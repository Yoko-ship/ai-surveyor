"""
Регистрация в мини-приложении, раздел «Пользователи» и выгрузки анализа.

Запуск из корня проекта (pytest и httpx в sandbox\\.venv не установлены — свой ASGI-клиент,
живой сервер не трогаем и не перезапускаем):
    set PYTHONIOENCODING=utf-8
    sandbox\\.venv\\Scripts\\python.exe tests\\test_registration.py

Сеть не используется: initData подписываем вымышленным токеном, отправка в Telegram подменяется
сборщиком исходящих. Все тестовые записи удаляются в конце (рабочая база остаётся чистой).

Что проверяется:
   1  код приходит в чат того же telegram_id, в ответе кода нет;
   2  повторная отправка раньше минуты запрещена;
   3  неверный код и исчерпание пяти попыток;
   4  бот не может написать («не нажимал Start») — честная ошибка, а не «код отправлен»;
   5  регистрация: пользователь сразу активен, роль «сотрудник», телефон в users.phone,
      два согласия в pd_consents (основное и телефон), подсказки департаментов;
   6  «Сделать админом» работает только у админа, последнего админа снять нельзя;
   7  чужой человек не может принять решение по запросу;
   8  выгрузки: чужому 403, участнику и админу 200; openpyxl читает созданный xlsx обратно;
   8а вероятность подтверждения: считается при отправке на согласование, приходит в API, попадает
      в PDF и XLSX, факт решения ложится рядом с прогнозом;
   9  телефон не попадает ни в audit, ни в tg_messages, ни в llm_calls; mask_pd его маскирует.
"""
import asyncio
import hashlib
import hmac
import json as _json
import os
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlencode

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Тест ходит в приложение напрямую с адреса 127.0.0.1 (см. tests/test_tgbot.py).
os.environ.setdefault("SURVEYOR_DEV", "1")

from tmpdb import temp_db  # noqa: E402  (tests/tmpdb.py)
from app import approvals, auth, db, exports, llm, outcomes, registration, telegram, tgbot  # noqa: E402
from app.main import app                                                                    # noqa: E402

TOKEN = "123456:TEST"                      # вымышленный токен бота, сеть не трогаем
BRANCH = "тест-рег"
TG_NEW, TG_ADMIN1, TG_ADMIN2, TG_STRANGER, TG_NOSTART = "9910001", "9910002", "9910003", "9910004", "9910005"
ALL_TG = [TG_NEW, TG_ADMIN1, TG_ADMIN2, TG_STRANGER, TG_NOSTART]
LOGINS = ["тест-рег-админ1", "тест-рег-админ2", "тест-рег-чужой"]
PHONE = "+998900000001"                    # вымышленный номер: не принадлежит никому
# с 21.09.2026 в анкете одно поле «Имя»: одного слова достаточно
PD_NAME = "Регистратов"
DEPARTMENT = "Тестовый департамент андеррайтинга"

OUT = []                                   # собранные исходящие: (метод, payload)
NO_CHAT = set()                            # чаты, куда бот «не может писать»
MARK = {"audit_id": 0}
TOKENS = {}                                # логин -> токен сессии


# ---------- минимальный ASGI-клиент ----------

def call(method: str, path: str, body=None, params=None, who=None, raw=False):
    query = urlencode(params or {}, encoding="utf-8")
    payload = _json.dumps(body, ensure_ascii=False).encode() if body is not None else b""
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method,
             "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "",
             "query_string": query.encode(),
             "headers": [(b"host", b"test"), (b"content-type", b"application/json"),
                         (b"content-length", str(len(payload)).encode())]
                        + ([(b"cookie", f"sid={TOKENS[who]}".encode())] if who else []),
             "client": ("127.0.0.1", 0), "server": ("test", 80)}
    out = {"status": None, "chunks": []}

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"] = msg["status"]
        elif msg["type"] == "http.response.body":
            out["chunks"].append(msg.get("body") or b"")

    asyncio.run(app(scope, receive, send))
    blob = b"".join(out["chunks"])
    if raw:
        return out["status"], blob
    try:
        return out["status"], _json.loads(blob.decode("utf-8"))
    except (ValueError, UnicodeDecodeError):
        return out["status"], blob


# ---------- подделки Telegram ----------

def init_data(tg_id: str, auth_date=None) -> str:
    """Собирает initData так, как это делает Telegram (тот же алгоритм, что в tests/test_telegram.py)."""
    user = {"id": int(tg_id), "first_name": "Тест", "username": "t" + tg_id}
    fields = {"query_id": "AAEreg", "user": _json.dumps(user, ensure_ascii=False),
              "auth_date": str(auth_date or int(time.time()))}
    dcs = telegram.data_check_string(list(fields.items()))
    secret = hmac.new(b"WebAppData", TOKEN.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, dcs.encode(), hashlib.sha256).hexdigest()
    return urlencode(fields)


def fake_deliver(method, payload):
    OUT.append((method, payload))
    if str(payload.get("chat_id")) in NO_CHAT:
        return {"ok": False, "description": "Forbidden: bot can't initiate conversation with a user"}
    return {"ok": True, "result": {"message_id": len(OUT)}}


def fake_deliver_file(method, fields, field, filename, blob, mime):
    OUT.append((method, dict(fields, filename=filename, bytes=len(blob), mime=mime)))
    return {"ok": True, "result": {"message_id": len(OUT)}}


def texts_for(tg_id) -> list:
    return [p.get("text", "") for m, p in OUT if str(p.get("chat_id")) == str(tg_id)]


def code_from_chat(tg_id) -> str:
    for t in reversed(texts_for(tg_id)):
        m = re.search(r"(\d{6})", t)
        if m:
            return m.group(1)
    raise AssertionError("код в чат не приходил: " + str(texts_for(tg_id)))


# ---------- подготовка и уборка ----------

def setup():
    db.ensure_schema()
    with db.tx() as con:
        MARK["audit_id"] = con.execute("SELECT COALESCE(MAX(id), 0) FROM audit").fetchone()[0]
        for tg in ALL_TG:
            con.execute("DELETE FROM users WHERE telegram_id=?", (tg,))
            con.execute("DELETE FROM reg_codes WHERE telegram_id=?", (tg,))
            con.execute("DELETE FROM pd_consents WHERE telegram_id=?", (tg,))
        ts = db.now()
        ids = {}
        for login, role, tg in ((LOGINS[0], "админ", TG_ADMIN1), (LOGINS[1], "админ", TG_ADMIN2),
                                (LOGINS[2], "андеррайтер", TG_STRANGER)):
            con.execute("DELETE FROM users WHERE login=?", (login,))
            cur = con.execute("INSERT INTO users (login, full_name, role, branch, department, position,"
                              " password_hash, salt, status, telegram_id, created_at)"
                              " VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                              (login, f"Тестовый {login}", role, BRANCH, "Тестовый департамент",
                               "менеджер", "x", "y", "активен", tg, ts))
            ids[login] = cur.lastrowid
    return ids


def make_request(author_id: int) -> int:
    """Запрос с объектом, расчётом, оценкой стоимости и документом — на нём проверяем выгрузки."""
    with db.tx() as con:
        cur = con.execute("INSERT INTO requests (external_no, branch, product_code, policyholder,"
                          " created_by_user_id, created_at, status) VALUES (?,?,?,?,?,?,?)",
                          ("ТЕСТ-РЕГ-1", BRANCH, "0311", "ООО «Тест-Рег»", author_id, db.now(), "посчитан"))
        rid = cur.lastrowid
        cur = con.execute("INSERT INTO objects (request_id, object_type, address, region, value_amount,"
                          " sum_insured, attributes) VALUES (?,?,?,?,?,?,?)",
                          (rid, "Легковой автомобиль", "г. Ташкент", "Ташкент", 100_000_000.0,
                           100_000_000.0, _json.dumps({"factors": {}}, ensure_ascii=False)))
        oid = cur.lastrowid
        chain = [{"name": "Базовая ставка", "value_pct": 1.5},
                 {"name": "Возраст автомобиля", "mult": 1.2}]
        con.execute("INSERT INTO calculations (request_id, object_id, net_rate_pct, risk_load_pct,"
                    " cat_load_pct, gross_rate_pct, min_rate_pct, applied_rate_pct, premium, verdict,"
                    " explanation, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (rid, oid, 1.5, 0.1, 0.0, 1.8, 2.0, 2.0, 2_000_000.0, "можно принять",
                     _json.dumps({"chain": chain, "manual": False}, ensure_ascii=False), db.now()))
        con.execute("INSERT INTO documents (request_id, doc_name, received) VALUES (?,?,1)",
                    (rid, "Техпаспорт"))
        cur = con.execute("INSERT INTO valuations (created_at, request_id, object_type, declared_value,"
                          " ai_value, method, method_version) VALUES (?,?,?,?,?,?,?)",
                          (db.now(), rid, "Легковой автомобиль", 100_000_000.0, 98_000_000.0,
                           "объявления", "v1"))
        vid = cur.lastrowid
        con.execute("INSERT INTO valuation_sources (valuation_id, source, status, url, ads_count, median,"
                    " fetched_at) VALUES (?,?,?,?,?,?,?)",
                    (vid, "avtoelon.uz", "ок", "https://avtoelon.uz/", 5, 98_000_000.0, db.now()))
    return rid


def teardown(rid):
    with db.tx() as con:
        con.execute("DELETE FROM request_reviewers WHERE request_id=?", (rid,))
        con.execute("DELETE FROM decision_outcomes WHERE request_id=?", (rid,))
        con.execute("DELETE FROM valuation_sources WHERE valuation_id IN"
                    " (SELECT id FROM valuations WHERE request_id=?)", (rid,))
        con.execute("DELETE FROM valuations WHERE request_id=?", (rid,))
        con.execute("DELETE FROM documents WHERE request_id=?", (rid,))
        con.execute("DELETE FROM calculations WHERE request_id=?", (rid,))
        con.execute("DELETE FROM object_perils WHERE object_id IN"
                    " (SELECT id FROM objects WHERE request_id=?)", (rid,))
        con.execute("DELETE FROM objects WHERE request_id=?", (rid,))
        con.execute("DELETE FROM requests WHERE id=?", (rid,))
        con.execute("DELETE FROM audit WHERE entity=?", (f"request:{rid}",))
        for tg in ALL_TG:
            con.execute("DELETE FROM sessions WHERE user_id IN (SELECT id FROM users WHERE telegram_id=?)", (tg,))
            con.execute("DELETE FROM audit WHERE who IN (SELECT login FROM users WHERE telegram_id=?)", (tg,))
            con.execute("DELETE FROM audit WHERE who=? OR entity=?", (f"tg{tg}", f"telegram:{tg}"))
            con.execute("DELETE FROM pd_consents WHERE telegram_id=?", (tg,))
            con.execute("DELETE FROM reg_codes WHERE telegram_id=?", (tg,))
            con.execute("DELETE FROM tg_messages WHERE telegram_id=?", (tg,))
            con.execute("DELETE FROM users WHERE telegram_id=?", (tg,))
        for login in LOGINS:
            con.execute("DELETE FROM audit WHERE who=?", (login,))
            con.execute("DELETE FROM users WHERE login=?", (login,))
        left_u = con.execute("SELECT COUNT(*) FROM users WHERE branch=?", (BRANCH,)).fetchone()[0]
        left_r = con.execute("SELECT COUNT(*) FROM requests WHERE branch=?", (BRANCH,)).fetchone()[0]
        left_c = con.execute("SELECT COUNT(*) FROM reg_codes").fetchone()[0]
    print(f"  очищено; тестовых людей осталось: {left_u}; запросов branch='{BRANCH}': {left_r}; "
          f"строк в reg_codes: {left_c}")


def age_code(tg_id, seconds: int):
    """Состаривает отправленный код, чтобы не ждать минуту в тесте."""
    from datetime import datetime, timedelta
    with db.tx() as con:
        con.execute("UPDATE reg_codes SET created_at=? WHERE telegram_id=?",
                    ((datetime.now() - timedelta(seconds=seconds)).isoformat(timespec="seconds"), tg_id))


# ---------- проверки ----------

def check_send_code():
    st, b = call("POST", "/tg/register/send-code", {"initData": init_data(TG_NEW), "phone": PHONE})
    assert st == 200 and b["sent"] is True, (st, b)
    code = code_from_chat(TG_NEW)
    assert code not in _json.dumps(b, ensure_ascii=False), b        # кода в ответе сервера нет
    with db.tx() as con:
        row = db.rows(con, "SELECT * FROM reg_codes WHERE telegram_id=?", TG_NEW)[0]
    assert code not in _json.dumps(row, ensure_ascii=False), row    # и в базе он только отпечатком
    assert PHONE not in _json.dumps(row, ensure_ascii=False), row   # номер в reg_codes не хранится
    assert row["salt"] and len(row["salt"]) >= 16, row              # своя соль на каждый код
    assert row["code_hash"] != hashlib.sha256((code + TG_NEW).encode()).hexdigest(), row
    print(f"1. код отправлен в чат {TG_NEW}; в ответе и в базе кода нет, номера тоже, "
          f"хэш кода с индивидуальной солью — ок")

    st, b = call("POST", "/tg/register/send-code", {"initData": init_data(TG_NEW), "phone": PHONE})
    assert st == 429 and "не чаще одного раза в минуту" in b["detail"], (st, b)
    print("2. повторная отправка раньше минуты — отказ 429 — ок")
    return code


def check_wrong_code(code: str):
    body = {"initData": init_data(TG_NEW), "phone": PHONE, "code": "000000" if code != "000000" else "111111"}
    for n in range(4):
        st, b = call("POST", "/tg/register/verify-code", body)
        assert st == 400 and "Осталось попыток" in b["detail"], (n, st, b)
    st, b = call("POST", "/tg/register/verify-code", body)
    assert st == 429 and "Попытки исчерпаны" in b["detail"], (st, b)
    # даже верный код после исчерпания попыток не принимается
    st, b = call("POST", "/tg/register/verify-code", {"initData": init_data(TG_NEW), "phone": PHONE, "code": code})
    assert st == 429, (st, b)
    print("3. неверный код: пять попыток, потом код сгорает — верный уже не принимается — ок")


def check_daily_limit():
    """Не больше пяти кодов в сутки на один telegram_id (раздел 7 документа юриста)."""
    with db.tx() as con:
        con.execute("DELETE FROM reg_codes WHERE telegram_id=?", (TG_NEW,))
    for n in range(registration.DAILY_MAX_CODES):
        age_code(TG_NEW, 120)
        st, b = call("POST", "/tg/register/send-code", {"initData": init_data(TG_NEW), "phone": PHONE})
        assert st == 200, (n, st, b)
    age_code(TG_NEW, 120)
    st, b = call("POST", "/tg/register/send-code", {"initData": init_data(TG_NEW), "phone": PHONE})
    assert st == 429 and "не больше" in b["detail"], (st, b)
    with db.tx() as con:            # для следующих проверок начинаем сутки заново
        con.execute("DELETE FROM reg_codes WHERE telegram_id=?", (TG_NEW,))
    print(f"2а. суточный лимит {registration.DAILY_MAX_CODES} кодов на telegram_id — шестой отклонён — ок")


def check_no_start():
    NO_CHAT.add(TG_NOSTART)
    st, b = call("POST", "/tg/register/send-code", {"initData": init_data(TG_NOSTART), "phone": "+998900000009"})
    assert st == 409 and "Start" in b["detail"], (st, b)
    with db.tx() as con:
        assert not db.rows(con, "SELECT 1 FROM reg_codes WHERE telegram_id=?", TG_NOSTART)
    print("4. бот не может написать: честная ошибка «нажмите Start», код не заводится — ок")


def check_register():
    age_code(TG_NEW, 120)                      # минута прошла — можно слать заново
    st, b = call("POST", "/tg/register/send-code", {"initData": init_data(TG_NEW), "phone": PHONE})
    assert st == 200, (st, b)
    code = code_from_chat(TG_NEW)
    st, b = call("POST", "/tg/register/verify-code", {"initData": init_data(TG_NEW), "phone": PHONE, "code": code})
    assert st == 200 and b["verified"] is True, (st, b)

    form = {"initData": init_data(TG_NEW), "phone": PHONE, "name": PD_NAME,
            "department": DEPARTMENT, "position": "менеджер", "branch": BRANCH,
            "consent": False, "consent_phone": True}
    st, b = call("POST", "/tg/register/submit", form)
    assert st == 422 and "согласия" in b["detail"], (st, b)         # без согласия — отказ

    st, b = call("POST", "/tg/register/submit", form | {"consent": True, "name": "  Ж  "})
    assert st == 422 and "имя" in b["detail"].lower(), (st, b)      # имя короче 2 символов — отказ

    form["consent"] = True
    st, b = call("POST", "/tg/register/submit", form)
    assert st == 200 and b["role"] == "сотрудник" and b["status"] == "активен", (st, b)
    assert b.get("token"), b
    TOKENS["новый"] = b["token"]
    with db.tx() as con:
        u = db.rows(con, "SELECT * FROM users WHERE telegram_id=?", TG_NEW)[0]
        cons = db.rows(con, "SELECT * FROM pd_consents WHERE telegram_id=? ORDER BY id", TG_NEW)
        left = db.rows(con, "SELECT 1 FROM reg_codes WHERE telegram_id=?", TG_NEW)
    assert u["status"] == "активен" and u["role"] == "сотрудник", u
    assert u["phone"] == PHONE and u["full_name"] == PD_NAME, u
    assert u["department"] == DEPARTMENT and u["position"] == "менеджер", u
    assert {c["scope"] for c in cons} == {"основное", "телефон"}, cons
    assert all(c["version"] and c["consent_text_hash"] and c["user_id"] == u["id"] for c in cons), cons
    assert not left, left                                            # код после регистрации удалён
    print(f"5. регистрация: пользователь № {u['id']} сразу «{u['status']}», роль «{u['role']}», "
          f"согласий в pd_consents: {len(cons)} ({', '.join(c['scope'] for c in cons)}) — ок")

    st, deps = call("GET", "/tg/register/departments")
    assert st == 200 and DEPARTMENT in deps["items"], deps
    st, pos = call("GET", "/tg/register/positions")
    assert st == 200 and "ассистент менеджера" in pos["items"], pos
    st, c = call("GET", "/tg/consent", params={"scope": "телефон"})
    assert st == 200 and c["scope"] == "телефон" and c["hash"], c
    st, me = call("GET", "/tg/me", who="новый")
    assert st == 200 and me["status"] == "активен" and "users" in [n["key"] for n in me["nav"]], me
    assert "расчёт" in me["rights"] and "все запросы" not in me["rights"], me
    print("6. подсказки департаментов, список должностей, текст согласия и разделы сотрудника — ок")
    return u["id"]


def check_users(ids, new_uid):
    st, b = call("GET", "/tg/users", who="новый")
    assert st == 200 and b["can_manage"] is False, (st, b)
    mine = [x for x in b["items"] if x["id"] == new_uid][0]
    assert mine["phone"] == PHONE and mine["department"] == DEPARTMENT and mine["is_admin"] is False, mine
    assert any(x["is_admin"] for x in b["items"]), b["items"][:3]

    st, b = call("POST", f"/tg/users/{new_uid}/make-admin", who="новый")
    assert st == 403, (st, b)                                        # не админ — нельзя
    st, b = call("POST", f"/tg/users/{new_uid}/make-admin", who=LOGINS[0])
    assert st == 200 and b["role"] == "админ", (st, b)
    with db.tx() as con:
        assert db.rows(con, "SELECT role FROM users WHERE id=?", new_uid)[0]["role"] == "админ"
    assert any("администратора" in t for t in texts_for(TG_NEW)), texts_for(TG_NEW)

    st, b = call("POST", f"/tg/users/{new_uid}/revoke-admin", who=LOGINS[0])
    assert st == 200 and b["role"] == "сотрудник", (st, b)
    with db.tx() as con:
        admins = db.rows(con, "SELECT id, login FROM users WHERE role='админ' AND status='активен'")
        others = [a["id"] for a in admins if a["login"] not in LOGINS[:2]]
        # оставляем ровно одного админа из тестовых и проверяем защиту от снятия последнего
        con.execute("UPDATE users SET role='андеррайтер' WHERE login=?", (LOGINS[1],))
    if not others:
        st, b = call("POST", f"/tg/users/{ids[LOGINS[0]]}/revoke-admin", who=LOGINS[0])
        assert st == 422 and "единственный" in b["detail"], (st, b)
        print("7. «Сделать админом» — только у админа; последнего админа снять нельзя — ок")
    else:
        print(f"7. «Сделать админом» — только у админа; кроме тестовых есть ещё {len(others)} "
              f"администратор(ов), проверку «последнего» пропускаем — ок")
    with db.tx() as con:                                              # возвращаем роль второму админу
        con.execute("UPDATE users SET role='админ' WHERE login=?", (LOGINS[1],))
    # в журнале — только логин и событие: ни ФИО, ни телефона
    with db.tx() as con:
        det = " ".join((r["detail"] or "") + (r["who"] or "")
                       for r in db.rows(con, "SELECT who, detail FROM audit WHERE id>?", MARK["audit_id"]))
    assert PD_NAME not in det and PHONE not in det, det[:200]


def check_decide_and_exports(ids, rid, new_uid):
    # кого можно выбрать рассматривающими: только действующие админы, без самого инициатора
    st, cand = call("GET", f"/requests/{rid}/reviewer-candidates", who="новый")
    assert st == 200 and cand["min"] == 1 and cand["max"] == 3, (st, cand)
    got = {c["id"] for c in cand["items"]}
    assert ids[LOGINS[0]] in got and new_uid not in got, cand
    assert all(c.get("full_name") and "position" in c and "department" in c for c in cand["items"]), cand
    print(f"7а. список рассматривающих: {len(cand['items'])} администратор(ов), "
          f"инициатора в списке нет — ок")

    with db.tx() as con:
        out = approvals.assign(con, rid, [ids[LOGINS[0]]], who=LOGINS[0])
    assert len(out["reviewers"]) == 1 and out["approval_status"] == "на согласовании", out
    print("8. отправка на согласование с одним рассматривающим (от 1 до 3) — ок")

    st, b = call("POST", f"/requests/{rid}/decide", {"decision": "одобрил"}, who=LOGINS[2])
    assert st == 403, (st, b)                                        # чужой решать не может
    print("9. чужой человек по запросу решение принять не может — 403 — ок")

    for path in (f"/requests/{rid}/analysis.pdf", f"/requests/{rid}/analysis.xlsx"):
        st, b = call("GET", path, who=LOGINS[2])
        assert st == 403, (path, st, b)
    print("10. чужому выгрузки анализа недоступны — 403 — ок")

    st, blob = call("GET", f"/requests/{rid}/analysis.pdf", who="новый", raw=True)
    assert st == 200 and blob[:4] == b"%PDF", (st, blob[:40])
    st, blob = call("GET", f"/requests/{rid}/analysis.pdf", who=LOGINS[0], raw=True)
    assert st == 200 and blob[:4] == b"%PDF", (st, blob[:40])
    print(f"11. analysis.pdf отдаётся инициатору и админу ({len(blob)} байт, начинается с %PDF) — ок")

    st, blob = call("GET", f"/requests/{rid}/analysis.xlsx", who="новый", raw=True)
    assert st == 200 and blob[:2] == b"PK", (st, blob[:40])
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob))
    assert wb.sheetnames == ["Запрос", "Ставка", "Проверки", "Стоимость", "Документы", "Решения",
                             "Вероятность"], wb.sheetnames
    values = [str(c.value) for row in wb["Запрос"].iter_rows() for c in row]
    assert any("ООО «Тест-Рег»" in v for v in values), values[:20]
    links = [str(c.value) for row in wb["Стоимость"].iter_rows() for c in row]
    assert any("avtoelon.uz" in v for v in links), links
    solutions = [str(c.value) for row in wb["Решения"].iter_rows() for c in row]
    assert any("ожидает" in v for v in solutions), solutions
    print(f"12. analysis.xlsx читается openpyxl обратно: листы {', '.join(wb.sheetnames)}; "
          f"ссылка на источник стоимости на месте — ок")

    # вероятность подтверждения: посчитана при отправке на согласование (пункт 8) и лежит в базе
    with db.tx() as con:
        row = outcomes.latest(con, rid)
        s = outcomes.summary(con, rid)
        stored = db.rows(con, "SELECT * FROM decision_outcomes WHERE request_id=?", rid)
    assert row and row["model_version"] == "prob-1", row
    assert len(stored) == 1, stored
    assert s["ready"] and 3 <= s["probability"] <= 97, s
    assert s["probability"] == int(round(row["probability"])), (s, row)
    assert s["calibrated"] == 0 and s["summary"] and "Вероятность подтверждения" in s["summary"], s
    assert row["product_code"] == "0311" and row["branch"] == BRANCH and row["calculation_id"], row
    print(f"13. вероятность подтверждения посчитана при отправке на согласование: "
          f"{s['probability']} % — «{s['verdict']}»; минусов {len(s['minus'])}, плюсов {len(s['plus'])}, "
          f"советов {len(s['how_to_raise'])} — ок")

    # то же число приходит в API: карточка запроса, «Ждут меня», «Мои запросы»
    st, card = call("GET", f"/requests/{rid}", who=LOGINS[0])
    assert st == 200 and card["probability"]["probability"] == s["probability"], card.get("probability")
    st, box = call("GET", "/approvals/inbox", who=LOGINS[0])
    mine_box = [i for i in box["items"] if i["request_id"] == rid]
    assert mine_box and mine_box[0]["probability"]["probability"] == s["probability"], box
    assert mine_box[0]["probability"]["summary"], mine_box[0]
    st, appr = call("GET", f"/requests/{rid}/approvals", who=LOGINS[0])
    assert st == 200 and appr["probability"]["probability"] == s["probability"], appr
    print(f"13а. вероятность приходит в GET /requests/{rid} (ключ probability), /approvals/inbox "
          f"и /requests/{rid}/approvals — одно и то же число {s['probability']} % — ок")

    # 21.09.2026: мини-апп только для аналитики — карточка согласующему в Telegram не уходит
    # (ни текста с вероятностью, ни PDF анализа во вложении), хотя запрос отправлен на согласование
    cards = [p for m, p in OUT if m in ("sendDocument", "sendPhoto")
             or "Вероятность подтверждения" in str(p.get("text", ""))]
    assert not cards, cards[:2]
    print("13б. отправка на согласование карточку и вложения в Telegram не шлёт — ок")

    # выгрузки: число и разбор вместо надписи «вероятность ещё не рассчитана»
    st, blob = call("GET", f"/requests/{rid}/analysis.xlsx", who=LOGINS[0], raw=True)
    assert st == 200
    import io
    from openpyxl import load_workbook
    wb2 = load_workbook(io.BytesIO(blob))
    assert "Вероятность" in wb2.sheetnames, wb2.sheetnames
    vals = [c.value for row in wb2["Вероятность"].iter_rows() for c in row if c.value is not None]
    assert s["probability"] in vals, vals[:10]
    assert any(isinstance(v, str) and "не калибрована" in v for v in vals), vals
    assert not any(isinstance(v, str) and outcomes.NOT_READY in v for v in vals), vals
    head = [c.value for row in wb2["Запрос"].iter_rows() for c in row]
    assert s["probability"] in head, head
    st, pdf_blob = call("GET", f"/requests/{rid}/analysis.pdf", who=LOGINS[0], raw=True)
    assert st == 200 and pdf_blob[:4] == b"%PDF" and len(pdf_blob) > len(blob) // 4, len(pdf_blob)
    assert outcomes.NOT_READY.encode("utf-8") not in pdf_blob        # текст PDF сжат, но проверка дешёвая
    print(f"13в. в analysis.xlsx появился лист «Вероятность» ({s['probability']} %, пометка про "
          f"калибровку), в analysis.pdf — раздел с разбором ({len(pdf_blob)} байт) — ок")

    # решение согласующего: факт ложится в ту же строку, где лежит прогноз
    with db.tx() as con:
        approvals.decide(con, rid, LOGINS[0], "одобрил", None)
        fact = outcomes.latest(con, rid)
        stored = db.rows(con, "SELECT * FROM decision_outcomes WHERE request_id=?", rid)
    assert fact["decision"] == "согласован" and fact["decided_at"], fact
    assert fact["decided_by"] == LOGINS[0], fact
    assert len(stored) == 1, stored                                  # факт не задвоился
    assert int(round(fact["probability"])) == s["probability"], fact
    print(f"13г. факт решения «{fact['decision']}» записан рядом с прогнозом {s['probability']} % "
          f"(строк в decision_outcomes по запросу: {len(stored)}) — ок")


def check_no_pd():
    needles = [PHONE, PD_NAME, "Регистратов"]
    bad = []
    with db.tx() as con:
        for table, where in (("tg_messages", ""), ("llm_calls", ""),
                             ("audit", f" WHERE id > {MARK['audit_id']}")):
            try:
                cols = [r[1] for r in con.execute(f"PRAGMA table_info({table})")]
                rows = db.rows(con, f"SELECT * FROM {table}{where}")
            except Exception:
                continue
            for row in rows:
                for c in cols:
                    v = row.get(c)
                    if isinstance(v, str) and any(n in v for n in needles):
                        bad.append((table, c, v[:80]))
    assert not bad, bad
    masked = llm.mask_pd(f"Позвоните на {PHONE}")
    assert PHONE not in masked and "[ТЕЛЕФОН]" in masked, masked
    print(f"14. телефон и ФИО не попали ни в audit, ни в tg_messages, ни в llm_calls; "
          f"маскировка даёт «{masked}» — ок")


def main():
    ids = setup()
    rid = make_request(0)                    # автор проставится после регистрации
    real = (telegram.bot_token, tgbot.bot_token, tgbot._deliver, tgbot._deliver_file)
    try:
        telegram.bot_token = lambda: TOKEN
        tgbot.bot_token = lambda: TOKEN
        tgbot._deliver = fake_deliver
        tgbot._deliver_file = fake_deliver_file
        code = check_send_code()
        check_wrong_code(code)
        check_daily_limit()
        check_no_start()
        new_uid = check_register()
        with db.tx() as con:                 # запрос подан только что зарегистрировавшимся человеком
            con.execute("UPDATE requests SET created_by_user_id=? WHERE id=?", (new_uid, rid))
        with db.tx() as con:
            for login in LOGINS:
                u = db.rows(con, "SELECT * FROM users WHERE login=?", login)[0]
                TOKENS[login], _ = auth.create_session(con, u, "127.0.0.1", "test")
        check_users(ids, new_uid)
        check_decide_and_exports(ids, rid, new_uid)
        check_no_pd()
        print("\nВсе проверки регистрации, пользователей и выгрузок пройдены.")
    finally:
        telegram.bot_token, tgbot.bot_token, tgbot._deliver, tgbot._deliver_file = real
        teardown(rid)


if __name__ == "__main__":
    with temp_db("surveyor-registration.db"):  # рабочая data/surveyor.db не меняется
        main()
