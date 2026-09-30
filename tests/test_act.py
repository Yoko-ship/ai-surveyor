"""
Сюрвейерский акт, лёгкая версия (app/act.py, app/act_engine.py, app/docx_lite.py) — ТЗ 2.0 от 29.09.2026.

Запуск из корня (свой ASGI-клиент, живой сервер не трогаем):
    set PYTHONIOENCODING=utf-8
    sandbox\\.venv\\Scripts\\python.exe tests\\test_act.py

Всё — во временной копии базы (tests/tmpdb.py) и во временной папке файлов.
Дополнения 29.09.2026 (проверки 21–27): сценарии PML/EML/MFL, разбор документов, франшиза, рекомендации. Сеть и модель подменяются:
llm.chat_raw отдаёт заготовленный ответ, llm._post бросает исключение (любой выход в сеть = ошибка теста).
Проверки 37а–37г (30.09.2026): запрос филиала — разбор DOCX/XLSX/PDF с 16 строками (оба образца заказчика),
ответ модели по скану, сверка с расчётом акта (продукт 0832), многолетний срок, физлицо, три языка, Word и PDF.
Проверки 38а–38ж (30.09.2026): договор страхования — учебный договор, договоры на узбекской кириллице и латинице,
русском и английском (DOCX и PDF с текстом), длинный договор, скан (ответ модели подменён), дочитывание текста
моделью с маскировкой ПД, сверка с расчётом акта (график, существенные условия ГК ст. 929), запрос филиала + договор.
Проверки 39а–39з (30.09.2026, вечер): источник условий решает сервер (правки «было → стало» в акте), вид документа по
заголовку (полис, заявление), существенные условия без категоричности, сверка «запрос ↔ договор» как при загрузке,
границы ввода, счёт не уходит в модель, стороны-юрлица, отрицательные суммы, срок разбора PDF.
Проверки 40а–40в (30.09.2026, шаблон договора 0102): чтение DOCX деревом XML (прогоны, w:tab, w:br, w:sdt) и новые
пределы (10 000 ячеек, 3 000 абзацев, 4 000 знаков в строке таблицы DOCX); бланк договора личного страхования на
выдуманных данных (заголовок в две строки, подчёркивания, таблица приложения 1, пп. 2.6 и 5.4) — is_template, «не
заполнено», без ст. 929 и сверки, модель не вызывается; тот же шаблон заполненный — полноценный договор со сверкой.
Ставки в проверках берутся из справочника копии базы (engine.rate_for / engine.min_rate), а не из головы.
"""
import asyncio
import io
import json as _json
import os
import re
import secrets
import shutil
import sys
import tempfile
import zipfile
from datetime import date, datetime, timedelta
from time import monotonic
from pathlib import Path
from urllib.parse import urlencode
from xml.dom import minidom

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

os.environ.pop("SURVEYOR_DEV", None)      # guard проверяем целиком, без режима разработчика
os.environ["SURVEYOR_NO_BACKGROUND"] = "1"

import pymupdf                            # noqa: E402

from tmpdb import temp_db                 # noqa: E402
from app import act, act_engine as ae, db, guest, llm   # noqa: E402
from app import act_market as am, act_texts as tx       # noqa: E402
from app.engine import Input, min_rate, premium_of, rate_for   # noqa: E402
from app.main import app                  # noqa: E402

passed, failed = 0, 0
COOKIES = {}
HEADERS = []                              # дополнительные заголовки (Bearer администратора)

CRANE_MUST = {"product_code": "0318", "sum_insured": 2_945_000_000, "object_value": 3_100_000_000,
              "region": "Ташкентская область"}
CRANE_OPT = {"location": "open_area", "losses_3y": {"count": 0, "small_count": 0, "amount": 0}}
CRANE_TYPE = "Спецтехника — автокран"


def ok(name, cond, extra=""):
    global passed, failed
    if cond:
        passed += 1
        print("  ок  ", name)
    else:
        failed += 1
        print("  ПЛОХО", name, str(extra)[:400])


# ------------------------------------------------------------------ ASGI-клиент

def _send(method, path, params, headers, payload, raw=False):
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method,
             "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "",
             "query_string": urlencode(params or {}, encoding="utf-8").encode(),
             "headers": headers + list(HEADERS), "client": ("203.0.113.9", 0), "server": ("test", 80)}
    out = {"status": None, "chunks": [], "headers": []}

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"], out["headers"] = msg["status"], msg.get("headers") or []
        elif msg["type"] == "http.response.body":
            out["chunks"].append(msg.get("body") or b"")

    asyncio.run(app(scope, receive, send))
    for k, v in out["headers"]:
        if k.lower() == b"set-cookie":
            pair = v.decode("latin-1").split(";")[0]
            name, _, value = pair.partition("=")
            COOKIES[name.strip()] = value.strip()
    body = b"".join(out["chunks"])
    if raw:
        return out["status"], body, dict((k.decode().lower(), v.decode("latin-1")) for k, v in out["headers"])
    text = body.decode("utf-8", "replace")
    try:
        return out["status"], _json.loads(text)
    except ValueError:
        return out["status"], text


def _cookie_hdr():
    return [(b"cookie", "; ".join(f"{k}={v}" for k, v in COOKIES.items()).encode())] if COOKIES else []


def call(method, path, body=None, params=None, raw=False):
    payload = _json.dumps(body, ensure_ascii=False).encode() if body is not None else b""
    hdrs = [(b"host", b"test"), (b"content-type", b"application/json"),
            (b"content-length", str(len(payload)).encode())] + _cookie_hdr()
    return _send(method, path, params, hdrs, payload, raw)


def upload(files, fields=None):
    boundary = "----insonact"
    parts = []
    for k, v in (fields or {}).items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
    for name, mime, blob in files:
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="files"; '
                     f'filename="{name}"\r\nContent-Type: {mime}\r\n\r\n'.encode() + blob + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    payload = b"".join(parts)
    hdrs = [(b"host", b"test"), (b"content-type", f"multipart/form-data; boundary={boundary}".encode()),
            (b"content-length", str(len(payload)).encode())] + _cookie_hdr()
    return _send("POST", "/act/photos", None, hdrs, payload)


def image(color=(200, 150, 0), kind="png", w=320, h=240) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.set_rect(pix.irect, color)
    return pix.tobytes("jpg" if kind == "jpg" else "png")


# ------------------------------------------------------------------ подмена модели

CALLS = []
REPLY = {"text": None}


def fake_chat_raw(purpose, messages, max_tokens=700, temperature=0.2, files=None, timeout=None, retries=None):
    CALLS.append({"purpose": purpose, "messages": messages, "files": files, "temperature": temperature,
                  "timeout": timeout, "retries": retries})
    if REPLY.get("sleep"):
        import time as _t
        _t.sleep(REPLY["sleep"])
    text = REPLY["text"](messages) if callable(REPLY["text"]) else REPLY["text"]
    return {"text": text, "ok": bool(text), "notes": [], "ms": 1, "reason": None if text else "пустой ответ"}


def no_network(*a, **kw):
    raise AssertionError("тест не должен ходить в сеть")


ORIG = {}


def model_on(on=True, files=True):
    llm.enabled = lambda: on
    llm.supports_files = lambda: on and files


CRANE_REPLY = "```json\n" + _json.dumps({
    "files": [{"n": 1, "view": "front"}, {"n": 2, "view": "left"}, {"n": 3, "view": "plate"},
              {"n": 4, "view": "document", "document_kind": "лист технических параметров"}],
    "object_kind": "truck_crane", "class_hint": "special_machinery", "condition": "new",
    "fields": [
        {"key": "object_type", "value": "автокран", "source": "photo", "file": 1},
        {"key": "brand", "value": "XCMG", "source": "marking", "file": 2},
        {"key": "model", "value": "QY50K5D", "source": "marking", "file": 2},
        {"key": "model", "value": "XCMG QY50K5D", "source": "document", "file": 4},
        {"key": "year", "value": "2026", "source": "document", "file": 4},
        {"key": "manufacture_date", "value": "2026-03", "source": "plate", "file": 3},
        {"key": "serial_no", "value": "LXGCPA393TA006921", "source": "plate", "file": 3},
        {"key": "serial_no", "value": "LXGCPA393TA006921", "source": "document", "file": 4},
        {"key": "manufacturer", "value": "Xuzhou Construction Machinery Group Co., Ltd.", "source": "plate",
         "file": 3},
        {"key": "curb_mass", "value": "36 170 кг", "source": "plate", "file": 3},
        {"key": "curb_mass", "value": "38 600 кг", "source": "document", "file": 4},
        {"key": "engine_power", "value": "248 кВт", "source": "plate", "file": 3},
        {"key": "engine_power", "value": "251 кВт", "source": "document", "file": 4},
        {"key": "location", "value": "открытая площадка", "source": "photo", "file": 1},
        {"key": "color", "value": None, "source": "photo", "file": 1},
        {"key": "owner", "value": "что-то чужое", "source": "document", "file": 4},
    ],
    "damages": []}, ensure_ascii=False) + "\n```"

CRANE_FILES = [("front.jpg", "image/jpeg", image(kind="jpg")), ("left.png", "image/png", image((250, 200, 0))),
               ("plate.png", "image/png", image((90, 90, 90))), ("sheet.png", "image/png", image((255, 255, 255)))]


def expected_rate(con, level_adj: float, product="0318", cls="3", otype=CRANE_TYPE):
    """Ручной расчёт по справочнику копии базы: ставка продукта по тарифной политике (если её нет —
    техническая ставка без коэффициентов), поправка по уровню, минимум."""
    ref = db.load_reference(con)
    mr = min_rate(ref, product)
    base = mr.get("company")
    if base is None:
        base = rate_for(ref, Input(product_code=product, class_code=cls, object_type=otype, value_amount=1,
                                   sum_insured=1, factors={}))["gross_pct"]
    floor = mr["floor"]
    applied = round(max(base * (1 + level_adj / 100), floor or 0), 4)
    return round(base, 4), floor, applied


def all_text(a: dict) -> str:
    out = [a["title"], a["footer"]]
    for s in a["sections"]:
        out += [s["title"]] + s["paragraphs"]
        out += [f"{r['label']} {r['value']} {r.get('note') or ''}" for r in s["rows"]]
        for li in s.get("lists") or []:
            out += [li["title"]] + li["items"]
    return "\n".join(out)


# ------------------------------------------------------------------ 1. пример заказчика

def check_crane():
    print("1. Пример заказчика: автокран XCMG QY50K5D")
    REPLY["text"] = CRANE_REPLY
    model_on(True)
    CALLS.clear()
    st, b = upload(CRANE_FILES, {"lang": "ru", "product_code": "0318"})
    ok("фото приняты", st == 200 and b.get("ok"), (st, b))
    ok("гостю выдана cookie", bool(COOKIES.get("gid")))
    ok("один запрос к модели со всеми файлами", len(CALLS) == 1 and len(CALLS[0]["files"]) == 4, len(CALLS))
    ok("низкая температура", CALLS[0]["temperature"] <= 0.2)
    sys_prompt = CALLS[0]["messages"][0]["content"] + CALLS[0]["messages"][1]["content"]
    ok("в инструкции «не выдумывай» и запрет на данные людей",
       "не выдумывай" in sys_prompt and "не извлекай" in sys_prompt and "посторонних машин" in sys_prompt)
    ok("маскировка ПД не портит инструкцию", llm.mask_pd(sys_prompt) == sys_prompt)
    ok("имена файлов в модель не уходят",
       all(f["name"].startswith("file ") for f in CALLS[0]["files"]) and "front.jpg" not in sys_prompt)
    ok("PNG пережат в JPEG перед отправкой", all(f["mime"] == "image/jpeg" for f in CALLS[0]["files"]))
    views = {f["index"]: f["view"] for f in b["files"]}
    ok("ракурсы распознаны (сопоставление по номеру файла в запросе)",
       views == {1: "front", 2: "left", 3: "plate", 4: "document"}, views)
    ok("распознавание — одна попытка с таймаутом 20 с", CALLS[0]["retries"] == 0 and CALLS[0]["timeout"] == 20,
       (CALLS[0]["retries"], CALLS[0]["timeout"]))
    keys = [(r["key"], r["source"]) for r in b["recognized"]]
    ok("оба значения массы вернулись (табличка и документ)",
       ("curb_mass", "plate") in keys and ("curb_mass", "document") in keys, keys)
    ok("пустые и неизвестные поля отброшены",
       not any(k == "color" for k, _ in keys) and not any(k == "owner" for k, _ in keys))
    ok("название изготовителя не принято за ФИО",
       any(r["key"] == "manufacturer" for r in b["recognized"]), keys)
    ok("у каждого значения источник словами и «проверьте»",
       all(r["source_label"] and r["check_label"] == "проверьте" for r in b["recognized"]))
    miss = {v["code"] for v in b["missing_views"]}
    ok("не хватает ракурсов: сзади, правый борт, счётчик", miss == {"back", "right", "odometer"}, miss)
    ok("предупреждение про данные людей и тестовый сервер",
       "данными людей" in b["warning"] and "тестовый" in b["warning"])
    ok("ai = true", b["ai"] is True)
    ok("вид объекта — автокран", (b.get("object_kind") or {}).get("code") == "truck_crane")
    sid = b["session"]

    CALLS.clear()
    st, a = call("POST", "/act/make", {"session": sid, "lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT,
                                       "recognized": b["recognized"]})
    ok("/act/make к модели не обращается", not CALLS, len(CALLS))
    ok("акт сформирован", st == 200 and a.get("ok"), (st, a))
    ok("пять разделов", [s["n"] for s in a["sections"]] == [1, 2, 3, 4, 5])
    ok("строка о подтверждении андеррайтером",
       a["footer"] == "Акт сформирован ИИ-сюрвейером, подлежит подтверждению андеррайтером")
    ok("шапка акта", a["title"] == "СЮРВЕЙЕРСКИЙ АКТ ПРЕДСТРАХОВОГО ОСМОТРА" and a["number"] and a["date"])
    ok("уровень риска по правилу — низкий (1 повышает, 3 снижают)",
       a["risk"]["level"] == "low" and a["risk"]["up"] == 1 and a["risk"]["down"] == 3, a["risk"])
    ok("признаки с объяснением", len(a["risk"]["factors"]) == 4 and all(f["text"] for f in a["risk"]["factors"]))
    with db.tx() as con:
        base, floor, applied = expected_rate(con, 0)
    ok("базовая ставка = ставка продукта по тарифной политике",
       abs(a["rate"]["base_pct"] - base) < 1e-9, (a["rate"]["base_pct"], base))
    ok("поправка низкого уровня 0 %", a["rate"]["adj_pct"] == 0)
    ok("итоговая ставка совпадает с ручным расчётом", abs(a["rate"]["applied_pct"] - applied) < 1e-9,
       (a["rate"]["applied_pct"], applied))
    ok("минимум продукта из справочника", a["rate"]["min_pct"] == floor, (a["rate"]["min_pct"], floor))
    prem = round(applied / 100 * CRANE_MUST["sum_insured"])
    ok("премия = сумма × ставка", a["premium"]["amount"] == prem, (a["premium"]["amount"], prem))
    ok("калибровка помечена", a["rate"]["calibrated"] == 0 and a["risk"]["calibrated"] == 0)
    ok("тип объекта для ставки — автокран", a["rate"]["object_type"] == CRANE_TYPE, a["rate"]["object_type"])
    ok("версия тарифа сохранена", a["rate"]["tariff_version_id"] is not None)
    ok("сумма ≈95 % стоимости — «в норме»",
       a["value"]["verdict"] == "normal" and abs(a["value"]["ratio_pct"] - 95.0) < 0.01, a["value"])
    ok("текст раздела 3 «В норме»", "В норме" in a["value"]["text"], a["value"]["text"])
    ok("«Франшиза не требуется»", a["franchise"]["needed"] is False and a["franchise"]["text"] ==
       "Франшиза не требуется", a["franchise"])
    codes = {c["code"] for c in a["clauses"]}
    ok("оговорки спецтехники", {"sp_attachments_storage", "sp_territory", "sp_reinspection"} <= codes, codes)
    ok("оговорки помечены экспертными", all(c["expert"] for c in a["clauses"]))
    ok("повторный осмотр через 12 месяцев или 2 000 моточасов",
       any("12 месяцев" in c["text"] and "2 000 моточасов" in c["text"] for c in a["clauses"]))
    dk = {d["key"]: d for d in a["discrepancies"]}
    ok("расхождение по массе и мощности", set(dk) == {"curb_mass", "engine_power"}, list(dk))
    ok("серийный номер и модель (с маркой в документе) не считаются расхождением",
       "serial_no" not in dk and "model" not in dk)
    s5 = all_text({"title": "", "footer": "", "sections": [a["sections"][4]]})
    ok("в заключении: масса 36 170 кг на табличке и 38 600 кг в документе",
       "на табличке 36 170 кг" in s5 and "в документе 38 600 кг" in s5, s5[:600])
    ok("в заключении: мощность 248 кВт и 251 кВт", "248 кВт" in s5 and "251 кВт" in s5)
    ok("приоритет у документа с печатью производителя", "печатью производителя" in s5)
    ok("решение: принять с оговорками", a["decision"]["code"] == "accept_with_clauses", a["decision"])
    ok("андеррайтеру: снять расхождение и дозапросить фото",
       any("Снаряжённая масса" in c for c in a["decision"]["checks"])
       and any("правый борт" in c for c in a["decision"]["checks"]), a["decision"]["checks"])
    s1 = {r["label"]: r for r in a["sections"][0]["rows"]}
    ok("в разделе 1 заводской номер из документа",
       s1["Заводской (серийный) номер"]["value"] == "LXGCPA393TA006921"
       and "из документа" in s1["Заводской (серийный) номер"]["note"], s1.get("Заводской (серийный) номер"))
    ok("в разделе 1 масса из документа и вариант с таблички",
       s1["Снаряжённая масса"]["value"] == "38 600 кг" and "36 170 кг" in s1["Снаряжённая масса"]["note"])
    # 30.09.2026: «данные недоступны» — не больше трёх строк, остальное одной строкой «Не указано: …»
    na_rows = [r for r in a["sections"][0]["rows"] if r["value"] == "данные недоступны"]
    ok("чего нет — «данные недоступны» не больше трёх строк, остальное — «Не указано»",
       len(na_rows) <= 3 and "Цвет" not in s1 and "цвет" in (s1.get("Не указано") or {}).get("value", ""),
       ([r["label"] for r in na_rows], s1.get("Не указано")))
    ok("цвет в списке недоступных данных", "Цвет" in a["missing"], a["missing"])
    text = all_text(a)
    ok("цифры в тексте совпадают с расчётом",
       act.pct(applied, "ru") in text and act.money(prem, "ru") in text)
    return a, sid


# ------------------------------------------------------------------ 2. сумма и стоимость

def check_value():
    print("2. Недострахование и превышение")
    st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(CRANE_MUST, sum_insured=2_000_000_000),
                                       "optional": CRANE_OPT})
    ok("< 90 % — недострахование, ГК ст. 936",
       st == 200 and a["value"]["verdict"] == "under" and a["value"]["legal_ref"] == "ГК РУз, ст. 936"
       and "пропорциональной" in a["value"]["text"], a.get("value"))
    ok("андеррайтеру — про пропорциональную выплату",
       any("пропорциональн" in c for c in a["decision"]["checks"]))
    st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(CRANE_MUST, sum_insured=3_500_000_000),
                                       "optional": CRANE_OPT})
    ok("> 100 % — превышение, ГК ст. 938",
       st == 200 and a["value"]["verdict"] == "over" and a["value"]["legal_ref"] == "ГК РУз, ст. 938"
       and "снизить" in a["value"]["text"], a.get("value"))
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, price_new=3_000_000_000, purchase_year=
                                                        date.today().year - 2)})
    dep = a["value"]["depreciated"]
    ok("ориентир с износом: транспорт 20 % в год", dep and dep["wear_pct_per_year"] == 20
       and dep["value"] == round(3_000_000_000 * (1 - 0.2 * 2)), dep)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(CRANE_MUST, sum_insured="12abc"),
                                       "optional": dict(CRANE_OPT, year=1800)})
    ok("мусор во вводе — 422 с полями", st == 422 and "sum_insured" in a.get("errors", {})
       and "year" in a.get("errors", {}), a)
    with db.tx() as con:
        row = db.rows(con, "SELECT detail FROM audit WHERE action='акт: ошибка ввода' ORDER BY id DESC LIMIT 1")
    ok("ошибка ввода — в журнале, без значений", row and "sum_insured" in row[0]["detail"]
       and "12abc" not in row[0]["detail"], row)


# ------------------------------------------------------------------ 3. франшиза

def check_franchise():
    print("3. Франшиза — только при основании")
    from app.risk_analytics import load_thresholds
    with db.tx() as con:
        th = load_thresholds(con)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, losses_3y={"count": 2, "small_count": 2})})
    fr = a["franchise"]
    band = th["franchise_by_level"][ae.RA_LEVEL[a["risk"]["level"]]]
    cap = th["franchise_class_caps"].get("3")
    hi = min(band[1], cap) if cap is not None else band[1]
    ok("2 мелких убытка — франшиза советуется", fr["needed"] and any(g["code"] == "fr_g_small_losses"
                                                                     for g in fr["grounds"]), fr)
    ok("размер — вилка из порогов (не выдуман)", fr.get("size") and fr["size"]["to_pct"] == hi
       and fr["size"]["from_pct"] == min(band[0], hi), (fr.get("size"), band, cap))
    ok("размер определяет андеррайтер", "андеррайтер" in fr["text"], fr["text"])
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, want_lower_premium=True)})
    ok("просьба клиента — основание", a["franchise"]["needed"] and
       [g["code"] for g in a["franchise"]["grounds"]] == ["fr_g_client"], a["franchise"])
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})
    ok("без оснований — «Франшиза не требуется»", a["franchise"]["text"] == "Франшиза не требуется")
    # обязательный вид: строительно-монтажные риски, обязательное (0820) — тариф по ПКМ, без поправок
    must = {"product_code": "0820", "sum_insured": 1_000_000_000, "object_value": 1_000_000_000,
            "region": "Ташкент"}
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must,
                                       "optional": {"losses_3y": {"count": 3, "small_count": 3},
                                                    "want_lower_premium": True, "location": "construction"}})
    with db.tx() as con:
        rt = db.rows(con, "SELECT rate_text FROM products WHERE code='0820'")[0]["rate_text"]
    act_rate = float(re.search(r"(\d+(?:,\d+)?)\s*%", rt).group(1).replace(",", "."))
    ok("обязательный вид: ставка по акту без поправок",
       st == 200 and a["rate"]["mode"] == "statutory" and a["rate"]["applied_pct"] == act_rate
       and a["rate"]["adj_pct"] == 0, a.get("rate"))
    ok("обязательный вид: премия по ставке акта",
       a["premium"]["amount"] == round(act_rate / 100 * 1_000_000_000), a["premium"])
    ok("обязательный вид: франшиза не применяется",
       a["franchise"]["needed"] is False and "обязательный вид" in a["franchise"]["text"], a["franchise"])
    must["product_code"] = "1002"            # ОСГО владельцев ТС: в тексте тарифа числа нет
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must, "optional": {}})
    ok("обязательный вид без числа в справочнике — ставка не выдумана",
       st == 200 and a["rate"]["applied_pct"] is None and a["premium"]["amount"] is None, a.get("rate"))


# ------------------------------------------------------------------ 4. продукт без ставки

def check_no_rate():
    print("4. Продукт «по программе» — ставка не определена")
    must = dict(CRANE_MUST, product_code="0321")
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must, "optional": CRANE_OPT})
    ok("режим undefined, цифр нет", st == 200 and a["rate"]["mode"] == "undefined"
       and a["rate"]["applied_pct"] is None and a["rate"]["base_pct"] is None
       and a["premium"]["amount"] is None, a.get("rate"))
    rows = {r["label"]: r["value"] for r in a["sections"][3]["rows"]}
    ok("в акте: «ставка не определена — нужен расчёт андеррайтера»",
       rows.get("Рекомендуемый тариф") == "ставка не определена — нужен расчёт андеррайтера", rows)
    ok("премия — «данные недоступны»", rows.get("Страховая премия") == "данные недоступны")
    ok("андеррайтеру — определить ставку", any("Определить ставку" in c for c in a["decision"]["checks"]))
    ok("не «принять» без оговорок", a["decision"]["code"] != "accept")


# ------------------------------------------------------------------ 5. модель недоступна / мусор

def check_no_model():
    print("5. Модель недоступна или вернула мусор")
    model_on(False)
    CALLS.clear()
    st, b = upload(CRANE_FILES[:2], {"lang": "ru"})
    ok("фото сохранены, ai = false", st == 200 and b["ai"] is False and b["session"], b)
    ok("к модели не обращались", not CALLS)
    ok("честное сообщение", "недоступно" in b["message"], b["message"])
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru", "must": CRANE_MUST,
                                       "optional": CRANE_OPT})
    s2 = all_text({"title": "", "footer": "", "sections": [a["sections"][1]]})
    ok("акт сформирован без модели", st == 200 and len(a["sections"]) == 5)
    ok("раздел осмотра честный: распознать не удалось, осмотр не проводился",
       "не удалось" in s2 and "осмотр не проводился" in s2, s2)
    ok("андеррайтеру — проверить фото вручную", any("вручную" in c for c in a["decision"]["checks"]))

    model_on(True)
    REPLY["text"] = "Извините, я не могу помочь с этим запросом."
    st, b = upload(CRANE_FILES[:1], {"lang": "ru"})
    ok("не-JSON — ai = false и причина", st == 200 and b["ai"] is False and "не по схеме" in b["message"], b)
    REPLY["text"] = '{"files": "не список", "fields": 5}'
    st, b = upload(CRANE_FILES[:1], {"lang": "ru"})
    ok("JSON не по схеме — ai = false", b["ai"] is False, b)
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru", "must": CRANE_MUST,
                                       "optional": CRANE_OPT})
    ok("акт после мусора формируется", st == 200 and a["inspection"]["done"] is False)

    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})
    s2 = all_text({"title": "", "footer": "", "sections": [a["sections"][1]]})
    ok("без фото — «осмотр не проводился»", "Осмотр не проводился" in s2 and "осмотр не проводился" in s2, s2)
    ok("без фото — документы не представлены повышают риск",
       any(f["code"] == "f_docs_up" for f in a["risk"]["factors"]))

    st, b = upload([("x.webp", "image/webp", b"RIFF\x00\x00\x00\x00WEBPVP8 ")], {"lang": "ru"})
    ok("WEBP вежливо отклонён", st == 422 and b["rejected"] and "JPG, PNG или PDF" in b["rejected"][0]["error"], b)
    st, b = upload([(f"{i}.png", "image/png", image()) for i in range(11)], {"lang": "ru"})
    ok("больше 10 файлов — 413", st == 413, st)


# ------------------------------------------------------------------ 6. ПД

def check_pd():
    print("6. Персональные данные")
    model_on(True)
    name = "Иванов Иван Иванович"
    REPLY["text"] = _json.dumps({
        "files": [{"n": 1, "view": "document"}],
        "fields": [{"key": "location", "value": name, "source": "document", "file": 1},
                   {"key": "manufacturer", "value": "Петров П. С.", "source": "document", "file": 1},
                   {"key": "serial_no", "value": "AA1234567", "source": "document", "file": 1},
                   {"key": "model", "value": "QY50K5D", "source": "document", "file": 1}],
        "damages": [{"what": "вмятина, владелец Сидоров Сидор", "file": 1}]}, ensure_ascii=False)
    st, b = upload([("Иванов Иван.jpg", "image/jpeg", image(kind="jpg"))], {"lang": "ru"})
    dump = _json.dumps(b, ensure_ascii=False)
    ok("ФИО из ответа модели отброшено", name not in dump and "Петров" not in dump and "Сидоров" not in dump, dump)
    ok("паспорт отброшен", "AA1234567" not in dump)
    ok("остальное сохранилось", any(r["key"] == "model" for r in b["recognized"]))
    ok("имя файла замаскировано", "Иванов" not in b["files"][0]["name"], b["files"])
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru", "must": CRANE_MUST,
                                       "optional": CRANE_OPT,
                                       "recognized": b["recognized"] + [
                                           {"key": "location", "value": "Сидоров Сидор Сидорович",
                                            "source": "input"}]})
    ok("ФИО во вводе сотрудника тоже отброшено", "Сидоров" not in _json.dumps(a, ensure_ascii=False))
    with db.tx() as con:
        rows = db.rows(con, "SELECT result_json, files_json FROM act_uploads WHERE id=?", b["session"])
        audit = db.rows(con, "SELECT who, action, entity, detail FROM audit WHERE action LIKE 'акт%'")
        calls = db.rows(con, "SELECT purpose, error FROM llm_calls")
        acts = db.rows(con, "SELECT act_json FROM acts")
    stored = _json.dumps(rows, ensure_ascii=False) + _json.dumps(acts, ensure_ascii=False)
    ok("в базе загрузок и актов ФИО нет", name not in stored and "Петров" not in stored
       and "Сидоров" not in stored and "Иванов" not in stored)
    journal = _json.dumps(audit, ensure_ascii=False) + _json.dumps(calls, ensure_ascii=False)
    ok("в журнал ФИО не попадает", "Иванов" not in journal and "Сидоров" not in journal and "Петров" not in journal)
    ok("в журнал не попадают распознанные значения", "QY50K5D" not in journal and "LXGCPA" not in journal)
    ok("в журнале есть счётчик отброшенного", any('"dropped_pd"' in (r["detail"] or "") for r in audit))


# ------------------------------------------------------------------ 7. доступ и лимит

def check_access(aid):
    print("7. Чужой акт и гостевой лимит")
    st, a = call("GET", f"/act/{aid}")
    ok("свой акт открывается", st == 200 and a["id"] == aid, st)
    saved = dict(COOKIES)
    COOKIES.clear()
    st, a = call("GET", f"/act/{aid}")
    ok("чужой гость — 404", st == 404, (st, a))
    st, _b, _h = call("GET", f"/act/{aid}.pdf", raw=True)
    ok("чужой PDF — 404", st == 404, st)
    st, _b, _h = call("GET", f"/act/{aid}.docx", raw=True)
    ok("чужой DOCX — 404", st == 404, st)
    # администратор видит любой
    now = datetime.now().isoformat(timespec="seconds")
    token = secrets.token_urlsafe(32)
    with db.tx() as con:
        cur = con.execute("INSERT INTO users (login, full_name, role, branch, password_hash, salt, status, created_at,"
                          " approved_by, approved_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                          ("act_test_admin", "Test Admin", "админ", "тест", secrets.token_hex(32),
                           secrets.token_hex(16), "активен", now, "test", now))
        con.execute("INSERT INTO sessions (token, user_id, created_at, expires_at, ip, user_agent) VALUES (?,?,?,?,?,?)",
                    (token, cur.lastrowid, now, (datetime.now() + timedelta(hours=2)).isoformat(timespec="seconds"),
                     "127.0.0.1", "test_act"))
    HEADERS.append((b"authorization", f"Bearer {token}".encode()))
    st, a = call("GET", f"/act/{aid}")
    ok("администратор открывает любой акт", st == 200, st)
    st, s = call("PUT", "/act/settings", {"settings": {"adj_pct": {"moderate": 25}}, "note": "тест"})
    ok("администратор меняет поправки", st == 200 and s["settings"]["adj_pct"]["moderate"] == 25, (st, s))
    st, s = call("PUT", "/act/settings", {"settings": {"adj_pct": {"low": -5}}})
    ok("кривые настройки — 422", st == 422, (st, s))
    HEADERS.clear()
    st, s = call("PUT", "/act/settings", {"settings": {"adj_pct": {"moderate": 30}}})
    ok("гость настройки не меняет", st in (401, 403), st)
    st, s = call("GET", "/act/settings")
    ok("настройки читаются", st == 200 and s["settings"]["adj_pct"]["moderate"] == 25, (st, s))
    with db.tx() as con:
        con.execute("DELETE FROM act_settings")
    COOKIES.clear()
    COOKIES.update(saved)

    guest.reset()
    old = guest.LIMITS["analysis"]
    guest.LIMITS["analysis"] = 2
    try:
        codes = [call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})[0]
                 for _ in range(3)]
    finally:
        guest.LIMITS["analysis"] = old
        guest.reset()
    ok("гостевой лимит на акты срабатывает (429)", codes == [200, 200, 429], codes)


# ------------------------------------------------------------------ 8. Word и PDF

def pdf_text(doc) -> str:
    """Текст PDF: извлечение отдаёт пробел как NBSP, а дефис как мягкий перенос — приводим к обычным."""
    raw = "".join(p.get_text() for p in doc).replace(" ", " ").replace("­", "-")
    return re.sub(r"\s+", " ", raw)


RU_TITLES = ["Объект и идентификация", "Результаты осмотра", "Стоимость и страховая сумма",
             "Риск-факторы и франшиза", "Заключение и рекомендация"]


def check_files(aid):
    print("8. Word и PDF")
    st, blob, h = call("GET", f"/act/{aid}.docx", raw=True)
    ok("DOCX отдаётся", st == 200 and "wordprocessingml" in h.get("content-type", ""), (st, h))
    try:
        z = zipfile.ZipFile(io.BytesIO(blob))
        names = set(z.namelist())
        xml = z.read("word/document.xml").decode("utf-8")
        minidom.parseString(xml.encode("utf-8"))
        good = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"} <= names
        for n in names:
            if n.endswith(".xml") or n.endswith(".rels"):
                minidom.parseString(z.read(n))
    except Exception as e:
        good, xml = False, str(e)
    ok("DOCX — zip с корректным XML", good, xml[:200])
    plain = re.sub(r"<[^>]+>", "", xml)
    ok("в DOCX пять разделов", all(f"{i}. {tt}" in plain for i, tt in enumerate(RU_TITLES, 1)))
    ok("в DOCX строка о подтверждении андеррайтером",
       "Акт сформирован ИИ-сюрвейером, подлежит подтверждению андеррайтером" in plain)
    ok("в DOCX шапка", "СЮРВЕЙЕРСКИЙ АКТ ПРЕДСТРАХОВОГО ОСМОТРА" in plain)
    st, blob, h = call("GET", f"/act/{aid}.pdf", raw=True)
    ok("PDF отдаётся", st == 200 and h.get("content-type") == "application/pdf", (st, h))
    try:
        doc = pymupdf.open(stream=blob, filetype="pdf")
        text = pdf_text(doc)
        pages = doc.page_count
    except Exception as e:
        text, pages = str(e), 0
    ok("PDF открывается pymupdf", pages >= 1, text[:200])
    ok("в PDF пять разделов", all(f"{i}. {tt}" in text for i, tt in enumerate(RU_TITLES, 1)), text[:300])
    ok("в PDF строка о подтверждении андеррайтером",
       "Акт сформирован ИИ-сюрвейером, подлежит подтверждению андеррайтером" in text)
    st, blob, h = call("GET", f"/act/{aid}.pdf", params={"lang": "uz"}, raw=True)
    doc = pymupdf.open(stream=blob, filetype="pdf")
    text = pdf_text(doc)
    ok("PDF на узбекском", "Koʻzdan kechirish natijalari" in text or "Ko'zdan kechirish natijalari" in text,
       text[:300])


# ------------------------------------------------------------------ 9. три языка

def check_langs(aid):
    print("9. Три языка")
    want = {"ru": RU_TITLES,
            "uz": ["Obyekt va uni identifikatsiya qilish", "Koʻzdan kechirish natijalari", "Qiymat va sugʻurta summasi",
                   "Xavf omillari va franshiza", "Xulosa va tavsiya"],
            "en": ["Object and identification", "Inspection results", "Value and sum insured",
                   "Risk factors and deductible", "Conclusion and recommendation"]}
    for lang, titles in want.items():
        st, a = call("GET", f"/act/{aid}", params={"lang": lang})
        ok(f"{lang}: заголовки разделов", st == 200 and [s["title"] for s in a["sections"]] == titles,
           [s.get("title") for s in a.get("sections") or []])
        if lang != "ru":
            labels = [a["title"], a["footer"]] + [r["label"] for s in a["sections"] for r in s["rows"]] + \
                     [li["title"] for s in a["sections"] for li in s.get("lists") or []] + \
                     [f["text"] for f in a["risk"]["factors"]] + a["decision"]["checks"] + [a["franchise"]["text"]] + \
                     a["rate"]["how"] + [a["value"]["text"], a["risk"]["rule"]] + \
                     [p for s in a["sections"] for p in s["paragraphs"]]
            cyr = [x for x in labels if re.search(r"[А-Яа-яЁё]", x or "")]
            ok(f"{lang}: в подписях и выводах нет кириллицы", not cyr, cyr[:5])
    st, a = call("POST", "/act/make", {"lang": "en", "must": CRANE_MUST, "optional": CRANE_OPT})
    ok("en: сумма в формате UZS", "2,945,000,000 UZS" in all_text(a), all_text(a)[:300])


# ------------------------------------------------------------------ 10. движок: чистые функции

def check_engine():
    print("10. Лёгкий движок (чистые функции)")
    today = date(2026, 9, 29)
    r = ae.risk_level({"inspected": True, "damages": ["вмятина"], "location": "construction", "losses_count": 3,
                       "documents": False, "today": today})
    ok("все четыре признака «повышают» — высокий", r["level"] == "high" and r["up"] == 4, r)
    r = ae.risk_level({"inspected": False, "documents": True, "today": today})
    ok("мало известного — умеренный", r["level"] == "moderate", r)
    r = ae.risk_level({"inspected": True, "year": 2026, "location": "guarded", "losses_count": 0,
                       "documents": True, "today": today})
    ok("всё снижает — низкий", r["level"] == "low" and r["down"] == 4, r)
    r = ae.risk_level({"inspected": True, "year": 2015, "location": "open_area", "guard": True,
                       "losses_count": 1, "documents": True, "today": today})
    ok("открытая площадка под охраной и один убыток — не влияют", r["up"] == 0 and r["down"] == 1, r)
    r = ae.risk_level({"inspected": True, "documents": True, "today": today})
    ok("год неизвестен — не пишем «не новый»", r["factors"][0]["code"] == "f_cond_no_year", r["factors"][0])
    d = ae.discrepancies([{"key": "curb_mass", "value": "36 170 кг", "source": "plate"},
                          {"key": "curb_mass", "value": "36,17 т", "source": "document"}])
    ok("масса в тоннах и килограммах — одно и то же", not d, d)
    d = ae.discrepancies([{"key": "serial_no", "value": "LXGCPA393TA006921", "source": "plate"},
                          {"key": "serial_no", "value": "LXGCPA393TA006927", "source": "document"}])
    ok("номер отличается одним знаком — расхождение", len(d) == 1 and d[0]["priority"] == "document", d)
    d = ae.discrepancies([{"key": "year", "value": "2025", "source": "document"}], {"year": 2026})
    ok("год во вводе и в документе — расхождение", len(d) == 1 and d[0]["key"] == "year", d)
    d = ae.discrepancies([{"key": "model", "value": "QY50K5D", "source": "document"},
                          {"key": "model", "value": "QY50K", "source": "marking"}])
    ok("модель на стреле и в документе — расхождение", len(d) == 1, d)
    ok("ракурсы здания", ae.required_views("property") == ["facade", "roof", "interior", "electrical",
                                                           "fire_safety"])
    ok("группа по классу и типу", ae.object_group("3", "Спецтехника — автокран") == "special"
       and ae.object_group("8", "Склад") == "property" and ae.object_group("7") == "cargo")
    ok("ошибки настроек ловятся", ae.check_settings({"level_rule": {"low_max_net": 3, "high_min_net": 1}}))
    ok("верные настройки проходят", not ae.check_settings({"adj_pct": {"moderate": 25}}))
    ok("число из «38,600 kg» и «2 945 000 000»", ae.to_number("38,600 kg") == 38600
       and ae.to_number("2 945 000 000") == 2945000000)


# ------------------------------------------------------------------ 11. связка текста моделью

def check_polish():
    print("11. Литературная связка моделью убрана: решение и списки — только из шаблонов")
    model_on(True)
    REPLY["text"] = "Объект осмотрен, выявлено 999 замечаний. Рекомендация: отказать."
    CALLS.clear()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT, "polish": True})
    ok("polish в теле игнорируется: к модели не обращались", st == 200 and not CALLS, len(CALLS))
    ok("ни один раздел не переписан моделью",
       not any(s.get("polished") for s in a["sections"]) and "999" not in all_text(a))
    ok("решение — из шаблона", a["decision"]["text"] == act.t("d_" + a["decision"]["code"], "ru"), a["decision"])
    ok("функции polish больше нет", not hasattr(act, "polish"))
    model_on(False)


# ------------------------------------------------------------------ 12. хранение

def check_cleanup(sid, aid):
    print("12. Хранение: фото 24 часа, акт 7 дней")
    with db.tx() as con:
        up = db.rows(con, "SELECT created_at, expires_at, files_json FROM act_uploads WHERE id=?", sid)[0]
        ac = db.rows(con, "SELECT created_at, expires_at FROM acts WHERE id=?", aid)[0]
    h = (datetime.fromisoformat(up["expires_at"]) - datetime.fromisoformat(up["created_at"])).total_seconds() / 3600
    d = (datetime.fromisoformat(ac["expires_at"]) - datetime.fromisoformat(ac["created_at"])).days
    ok("фото — 24 часа", h == 24, h)
    ok("акт — 7 дней", d == 7, d)
    files = _json.loads(up["files_json"])
    ok("в базе у файла нет имени — только номер, формат и путь",
       all("name" not in f and "orig_name" not in f and f.get("index") and f.get("fmt") for f in files)
       and "front.jpg" not in up["files_json"], files[:1])
    ok("путь файла — через db.stored_path", all(f["path"] == db.stored_path(act.DIR / sid / Path(f["path"]).name)
                                                for f in files), files[:1])
    ok("файлы на диске", (act.DIR / sid).is_dir())
    with db.tx() as con:
        con.execute("UPDATE act_uploads SET expires_at='2020-01-01T00:00:00' WHERE id=?", (sid,))
        con.execute("UPDATE acts SET expires_at='2020-01-01T00:00:00' WHERE id=?", (aid,))
    removed = act.cleanup()
    ok("просроченное удаляется", removed >= 2, removed)
    ok("папка с фото удалена", not (act.DIR / sid).exists())
    st, a = call("GET", f"/act/{aid}")
    ok("просроченный акт не открывается", st == 404, st)


# ------------------------------------------------------------------ 13–20. замечания контролёра 29.09.2026

def set_limits(**kw):
    """Версия настроек акта с другими пределами (как правка администратора), без API."""
    with db.tx() as con:
        con.execute("INSERT INTO act_settings (created_at, created_by, settings_json, calibrated, note) "
                    "VALUES (?,?,?,?,?)", (db.now(), "тест", _json.dumps({"limits": kw}), 0, "тест"))


def clear_settings():
    with db.tx() as con:
        con.execute("DELETE FROM act_settings")


def png_bomb(w=14000, h=14000) -> bytes:
    """Настоящая «бомба»: PNG 14000×14000 оттенков серого из нулей — сотни килобайт в файле, ~200 МБ в памяти."""
    import struct
    import zlib

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    comp = zlib.compressobj(9)
    row = b"\x00" * (w + 1)
    parts = [comp.compress(row * 500) for _ in range(h // 500)] + [comp.flush()]
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 0, 0, 0, 0))
            + chunk(b"IDAT", b"".join(parts)) + chunk(b"IEND", b""))


def jpeg_claiming(w, h) -> bytes:
    """JPEG, у которого в SOF0 записаны огромные размеры (данные кадра — мусор)."""
    import struct
    app0 = b"\xff\xe0" + struct.pack(">H", 16) + b"JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    sof = b"\xff\xc0" + struct.pack(">HBHHB", 17, 8, h, w, 3) + b"\x01\x11\x00\x02\x11\x01\x03\x11\x01"
    return b"\xff\xd8" + app0 + sof + b"\xff\xda\x00\x02" + b"\x00" * 64 + b"\xff\xd9"


def pdf_pages(n) -> bytes:
    doc = pymupdf.open()
    for _ in range(n):
        doc.new_page()
    return doc.tobytes()


def noise(w, h, kind="png") -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, w, h, os.urandom(w * h * 3), False)
    return pix.tobytes("jpg" if kind == "jpg" else "png")


def check_bomb():
    print("13. Картинка-бомба и PDF: размеры из заголовка, без раскрытия")
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    act.reset_limits()
    bomb = png_bomb()
    opened = []
    orig_pix = pymupdf.Pixmap

    class spy(orig_pix):                     # подкласс: внутренние проверки isinstance в pymupdf не ломаются
        def __init__(self, *a, **kw):
            opened.append(a[:1] if a and isinstance(a[0], (bytes, bytearray)) else "other")
            super().__init__(*a, **kw)
    ok("бомба маленькая на диске", len(bomb) < 1024 * 1024, len(bomb))
    ok("размеры PNG читаются из IHDR", act.image_size(bomb, "png") == (14000, 14000))
    ok("размеры JPEG — из SOF", act.image_size(image(kind="jpg", w=321, h=123), "jpg") == (321, 123))
    ok("размеры JPEG с мусором в SOF", act.image_size(jpeg_claiming(12000, 9000), "jpg") == (12000, 9000))
    ok("битый заголовок — размеров нет", act.image_size(b"\x89PNG\r\n\x1a\n" + b"\x00" * 30, "png") is None
       and act.image_size(b"\xff\xd8\xff\xda\x00\x02", "jpg") is None)
    pymupdf.Pixmap = spy
    try:
        CALLS.clear()
        st, b = upload([("bomb.png", "image/png", bomb), ("big.jpg", "image/jpeg", jpeg_claiming(12000, 9000)),
                        ("broken.png", "image/png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 64),
                        ("ok.png", "image/png", image())], {"lang": "ru"})
    finally:
        pymupdf.Pixmap = orig_pix
    rej = {r["index"]: r["error"] for r in b.get("rejected") or []}
    ok("бомба 14000×14000 отклонена с понятным сообщением",
       st == 200 and "14000×14000" in rej.get(1, "") and "50 Мп" in rej.get(1, ""), rej)
    ok("JPEG 12000×9000 (108 Мп) отклонён", "Мп" in rej.get(2, ""), rej)
    ok("размеры не прочитались — отклонено", "размеры" in rej.get(3, ""), rej)
    ok("нормальный снимок принят (index 4)", [f["index"] for f in b["files"]] == [4], b.get("files"))
    ok("отклонённые картинки не раскрывались (Pixmap только для принятой)",
       len([x for x in opened if x != "other"]) == 1, len(opened))
    ok("к модели ушёл один файл", len(CALLS) == 1 and len(CALLS[0]["files"]) == 1)
    st, b = upload([("many.pdf", "application/pdf", pdf_pages(11)), ("one.pdf", "application/pdf", pdf_pages(1))],
                   {"lang": "ru"})
    rej = {r["index"]: r["error"] for r in b.get("rejected") or []}
    ok("PDF больше 10 страниц отклонён", "10 страниц" in rej.get(1, ""), rej)
    ok("PDF в 1 страницу принят и не растеризуется (уходит модели как PDF)",
       [f["index"] for f in b["files"]] == [2] and CALLS[-1]["files"][0]["mime"] == "application/pdf")
    set_limits(max_image_mp=0.05)             # 320×240 = 0,077 Мп — теперь больше предела
    try:
        st, b = upload([("ok.png", "image/png", image())], {"lang": "ru"})
        ok("предел мегапикселей — из настроек", st == 422 and "Мп" in b["rejected"][0]["error"], b)
    finally:
        clear_settings()


async def _asend(method, path, headers, payload, out):
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method,
             "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "", "query_string": b"",
             "headers": headers, "client": ("203.0.113.9", 0), "server": ("test", 80)}
    res = {"status": None}

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            res["status"] = msg["status"]
    await app(scope, receive, send)
    import time as _t
    out.append((path, res["status"], _t.monotonic()))


def multipart(files, fields=None):
    boundary = "----insonact"
    parts = [f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode()
             for k, v in (fields or {}).items()]
    for name, mime, blob in files:
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="files"; '
                     f'filename="{name}"\r\nContent-Type: {mime}\r\n\r\n'.encode() + blob + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    return boundary, b"".join(parts)


def check_threadpool():
    print("14. Распознавание не останавливает сервер; срок 25 с; одна попытка")
    import inspect
    import time as _t
    from app import act as act_mod
    eps = {}
    for r in act_mod.router.routes:
        eps[(r.path, tuple(sorted(r.methods)))] = r.endpoint
    heavy = [eps[("/act/photos", ("POST",))], eps[("/act/make", ("POST",))], eps[("/act/{aid}.docx", ("GET",))],
             eps[("/act/{aid}.pdf", ("GET",))], eps[("/act/{aid}/send", ("POST",))]]
    ok("обработчики фото, акта, выгрузки и отправки — обычные def (пул потоков)",
       not any(inspect.iscoroutinefunction(f) for f in heavy), [f.__name__ for f in heavy
                                                                if inspect.iscoroutinefunction(f)])
    # пока идёт медленное распознавание (1,5 с), другой запрос отвечает сразу
    model_on(True)
    REPLY["text"], REPLY["sleep"] = CRANE_REPLY, 1.5
    act.reset_limits()
    boundary, payload = multipart(CRANE_FILES[:1], {"lang": "ru"})
    up_hdrs = [(b"host", b"test"), (b"content-type", f"multipart/form-data; boundary={boundary}".encode()),
               (b"content-length", str(len(payload)).encode())] + _cookie_hdr()
    get_hdrs = [(b"host", b"test")] + _cookie_hdr()
    done = []

    async def both():
        import asyncio as _a
        t0 = _t.monotonic()
        up = _a.ensure_future(_asend("POST", "/act/photos", up_hdrs, payload, done))
        await _a.sleep(0.2)
        await _asend("GET", "/act/settings", get_hdrs, b"", done)
        await up
        return t0
    try:
        t0 = asyncio.run(both())
    finally:
        REPLY["sleep"] = 0
    order = [p for p, _s, _ in done]
    ok("GET /act/settings ответил раньше, чем закончилось распознавание",
       order == ["/act/settings", "/act/photos"] and done[0][2] - t0 < 1.2, [(p, round(x - t0, 2)) for p, _s, x in done])
    ok("загрузка при этом прошла", done[-1][1] == 200, done)
    # общий срок: модель «висит» дольше срока — ai=false, честная причина
    REPLY["sleep"] = 1.5
    try:
        t1 = _t.monotonic()
        rec = act.recognize([{"blob": image(), "fmt": "png"}], "ru", {"ai_deadline_sec": 0.5})
        spent = _t.monotonic() - t1
    finally:
        REPLY["sleep"] = 0
    ok("не уложились в срок — ai=false и сообщение", rec["ok"] is False and "не ответила" in rec["reason"], rec)
    ok("ответ не ждёт модель дольше срока", spent < 1.2, spent)
    # сервер-подобный путь: модель падает — фото сохранены, акт формируется
    def boom(*a, **kw):
        raise RuntimeError("сеть")
    fake_raw = llm.chat_raw
    llm.chat_raw = boom
    try:
        st, b = upload(CRANE_FILES[:2], {"lang": "ru"})
    finally:
        llm.chat_raw = fake_raw
    ok("ошибка модели — фото сохранены, ai=false", st == 200 and b["ai"] is False and b["session"]
       and (act.DIR / b["session"]).is_dir(), b)
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru", "must": CRANE_MUST,
                                       "optional": CRANE_OPT})
    ok("акт после сбоя модели формируется", st == 200 and a["inspection"]["done"] is False, st)
    # llm.chat_raw: retries=0 — ровно один запрос к сети, таймаут передан
    posts = []

    def fake_post(url, body, headers, timeout=None):
        posts.append(timeout)
        raise TimeoutError("медленно")
    old = (llm._post, llm.enabled, llm.provider)
    llm._post, llm.enabled, llm.provider = fake_post, (lambda: True), (lambda: "openai")
    try:
        res = ORIG["chat_raw"]("тест", [{"role": "user", "content": "x"}], timeout=20, retries=0)
    finally:
        llm._post, llm.enabled, llm.provider = old
    ok("llm.chat_raw(retries=0): одна попытка, таймаут 20 с", posts == [20] and res["text"] is None, posts)


def check_body_limit():
    print("15. Размер тела /act/photos — по Content-Length до чтения")
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    CALLS.clear()
    before = set(p.name for p in act.DIR.iterdir()) if act.DIR.exists() else set()
    boundary, payload = multipart(CRANE_FILES[:1], {"lang": "ru"})
    read = []

    async def run(headers):
        out = {}

        async def receive():
            read.append(1)
            return {"type": "http.request", "body": payload, "more_body": False}

        async def send(msg):
            if msg["type"] == "http.response.start":
                out["status"] = msg["status"]
        scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "POST",
                 "scheme": "http", "path": "/act/photos", "raw_path": b"/act/photos", "root_path": "",
                 "query_string": b"", "headers": headers, "client": ("203.0.113.9", 0), "server": ("test", 80)}
        await app(scope, receive, send)
        return out.get("status")
    ctype = (b"content-type", f"multipart/form-data; boundary={boundary}".encode())
    st = asyncio.run(run([(b"host", b"test"), ctype, (b"content-length", str(act.MAX_BODY + 1).encode())]
                         + _cookie_hdr()))
    ok("больше 10 × 15 МБ + запас — 413", st == 413, st)
    ok("тело при этом не читалось", not read, len(read))
    st = asyncio.run(run([(b"host", b"test"), ctype] + _cookie_hdr()))
    ok("без Content-Length — 411", st == 411, st)
    after = set(p.name for p in act.DIR.iterdir()) if act.DIR.exists() else set()
    ok("ни файлов, ни обращений к модели", after == before and not CALLS)
    st = asyncio.run(run([(b"host", b"test"), ctype, (b"content-length", str(len(payload)).encode())]
                         + _cookie_hdr()))
    ok("обычный размер проходит", st == 200, st)


def check_limits():
    print("16. Лимиты: фото гостя по числу файлов, распознавания на сервер")
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    act.reset_limits()
    guest.reset()
    set_limits(guest_photos_per_hour=3)
    try:
        codes = [upload(CRANE_FILES[:2], {"lang": "ru"})[0], upload(CRANE_FILES[:2], {"lang": "ru"}),
                 upload(CRANE_FILES[:1], {"lang": "ru"})[0]]
    finally:
        clear_settings()
        act.reset_limits()
    st2, b2 = codes[1]
    ok("2 фото + 2 фото при пределе 3 — второй запрос 429, третий (1 фото) проходит",
       [codes[0], st2, codes[2]] == [200, 429, 200], [codes[0], st2, codes[2]])
    ok("429 с понятным сообщением и Retry-After", "фото в час" in b2.get("detail", "") and b2.get("limit") == 3
       and b2.get("retry_after_sec"), b2)
    ok("по умолчанию 60 фото в час", ae.DEFAULT_SETTINGS["limits"]["guest_photos_per_hour"] == 60
       and ae.DEFAULT_SETTINGS["limits"]["ai_calls_per_hour"] == 120)
    # guard больше не считает /act/photos запросами (иначе 30 запросов в час перекрыли бы счёт по фото)
    from app import guard
    ok("guard не считает /act/photos запросами", ("POST", "/act/photos") not in guard.GUEST_BUCKET)
    set_limits(ai_calls_per_hour=1)
    CALLS.clear()
    try:
        _s1, b1 = upload(CRANE_FILES[:1], {"lang": "ru"})
        _s2, b2 = upload(CRANE_FILES[:1], {"lang": "ru"})
    finally:
        clear_settings()
        act.reset_limits()
    ok("общий предел распознаваний: второй раз ai=false с честным сообщением",
       b1["ai"] is True and b2["ai"] is False and "лимит" in b2["message"] and len(CALLS) == 1, (b2["message"],
                                                                                                  len(CALLS)))
    ok("фото при этом сохранены", _s2 == 200 and b2["session"])
    ok("кривые пределы в настройках ловятся",
       ae.check_settings({"limits": {"max_image_mp": 0}}) and ae.check_settings({"limits": {"x": 1}})
       and not ae.check_settings({"limits": {"max_image_mp": 40}}))


def check_budget():
    print("17. Вложения в модель — не больше 10 МБ (настройка), о непрочитанном — честно")
    model_on(True)
    REPLY["text"] = _json.dumps({"files": [], "fields": [], "damages": []})
    act.reset_limits()
    big = [noise(1200, 1200, "png") for _ in range(3)]
    set_limits(ai_max_mb=1)
    CALLS.clear()
    try:
        st, b = upload([("a.png", "image/png", big[0]), ("b.png", "image/png", big[1]),
                        ("c.png", "image/png", big[2]), ("d.pdf", "application/pdf", pdf_pages(1))], {"lang": "ru"})
    finally:
        clear_settings()
    sent = CALLS[0]["files"] if CALLS else []
    total = sum(len(f["data"]) for f in sent)
    ok("суммарно в модель не больше предела", st == 200 and 0 < total <= 1024 * 1024, total)
    ok("PDF (документ) — в приоритете", any(f["mime"] == "application/pdf" for f in sent))
    ok("непрочитанные перечислены номерами", b.get("not_sent") and b["notes"]
       and all(str(n) in b["notes"][0] for n in b["not_sent"]), (b.get("not_sent"), b.get("notes")))
    ok("у непрочитанных read_by_ai = false",
       all(not f["read_by_ai"] for f in b["files"] if f["index"] in b["not_sent"]))
    ok("по умолчанию предел 10 МБ", ae.DEFAULT_SETTINGS["limits"]["ai_max_mb"] == 10)
    # сильнее пережать: 4 шума 2000×2000 не влезают в 10 МБ при обычном сжатии, но влезают после второго шага
    many = [{"blob": noise(2000, 2000, "jpg"), "fmt": "jpg"} for _ in range(4)]
    first = sum(len(act._for_model(f["blob"], "jpg")[0]) for f in many)
    payload, sent_i, left = act.pick_for_model(many, 10 * 1024 * 1024)
    ok("не поместились — пережаты сильнее, а не выброшены" if first > 10 * 1024 * 1024 else
       "поместились при обычном сжатии", len(sent_i) == 4 and not left
       and sum(len(p["data"]) for p in payload) <= 10 * 1024 * 1024, (first, [len(p["data"]) for p in payload]))


def check_sources():
    print("18. Источник значения нельзя подменить")
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    act.reset_limits()
    st, b = upload(CRANE_FILES, {"lang": "ru", "product_code": "0318"})
    sid = b["session"]
    fake = [{"key": "serial_no", "value": "FAKE0000000000001", "source": "document", "file": "f4"}]
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT,
                                       "recognized": fake})
    rec = {(r["key"], r["value"]): r for r in a["recognized"]}
    ok("без загрузки «из документа» превращается во «введено сотрудником»",
       rec[("serial_no", "FAKE0000000000001")]["source"] == "input", a["recognized"])
    s1 = {r["label"]: r for r in a["sections"][0]["rows"]}
    ok("в разделе 1 — «введено сотрудником»", "введено сотрудником" in (s1["Заводской (серийный) номер"]["note"]
                                                                         or ""), s1["Заводской (серийный) номер"])
    edited = [dict(r) for r in b["recognized"]]
    for r in edited:
        if r["key"] == "serial_no" and r["source"] == "document":
            r["value"] = "LXGCPA393TA000000"            # сотрудник исправил значение, источник оставил
    st, a = call("POST", "/act/make", {"session": sid, "lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT,
                                       "recognized": edited + fake})
    rec = {(r["key"], r["value"]): r["source"] for r in a["recognized"]}
    ok("изменённое значение — input", rec.get(("serial_no", "LXGCPA393TA000000")) == "input", rec)
    ok("чужое значение с source=document в живой сессии — input",
       rec.get(("serial_no", "FAKE0000000000001")) == "input", rec)
    ok("неизменённые значения сохраняют источник", rec.get(("serial_no", "LXGCPA393TA006921")) == "plate"
       and rec.get(("curb_mass", "38 600 кг")) == "document", rec)
    ok("расхождение по номеру не выдумано из подмены",
       not any(d["key"] == "serial_no" and any("document" in g["sources"] and "FAKE" in g["value"]
                                               for g in d["values"]) for d in a["discrepancies"]), a["discrepancies"])
    with db.tx() as con:
        row = db.rows(con, "SELECT detail FROM audit WHERE action='акт сформирован' ORDER BY id DESC LIMIT 1")
    ok("в журнале — счётчик понижённых источников", '"sources_downgraded": 2' in row[0]["detail"], row)
    saved = dict(COOKIES)
    COOKIES.clear()
    st, a = call("POST", "/act/make", {"session": sid, "lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT,
                                       "recognized": b["recognized"]})
    COOKIES.clear()
    COOKIES.update(saved)
    ok("чужая сессия — все источники input", st == 200 and all(r["source"] == "input" for r in a["recognized"]),
       {r["source"] for r in a["recognized"]})


def init_data(tg_id: int, token: str, auth_date=None) -> str:
    import hashlib
    import hmac
    import time as _t
    from urllib.parse import urlencode as _ue
    pairs = [("auth_date", str(int(auth_date or _t.time()))), ("query_id", "AAHtest"),
             ("user", _json.dumps({"id": tg_id, "first_name": "Test"}, separators=(",", ":")))]
    dcs = "\n".join(f"{k}={v}" for k, v in sorted(pairs))
    secret = hmac.new(b"WebAppData", token.encode(), hashlib.sha256).digest()
    h = hmac.new(secret, dcs.encode(), hashlib.sha256).hexdigest()
    return _ue(pairs + [("hash", h)])


def check_send(aid):
    print("19. Отправка акта ботом: POST /act/{id}/send")
    from app import telegram, tgbot
    token = "123456:TEST-token-не-для-журнала"
    sent = []
    reply = {"ok": True, "result": {"message_id": 5}}

    def fake_file(method, fields, field, filename, blob, mime):
        sent.append({"method": method, "chat_id": fields.get("chat_id"), "field": field, "filename": filename,
                     "head": blob[:4], "mime": mime, "caption": fields.get("caption")})
        return dict(reply)
    old = (tgbot.bot_token, telegram.bot_token, tgbot._deliver_file, tgbot._deliver)
    tgbot.bot_token = telegram.bot_token = lambda: token
    tgbot._deliver_file = fake_file
    tgbot._deliver = no_network
    act.reset_limits()
    try:
        good = init_data(777000111, token)
        st, r = call("POST", f"/act/{aid}/send", {"format": "pdf"})
        ok("гость без Telegram — 409 с понятным сообщением",
           st == 409 and r.get("code") == "no_telegram" and "Telegram" in r["detail"]
           and "компьютере" in r["detail"], (st, r))
        st, r = call("POST", f"/act/{aid}/send", {"format": "pdf", "initData": init_data(777000111, "чужой:токен")})
        ok("поддельный initData — 403", st == 403 and r.get("code") == "bad_init_data", (st, r))
        st, r = call("POST", f"/act/{aid}/send", {"format": "xls", "initData": good})
        ok("неизвестный формат — 422", st == 422, (st, r))
        st, r = call("POST", f"/act/{aid}/send", {"format": "pdf", "lang": "uz", "initData": good})
        ok("PDF отправлен: 200 и sendDocument этому пользователю",
           st == 200 and r.get("sent") and sent and sent[-1]["method"] == "sendDocument"
           and sent[-1]["chat_id"] == "777000111" and sent[-1]["head"] == b"%PDF"
           and sent[-1]["filename"].endswith(".pdf"), (st, r, sent[-1:]))
        ok("язык файла — из запроса", r.get("lang") == "uz" and "Syurveyer" in (sent[-1]["caption"] or ""), r)
        st, r = call("POST", f"/act/{aid}/send", {"format": "docx", "initData": good})
        ok("DOCX отправлен документом", st == 200 and sent[-1]["head"][:2] == b"PK"
           and sent[-1]["filename"].endswith(".docx") and sent[-1]["method"] == "sendDocument", (st, sent[-1:]))
        saved = dict(COOKIES)
        COOKIES.clear()
        n_before = len(sent)
        st, r = call("POST", f"/act/{aid}/send", {"format": "pdf", "initData": good})
        COOKIES.clear()
        COOKIES.update(saved)
        ok("не владелец — 404 и ничего не отправлено", st == 404 and len(sent) == n_before, (st, r))
        reply.update(ok=False, description="Forbidden: bot can't initiate conversation with a user")
        st, r = call("POST", f"/act/{aid}/send", {"format": "pdf", "initData": good})
        ok("бот не может написать первым — 502 и подсказка нажать Start",
           st == 502 and r.get("code") == "start_bot" and "Start" in r["detail"], (st, r))
        reply.clear()
        reply.update(ok=True, result={"message_id": 6})
        codes = [call("POST", f"/act/{aid}/send", {"format": "pdf", "initData": good})[0] for _ in range(8)]
        ok("лимит: 10 отправок в час, 11-я — 429", codes[:7] == [200] * 7 and codes[7] == 429, codes)
        tgbot.bot_token = telegram.bot_token = lambda: ""
        act.reset_limits()
        st, r = call("POST", f"/act/{aid}/send", {"format": "pdf", "initData": good})
        ok("бот не подключён — 503", st == 503 and r.get("code") == "bot_off", (st, r))
        with db.tx() as con:
            dump = _json.dumps(db.rows(con, "SELECT * FROM audit WHERE action LIKE 'акт%'"), ensure_ascii=False) + \
                _json.dumps(db.rows(con, "SELECT * FROM tg_messages"), ensure_ascii=False)
        ok("токен бота не попал ни в журнал, ни в журнал бота", "TEST-token" not in dump)
    finally:
        tgbot.bot_token, telegram.bot_token, tgbot._deliver_file, tgbot._deliver = old
        act.reset_limits()


def check_misc():
    print("20. Мелочи: ФИО и номера, регион, несколько классов, папки-сироты")
    pl = act.pd_like
    ok("изготовитель — название компании целиком: не ФИО",
       not pl("manufacturer", "Xuzhou Construction Machinery Group Co., Ltd.")
       and not pl("manufacturer", "XUZHOU CONSTRICTION MACHINERY GROUP IMPORT & EXPORT CO., LTD (XCMG)"))
    ok("компания + ФИО вне названия — ФИО", pl("manufacturer", "Xuzhou Machinery Group Co., Ltd., Ivan Petrov")
       and pl("manufacturer", "ООО «Техника», директор Иванов Иван Иванович")
       and pl("brand", "Иванов И. И. Group"))
    ok("9 цифр без букв в номере — отброшено (возможный ИНН/телефон)", pl("serial_no", "123456789")
       and pl("engine_no", "301 234 567".replace(" ", "")))
    ok("номер с буквами — проходит", not pl("engine_no", "D123456789") and not pl("serial_no", "A123456789"))
    ok("лист параметров заказчика: «Паспорт № 20102600523» и «D9264002842» проходят",
       not pl("serial_no", "Паспорт № 20102600523") and not pl("engine_no", "D9264002842")
       and not pl("engine_model", "SC9DF340Q6") and not pl("serial_no", "LXGCPA393TA006921"))
    # регион кодом и названием
    st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(CRANE_MUST, region="tashkent_region"),
                                       "optional": CRANE_OPT})
    reg = {r["label"]: r["value"] for r in a["sections"][0]["rows"]}.get("Регион")
    ok("регион кодом печатается по-русски", reg == "Ташкентская область", reg)
    st, u = call("GET", f"/act/{a['id']}", params={"lang": "uz"})
    reg_uz = [r["value"] for r in u["sections"][0]["rows"]][-1] if u.get("sections") else None
    ok("тот же акт на узбекском — регион по-узбекски",
       any(r["value"] == "Toshkent viloyati" for r in u["sections"][0]["rows"]), reg_uz)
    st, a = call("POST", "/act/make", {"lang": "en", "must": dict(CRANE_MUST, region="Ташкентская область"),
                                       "optional": CRANE_OPT})
    ok("каноническое название → перевод на язык акта",
       any(r["value"] == "Tashkent region" for r in a["sections"][0]["rows"]))
    st, a = call("POST", "/act/make", {"lang": "en", "must": dict(CRANE_MUST, region="Чирчик, промзона"),
                                       "optional": CRANE_OPT})
    ok("незнакомый регион — как ввели", any(r["value"] == "Чирчик, промзона" for r in a["sections"][0]["rows"]))
    # продукт с несколькими классами
    with db.tx() as con:
        multi = db.rows(con, "SELECT pc.product_code AS code FROM product_classes pc JOIN products p "
                             "ON p.code = pc.product_code WHERE p.pricing_mode NOT IN ('по программе') "
                             "GROUP BY pc.product_code HAVING COUNT(DISTINCT pc.class_code) > 1 LIMIT 1")
    if multi:
        st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(CRANE_MUST, product_code=multi[0]["code"]),
                                           "optional": CRANE_OPT})
        text = all_text(a) if st == 200 else str(a)
        ok(f"продукт {multi[0]['code']} с несколькими классами — пометка в акте и разбор по частям (30.09.2026)",
           st == 200 and a["rate"]["multi_class"] and "нескольким классам" in text
           and "договор разобран по частям" in text and a["parts"]["mode"] == "multi", text[:300])
    else:
        ok("в справочнике нашёлся продукт с несколькими классами", False)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})
    ok("у продукта с одним классом пометки нет", not a["rate"]["multi_class"]
       and "нескольким классам" not in all_text(a))
    # папка-сирота: запись в базу упала — папки нет
    model_on(False)
    act.reset_limits()
    before = set(p.name for p in act.DIR.iterdir()) if act.DIR.exists() else set()
    orig_cleanup = act.cleanup

    def broken(con=None):
        raise RuntimeError("база недоступна")
    act.cleanup = broken
    try:
        try:
            st, _b = upload(CRANE_FILES[:1], {"lang": "ru"})
        except Exception as e:                     # ASGI-приложение пробрасывает исключение после 500
            st = type(e).__name__
    finally:
        act.cleanup = orig_cleanup
    after = set(p.name for p in act.DIR.iterdir()) if act.DIR.exists() else set()
    ok("запись в базу упала — папка с фото удалена", after == before, (st, after - before))
    old_dir = act.DIR / "deadbeefdeadbeefdeadbeef"
    new_dir = act.DIR / "cafecafecafecafecafecafe"
    old_dir.mkdir(parents=True, exist_ok=True)
    new_dir.mkdir(parents=True, exist_ok=True)
    (old_dir / "f1.png").write_bytes(image())
    past = datetime.now().timestamp() - 3 * 3600
    os.utime(old_dir, (past, past))
    act.cleanup()
    ok("очистка: старая папка без записи удалена, свежая оставлена", not old_dir.exists() and new_dir.exists())
    shutil.rmtree(new_dir, ignore_errors=True)


# ------------------------------------------------------------------ 21–27. дополнения 29.09.2026

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CONTRACT = Path(__file__).resolve().parent.parent / "sandbox" / "flow150_contract.docx"
WH_MUST = {"product_code": "0808", "sum_insured": 1_000_000_000, "object_value": 1_200_000_000,
           "region": "Ташкент"}
WH_OPT = {"object_kind": "warehouse", "protection": "alarm", "losses_3y": {"count": 0, "small_count": 0}}


def fresh():
    guest.reset()
    act.reset_limits()


def docx_bytes(lines, extra_parts=None) -> bytes:
    """Минимальный DOCX: абзацы из строк (и лишние части архива, если нужны)."""
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = "".join(f"<w:p><w:r><w:t>{ln}</w:t></w:r></w:p>" for ln in lines)
    xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{ns}"><w:body>{body}</w:body></w:document>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        z.writestr("word/document.xml", xml)
        for name, data in (extra_parts or {}).items():
            z.writestr(name, data)
    return buf.getvalue()


def docx_bomb(mb=60) -> bytes:
    """DOCX с частью из нулей: в файле — десятки килобайт, после распаковки — mb мегабайт."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", '<?xml version="1.0"?><w:document/>')
        z.writestr("word/media/zero.bin", b"\x00" * (mb * 1024 * 1024))
    return buf.getvalue()


def ra_same(con, must, optional):
    """Расчёт risk_analytics напрямую — тем же быстрым режимом, что и акт."""
    from app import risk_analytics as ra
    m, o, a = ra.apply_defaults(con, must, optional)
    return ra.analyze(con, m, o, assumptions=a)


# сценарий акта ← сценарий risk_analytics (там названия EML и PML переставлены, модуль не меняем)
MAP_RA = {"PML": "EML", "EML": "PML", "MFL": "MFL"}
SCEN_REPORT = {}                          # цифры сценариев для отчёта (печатаются в конце)


def order_ok(sc: dict) -> bool:
    return sc["pml"]["amount"] <= sc["eml"]["amount"] <= sc["mfl"]["amount"]


def check_scenarios():
    print("21. Сценарии PML / EML / MFL — сверка с risk_analytics (порядок заказчика PML ≤ EML ≤ MFL)")
    fresh()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, object_kind="truck_crane")})
    sc = a.get("scenarios") or {}
    ok("автокран: блок сценариев есть", st == 200 and sc.get("available") and sc.get("rule") == "vehicle", sc)
    with db.tx() as con:
        an = ra_same(con, {"class_code": "3", "product_code": "0318", "object_type": CRANE_TYPE,
                           "sum_insured": CRANE_MUST["sum_insured"], "object_value": CRANE_MUST["object_value"],
                           "region": CRANE_MUST["region"], "vehicle_type": "special"},
                     {"losses_3y": {"count": 0, "amount": 0, "small_count": 0}})
    # порядок заказчика PML ≤ EML ≤ MFL: PML акта = EML модуля («защита сработала»), EML акта = PML модуля
    for s in ("PML", "EML", "MFL"):
        src = MAP_RA[s]
        ok(f"автокран: {s} акта = {src} risk_analytics", sc[s.lower()]["amount"] == round(an["scenarios"][src]["amount"])
           and sc[s.lower()]["pct"] == round(an["scenarios"][src]["amount"] / CRANE_MUST["sum_insured"] * 100, 1)
           and sc[s.lower()]["source_scenario"] == src, (sc.get(s.lower()), an["scenarios"][src]["amount"]))
    ok("автокран: PML ≤ EML ≤ MFL", order_ok(sc), [sc[k]["amount"] for k in ("pml", "eml", "mfl")])
    ok("автокран: PML — крупная авария с ремонтом (50 %), без слов про отсутствие защиты",
       sc["pml"]["pct"] == 50.0 and "авария" in sc["pml"]["what"] and "нет" not in sc["pml"]["what"].split("—")[0]
       and "противоугон" not in sc["pml"]["what"], sc["pml"])
    ok("автокран без противоугонной: EML 100 % — угон или гибель, подпись про отсутствие системы",
       sc["eml"]["pct"] == 100.0 and "угон" in sc["eml"]["what"] and "нет" in sc["eml"]["what"], sc["eml"])
    ok("автокран: MFL 100 % — защита не сработала", sc["mfl"]["pct"] == 100.0 and "не сработала" in sc["mfl"]["what"])
    ret = an["retention"]
    ok("лимит удержания = risk_analytics", sc["retention"]["limit"] == ret["retention_limit"]
       and sc["retention"]["known"] is (ret["retention_limit"] is not None), (sc["retention"], ret.get("retention_limit")))
    lim = sc["retention"]["limit"]
    ok("удержание сравнивается с EML: compared_with = eml, превышение = EML − лимит",
       sc["retention"]["compared_with"] == "eml"
       and (lim is None or sc["retention"]["eml_excess"] == max(sc["eml"]["amount"] - lim, 0)), sc["retention"])
    ok("MFL сверх удержания — отдельной справкой (mfl_excess)",
       lim is None or sc["retention"]["mfl_excess"] == max(sc["mfl"]["amount"] - lim, 0), sc["retention"])
    ok("что взято по умолчанию — в assumptions с пометкой (противоугонная влияет на EML)",
       any(x["code"] == "as_protection_veh_eml" and "EML" in x["text"] for x in sc["assumptions"])
       and all("по умолчанию" in x["text"] for x in sc["assumptions"]), sc["assumptions"])
    ok("how: честно о перестановке названий в risk_analytics",
       any("переставлены" in h and "PML акта = EML модуля" in h for h in sc["how"]), sc["how"])
    ok("calibrated = 0", sc["calibrated"] == 0 and sc["pml"]["calibrated"] == 0)
    rows = {r["label"]: r for r in a["sections"][3]["rows"]}
    ok("раздел 4: строки PML, EML, MFL и лимит удержания",
       {"PML — вероятный максимальный убыток", "EML — оценочный максимальный убыток",
        "MFL — максимально возможный убыток", "Лимит собственного удержания"} <= set(rows), list(rows))
    ok("раздел 4: определения заказчика (штатно / частично / отказ; удержание — с EML)",
       any("сработала штатно" in p and "сработала частично" in p and "отказе защиты" in p
           and "сравнивается лимит собственного удержания" in p for p in a["sections"][3]["paragraphs"]))
    ok("раздел 4: порядок строк PML, EML, MFL и плитки для экрана в том же порядке",
       [r["label"][:3] for r in a["sections"][3]["rows"] if r["label"][:3] in ("PML", "EML", "MFL")] == ["PML", "EML", "MFL"]
       and [x["name"] for x in sc["tiles"]] == ["PML", "EML", "MFL"] and sc["tiles"][0]["label"].startswith("PML —"))
    if lim is not None:
        note = rows["Лимит собственного удержания"]["note"]
        ok("строка удержания: сравнение с EML и справка по MFL", "EML" in note and "MFL" in note, note)
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                        "optional": dict(CRANE_OPT, object_kind="truck_crane", protection="tracker")})
    sc2 = a2["scenarios"]
    ok("спутниковый поиск — EML ниже (75 %), PML 50 %, MFL 100 %",
       sc2["eml"]["pct"] == 75.0 and sc2["pml"]["pct"] == 50.0 and sc2["mfl"]["pct"] == 100.0
       and "противоугонная система есть" in sc2["eml"]["what"], sc2["eml"])
    ok("спутниковый поиск: PML ≤ EML ≤ MFL", order_ok(sc2))
    SCEN_REPORT["crane"] = {k: (sc[k]["amount"], sc[k]["pct"]) for k in ("pml", "eml", "mfl")}
    SCEN_REPORT["crane_tracker"] = {k: (sc2[k]["amount"], sc2[k]["pct"]) for k in ("pml", "eml", "mfl")}

    st, a = call("POST", "/act/make", {"lang": "ru", "must": WH_MUST, "optional": WH_OPT})
    sc = a["scenarios"]
    with db.tx() as con:
        an = ra_same(con, {"class_code": "9", "product_code": "0808", "object_type": "Склад",
                           "sum_insured": WH_MUST["sum_insured"], "object_value": WH_MUST["object_value"],
                           "region": WH_MUST["region"], "activity": "warehouse"},
                     {"protection": "alarm", "losses_3y": {"count": 0, "small_count": 0}})
    ok("склад класса 9: правило класса 9", sc["available"] and sc["rule"] == "property9", sc)
    ok("склад класса 9: PML/EML/MFL акта = EML/PML/MFL risk_analytics",
       all(sc[s.lower()]["amount"] == round(an["scenarios"][MAP_RA[s]]["amount"]) for s in ("PML", "EML", "MFL")),
       ([sc[s.lower()]["amount"] for s in ("PML", "EML", "MFL")],
        [an["scenarios"][s]["amount"] for s in ("PML", "EML", "MFL")]))
    ok("склад класса 9: PML ≤ EML ≤ MFL", order_ok(sc), [sc[k]["amount"] for k in ("pml", "eml", "mfl")])
    ok("склад класса 9: подписи — состояние защиты и «помещения не указаны»",
       sc["pml"]["what"].startswith("защита сработала штатно") and sc["eml"]["what"].startswith("защита сработала частично")
       and sc["mfl"]["what"].startswith("защита не сработала") and "не указаны" in sc["pml"]["what"],
       [sc[k]["what"] for k in ("pml", "eml", "mfl")])
    ok("склад: сумма ниже стоимости — доля в объяснении", any("0,8333" in h for h in sc["how"]), sc["how"])
    SCEN_REPORT["wh9"] = {k: (sc[k]["amount"], sc[k]["pct"]) for k in ("pml", "eml", "mfl")}
    must8 = {"product_code": "0807", "sum_insured": 5_000_000_000, "object_value": 5_000_000_000, "region": "Ташкент"}
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must8,
                                       "optional": {"object_kind": "warehouse", "seismic_zone": 9,
                                                    "construction": "reinforced"}})
    with db.tx() as con:
        an = ra_same(con, {"class_code": "8", "product_code": "0807", "object_type": "Склад", "sum_insured": 5e9,
                           "object_value": 5e9, "region": "Ташкент", "construction": "reinforced",
                           "activity": "warehouse"}, {"seismic_zone": 9})
    sc = a["scenarios"]
    ok("склад класса 8 в 9-балльной зоне: сценарии = risk_analytics (со сменой названий)",
       sc["rule"] == "property8"
       and all(sc[s.lower()]["amount"] == round(an["scenarios"][MAP_RA[s]]["amount"]) for s in ("PML", "EML", "MFL")),
       sc)
    ok("склад класса 8 (9 баллов): PML ≤ EML ≤ MFL", order_ok(sc), [sc[k]["amount"] for k in ("pml", "eml", "mfl")])
    ok("склад класса 8 (9 баллов): формула названа сценарием акта",
       all((sc[k]["formula"] or "").startswith(k.upper() + " =") for k in ("pml", "eml", "mfl")),
       [sc[k]["formula"] for k in ("pml", "eml", "mfl")])
    ok("склад класса 8 (9 баллов): MFL — пожар всего объекта (отсеки не указаны)",
       "пожар" in sc["mfl"]["what"] and "отсеки не указаны" in sc["mfl"]["what"], sc["mfl"]["what"])
    SCEN_REPORT["wh8_zone9"] = {k: (sc[k]["amount"], sc[k]["pct"]) for k in ("pml", "eml", "mfl")}
    # сейсмозона не указана: MFL — полное уничтожение, подпись говорит именно это
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must8,
                                       "optional": {"object_kind": "warehouse", "construction": "reinforced",
                                                    "protection": "sprinkler"}})
    sc = a["scenarios"]
    ok("склад класса 8 без сейсмозоны: PML ≤ EML ≤ MFL", order_ok(sc), [sc[k]["amount"] for k in ("pml", "eml", "mfl")])
    ok("склад класса 8 без сейсмозоны: MFL — «полное уничтожение: сейсмозона не указана», 100 %",
       "полное уничтожение" in sc["mfl"]["what"] and "сейсмозона не указана" in sc["mfl"]["what"]
       and sc["mfl"]["pct"] == 100.0 and "полное уничтожение" in sc["mfl"]["formula"], sc["mfl"])
    ok("склад класса 8 без сейсмозоны: PML и EML — пожар, отсеки не указаны (весь объект)",
       all("пожар" in sc[k]["what"] and "отсеки не указаны" in sc[k]["what"] for k in ("pml", "eml")),
       [sc[k]["what"] for k in ("pml", "eml")])
    SCEN_REPORT["wh8_sprinkler"] = {k: (sc[k]["amount"], sc[k]["pct"]) for k in ("pml", "eml", "mfl")}
    pcts = [sc[k]["pct_text"] for k in ("pml", "eml", "mfl")]
    ok("проценты сценариев — одинаковое число знаков", len({len(re.sub(r"[^\d,.]", "", p).partition(",")[2]) for p in pcts}) == 1,
       pcts)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(WH_MUST, product_code="0701"), "optional": {}})
    # с 30.09.2026 у класса 7 — простое правило шаблона класса (одна отправка / накопление), а не «не считается»
    ok("груз (класс 7): сценарий по правилу шаблона — одна отправка; без полей класса — страховая сумма",
       st == 200 and a["scenarios"]["available"] is True and a["scenarios"]["source"] == "template"
       and a["scenarios"]["rule"] == "shipment" and a["scenarios"]["pml"]["amount"] == WH_MUST["sum_insured"]
       and {x["code"] for x in a["scenarios"]["assumptions"]} == {"as_tpl_shipment", "as_tpl_accumulation"},
       a.get("scenarios"))
    ok("груз: в разделе 4 строки PML, EML, MFL (подпись из шаблона), строки «не считается» нет",
       any(r["label"].startswith("PML") and r.get("note") == "одна отправка" for r in a["sections"][3]["rows"])
       and not any(r["value"] == "не считается" for r in a["sections"][3]["rows"]), a["sections"][3]["rows"][:6])
    # собственных средств нет — удержание «не задан»
    with db.tx() as con:
        saved = db.rows(con, "SELECT * FROM company_financials")
        con.execute("DELETE FROM company_financials")
    try:
        st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})
    finally:
        with db.tx() as con:
            for r in saved:
                cols = list(r)
                con.execute(f"INSERT INTO company_financials ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                            [r[c] for c in cols])
    ret = a["scenarios"]["retention"]
    ok("без собственных средств и резервов — лимит «не задан»",
       ret["known"] is False and ret["limit"] is None and "не задан" in ret["basis"], ret)


def check_documents():
    print("22. Разбор документов DOCX/XLSX/PDF с текстом — без модели")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    CALLS.clear()
    st, b = upload([("contract.docx", DOCX_MIME, CONTRACT.read_bytes())], {"lang": "ru", "product_code": "0808"})
    ok("учебный договор принят", st == 200 and b.get("ok"), (st, b))
    ok("модель не вызывалась", not CALLS, len(CALLS))
    rec = {(r["key"], r["value"]): r for r in b["recognized"]}
    ok("recognized: страховая сумма, стоимость, срок, регион, тип, конструкция, год",
       {("sum_insured", "4 200 000 000"), ("object_value", "5 000 000 000"), ("term_days", "365"),
        ("region", "город Ташкент"), ("object_type", "склад готовой продукции"), ("construction", "кирпич"),
        ("year", "2012")} <= set(rec), list(rec))
    ok("у каждого значения source = document и «проверьте»",
       all(r["source"] == "document" and r["check_label"] == "проверьте" and "из документа" in r["note"]
           for r in b["recognized"]))
    pf = b.get("prefill") or {}
    ok("prefill: сумма, стоимость, регион, срок — с источником",
       pf.get("sum_insured", {}).get("value") == 4_200_000_000 and pf.get("object_value", {}).get("value") == 5e9
       and pf.get("region", {}).get("code") == "tashkent_city" and pf.get("term_days", {}).get("value") == 365
       and all(v["source"] == "document" and v["check_label"] == "из документа, проверьте" for v in pf.values()), pf)
    ok("файл помечен как разобранный, в модель не ушёл",
       b["files"][0]["parsed"] and not b["files"][0]["read_by_ai"] and b["files"][0]["view"] == "document")
    ok("ИНН организации из договора не взят", "301234567" not in _json.dumps(b, ensure_ascii=False))
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru",
                                       "must": dict(WH_MUST, sum_insured=4_000_000_000, object_value=5_000_000_000),
                                       "optional": {}})
    dk = {d["key"]: d for d in a["discrepancies"]}
    ok("расхождение: страховая сумма в документе и во вводе", "sum_insured" in dk
       and dk["sum_insured"]["priority"] == "document" and "4 200 000 000" in dk["sum_insured"]["text"], dk)
    ok("стоимость совпала — расхождения нет", "object_value" not in dk)
    ok("документ засчитан как представленный", a["inspection"]["documents"] is True)
    s1 = {r["label"]: r for r in a["sections"][0]["rows"]}
    ok("раздел 1: конструкция из документа", s1.get("Конструкция, материал стен", {}).get("value") == "кирпич", s1)

    # ФИО, паспорт, ПИНФЛ, адрес проживания — не извлекаются и не хранятся
    lines = ["ЗАЯВЛЕНИЕ НА СТРАХОВАНИЕ ИМУЩЕСТВА", "Страхователь: Иванов Иван Иванович",
             "Паспорт: AA1234567", "ПИНФЛ: 31234567890123", "Адрес проживания: г. Ташкент, ул. Навои, 5",
             "Телефон: +998 90 123 45 67", "Страховая сумма: 1 000 000 000 сум", "Материал стен: кирпич"]
    st, b = upload([("zayavlenie.docx", DOCX_MIME, docx_bytes(lines))], {"lang": "ru", "product_code": "0808"})
    dump = _json.dumps(b, ensure_ascii=False)
    ok("заявление разобрано: сумма взята", any(r["key"] == "sum_insured" for r in b.get("recognized") or []), dump[:400])
    ok("ФИО, паспорт, ПИНФЛ, адрес проживания и телефон не извлечены",
       not any(x in dump for x in ("Иванов", "AA1234567", "31234567890123", "Навои", "123 45 67")), dump[:600])
    with db.tx() as con:
        stored = _json.dumps(db.rows(con, "SELECT result_json, files_json FROM act_uploads WHERE id=?", b["session"]),
                             ensure_ascii=False)
        journal = _json.dumps(db.rows(con, "SELECT detail FROM audit WHERE entity=?", f"act_upload:{b['session']}"),
                              ensure_ascii=False)
    ok("в базе и в журнале данных людей нет",
       not any(x in stored + journal for x in ("Иванов", "AA1234567", "31234567890123", "Навои")))
    ok("в журнале — только счётчики", "parsed_docs" in journal and "1 000 000 000" not in journal, journal[:300])

    # XLSX-выгрузка по технике: марка, модель, год, VIN, госномер
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for row in (["Заявление на страхование спецтехники"], ["Страховая сумма", "2 945 000 000 сум"],
                ["Стоимость имущества", "3 100 000 000 сум"], ["Марка", "XCMG"], ["Модель", "QY50K5D"],
                ["Год выпуска", "2026"], ["Заводской номер (VIN)", "LXGCPA393TA006921"],
                ["Государственный номер", "01 A 123 BC"]):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    CALLS.clear()
    st, b = upload([("vygruzka.xlsx", XLSX_MIME, buf.getvalue())], {"lang": "ru", "product_code": "0318"})
    got = {(r["key"], r["value"]) for r in b.get("recognized") or []}
    ok("XLSX: марка, модель, год, VIN, госномер, сумма — из документа",
       {("brand", "XCMG"), ("model", "QY50K5D"), ("year", "2026"), ("serial_no", "LXGCPA393TA006921"),
        ("reg_no", "01 A 123 BC"), ("sum_insured", "2 945 000 000")} <= got and not CALLS, got)

    # PDF с текстовым слоем — парсером; скан — модели
    doc = pymupdf.open()
    page = doc.new_page()
    font = act._fonts()[0]
    tw = pymupdf.TextWriter(page.rect)
    for k, ln in enumerate(["ДОГОВОР СТРАХОВАНИЯ ИМУЩЕСТВА № 7/2026", "Страховая сумма: 1 000 000 000 сум",
                            "Стоимость имущества: 1 100 000 000 сум",
                            "Адрес объекта: Самаркандская область, г. Самарканд", "Срок страхования: 12 месяцев"]):
        tw.append((72, 72 + 18 * k), ln, font=font, fontsize=11)
    tw.write_text(page)
    CALLS.clear()
    st, b = upload([("dogovor.pdf", "application/pdf", doc.tobytes())], {"lang": "ru", "product_code": "0808"})
    ok("PDF с текстом разобран без модели", st == 200 and not CALLS and b["files"][0]["parsed"]
       and (b.get("prefill") or {}).get("region", {}).get("code") == "samarkand", b.get("prefill"))
    CALLS.clear()
    st, b = upload([("scan.pdf", "application/pdf", pdf_pages(1))], {"lang": "ru"})
    ok("скан PDF без текста — читает модель", st == 200 and len(CALLS) == 1 and not b["files"][0]["parsed"])

    # ограничения: zip-бомба, DTD, макросы
    bomb = docx_bomb(60)
    st, b = upload([("bomb.docx", DOCX_MIME, bomb)], {"lang": "ru"})
    ok("zip-бомба в DOCX отклонена до распаковки", st == 422 and b["rejected"]
       and "слишком большой" in b["rejected"][0]["error"] and len(bomb) < 1024 * 1024, (st, b, len(bomb)))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><w:document/>')
    st, b = upload([("dtd.docx", DOCX_MIME, buf.getvalue())], {"lang": "ru"})
    ok("DOCX с DTD/ENTITY отклонён", st == 422 and "DTD" in b["rejected"][0]["error"], b)
    st, b = upload([("macro.docx", DOCX_MIME, docx_bytes(["Страховая сумма: 1 000 000 000 сум"],
                                                          {"word/vbaProject.bin": b"\x00" * 100}))], {"lang": "ru"})
    ok("макросы не исполняются — только пометка", st == 200 and any("макрос" in n for n in b["notes"]), b.get("notes"))
    set_limits(doc_max_unzip_mb=1)
    try:
        st, b = upload([("contract.docx", DOCX_MIME, docx_bomb(2))], {"lang": "ru"})
        ok("предел распаковки — из настроек (1 МБ)", st == 422, st)
    finally:
        clear_settings()
    st, b = upload([("x.zip", "application/zip", docx_bomb(1).replace(b"word/document.xml", b"other/documen.xml"))],
                   {"lang": "ru"})
    ok("прочий zip — формат не принимается", st == 422 and "DOCX или XLSX" in b["rejected"][0]["error"], b)


def check_franchise_apply():
    print("23. Франшиза: предложение по основанию, франшиза сотрудника, обязательный вид")
    from app import franchise as frm
    from app.risk_analytics import load_thresholds
    fresh()
    with db.tx() as con:
        th = load_thresholds(con)
        base, floor, applied = expected_rate(con, 20)            # умеренный уровень: без фото, открытая площадка
    S = CRANE_MUST["sum_insured"]
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})
    fr = a["franchise"]
    ok("без оснований: «Франшиза не требуется», статус none, премия не меняется",
       fr["text"] == "Франшиза не требуется" and fr["status"] == "none" and not fr["applied"]
       and fr["premium_after"] == fr["premium_before"] == a["premium"]["amount"] and fr["delta"] == 0, fr)
    ok("без оснований: альтернативы есть (мероприятия, сумма к стоимости)",
       {x["code"] for x in fr["alternatives"]} >= {"alt_sum_up"} and fr["how"], fr["alternatives"])

    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, losses_3y={"count": 2, "small_count": 2})})
    fr = a["franchise"]
    ok("основание есть: статус proposed, тип безусловная", fr["status"] == "proposed"
       and fr["type"] == "unconditional" and fr["needed"], fr)
    ok("размер — внутри вилки порогов", fr["size"]["from_pct"] <= fr["size_pct"] <= fr["size"]["to_pct"]
       and fr["size_amount"] == round(S * fr["size_pct"] / 100), fr)
    ok("множитель — из franchise.what_if (экспертная кривая)", abs(fr["multiplier"] - frm._mult(th, fr["size_pct"])) < 1e-4,
       (fr["multiplier"], frm._mult(th, fr["size_pct"])))
    rate = a["rate"]["applied_pct"]
    want_rate = max(round(rate * fr["multiplier"], 4), a["rate"]["min_pct"] or 0)
    ok("премия с франшизой = ставка акта × множитель, не ниже минимума",
       fr["premium_after"] == round(want_rate / 100 * S) and fr["rate_after"] == want_rate, (fr, want_rate))
    ok("предложенная франшиза в премию акта не включена", a["premium"]["amount"] == fr["premium_before"]
       and a["premium"]["franchise_applied"] is False and "андеррайтер" in fr["text"])

    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, deductible={"pct": 1, "type": "unconditional"})})
    fr = a["franchise"]
    hand_rate = max(round(applied * 0.85, 4), floor)          # 1 % по справочнику коэффициентов: ×0,85
    hand = round(hand_rate / 100 * S)
    ok("своя франшиза сотрудника 1 %: премия пересчитана руками", a["premium"]["amount"] == hand
       and fr["premium_after"] == hand and fr["premium_before"] == round(applied / 100 * S), (a["premium"], hand))
    ok("текст «Франшиза применена по решению сотрудника»",
       fr["text"].startswith("Франшиза применена по решению сотрудника") and fr["applied"]
       and fr["applied_by"] == "employee", fr["text"])
    rows = {r["label"]: r for r in a["sections"][3]["rows"]}
    ok("раздел 4: та же премия и ставка с франшизой",
       rows["Страховая премия"]["value"] == act.money(hand, "ru")
       and rows["Ставка с учётом франшизы"]["value"] == act.pct(hand_rate, "ru")
       and act.money(round(applied / 100 * S), "ru") in rows["Страховая премия"]["note"], rows.get("Страховая премия"))
    ok("андеррайтеру — подтвердить франшизу", any("франшизу, применённую" in c for c in a["decision"]["checks"]))
    # служебного имени функции (what_if) в тексте нет — проверяем смысл: множитель модуля франшизы и ставка акта
    mult_txt = act._mult(fr["multiplier"], "ru")
    ok("how: множитель модуля франшизы и ставка акта × множитель, премия модуля расчёта ставок не переносится",
       any("из модуля франшизы" in h and "множитель " + mult_txt in h for h in fr["how"])
       and any(h.startswith("Ставка акта") and act.pct(applied, "ru") + " × " + mult_txt in h
               and act.pct(hand_rate, "ru") in h and "не переносится" in h for h in fr["how"])
       and not any("what_if" in h or "движ" in h for h in fr["how"]), fr["how"])
    applied_id = a["id"]
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                        "optional": dict(CRANE_OPT, deductible={"amount": 29_450_000})})
    ok("франшиза суммой = 1 %", a2["franchise"]["size_pct"] == 1.0 and a2["premium"]["amount"] == hand, a2["franchise"])
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                        "optional": dict(CRANE_OPT, deductible={"pct": 7})})
    ok("выше потолка класса 3 (5 %) — предупреждение, решение андеррайтера",
       a2["franchise"]["warning"] and "потолка" in a2["franchise"]["warning"] and a2["franchise"]["applied"], a2["franchise"])
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                        "optional": dict(CRANE_OPT, deductible={"pct": 80})})
    ok("франшиза 80 % — ошибка ввода 422", st == 422 and "deductible" in a2.get("errors", {}), a2)
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                        "optional": dict(CRANE_OPT, deductible={"pct": 1, "amount": 5})})
    ok("pct и amount вместе — 422", st == 422, a2)
    must = {"product_code": "0820", "sum_insured": 1_000_000_000, "object_value": 1_000_000_000, "region": "Ташкент"}
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": must,
                                        "optional": {"deductible": {"pct": 1}, "location": "construction"}})
    fr = a2["franchise"]
    ok("обязательный вид: франшиза сотрудника не применена, премия по акту",
       st == 200 and fr["status"] == "statutory" and not fr["applied"]
       and a2["premium"]["amount"] == a2["premium"]["before_franchise"] and "обязательным видам" in fr["warning"], fr)
    # низкий уровень: ставка уже на минимуме продукта — франшиза премию не снижает, и акт говорит об этом
    low = dict(CRANE_OPT, object_kind="truck_crane", year=date.today().year, condition="new",
               documents_provided=True, deductible={"pct": 1})
    st, a2 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": low})
    fr = a2["franchise"]
    ok("ставка на минимуме: премия с франшизой не ниже минимальной ставки",
       a2["risk"]["level"] == "low" and fr["floor_applied"] and a2["premium"]["amount"] == round(floor / 100 * S)
       and "минимальную ставку" in fr["text"], (a2["risk"]["level"], fr))
    return applied_id, hand


def check_measures():
    print("24. Рекомендации страхователю")
    fresh()
    with db.tx() as con:
        _b, floor, applied = expected_rate(con, 20)
    S = CRANE_MUST["sum_insured"]
    P = round(applied / 100 * S)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, object_kind="truck_crane")})
    ms = {m["code"]: m for m in a["measures"]}
    ok("спецтехника на открытой площадке: охраняемая стоянка и спутниковый мониторинг",
       {"sp_parking_guarded", "sp_gps", "sp_crane_setup", "sp_operator"} <= set(ms), list(ms))
    ok("эффект мероприятия — из справочника коэффициентов (−10 %)",
       ms["sp_parking_guarded"]["effect_pct"] == -10.0 and ms["sp_gps"]["effect_pct"] == -10.0, ms["sp_gps"])
    ok("без эффекта — «на ставку не влияет, снижает вероятность убытка»",
       ms["sp_operator"]["premium_delta"] is None and "снижает вероятность" in ms["sp_operator"]["effect_text"])
    want = max(round(P * 0.9 * 0.9), round(floor / 100 * S))
    summ = a["measures_summary"]
    ok("скидки перемножаются и не ниже минимальной ставки продукта",
       summ["premium_after"] == want and summ["premium_before"] == P, (summ, want))
    ok("каждое мероприятие: что, зачем, срок, calibrated=0",
       all(m["text"] and m["why"] and m["deadline_days"] and m["calibrated"] == 0 for m in a["measures"]))
    s5 = a["sections"][4]
    li = next((x for x in s5["lists"] if x["title"] == "Рекомендации страхователю"), None)
    ok("раздел 5: подраздел «Рекомендации страхователю»", li and len(li["items"]) == len(a["measures"])
       and "Зачем:" in li["items"][0] and "Срок:" in li["items"][0], li)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, object_kind="truck_crane", guard=True)})
    ok("под охраной — охраняемая стоянка не нужна", "sp_parking_guarded" not in {m["code"] for m in a["measures"]})

    st, a = call("POST", "/act/make", {"lang": "ru", "must": WH_MUST, "optional": WH_OPT})
    ms = {m["code"]: m for m in a["measures"]}
    ok("склад класса 9: мероприятия таблицы preventive_measures и хранение на стеллажах",
       {"burglary_protection", "inventory_control", "wh_storage"} <= set(ms)
       and ms["burglary_protection"]["source"] == "preventive_measures", list(ms))
    ok("склад: сумма 1 млрд — охранная сигнализация обязательна, срок 60 дней",
       ms["burglary_protection"]["mandatory"] and ms["burglary_protection"]["deadline_days"] == 60,
       ms["burglary_protection"])
    st, u = call("GET", f"/act/{a['id']}", params={"lang": "uz"})
    ok("uz: мероприятия таблицы переведены", not re.search(r"[А-Яа-яЁё]", " ".join(
        m["text"] + m["why"] for m in u["measures"])), [m["text"] for m in u["measures"]])
    must8 = {"product_code": "0807", "sum_insured": 5_000_000_000, "object_value": 5_000_000_000, "region": "Ташкент"}
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must8,
                                       "optional": {"object_kind": "warehouse", "construction": "wood"}})
    ms = {m["code"]: m for m in a["measures"]}
    ok("склад класса 8 из дерева: огнезащитная обработка со скидкой (из движка)",
       "fire_treatment" in ms and ms["fire_treatment"]["effect_pct"] and ms["fire_treatment"]["effect_pct"] < 0
       and ms["fire_treatment"]["mandatory"], ms.get("fire_treatment"))
    must = {"product_code": "0820", "sum_insured": 1_000_000_000, "object_value": 1_000_000_000, "region": "Ташкент"}
    st, a = call("POST", "/act/make", {"lang": "ru", "must": must, "optional": {}})
    ok("обязательный вид: скидок на премию нет", all(m["premium_delta"] is None for m in a["measures"]))


def check_new_langs(aid):
    print("25. Новые блоки на трёх языках")
    for lang in ("uz", "en"):
        st, a = call("GET", f"/act/{aid}", params={"lang": lang})
        fr = a["franchise"]
        texts = ([r["label"] for s in a["sections"] for r in s["rows"]]
                 + [r.get("note") or "" for r in a["sections"][3]["rows"]]
                 + [li["title"] for s in a["sections"] for li in s["lists"]]
                 + [p for s in a["sections"] for p in s["paragraphs"]]
                 + [fr["text"], fr["type_label"] or ""] + fr["how"] + [x["text"] for x in fr["alternatives"]]
                 + [a["scenarios"][k]["what"] for k in ("pml", "eml", "mfl")] + a["scenarios"]["how"]
                 + [x["text"] for x in a["scenarios"]["assumptions"]] + [a["scenarios"]["retention"]["basis"]]
                 + [m["text"] + " " + m["why"] + " " + m["effect_text"] for m in a["measures"]]
                 + [a["measures_summary"]["text"] or ""])
        cyr = [x for x in texts if re.search(r"[А-Яа-яЁё]", x or "")]
        ok(f"{lang}: сценарии, франшиза, мероприятия без кириллицы", not cyr, cyr[:4])
    st, a = call("GET", f"/act/{aid}", params={"lang": "en"})
    ok("en: франшиза применена сотрудником", a["franchise"]["text"].startswith("Deductible applied by staff"))


def check_new_files(aid, hand):
    print("26. Word и PDF содержат новые блоки и ту же премию")
    st, blob, h = call("GET", f"/act/{aid}.docx", raw=True)
    xml = zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8")
    plain = re.sub(r"<[^>]+>", "", xml)
    need = ["PML — вероятный максимальный убыток", "EML — оценочный максимальный убыток",
            "MFL — максимально возможный убыток", "Лимит собственного удержания", "Рекомендации страхователю",
            "Франшиза применена по решению сотрудника", "Ставка с учётом франшизы", act.money(hand, "ru")]
    ok("DOCX: сценарии, франшиза, рекомендации, премия с франшизой", all(x in plain for x in need),
       [x for x in need if x not in plain])
    st, blob, h = call("GET", f"/act/{aid}.pdf", raw=True)
    text = pdf_text(pymupdf.open(stream=blob, filetype="pdf"))
    flat = [x.replace(" ", " ") for x in need]           # pdf_text приводит неразрывный пробел к обычному
    ok("PDF: те же блоки и премия", all(x in text for x in flat), [x for x in flat if x not in text])
    st, a = call("GET", f"/act/{aid}")
    ok("премия в JSON, в разделе 4 и в файлах одна", a["premium"]["amount"] == hand
       and a["premium"]["text"] == act.money(hand, "ru"))


def check_speed():
    print("27. Акт формируется быстро и без сети")
    import time as _t
    fresh()
    CALLS.clear()
    t0 = _t.monotonic()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, deductible={"pct": 1})})
    sec = _t.monotonic() - t0
    ok("акт со всеми блоками — меньше 2 секунд, без модели", st == 200 and sec < 2 and not CALLS, sec)
    with db.tx() as con:
        rows = db.rows(con, "SELECT detail FROM audit WHERE entity=?", f"act:{a['id']}")
    ok("в журнале — статус франшизы и счётчики, без значений",
       rows and '"franchise": "applied"' in rows[0]["detail"] and "29 450 000" not in rows[0]["detail"], rows)


# ------------------------------------------------------------------ 28–35. замечания контролёра (30.09.2026)

def xlsx_big(rows, cols, sheets=1) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    words = ["склад", "сумма", "Ташкент", "объект", "qiymat", "ombor", "value", "итого", "страхование"]
    for s in range(sheets):
        if s:
            ws = wb.create_sheet(f"Лист {s + 1}")
        for r in range(rows):
            ws.append([f"{words[(r + c) % len(words)]} {r}-{c} {words[(r * c) % len(words)]}" for c in range(cols)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def check_doc_limits():
    print("28. Разбор документов: пределы до разбора, срок внутри циклов, не больше двух одновременно")
    import threading
    import time as _t
    from app import act_extras as ax, docparse as D, ingest
    fresh()
    CALLS.clear()
    for rows, cols, sheets, name in ((500, 60, 1, "500 × 60"), (2000, 3, 3, "2000 строк × 3 листа")):
        blob = xlsx_big(rows, cols, sheets)
        t0 = _t.monotonic()
        st, b = upload([("big.xlsx", XLSX_MIME, blob)], {"lang": "ru", "product_code": "0808"})
        sec = _t.monotonic() - t0
        ok(f"XLSX {name} ({len(blob) // 1024} КБ) — быстрее 3 с", st == 200 and sec < 3, (st, round(sec, 2)))
        ok(f"XLSX {name}: честная пометка «прочитана только часть»",
           any("только часть документа" in n for n in b.get("notes") or []) and b["files"][0]["parsed"], b.get("notes"))
    ok("большие XLSX в модель не уходили", not CALLS, len(CALLS))
    # пределы работают до разбора: листов, строк, колонок, длина ячейки
    folder = Path(tempfile.mkdtemp(prefix="act-lim-"))
    try:
        p = folder / "w.xlsx"
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["x" * 900, "короткая"])
        ws.append([f"c{c}" for c in range(40)])
        for r in range(300):
            ws.append([f"r{r}", "a", "b"])
        for s in range(4):
            wb.create_sheet(f"S{s}").append(["лист", s])
        wb.save(p)
        got = ax.read_limited(p, {})
        t1 = got["tables"][0]["rows"]
        ok("предел: не больше 3 листов", len(got["tables"]) == 3, [t["name"] for t in got["tables"]])
        ok("предел: не больше 200 строк и 30 колонок с листа",
           len(t1) == 200 and max(len(r) for r in t1) == 30 and got["tables"][0]["обрезан"],
           (len(t1), max(len(r) for r in t1)))
        ok("предел: ячейка и строка не длиннее 500 знаков",
           all(len(" | ".join(r)) <= 500 for r in t1) and all(len(ln) <= 500 for ln in got["text"].splitlines()))
        ok("предел: отметка truncated", got["truncated"] is True)
        small = ax.read_limited(p, {"doc_max_cells": 100})
        ok("предел ячеек на файл — из настроек (100)", sum(len([c for c in r if c]) for t in small["tables"]
                                                             for r in t["rows"]) <= 100, small["tables"][0]["rows"][:2])
        # учебный договор: тот же prefill и те же значения, что у прежнего разбора без пределов
        with db.tx() as con:
            old = ax.parse_document(con, CONTRACT, "8")
            new = ax.parse_document_limited(con, CONTRACT, "8", {})
        ok("flow150_contract.docx: prefill и значения те же, что без пределов",
           old["prefill"] == new["prefill"] and old["items"] == new["items"] and not new["notes"], (old, new))
        r_old, r_new = ingest.read_file(CONTRACT), ax.read_limited(CONTRACT, {})
        ok("потоковое чтение DOCX = ingest.read_docx (текст и таблицы)",
           r_old["text"] == r_new["text"] and r_old["tables"] == r_new["tables"])
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    # кэш свёртки: поведение прежнее, список карты — свой у каждого вызова
    s = "Sugʻurta summasi: 1 000 000 сум"
    f1, i1 = D.fold_map(s)
    f2, i2 = D._fold_map_raw(s)
    i1.append(-1)
    ok("fold_map из кэша = прежний расчёт, карта не делится между вызовами",
       (f1, i1[:-1]) == (f2, i2) and D.fold_map(s)[1] == i2 and D.fold(s) == f2)
    # срок: проверка внутри цикла освобождает поток — разбор возвращается, поток не остаётся висеть
    orig = ingest.detect_language
    ingest.detect_language = lambda text: (_t.sleep(0.4), orig(text))[1]
    set_limits(doc_parse_sec=0.2)
    try:
        before = threading.active_count()
        own_before = [x.name for x in threading.enumerate() if not x.name.startswith("AnyIO worker")]
        t0 = _t.monotonic()
        st, b = upload([("contract.docx", DOCX_MIME, CONTRACT.read_bytes())], {"lang": "ru", "product_code": "0808"})
        sec = _t.monotonic() - t0
        ok("срок одного файла: «не разобран: слишком большой», запрос продолжается",
           st == 200 and not b["files"][0]["parsed"] and any("слишком большой" in n for n in b["notes"])
           and sec < 2, (st, b.get("notes"), sec))
        # ждём с пределом, а не проверяем мгновенно: потоки пула сервера (AnyIO worker) живут своим сроком и к
        # разбору не относятся; поток разбора должен закончиться за время ожидания
        def own():
            return [x.name for x in threading.enumerate() if not x.name.startswith("AnyIO worker")]
        deadline = _t.monotonic() + 5
        while len(own()) > len(own_before) and _t.monotonic() < deadline:
            _t.sleep(0.05)
        ok("поток разбора освобождён (лишних потоков нет)", len(own()) <= len(own_before),
           (before, threading.active_count(), sorted(set(own()) - set(own_before))))
        with db.tx() as con:
            j = db.rows(con, "SELECT detail FROM audit WHERE entity=? AND action='акт: документ не разобран'",
                        f"act_upload:{b['session']}")
        ok("превышение срока — в журнале", j and "timeout" in j[0]["detail"], j)
    finally:
        ingest.detect_language = orig
        clear_settings()
    # срок на все документы запроса: второй файл уже не разбирается
    ingest.detect_language = lambda text: (_t.sleep(0.35), orig(text))[1]
    set_limits(doc_parse_sec=5, doc_parse_total_sec=0.3)
    try:
        st, b = upload([("a.docx", DOCX_MIME, CONTRACT.read_bytes()), ("b.docx", DOCX_MIME, CONTRACT.read_bytes())],
                       {"lang": "ru", "product_code": "0808"})
        ok("общий срок запроса: оба файла помечены, ответ есть",
           st == 200 and not any(f["parsed"] for f in b["files"])
           and len(b["documents"]) == 2 and all(d["notes"] for d in b["documents"]), (st, b.get("documents")))
    finally:
        ingest.detect_language = orig
        clear_settings()
    # не больше двух документов одновременно: оба места заняты — файл честно «сервер занят»
    ok("мест для разбора — два", ax.PARSE_SLOTS == 2)
    took = [ax.parse_slot(0), ax.parse_slot(0)]
    ok("третий документ места не получает", not ax.parse_slot(0))
    set_limits(doc_parse_total_sec=0.3)
    try:
        st, b = upload([("contract.docx", DOCX_MIME, CONTRACT.read_bytes())], {"lang": "ru", "product_code": "0808"})
        ok("семафор занят: «сервер занят разбором», запрос не падает",
           st == 200 and any("сервер занят" in n for n in b["notes"]) and not b["files"][0]["parsed"], b.get("notes"))
    finally:
        for x in took:
            if x:
                ax.parse_slot_release()
        clear_settings()
    st, b = upload([("contract.docx", DOCX_MIME, CONTRACT.read_bytes())], {"lang": "ru", "product_code": "0808"})
    ok("места освобождены — договор снова разбирается", st == 200 and b["files"][0]["parsed"])
    errs = ae.check_settings({"limits": dict(ae.DEFAULT_SETTINGS["limits"], doc_max_rows=5)})
    ok("настройки: предел строк проверяется", any("doc_max_rows" in e for e in errs), errs)
    ok("настройки по умолчанию: 10 000 ячеек, 3 000 абзацев, 200 строк, 30 колонок, 3 листа, 500 знаков, "
       "4 000 знаков в строке DOCX, 200 000 знаков, 5 с и 8 с",
       {k: ae.DEFAULT_SETTINGS["limits"][k] for k in ax.DOC_LIMITS} == ax.DOC_LIMITS
       and not ae.check_settings({"limits": dict(ae.DEFAULT_SETTINGS["limits"])}))


def check_dtd_prolog():
    print("29. DTD/ENTITY ищется во всём прологе XML, а не в первых 4096 байтах")
    from app import act_extras as ax
    fresh()
    xml = ('<?xml version="1.0"?><!--' + "x" * 5000 + '--><!DOCTYPE x [<!ENTITY a "aaaa">]>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", xml)
    st, b = upload([("dtd5000.docx", DOCX_MIME, buf.getvalue())], {"lang": "ru"})
    ok("комментарий 5000 байт перед DOCTYPE — файл отклонён", st == 422 and "DTD" in b["rejected"][0]["error"], (st, b))
    long_c = b'<?xml version="1.0"?><!--' + b"-" * 3 + b"y" * 200_000 + b'-->\n<!ENTITY e "x"><r/>'
    ok("комментарий 200 КБ (через границы кусков) — ENTITY найдена", ax.xml_prolog_has_dtd(io.BytesIO(long_c), 10 ** 7))
    ok("обычный XML с комментарием и инструкцией — годится",
       not ax.xml_prolog_has_dtd(io.BytesIO(b'\xef\xbb\xbf<?xml version="1.0"?>\n<!-- c --><?pi x?><r><!-- <!DOCTYPE --></r>'),
                                 10 ** 6))
    ok("UTF-16 с DOCTYPE — отклоняется", ax.xml_prolog_has_dtd(
        io.BytesIO('<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE r><r/>'.encode("utf-16")), 10 ** 6))
    ok("DOCTYPE после корня (в тексте) не ищется — пролог закончился", not ax.xml_prolog_has_dtd(
        io.BytesIO(b"<r>" + b"z" * 10 + b"&lt;!DOCTYPE</r>"), 10 ** 6))
    st, b = upload([("contract.docx", DOCX_MIME, CONTRACT.read_bytes())], {"lang": "ru", "product_code": "0808"})
    ok("настоящий договор проверку проходит", st == 200 and b["files"][0]["parsed"])


def check_fr_proposed():
    print("30. Предложенная франшиза: минимум ставки и множитель выше 2 %")
    fresh()
    with db.tx() as con:
        _b, floor, _a = expected_rate(con, 0)
    S = CRANE_MUST["sum_insured"]
    low = dict(CRANE_OPT, object_kind="truck_crane", year=date.today().year, condition="new",
               documents_provided=True, want_lower_premium=True)
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": low})
    fr = a["franchise"]
    ok("низкий уровень + просьба клиента: франшиза предложена", fr["status"] == "proposed" and fr["size_pct"] > 0, fr)
    ok("предложенная упирается в минимум: floor_applied и то же пояснение, что у применённой",
       fr["floor_applied"] is True and fr["premium_after"] == round(floor / 100 * S)
       and "минимальную ставку продукта" in fr["text"], (fr["floor_applied"], fr["text"]))
    # высокий уровень (вилка 5–10 %, потолок класса 3 — 5 %): вилка 5 % — одно число, множитель — продолжение кривой
    fr_hi = {"code": "fr_advise_range", "needed": True, "status": "proposed", "type": "unconditional",
             "size": {"from_pct": 5.0, "to_pct": 5.0, "from_amount": S * 0.05, "to_amount": S * 0.05},
             "size_pct": 5.0, "size_amount": round(S * 0.05), "premium_before": 100, "premium_after": 90,
             "floor_applied": True, "engine": {"extrapolated": True},
             "how": [{"code": "frh_size", "params": {"from": 5.0, "to": 5.0, "pct": 5.0, "losses": 0}},
                     {"code": "frh_apply", "params": {}}]}
    text = act._fr_text(fr_hi, "ru")
    ok("вилка с равными границами — «5 %», а не «от 5 % до 5 %»", "франшизу 5 %" in text and "от 5" not in text, text)
    ok("предложенная выше 2 %: «экспертное продолжение», и пояснение о минимуме", "экспертное продолжение" in text
       and "минимальную ставку продукта" in text, text)
    how = act._fr_how_text(fr_hi["how"][0], "ru")
    ok("как посчитано: вилка «ровно 5 %»", "ровно 5 %" in how and "–" not in how, how)
    for lang in ("uz", "en"):
        tx_ = act._fr_text(fr_hi, lang)
        ok(f"{lang}: пояснения о минимуме и о 2 % переведены", not re.search(r"[А-Яа-яЁё]", tx_), tx_)


def check_alt_base():
    print("31. «Вместо франшизы можно»: все варианты — от премии без франшизы")
    fresh()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, object_kind="truck_crane",
                                                        deductible={"pct": 1, "type": "unconditional"})})
    fr = a["franchise"]
    P = a["premium"]["before_franchise"]
    alts = fr["alternatives"]
    ok("франшиза применена, премия акта с ней ниже базы", fr["applied"] and a["premium"]["amount"] < P, a["premium"])
    ok("у каждого варианта база — премия без франшизы",
       alts and all(x["base_premium"] == P for x in alts), [(x["code"], x.get("base_premium")) for x in alts])
    ok("premium_delta = premium − премия без франшизы",
       all(x["premium"] is None or x["premium_delta"] == x["premium"] - P for x in alts),
       [(x["code"], x["premium"], x["premium_delta"]) for x in alts])
    m = next((x for x in alts if x["code"] == "alt_measures"), None)
    want = a["measures_summary"]
    ok("мероприятия вместо франшизы: скидки от базы без франшизы",
       m is not None and m["premium"] == P + m["premium_delta"] and m["premium_delta"] < 0, (m, want))
    ok("в тексте каждого варианта — «Вместо франшизы»", all(x["text"].startswith("Вместо франшизы") for x in alts),
       [x["text"] for x in alts])
    st, u = call("GET", f"/act/{a['id']}", params={"lang": "en"})
    ok("en: «Instead of a deductible», ссылка на норму — по-английски",
       all(x["text"].startswith("Instead of a deductible") for x in u["franchise"]["alternatives"])
       and all(not re.search(r"[А-Яа-яЁё]", x["legal_ref"] or "") for x in u["franchise"]["alternatives"]),
       u["franchise"]["alternatives"])


def check_protection_other():
    print("32. Защита для класса без списка — не ошибка, а пометка")
    fresh()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": dict(WH_MUST, product_code="0701"),
                                       "optional": {"protection": "alarm"}})
    ok("груз (класс 7) с protection — 200, не ошибка ввода", st == 200 and a.get("ok"), (st, a.get("errors")))
    asm = [x["text"] for x in a["scenarios"]["assumptions"]]
    ok("пометка в «принято по умолчанию»", any("для этого класса" in x and "не учтено" in x for x in asm), asm)
    ok("в разделе 4 есть список «Принято по умолчанию»",
       any(li["title"] == "Принято по умолчанию (уточните)" for li in a["sections"][3]["lists"]))
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": dict(CRANE_OPT, protection="sprinkler")})
    ok("класс 3 с кодом защиты имущества (sprinkler) — по-прежнему 422", st == 422 and "protection" in a.get("errors", {}), a)


def check_lang_fields():
    print("33. uz/en: формула и ссылка на норму — не на чужом языке")
    fresh()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": dict(CRANE_OPT, object_kind="truck_crane")})
    ok("ru: формула есть, ссылка — «Положение № 1806, п. 15»", a["scenarios"]["pml"]["formula"]
       and a["scenarios"]["retention"]["legal_ref"] == "Положение № 1806, п. 15", a["scenarios"]["retention"])
    for lang, ref in (("uz", "1806-son Nizom, 15-band"), ("en", "Regulation No. 1806, para. 15")):
        st, u = call("GET", f"/act/{a['id']}", params={"lang": lang})
        sc = u["scenarios"]
        ok(f"{lang}: formula = null у всех трёх сценариев", all(sc[k]["formula"] is None for k in ("pml", "eml", "mfl")))
        ok(f"{lang}: legal_ref удержания — на языке акта", sc["retention"]["legal_ref"] == ref, sc["retention"]["legal_ref"])
        texts = [sc["definitions"], sc["retention"]["text"] or ""] + [x["label"] + " " + x["what"] + " " + x["pct_text"]
                                                                        for x in sc["tiles"]]
        ok(f"{lang}: определения, плитки и удержание без кириллицы", not any(re.search(r"[А-Яа-яЁё]", x) for x in texts),
           texts)


def check_minor():
    print("34. Мелочи: управляющие байты, проценты, подпись года")
    src = (Path(act.__file__)).read_bytes()
    bad = [b for b in src if (b < 32 and b not in (9, 10, 13)) or b == 127]
    ok("в app/act.py нет сырых управляющих символов", not bad, bad[:5])
    ok("сигнатура zip записана как b\"PK\\x03\\x04\"", b'b"PK\\x03\\x04"' in src)
    ok("_format_of по-прежнему узнаёт DOCX", act._format_of(docx_bytes(["Страховая сумма: 1 сум"])) == "docx")
    sc = {"available": True, "order": "classic", "class_code": "9", "rule": "property9", "k": 1,
          "items": {s: {"amount": p * 10, "pct": p, "what": "sc_w_c9", "state": "sc_state_" + s.lower()}
                    for s, p in (("PML", 15.0), ("EML", 37.5), ("MFL", 100.0))},
          "retention": {"known": False}, "assumptions": []}
    v = act._scenarios_view(sc, {"class_code": "9"}, "ru")["json"]
    ok("проценты рядом — одинаково: 15,0 % · 37,5 % · 100,0 %",
       [v[k]["pct_text"] for k in ("pml", "eml", "mfl")] == ["15,0 %", "37,5 %", "100,0 %"],
       [v[k]["pct_text"] for k in ("pml", "eml", "mfl")])
    sc["items"]["EML"]["pct"] = 40.0
    v = act._scenarios_view(sc, {"class_code": "9"}, "en")["json"]
    ok("все целые — без дроби: 15% · 40% · 100%", [v[k]["pct_text"] for k in ("pml", "eml", "mfl")] == ["15%", "40%", "100%"])
    fresh()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": WH_MUST, "optional": WH_OPT})
    labels = [r["label"] for r in a["sections"][0]["rows"]]
    ok("здание: «Год постройки»", "Год постройки" in labels and "Год выпуска" not in labels, labels)
    for lang, want in (("uz", "Qurilgan yili"), ("en", "Year built")):
        st, u = call("GET", f"/act/{a['id']}", params={"lang": lang})
        ok(f"{lang}: здание — «{want}»", want in [r["label"] for r in u["sections"][0]["rows"]])
    st, a = call("POST", "/act/make", {"lang": "en", "must": CRANE_MUST, "optional": CRANE_OPT})
    labels = [r["label"] for r in a["sections"][0]["rows"]]
    ok("техника: «Year of manufacture»", "Year of manufacture" in labels and "Year built" not in labels, labels)


def check_screen():
    print("35. Экран: плитки сценариев из ответа сервера, сброс пометок «из документа»")
    html = (Path(act.__file__).parent / "tg.html").read_text(encoding="utf-8")
    body = html[html.index("function actScenHtml"):html.index("function actMeasuresHtml")]
    ok("плитки берутся из scenarios.tiles (порядок и подписи сервера)", "s.tiles" in body and "x.label" in body
       and "x.pct_text" in body, body[:300])
    ok("удержание: текст сервера (сравнение с EML)", "r.text" in body)
    pre = html[html.index("function wzApplyPrefill"):html.index("function preTag")]
    ok("wzApplyPrefill: пометка остаётся только у значений из нового ответа",
       "delete CH.pre[k]" in pre and "CH.preDoc" in pre, pre[-300:])
    call_site = html[html.index("CH.preDoc = {};"):html.index("function wzApplyPrefill")]
    ok("новая загрузка без prefill тоже сбрасывает пометки", "wzApplyPrefill(d.prefill && typeof d.prefill" in call_site
       and ": {})" in call_site, call_site)


# ------------------------------------------------------------------ 36. оценка по объявлениям (снимки экрана)

def upload_to(path, files, fields=None):
    boundary = "----insonmk"
    parts = []
    for k, v in (fields or {}).items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
    for name, mime, blob in files:
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="files"; '
                     f'filename="{name}"\r\nContent-Type: {mime}\r\n\r\n'.encode() + blob + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    payload = b"".join(parts)
    hdrs = [(b"host", b"test"), (b"content-type", f"multipart/form-data; boundary={boundary}".encode()),
            (b"content-length", str(len(payload)).encode())] + _cookie_hdr()
    return _send("POST", path, None, hdrs, payload)


RATE = 12650.0                            # подменённый курс ЦБ в тесте (сеть не используется)
TODAY = date.today()


def _days_ago(n):
    return (TODAY - timedelta(days=n)).isoformat()


def _days_ahead(n):
    return (TODAY + timedelta(days=n)).isoformat()


# 7 объявлений автокрана XCMG QY50K5D: одно другой модели, один выброс, одно в долларах, у одного даты нет.
# Лишние поля продавца и телефон в названии — модель нарушила запрет, сервер должен их убрать.
MARKET_REPLY = _json.dumps({"listings": [
    {"file": 1, "title": "XCMG QY50K5D 2021, звоните +998 90 123-45-67", "price": 2_650_000_000, "currency": "UZS",
     "year": 2021, "hours": 4200, "region": "Ташкент", "posted": "Сегодня 10:15", "posted_date": None,
     "site": "olx", "relevant": True, "seller": "Иванов Иван Иванович", "phone": "+998 90 123-45-67"},
    {"file": 1, "title": "Автокран XCMG QY50K5D", "price": "2 500 000 000 сум", "currency": "UZS", "year": 2020,
     "region": "Каримов Алишер Бахтиёрович", "posted": _days_ago(18), "site": "olx", "relevant": True},
    {"file": 2, "title": "XCMG QY50K5D 50 тонн", "price": 205000, "currency": "USD", "year": 2022,
     "region": "Самарканд", "posted": "вчера", "site": "olx", "relevant": True},
    {"file": 2, "title": "XCMG QY50K5D", "price": 2_400_000_000, "currency": "UZS", "year": 2019,
     "region": "Навои", "posted_date": _days_ago(41), "site": "olx", "relevant": True},
    {"file": 2, "title": "XCMG QY25K5D 25 тонн", "price": 1_500_000_000, "currency": "UZS", "year": 2021,
     "posted": "вчера", "site": "olx", "relevant": False, "why_excluded": "другая модель: QY25K5D"},
    {"file": 3, "title": "XCMG QY50K5D срочно", "price": 900_000_000, "currency": "UZS", "year": 2021,
     "posted": _days_ago(5), "site": "olx", "relevant": True},
    {"file": 3, "title": "XCMG QY50K5D 2023", "price": 2_800_000_000, "currency": "UZS", "year": 2023,
     "posted": "3 дня назад", "site": "olx", "relevant": True},
    # дубль с перекрывающегося снимка — должен схлопнуться
    {"file": 3, "title": "XCMG QY50K5D 2023", "price": 2_800_000_000, "currency": "UZS", "year": 2023,
     "posted": None, "site": "olx", "relevant": True},
    # L8: даты публикации на снимке нет — в расчёт не берётся; телефон в названии с хвостом «тел.»
    {"file": 3, "title": "XCMG QY50K5D, тел. +998 93 555 66 77", "price": 2_550_000_000, "currency": "UZS",
     "year": 2020, "posted": None, "site": "olx", "relevant": True},
    # L9: «7 месяцев назад» — старше 6 месяцев
    {"file": 3, "title": "XCMG QY50K5D 2018", "price": 2_000_000_000, "currency": "UZS", "year": 2018,
     "posted": "7 месяцев назад", "site": "olx", "relevant": True},
    # L10: дата позже даты снимков — некорректна
    {"file": 3, "title": "XCMG QY50K5D 2022 новый", "price": 2_700_000_000, "currency": "UZS", "year": 2022,
     "posted_date": _days_ahead(10), "site": "olx", "relevant": True},
]}, ensure_ascii=False)

SHOT_FILES = [("shot1.png", "image/png", image((240, 240, 240), w=390, h=844)),
              ("shot2.jpg", "image/jpeg", image((230, 230, 230), kind="jpg", w=390, h=844)),
              ("shot3.png", "image/png", image((250, 250, 250), w=390, h=844))]
PD_BITS = ("Иванов", "123-45-67", "123 45 67", "Бахтиёрович", "Каримов", "555 66 77")


def hand_q(xs, p):
    """Перцентиль «руками»: позиция p × (n − 1), линейно между соседями."""
    xs = sorted(xs)
    pos = p * (len(xs) - 1)
    lo = int(pos)
    return xs[lo] + (xs[min(lo + 1, len(xs) - 1)] - xs[lo]) * (pos - lo)


class NetSpy:
    """Подмена сетевого слоя: любое обращение записывается и отклоняется (сервер не должен ходить в сеть)."""

    def __init__(self):
        self.hits = []

    def __enter__(self):
        import socket
        import urllib.request
        from app import valuation_sources as vs
        self.saved = (socket.socket.connect, socket.create_connection, urllib.request.urlopen, vs._http_get,
                      vs.robots_check, llm._post)

        def rec(kind):
            def f(*a, **kw):
                self.hits.append((kind, repr(a[:2])[:200]))
                raise AssertionError("сеть запрещена в тесте: " + kind)
            return f
        orig_connect = self.saved[0]

        def connect(s, addr, *a):
            # asyncio на Windows соединяет внутренний socketpair через 127.0.0.1 — это не сеть
            if isinstance(addr, tuple) and addr[0] in ("127.0.0.1", "::1"):
                return orig_connect(s, addr, *a)
            return rec("socket")(addr)
        socket.socket.connect = connect
        socket.create_connection = rec("create_connection")
        urllib.request.urlopen = rec("urlopen")
        vs._http_get = rec("valuation_sources._http_get")
        vs.robots_check = rec("valuation_sources.robots_check")
        llm._post = rec("llm._post")
        return self

    def __exit__(self, *exc):
        import socket
        import urllib.request
        from app import valuation_sources as vs
        (socket.socket.connect, socket.create_connection, urllib.request.urlopen, vs._http_get,
         vs.robots_check, llm._post) = self.saved


def check_market_engine():
    print("36а. Оценка по объявлениям: чистая функция")
    sd = date(2026, 9, 30)
    P = ae.parse_posted
    ok("даты публикации: ISO, ДД.ММ.ГГГГ, «сегодня», «вчера», «3 дня назад», «12 сентября», bugun/kecha",
       P("2026-09-12", sd) == date(2026, 9, 12) and P("12.09.2026", sd) == date(2026, 9, 12)
       and P("Сегодня 10:15", sd) == sd and P("вчера", sd) == date(2026, 9, 29)
       and P("3 дня назад", sd) == date(2026, 9, 27) and P("12 сентября", sd) == date(2026, 9, 12)
       and P("15 октября", sd) == date(2025, 10, 15) and P("bugun", sd) == sd and P("kecha", sd) == date(2026, 9, 29)
       and P("2 мая 2026 г.", sd) == date(2026, 5, 2) and P("5 March 2026", sd) == date(2026, 3, 5))
    ok("дата позже снимка и мусор — не дата", P("2026-10-05", sd) is None and P("срочно", sd) is None
       and P(None, sd) is None)
    PX = ae.parse_posted_ex
    ok("«N … назад» по-русски: минуты, часы, дни, недели, месяцы, годы",
       P("10 минут назад", sd) == sd and P("5 часов назад", sd) == sd and P("30 часов назад", sd) == date(2026, 9, 29)
       and P("неделю назад", sd) == date(2026, 9, 23) and P("2 недели назад", sd) == date(2026, 9, 16)
       and P("7 месяцев назад", sd) == date(2026, 2, 28) and P("месяц назад", sd) == date(2026, 8, 30)
       and P("2 года назад", sd) == date(2024, 9, 30) and P("5 лет назад", sd) == date(2021, 9, 30))
    ok("«N … oldin» по-узбекски и «N … ago» по-английски",
       P("15 daqiqa oldin", sd) == sd and P("3 soat oldin", sd) == sd and P("3 kun oldin", sd) == date(2026, 9, 27)
       and P("2 hafta oldin", sd) == date(2026, 9, 16) and P("2 oy oldin", sd) == date(2026, 7, 30)
       and P("1 yil oldin", sd) == date(2025, 9, 30) and P("20 minutes ago", sd) == sd
       and P("an hour ago", sd) == sd and P("3 days ago", sd) == date(2026, 9, 27)
       and P("2 weeks ago", sd) == date(2026, 9, 16) and P("3 months ago", sd) == date(2026, 6, 30)
       and P("a year ago", sd) == date(2025, 9, 30))
    ok("дата позже снимка — «некорректна» (future), не «не видна»",
       PX("2026-10-05", sd) == (None, "future") and PX("срочно", sd) == (None, "none") and PX(None, sd) == (None, "none"))
    L = [{"id": "a", "title": "x", "price": 100, "currency": "UZS", "posted": "2026-09-01", "relevant": True},
         {"id": "b", "title": "x", "price": 110, "currency": "UZS", "posted": "2026-03-01", "relevant": True},
         {"id": "c", "title": "x", "price": 10, "currency": "USD", "posted": "вчера", "relevant": True},
         {"id": "d", "title": "x", "price": None, "currency": "UZS", "posted": "2026-09-01", "relevant": True}]
    r = ae.market_estimate(L, declared=100, shot_date=sd)
    codes = {e["id"]: e["code"] for e in r["excluded"]}
    ok("старше 6 месяцев, без цены, в долларах без курса — исключены с причиной",
       codes == {"b": "mx_too_old", "c": "mx_no_rate", "d": "mx_no_price"}, codes)
    ok("одно объявление — оценка ориентировочная (few)", r["verdict"] == "few" and r["median"] == 100, r["verdict"])
    r = ae.market_estimate(L, declared=100, shot_date=sd, usd_rate=11)
    ok("с курсом доллар пересчитан и вошёл", r["used"] == 2 and r["median"] == 105 and not r["date_assumed"],
       (r["used"], r.get("date_assumed")))
    # правило проекта (valuation_sources): объявление без даты публикации в расчёт не берётся
    U = [dict(L[0]), {"id": "u", "title": "x", "price": 104, "currency": "UZS", "posted": None, "relevant": True},
         {"id": "f", "title": "x", "price": 102, "currency": "UZS", "posted_date": "2026-10-03", "relevant": True},
         {"id": "o", "title": "x", "price": 101, "currency": "UZS", "posted": "7 месяцев назад", "relevant": True}]
    r = ae.market_estimate(U, declared=100, shot_date=sd)
    codes = {e["id"]: e["code"] for e in r["excluded"]}
    ok("без даты — «дата публикации не видна»; позже снимка — «некорректна»; «7 месяцев назад» — старше 6 мес.",
       codes == {"u": "mx_no_date", "f": "mx_bad_date", "o": "mx_too_old"} and r["used"] == 1, codes)
    ok("причины словами", tx.t("mx_no_date", "ru") == "дата публикации не видна"
       and "некорректна" in tx.t("mx_bad_date", "ru", shot="30.09.2026"))
    r = ae.market_estimate(U, declared=100, shot_date=sd, settings={"market": {"allow_undated": True}})
    ok("настройка allow_undated = true: без даты — дата снимка с пометкой",
       r["used"] == 2 and r["date_assumed"] == ["u"] and r["how"][0]["code"] == "mh_filter_undated", r["date_assumed"])
    ok("allow_undated по умолчанию false и проверяется как true/false",
       ae.DEFAULT_SETTINGS["market"]["allow_undated"] is False and ae.check_settings({"market": {"allow_undated": 1}})
       and not ae.check_settings({"market": {"allow_undated": True}}))
    # выбросы при числе подходящих меньше min_listings не ищутся
    two = [{"id": "p", "price": 100, "posted": "вчера"}, {"id": "q", "price": 1000, "posted": "вчера"}]
    r = ae.market_estimate(two, shot_date=sd)
    ok("2 объявления < 3: выбросы не ищутся (1000 при медиане 550 не выброс)",
       r["used"] == 2 and not r["excluded"] and any(h["code"] == "mh_outliers_skipped" for h in r["how"]), r["how"])
    # пустые и нулевые цены не роняют расчёт
    bad = [{"id": "z0", "price": 0}, {"id": "z1", "price": ""}, {"id": "z2", "price": "abc"},
           {"id": "z3", "price": -5}, {"id": "z4"}, "мусор"]
    r = ae.market_estimate(bad, declared=10, shot_date=sd, usd_rate="x")
    ok("пустые, нулевые и мусорные цены — «цена не видна», без падения",
       r["verdict"] == "none" and {e["code"] for e in r["excluded"]} == {"mx_no_price"} and len(r["excluded"]) == 5, r)
    # округление медианы — «половина вверх» (как Math.round на экране), а не банковское round()
    r = ae.market_estimate([{"id": "h1", "price": 2, "posted": "вчера"}, {"id": "h2", "price": 3, "posted": "вчера"}],
                           shot_date=sd)
    ok("медиана 2,5 → 3 (половина вверх; round() дал бы 2)", r["median"] == 3 and ae.round_half_up(2_593_249_999.5)
       == 2_593_250_000 and ae.round_half_up(0.5) == 1, r["median"])
    r = ae.market_estimate([dict(L[0], relevant=False)], declared=100, shot_date=sd)
    ok("ни одного подходящего — none, оценки нет", r["verdict"] == "none" and not r["available"], r["verdict"])
    st = ae.merge_settings({"market": {"min_listings": 1, "diff_pct": 5}})
    r = ae.market_estimate(L[:1], declared=104, shot_date=sd, settings=st)
    ok("порог расхождения — настройка (4 % < 5 % — confirmed)", r["verdict"] == "confirmed", r)
    r = ae.market_estimate(L[:1], declared=120, sum_insured=120, shot_date=sd, settings=st)
    ok("расхождение 16,7 % > 5 % — refine; страховая сумма сверяется с уточнённой (ГК 938)",
       r["verdict"] == "refine" and r["refined_value"] == 100 and r["insured_check"]["verdict"] == "over"
       and r["insured_check"]["legal_ref"] == "ГК РУз, ст. 938", r)
    ok("настройки оценки проверяются", ae.check_settings({"market": {"min_listings": 0}})
       and ae.check_settings({"market": {"outlier_low": 1.5}}) and ae.check_settings({"market": {"zzz": 1}})
       and not ae.check_settings({"market": {"min_listings": 4, "diff_pct": 20}}))
    ok("по умолчанию: 3 объявления, 15 %, 6 месяцев, выбросы 0,5 и 2, без даты — нельзя",
       ae.DEFAULT_SETTINGS["market"] == {"min_listings": 3, "diff_pct": 15, "max_age_months": 6, "outlier_low": 0.5,
                                         "outlier_high": 2.0, "allow_undated": False})
    # чистка названия: после вырезанного телефона не остаётся «, звоните» / «тел.»
    ok("хвосты «звоните», «тел.», «звонить» убираются вместе с телефоном",
       am.clean_text("XCMG QY50K5D 2021, звоните +998 90 123-45-67", 160)[0] == "XCMG QY50K5D 2021"
       and am.clean_text("XCMG, тел. +998 91 555 44 33", 160)[0] == "XCMG"
       and am.clean_text("Автокран +998 90 111 22 33 звонить", 160)[0] == "Автокран"
       and am.clean_text("Кран Мотель", 160)[0] == "Кран Мотель")
    ok("ссылка объявления: только https на четырёх площадках",
       am.safe_url("https://www.olx.uz/d/obyavlenie/x") and am.safe_url("https://m.avtoelon.uz/a/1")
       and not am.safe_url("http://169.254.169.254/latest/meta-data") and not am.safe_url("https://olx.uz.evil.com/")
       and not am.safe_url("https://u:p@olx.uz/") and not am.safe_url("https://example.uz/ad/1")
       and not am.safe_url("https://olx.uz:8443/"))
    # сбой robots.txt помнится 10 минут, а не до перезапуска (app/valuation_sources.py)
    from app import valuation_sources as vs
    hits = []

    def refuse(url, timeout=0):
        hits.append(url)
        raise OSError("нет сети")
    saved = vs._http_get
    vs._http_get = refuse
    try:
        vs._robots_cache.pop("robots-test.example", None)
        a1 = vs.robots_check("https://robots-test.example/x")
        a2 = vs.robots_check("https://robots-test.example/y")
        rp, why, at = vs._robots_cache["robots-test.example"]
        vs._robots_cache["robots-test.example"] = (rp, why, at - vs.ROBOTS_FAIL_TTL_SEC - 1)
        vs.robots_check("https://robots-test.example/z")
        ok("robots.txt: сбой запомнен на 10 минут, потом повтор", not a1[0] and not a2[0] and len(hits) == 2
           and vs.ROBOTS_FAIL_TTL_SEC == 600, hits)
    finally:
        vs._http_get = saved
        vs._robots_cache.pop("robots-test.example", None)


def check_market_links():
    print("36б. Ссылки поиска: сервер только составляет адреса")
    st, b = call("GET", "/act/market/links", params={"brand": "XCMG", "model": "QY50K5D", "year": "2021",
                                                     "object_kind": "truck_crane", "lang": "ru"})
    urls = {(ln["site"], ln["url"]) for ln in b.get("links") or []}
    ok("OLX: общий поиск по марке и модели (раздела спецтехники нет в valuation_sources)",
       ("olx", "https://www.olx.uz/list/q-xcmg-qy50k5d/") in urls, urls)
    ok("OLX: вторая ссылка с годом", ("olx", "https://www.olx.uz/list/q-xcmg-qy50k5d-2021/") in urls)
    ok("avtoelon.uz: раздел автокранов по марке (valuation.SPEC_SECTIONS)",
       ("avtoelon", "https://avtoelon.uz/spectehnika/gruzovaja-tehnika/avtokran/xcmg/") in urls, urls)
    ok("у каждой ссылки подпись и подсказка", all(ln["label"] and ln["hint"] for ln in b["links"]))
    ok("подсказка: что снять — 5–10 объявлений, до 5 снимков", "5–10" in b["hint"] and b["max_shots"] == 5
       and len(b["shot_tips"]) == 4 and any("телефон" in x for x in b["shot_tips"]), b.get("hint"))
    ok("сервер по ссылкам не ходит — сказано явно", "не ходит" in b["note"])
    st, b = call("GET", "/act/market/links", params={"object_kind": "warehouse", "lang": "ru"})
    urls = [(ln["site"], ln["url"]) for ln in b["links"]]
    ok("недвижимость: OLX в разделе «Недвижимость», запрос по-русски в кодировке, uybor и joymee",
       urls[0] == ("olx", "https://www.olx.uz/nedvizhimost/q-%D1%81%D0%BA%D0%BB%D0%B0%D0%B4/")
       and {"uybor", "joymee"} <= {s for s, _ in urls} and "avtoelon" not in {s for s, _ in urls}, urls)
    st, b = call("GET", "/act/market/links", params={"brand": "Chevrolet", "model": "Cobalt", "object_kind": "car"})
    urls = {(ln["site"], ln["url"]) for ln in b["links"]}
    ok("легковой: OLX в разделе легковых, avtoelon /avto/марка/модель/",
       ("olx", "https://www.olx.uz/transport/legkovye-avtomobili/q-chevrolet-cobalt/") in urls
       and ("avtoelon", "https://avtoelon.uz/avto/chevrolet/cobalt/") in urls, urls)
    st, b = call("GET", "/act/market/links", params={"brand": "Шакман", "model": "SX3258", "object_kind": "truck"})
    urls = {(ln["site"], ln["url"]) for ln in b["links"]}
    ok("кириллица: OLX — кодирование, avtoelon — транслитерация",
       ("olx", "https://www.olx.uz/list/q-%D1%88%D0%B0%D0%BA%D0%BC%D0%B0%D0%BD-sx3258/") in urls
       and ("avtoelon", "https://avtoelon.uz/spectehnika/") in urls, urls)
    st, b = call("GET", "/act/market/links", params={"lang": "ru"})
    ok("искать не по чему — 422 с понятной причиной", st == 422 and "марку" in b["detail"], (st, b))
    st, b = call("GET", "/act/market/links", params={"year": "1800", "brand": "XCMG"})
    ok("год проверяется", st == 422 and "year" in b["errors"], (st, b))
    for lang in ("uz", "en"):
        st, b = call("GET", "/act/market/links", params={"brand": "XCMG", "model": "QY50K5D", "year": "2021",
                                                         "object_kind": "truck_crane", "lang": lang})
        txt = " ".join([b["hint"], b["note"]] + b["shot_tips"] + [ln["label"] + ln["hint"] for ln in b["links"]])
        ok(f"{lang}: подписи и подсказки без кириллицы", st == 200 and not re.search(r"[А-Яа-яЁё]", txt), txt[:300])


def market_setup(rate=RATE):
    from app import valuation_sources as vs
    act._FX_CACHE.clear()
    ORIG.setdefault("cbu", vs.cbu_usd_rate)
    vs.cbu_usd_rate = (lambda d: rate)


def check_market_shots():
    print("36в. Снимки объявлений: одно чтение моделью, валюта, ПД, лимиты")
    fresh()
    clear_settings()
    market_setup()
    model_on(True)
    REPLY["text"] = MARKET_REPLY
    CALLS.clear()
    with NetSpy() as spy:
        st, b = upload_to("/act/market/shots", SHOT_FILES, {"lang": "ru", "site": "olx", "brand": "XCMG",
                                                            "model": "QY50K5D", "object_kind": "truck_crane"})
        st_l, _ = call("GET", "/act/market/links", params={"brand": "XCMG", "model": "QY50K5D"})
    ok("снимки приняты", st == 200 and b.get("ok") and b.get("shots_session"), (st, b))
    ok("ни одного сетевого запроса (к olx.uz — тем более)", not spy.hits and st_l == 200, spy.hits)
    ok("один запрос к модели со всеми снимками, 20 с, без повторов",
       len(CALLS) == 1 and len(CALLS[0]["files"]) == 3 and CALLS[0]["timeout"] == 20 and CALLS[0]["retries"] == 0,
       [(c["purpose"], c["timeout"], c["retries"]) for c in CALLS])
    prompt = CALLS[0]["messages"][0]["content"] + CALLS[0]["messages"][1]["content"]
    ok("в инструкции: не выдумывай, null, запрет на имена и телефоны, строгая схема",
       "не выдумывай" in prompt and "null" in prompt and "телефоны" in prompt and '"listings"' in prompt)
    ok("в инструкции дата снимка и что ищем", TODAY.isoformat() in prompt and "qy50k5d" in prompt.lower())
    ok("маскировка ПД не портит инструкцию", llm.mask_pd(prompt) == prompt)
    L = {r["id"]: r for r in b["listings"]}
    ok("10 объявлений (дубль с перекрывающегося снимка схлопнут)", len(L) == 10, list(L))
    ok("доллары пересчитаны по курсу ЦБ", L["L3"]["price_uzs"] == round(205000 * RATE) and L["L3"]["currency"] == "USD",
       L["L3"])
    ok("курс показан явно: ЦБ РУз, с датой", b["fx"] and b["fx"]["by"] == "cbu" and b["fx"]["rate"] == RATE
       and "cbu.uz" in b["fx"]["text"], b.get("fx"))
    ok("цена «2 500 000 000 сум» строкой — число", L["L2"]["price"] == 2_500_000_000)
    ok("другая модель — relevant = false с причиной", L["L5"]["relevant"] is False and "QY25K5D" in
       (L["L5"]["why_excluded"] or ""))
    ok("«сегодня» и «вчера» — даты от даты снимка", L["L1"]["posted_date"] == TODAY.isoformat()
       and L["L3"]["posted_date"] == _days_ago(1))
    ok("«3 дня назад» — дата от даты снимка", L["L7"]["posted_date"] == _days_ago(3), L["L7"])
    ok("дата не видна — «в расчёт не берётся»", L["L8"]["date_assumed"] and not L["L8"]["used"]
       and L["L8"]["date_note"] == "дата публикации не видна — в расчёт не берётся", L["L8"])
    ok("дата позже снимка — «некорректна»", L["L10"]["date_status"] == "future" and "некорректна" in L["L10"]["date_note"],
       L["L10"])
    ok("«7 месяцев назад» — дата посчитана", L["L9"]["posted_date"] == ae.months_before(TODAY, 7).isoformat(), L["L9"])
    dump = _json.dumps(b, ensure_ascii=False)
    ok("имена и телефоны продавцов не возвращаются", not any(x in dump for x in PD_BITS),
       [x for x in PD_BITS if x in dump])
    ok("название осталось, телефон и «звоните» вырезаны", L["L1"]["title"] == "XCMG QY50K5D 2021", L["L1"])
    ok("«тел.» без номера тоже убрано", L["L8"]["title"] == "XCMG QY50K5D", L["L8"])
    ok("предупреждение правдивое: не извлекаются, хранятся 24 часа",
       "не извлекаются" in b["warning"] and "24 часа" in b["warning"] and "не сохраняются" not in b["warning"],
       b["warning"])
    ok("регион с ФИО — отброшен", L["L2"]["region"] is None)
    ok("сказано, что данные продавцов убраны", any("убрано" in n for n in b["notes"]), b["notes"])
    # предварительная оценка — руками: L5 другой модели, L6 (900 млн) — выброс
    cand = [2_650_000_000, 2_500_000_000, round(205000 * RATE), 2_400_000_000, 900_000_000, 2_800_000_000]
    med0 = (sorted(cand)[2] + sorted(cand)[3]) / 2
    used = [x for x in cand if 0.5 * med0 <= x <= 2 * med0]
    ok("выброс руками: 900 млн < 0,5 × медианы", 900_000_000 < 0.5 * med0 and len(used) == 5, (med0, used))
    e = b["estimate"]
    ok("медиана и вилка пересчитаны руками", e["median"] == round(hand_q(used, 0.5)) == 2_593_250_000
       and e["low"] == round(hand_q(used, 0.25)) == 2_500_000_000
       and e["high"] == round(hand_q(used, 0.75)) == 2_650_000_000, e)
    ok("исключены: L5 (другая модель), L6 (выброс), L8 (без даты), L9 (старше 6 мес.), L10 (дата позже снимка)",
       sorted((x["id"], x["code"]) for x in e["excluded"]) == [
           ("L10", "mx_bad_date"), ("L5", "mx_not_relevant"), ("L6", "mx_outlier_low"), ("L8", "mx_no_date"),
           ("L9", "mx_too_old")], e["excluded"])
    ok("без заявленной стоимости — verdict ready", e["verdict"] == "ready" and e["calibrated"] == 0)
    ok("ссылки поиска в ответе", any(ln["site"] == "olx" for ln in b["links"]))
    with db.tx() as con:
        row = db.rows(con, "SELECT result_json, expires_at, created_at FROM act_uploads WHERE id=?", b["shots_session"])
        j = db.rows(con, "SELECT detail FROM audit WHERE entity=?", "act_upload:" + b["shots_session"])
    stored = row[0]["result_json"]
    ok("в базе нет данных продавцов", not any(x in stored for x in PD_BITS), stored[:300])
    ok("снимки хранятся 24 часа", (datetime.fromisoformat(row[0]["expires_at"]) -
                                   datetime.fromisoformat(row[0]["created_at"])) == timedelta(hours=24))
    detail = j[0]["detail"] if j else ""
    ok("в журнале только счётчики", j and "XCMG" not in detail and "2650000000" not in detail
       and '"listings": 10' in detail, detail)
    sid = b["shots_session"]

    # курс ЦБ: свой пул (не пул модели), общий срок 5 с, неудача помнится 10 минут
    from app import valuation_sources as vs
    ok("курс — в своём пуле, срок 5 с, неудача — 10 минут", act._FX_POOL is not act._AI_POOL
       and act.FX_DEADLINE_SEC == 5 and act.FX_FAIL_TTL_SEC == 600)
    calls, names = [], []

    def cbu_down(d):
        import threading as _thr
        calls.append(d)
        names.append(_thr.current_thread().name)
        return None
    vs.cbu_usd_rate = cbu_down
    act._FX_CACHE.clear()
    r1 = act._fx_submit(TODAY).result(5)
    r2 = act._fx_submit(TODAY).result(5)
    ok("неудача курса закэширована: второй раз cbu.uz не спрашивается", r1["rate"] is None and r2["rate"] is None
       and len(calls) == 1, calls)
    res, until = act._FX_CACHE[TODAY.isoformat()]
    act._FX_CACHE[TODAY.isoformat()] = (res, until - act.FX_FAIL_TTL_SEC - 1)
    act._fx_submit(TODAY).result(5)
    ok("через 10 минут курс спрашивается снова", len(calls) == 2, calls)
    ok("курс запрашивается в потоке act-fx, а не в потоке модели", names and all(n.startswith("act-fx") for n in names),
       names)
    import threading as _th
    import time as _time
    slow_gate = _th.Event()

    def cbu_slow(d):
        slow_gate.wait(8)
        return RATE
    vs.cbu_usd_rate = cbu_slow
    act._FX_CACHE.clear()
    fresh()
    t0 = _time.monotonic()
    st, b2 = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru", "brand": "XCMG", "model": "QY50K5D"})
    spent = _time.monotonic() - t0
    slow_gate.set()
    act._FX_INFLIGHT[TODAY.isoformat()].result(10)        # запоздавший ответ ЦБ ложится в кэш — дождёмся его
    ok("медленный cbu.uz: загрузка не ждёт дольше 5 с, курс не выдуман", st == 200 and b2["fx"] is None
       and spent < 7.5 and b2["usd_rate_needed"], (st, spent, b2.get("fx")))
    market_setup()

    # курс не получен: доллары не считаются, нужен курс сотрудника; с ним — «введён сотрудником»
    market_setup(rate=None)
    fresh()
    st, b = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru", "brand": "XCMG", "model": "QY50K5D"})
    L = {r["id"]: r for r in b["listings"]}
    ok("без курса: цена в долларах не пересчитана, просьба указать курс", L["L3"]["price_uzs"] is None
       and b["usd_rate_needed"] and any("курс" in n for n in b["notes"]) and b["fx"] is None, b.get("notes"))
    ok("без курса: объявление в долларах — причина «курс не задан»",
       ("L3", "mx_no_rate") in [(x["id"], x["code"]) for x in b["estimate"]["excluded"]])
    st, b = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru", "usd_rate": "12 600"})
    ok("курс сотрудника принят и показан", b["fx"] and b["fx"]["by"] == "employee" and b["fx"]["rate"] == 12600
       and {r["id"]: r for r in b["listings"]}["L3"]["price_uzs"] == 205000 * 12600, b.get("fx"))
    st, b = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru", "usd_rate": "12"})
    ok("курс с опечаткой — 422", st == 422 and "usd_rate" in b["errors"], (st, b))
    market_setup()

    # пределы: больше 5 снимков, не картинка, гостевой лимит по файлам, модель выключена
    st, b = upload_to("/act/market/shots", SHOT_FILES * 2, {"lang": "ru"})
    ok("больше 5 снимков — 413", st == 413 and "5" in b["detail"], (st, b))
    st, b = upload_to("/act/market/shots", [("x.pdf", "application/pdf", pdf_pages(1))], {"lang": "ru"})
    ok("PDF как снимок не принимается", st == 422 and "JPG" in b["rejected"][0]["error"], (st, b))
    st, b = upload_to("/act/market/shots", [("big.png", "image/png", png_bomb(9000, 9000))], {"lang": "ru"})
    ok("размер в пикселях проверяется до раскрытия", st == 422 and "Мп" in b["rejected"][0]["error"], (st, b))
    st, b = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru", "site": "avito"})
    ok("площадка — из списка", st == 422 and "site" in b["errors"])
    fresh()
    set_limits(guest_photos_per_hour=4)
    try:
        st1, _ = upload_to("/act/market/shots", SHOT_FILES, {"lang": "ru"})
        st2, b2 = upload_to("/act/market/shots", SHOT_FILES[:2], {"lang": "ru"})
        ok("гостевой лимит считается по файлам (3 + 2 > 4)", st1 == 200 and st2 == 429, (st1, st2))
    finally:
        clear_settings()
        fresh()
    set_limits(ai_calls_per_hour=1)
    try:
        CALLS.clear()
        call_ok, _ = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru"})
        st, b = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru"})
        ok("общий предел обращений к модели — общий с распознаванием фото",
           st == 200 and not b["ai"] and "лимит распознаваний" in b["message"] and len(CALLS) == 1, b.get("message"))
    finally:
        clear_settings()
        fresh()
    model_on(False)
    st, b = upload_to("/act/market/shots", SHOT_FILES[:1], {"lang": "ru"})
    ok("модель выключена: снимки сохранены, честная причина, можно ввести вручную",
       st == 200 and not b["ai"] and "вручную" in b["message"] and b["listings"] == [], b.get("message"))
    model_on(True)
    return sid


def mk_make(listings, sid=None, must=None, lang="ru", usd_rate=None, optional=None, fx=None):
    opt = dict(optional or CRANE_OPT, object_kind="truck_crane")
    opt["market"] = {"listings": listings, "shots_session": sid}
    if usd_rate is not None:
        opt["market"]["usd_rate"] = usd_rate
    if fx is not None:
        opt["market"]["fx"] = fx
    return call("POST", "/act/make", {"lang": lang, "must": must or CRANE_MUST, "optional": opt,
                                       "recognized": [{"key": "brand", "value": "XCMG", "source": "input"},
                                                      {"key": "model", "value": "QY50K5D", "source": "input"}]})


def sec3(a):
    return a["sections"][2]


def flat(s):
    """Текст без неразрывных пробелов: так сравнивать суммы в строках акта удобнее."""
    return re.sub(r"\s+", " ", str(s or "").replace(" ", " "))


def docx_plain(aid, lang="ru"):
    st, blob, h = call("GET", f"/act/{aid}.docx", params={"lang": lang}, raw=True)
    return flat(re.sub(r"<[^>]+>", "", zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8")))


def pdf_plain(aid, lang="ru"):
    st, blob, h = call("GET", f"/act/{aid}.pdf", params={"lang": lang}, raw=True)
    return flat(pdf_text(pymupdf.open(stream=blob, filetype="pdf")))


def screen_rows(stored):
    """Объявления так, как их отдаёт экран в /act/make (mkBody): с датой, площадкой и происхождением."""
    out = []
    for r in stored["listings"]:
        o = {k: r.get(k) for k in ("id", "title", "price", "currency", "year", "relevant", "region")}
        if r.get("posted_date"):
            o["posted_date"] = r["posted_date"]
        o.update(date_assumed=not r.get("posted_date"), site=r.get("site") or "olx", source="shot")
        out.append(o)
    return out


def check_market_make(sid):
    print("36г. Акт: блок оценки по объявлениям, решение, правки сотрудника, языки, Word и PDF")
    fresh()
    market_setup()
    CALLS.clear()
    with db.tx() as con:
        stored = _json.loads(db.rows(con, "SELECT result_json FROM act_uploads WHERE id=?", sid)[0]["result_json"])
    listings = screen_rows(stored)
    with NetSpy() as spy:
        st, a = mk_make(listings, sid)
    ok("акт с оценкой сформирован без сети и без модели", st == 200 and not spy.hits and not CALLS, (st, spy.hits))
    mv = a["market_value"]
    used = [2_650_000_000, 2_500_000_000, round(205000 * RATE), 2_400_000_000, 2_800_000_000]
    med = hand_q(used, 0.5)
    diff = (3_100_000_000 - med) / 3_100_000_000 * 100
    ok("медиана, вилка, число объявлений — руками", mv["median"] == round(med) and mv["low"] == 2_500_000_000
       and mv["high"] == 2_650_000_000 and mv["count"] == 10 and mv["used"] == 5, mv)
    ok("расхождение 16,3 % > 15 % — refine, уточнённая стоимость = медиана",
       mv["verdict"] == "refine" and mv["diff_pct"] == round(diff, 1) == 16.3 and mv["refined_value"] == round(med),
       (mv["verdict"], mv["diff_pct"]))
    ic = mv["insured_check"]
    ok("страховая сумма сверяется с уточнённой: превышение, ГК ст. 938",
       ic["verdict"] == "over" and ic["diff"] == 2_945_000_000 - round(med) and "938" in ic["legal_ref"]
       and ic["ratio_pct"] == round(2_945_000_000 / round(med) * 100, 2) == 113.56, ic)
    ok("источник: OLX, дата снимков, загружены сотрудником",
       mv["source_label"] == f"Источник: OLX, объявления на {TODAY.strftime('%d.%m.%Y')}, снимки загружены "
                             f"сотрудником.", mv["source_label"])
    ok("ссылки поиска и «как посчитано», calibrated = 0", mv["links"] and len(mv["how"]) >= 5
       and mv["calibrated"] == 0 and any("перцентил" in h for h in mv["how"]), mv["how"])
    ok("отброшенные с причинами", sorted(e["id"] for e in mv["excluded"]) == ["L10", "L5", "L6", "L8", "L9"]
       and all(e["reason"] for e in mv["excluded"]))
    ok("без правок: «Правки сотрудника: правок нет», медиана без правок та же, проверки правок нет",
       "Правки сотрудника: правок нет." in mv["source_lines"] and mv["edits"]["items"] == []
       and mv["median_original"] == mv["median"]
       and not any(c.startswith("Проверить правки") for c in a["decision"]["checks"]), mv["source_lines"])
    s3 = sec3(a)
    labels = [r["label"] for r in s3["rows"]]
    ok("раздел 3: «Оценка по объявлениям», медиана, вилка, вывод, уточнённая стоимость",
       {"Оценка по объявлениям", "Медиана цен объявлений", "Вилка (25–75-й перцентиль)", "Вывод по объявлениям",
        "Уточнённая стоимость (по объявлениям)", "Страховая сумма к уточнённой стоимости"} <= set(labels), labels)
    ok("раздел 3: строка источника и ссылка поиска под оценкой",
       s3["source_lines"][0] == mv["source_label"] and any(x.startswith("Ссылка поиска: https://www.olx.uz/")
                                                           for x in s3["source_lines"]), s3["source_lines"])
    ok("раздел 3: прежние строки на месте", labels[:4] == ["Страховая сумма", "Стоимость объекта",
                                                           "Отношение суммы к стоимости", "Вывод"], labels)
    # один итоговый вывод вместо «в норме» рядом с «превышением» (п. 7 замечаний)
    want = ("К заявленной стоимости страховая сумма составляет 95 % — в норме. Но по объявлениям стоимость ниже "
            "заявленной на 16,3 %: к уточнённой стоимости страховая сумма составляет 113,56 % — превышение (ГК РУз, "
            "ст. 938). Итог: стоимость нужно уточнить.")
    ok("вывод раздела 3 — один: к заявленной, к уточнённой, итог", flat(s3["rows"][3]["value"]) == want,
       flat(s3["rows"][3]["value"]))
    ok("value.final_verdict = refine, text — тот же итог; поле verdict прежнее",
       a["value"]["final_verdict"] == "refine" and flat(a["value"]["text"]) == want and a["value"]["verdict"] == "normal"
       and a["value"]["refined_ratio_pct"] == 113.56, a["value"])
    ok("в разделе 3 нет второго, противоположного вывода «В норме: …»",
       not any(flat(r["value"]).startswith("В норме") for r in s3["rows"]), [r["value"] for r in s3["rows"]])
    chk = a["decision"]["checks"]
    ok("решение: «уточнить стоимость объекта: по объявлениям …» и снизить сумму",
       any(c.startswith("Уточнить стоимость объекта: по объявлениям (5 шт.)") for c in chk)
       and any("выше уточнённой стоимости" in c for c in chk) and a["decision"]["code"] != "accept", chk)
    aid = a["id"]

    # Word и PDF
    plain = docx_plain(aid)
    ok("DOCX: блок оценки, строка источника, правок нет", "Оценка по объявлениям" in plain
       and "Медиана цен объявлений" in plain and "снимки загружены сотрудником" in plain
       and "Уточнить стоимость объекта" in plain and "Правки сотрудника: правок нет" in plain and want in plain)
    text = pdf_plain(aid)
    ok("PDF: блок оценки, строка источника, правок нет", "Оценка по объявлениям" in text
       and "Медиана цен объявлений" in text and "снимки загружены сотрудником" in text
       and "Правки сотрудника: правок нет" in text, text[:200])

    # три языка
    for lang, word in (("uz", "Eʼlonlar boʻyicha baholash"), ("en", "Valuation by listings")):
        st, b = call("GET", f"/act/{aid}", params={"lang": lang})
        s = sec3(b)
        txt = [r["label"] for r in s["rows"]] + [str(r["value"]) for r in s["rows"][3:]] + s["source_lines"] + \
            [li["title"] for li in s["lists"]] + s["lists"][0]["items"] + b["decision"]["checks"] + \
            [b["market_value"]["verdict_text"], b["market_value"]["source_label"], b["value"]["text"]]
        cyr = [x for x in txt if re.search(r"[А-Яа-яЁё]", x or "")]
        ok(f"{lang}: блок оценки и итоговый вывод на языке акта, без кириллицы",
           word in [r["label"] for r in s["rows"]] and not cyr, cyr[:4])
    st, blob, h = call("GET", f"/act/{aid}.pdf", params={"lang": "en"}, raw=True)
    ok("PDF en: блок оценки", "Valuation by listings" in pdf_text(pymupdf.open(stream=blob, filetype="pdf")))

    # п. 1: сотрудник снял галочки с двух самых дорогих объявлений (L7 — 2,8 млрд, L1 — 2,65 млрд)
    off2 = [dict(r, relevant=False) if r["id"] in ("L1", "L7") else dict(r) for r in listings]
    st, a = mk_make(off2, sid)
    mv = a["market_value"]
    used2 = [2_500_000_000, round(205000 * RATE), 2_400_000_000]
    reasons = {e["id"]: e for e in mv["excluded"]}
    ok("снятые галочки: причина «снято сотрудником», а не «другое изделие»",
       reasons["L1"]["code"] == reasons["L7"]["code"] == "mx_unchecked_by_employee"
       and "снято сотрудником" in reasons["L1"]["reason"] and "другое изделие" not in reasons["L1"]["reason"]
       and reasons["L5"]["code"] == "mx_not_relevant" and "другое изделие" in reasons["L5"]["reason"], reasons)
    ok("медиана без двух дорогих — руками; медиана без правок — как прочитала модель",
       mv["median"] == round(hand_q(used2, 0.5)) == 2_500_000_000 and mv["median_original"] == round(med), mv)
    ok("market_value.edits: снято 2, остальное 0, список правок",
       {k: mv["edits"][k] for k in ("unchecked", "checked", "price_changed", "removed", "manual")}
       == {"unchecked": 2, "checked": 0, "price_changed": 0, "removed": 0, "manual": 0}
       and sorted((e["id"], e["what"]) for e in mv["edits"]["items"]) == [("L1", "unchecked"), ("L7", "unchecked")],
       mv["edits"])
    line = "Правки сотрудника: снято 2, включено 0, исправлено цен 0, убрано 0, добавлено вручную 0."
    ok("раздел 3: строка «Правки сотрудника: снято 2, …» и список правок",
       line in sec3(a)["source_lines"] and any(li["title"] == "Правки сотрудника в объявлениях"
                                                for li in sec3(a)["lists"]), sec3(a)["source_lines"])
    lv = {r["id"]: r for r in mv["listings"]}
    ok("у объявления видно «снято сотрудником»", lv["L1"]["off_by"] == "employee"
       and lv["L1"]["edit_note"] == "снято сотрудником", lv["L1"])
    chk = [flat(c) for c in a["decision"]["checks"]]
    ok("решение: «проверить правки сотрудника» с медианой без правок рядом с итоговой",
       any(c.startswith("Проверить правки сотрудника в объявлениях (снято 2") and "2 500 000 000 сум" in c
           and "2 593 250 000 сум" in c for c in chk), chk)
    med_row = next(r for r in sec3(a)["rows"] if r["label"] == "Медиана цен объявлений")
    ok("раздел 3: у медианы — «без правок сотрудника — 2 593 250 000»",
       "без правок сотрудника — 2 593 250 000 сум" in flat(med_row["note"]), med_row)
    plain, text = docx_plain(a["id"]), pdf_plain(a["id"])
    ok("DOCX: правки видны — строка, «снято сотрудником», проверка правок", line in plain
       and "снято сотрудником" in plain and "Проверить правки сотрудника" in plain, plain[-600:])
    ok("PDF: правки видны — строка и «снято сотрудником»", "Правки сотрудника: снято 2" in text
       and "снято сотрудником" in text, text[-600:])

    # правка цены (было → стало), включение исключённого моделью и объявление вручную
    ed = [dict(r) for r in listings]
    ed[0]["price"] = 1_000_000_000                        # L1: 2 650 000 000 → 1 000 000 000
    ed[4]["relevant"] = True                              # L5: модель исключила, сотрудник включил
    ed.append({"title": "XCMG QY50K5D 2021 (у дилера)", "price": 2_550_000_000, "currency": "UZS", "year": 2021,
               "posted_date": _days_ago(2), "url": "https://www.olx.uz/d/obyavlenie/xcmg-qy50k5d-ID1.html",
               "source": "manual"})
    st, a = mk_make(ed, sid)
    mv = a["market_value"]
    cand = [1_000_000_000, 2_500_000_000, round(205000 * RATE), 2_400_000_000, 1_500_000_000, 900_000_000,
            2_800_000_000, 2_550_000_000]
    m0 = hand_q(cand, 0.5)
    used3 = [x for x in cand if 0.5 * m0 <= x <= 2 * m0]
    ok("правки: медиана и вилка руками", mv["used"] == len(used3) == 6 and mv["median"] == round(hand_q(used3, 0.5))
       and mv["low"] == round(hand_q(used3, 0.25)) and mv["high"] == round(hand_q(used3, 0.75)), mv)
    src = {r["id"]: (r["source"], r["edited"]) for r in mv["listings"]}
    ok("исправленные помечены, ручное — «введено сотрудником»", src["L1"] == ("shot", True)
       and src["L5"] == ("shot", True) and src["L2"] == ("shot", False) and src["m11"] == ("manual", False), src)
    items = {(e["id"], e["what"]): flat(e["text"]) for e in mv["edits"]["items"]}
    ok("цена: «цена исправлена сотрудником: было 2 650 000 000, стало 1 000 000 000»",
       items.get(("L1", "price")) == "XCMG QY50K5D 2021 — цена исправлена сотрудником: было 2 650 000 000, "
                                     "стало 1 000 000 000", items)
    ok("включено сотрудником (модель исключила) — с причиной модели",
       "включено сотрудником" in items.get(("L5", "checked"), "") and "QY25K5D" in items.get(("L5", "checked"), ""),
       items)
    ok("счёт правок: включено 1, исправлено цен 1, добавлено вручную 1",
       (mv["edits"]["checked"], mv["edits"]["price_changed"], mv["edits"]["manual"], mv["edits"]["unchecked"]) ==
       (1, 1, 1, 0), mv["edits"])
    ok("строки источника: снимки + введено сотрудником + строка правок",
       any("снимки загружены" in x for x in mv["source_lines"]) and any("введено сотрудником" in x
                                                                         for x in mv["source_lines"])
       and "Правки сотрудника: снято 0, включено 1, исправлено цен 1, убрано 0, добавлено вручную 1." in
       mv["source_lines"], mv["source_lines"])
    ok("адрес объявления сохранён, но не открывается сервером",
       mv["listings"][-1]["url"] == "https://www.olx.uz/d/obyavlenie/xcmg-qy50k5d-ID1.html")
    plain = docx_plain(a["id"])
    ok("DOCX: правка цены — было и стало", "было 2 650 000 000, стало 1 000 000 000" in plain)

    # правка валюты, года и даты публикации — тоже «было → стало»
    ed2 = [dict(r) for r in listings]
    ed2[1].update(year=2015, posted_date=_days_ago(20))
    st, a = mk_make(ed2, sid)
    whats = {(e["id"], e["what"]) for e in a["market_value"]["edits"]["items"]}
    ok("правка года и даты публикации видна", {("L2", "year"), ("L2", "posted_date")} <= whats, whats)

    # убранные из списка объявления загрузки: «убрано сотрудником», в «Не вошли в расчёт»
    cut = [dict(r) for r in listings if r["id"] not in ("L2", "L4")]
    st, a = mk_make(cut, sid)
    mv = a["market_value"]
    reasons = {e["id"]: e for e in mv["excluded"]}
    excl = next(li for li in sec3(a)["lists"] if li["title"] == "Не вошли в расчёт")
    ok("убранные: «убрано сотрудником», в списке «Не вошли в расчёт», счёт правок removed = 2",
       reasons["L2"]["code"] == reasons["L4"]["code"] == "mx_removed_by_employee" and mv["edits"]["removed"] == 2
       and mv["count"] == 10 and any("убрано сотрудником" in x for x in excl["items"]), (mv["edits"], excl))
    ok("убранные: проверка правок с медианой без правок",
       any(c.startswith("Проверить правки сотрудника") for c in a["decision"]["checks"]), a["decision"]["checks"])

    # п. 2: загрузка снимков недоступна — даты с экрана учитываются, всё «введено сотрудником», курс ЦБ не «сотрудника»
    act._FX_CACHE.clear()
    act._fx_lookup(TODAY)                                 # сервер уже выдавал курс ЦБ на сегодня
    dead = "0" * 24
    st, a = mk_make(listings, dead, fx={"rate": RATE, "by": "cbu", "as_of": TODAY.isoformat()})
    mv = a["market_value"]
    miss = ("Снимки объявлений недоступны (прошло больше 24 часов или сменилась сессия) — объявления учтены как "
            "введённые сотрудником.")
    codes = {e["id"]: e["code"] for e in mv["excluded"]}
    ok("недоступная загрузка: shots_missing и строка в акте", mv["shots_missing"] is True
       and mv["shots_missing_text"] == miss and miss in sec3(a)["source_lines"], mv["source_lines"])
    ok("недоступная загрузка: все «введено сотрудником», даты экрана учтены (6 месяцев работают)",
       all(r["source"] == "manual" for r in mv["listings"]) and codes.get("L9") == "mx_too_old"
       and codes.get("L8") == "mx_no_date" and mv["median"] == round(med) and mv["used"] == 5, codes)
    ok("курс ЦБ, полученный экраном от сервера и сверенный, — «ЦБ РУз», не «введён сотрудником»",
       mv["fx"]["by"] == "cbu" and "cbu.uz" in mv["fx"]["text"] and "сотрудник" not in mv["fx"]["text"], mv["fx"])
    plain = docx_plain(a["id"])
    ok("DOCX: строка о недоступных снимках", "Снимки объявлений недоступны" in plain)
    act._FX_CACHE.clear()
    st, a = mk_make(listings, dead, fx={"rate": RATE, "by": "cbu", "as_of": TODAY.isoformat()})
    fx = a["market_value"]["fx"]
    ok("курс ЦБ не с чем сверить — «ЦБ РУз, повторно не сверен», всё равно не «сотрудник»",
       fx["by"] == "cbu_unverified" and "не сверен" in fx["text"] and "введён сотрудником" not in fx["text"], fx)
    st, a = mk_make(listings, dead, fx={"rate": 12_700, "by": "employee", "as_of": TODAY.isoformat()})
    fx = a["market_value"]["fx"]
    ok("курс, введённый сотрудником, — «введён сотрудником»", fx["by"] == "employee" and fx["rate"] == 12_700
       and "введён сотрудником" in fx["text"], fx)
    st, b = mk_make(listings, dead, fx={"rate": RATE, "by": "robot"})
    ok("источник курса проверяется", st == 422 and "fx.by" in b["errors"]["market"], (st, b))
    market_setup()

    # п. 5: стоимость объекта заменена медианой — в акте видно, что заявил клиент
    must = dict(CRANE_MUST, object_value=round(med))
    st, a = mk_make(listings, sid, must=must, optional=dict(CRANE_OPT, declared_value_original=3_100_000_000))
    rows = {r["label"]: r for r in sec3(a)["rows"]}
    ok("замена медианой: «Заявлено клиентом: …; стоимость принята по объявлениям: …»",
       "Заявлено клиентом" in rows and flat(rows["Заявлено клиентом"]["note"]) ==
       "Заявлено клиентом: 3 100 000 000 сум; стоимость принята по объявлениям: 2 593 250 000 сум."
       and a["value"]["value_source"] == "listings" and a["value"]["declared_original"] == 3_100_000_000, rows)
    ok("после замены: объявления подтверждают стоимость, итог — по сумме к стоимости (превышение)",
       a["market_value"]["verdict"] == "confirmed" and a["value"]["final_verdict"] == "over", a["value"])
    ok("DOCX: «Заявлено клиентом» и «стоимость принята по объявлениям»",
       "стоимость принята по объявлениям: 2 593 250 000" in docx_plain(a["id"]))
    st, b = mk_make(listings, sid, optional=dict(CRANE_OPT, declared_value_original="abc"))
    ok("declared_value_original проверяется", st == 422 and "declared_value_original" in b["errors"], (st, b))

    # только ручной ввод, без снимков: источник — «введено сотрудником»; курс сотрудника
    manual = [{"title": "XCMG QY50K5D", "price": 2_900_000_000, "posted_date": _days_ago(4)},
              {"title": "XCMG QY50K5D", "price": 3_000_000_000, "posted_date": _days_ago(9)},
              {"title": "XCMG QY50K5D", "price": 240_000, "currency": "USD", "posted_date": _days_ago(1)}]
    market_setup(rate=None)
    st, a = mk_make(manual, None, usd_rate=12_500)
    mv = a["market_value"]
    ok("ручной ввод: «Источник: введено сотрудником», курс сотрудника, confirmed",
       mv["source_label"] == "Источник: введено сотрудником (объявлений: 3)." and mv["fx"]["by"] == "employee"
       and mv["median"] == 3_000_000_000 and mv["verdict"] == "confirmed"
       and not any(c.startswith("Уточнить стоимость") for c in a["decision"]["checks"]), mv)
    st, a = mk_make(manual, None)
    ok("ручной ввод без курса: доллары не вошли, курс не выдуман", a["market_value"]["used"] == 2
       and a["market_value"]["fx"] is None, a["market_value"]["fx"])
    st, a = mk_make([dict(r, posted_date=None) for r in manual[:2]], None)
    ok("ручной ввод без даты — не в расчёте («дата публикации не видна»)",
       not a["market_value"]["available"] and {e["code"] for e in a["market_value"]["excluded"]} == {"mx_no_date"})
    market_setup()

    # мало объявлений и ни одного
    st, a = mk_make(listings[:2], sid)
    mv = a["market_value"]
    ok("два объявления — «мало, ориентировочно», стоимость не уточняется, выбросы не искались",
       mv["verdict"] == "few" and mv["refined_value"] is None and "ориентировочная" in mv["verdict_text"]
       and not any(c.startswith("Уточнить стоимость") for c in a["decision"]["checks"])
       and any("выбросы не искались" in h for h in mv["how"]), mv["verdict"])
    none = [dict(r, relevant=False) for r in listings]
    st, a = mk_make(none, sid)
    st0, a0 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                         "optional": dict(CRANE_OPT, object_kind="truck_crane")})
    n_rel = sum(1 for r in stored["listings"] if r.get("relevant") is not False)
    ok("ни одного — оценки нет, строки раздела 3 как без неё, но правки видны",
       not a["market_value"]["available"] and a["market_value"]["verdict"] == "none"
       and [r["label"] for r in sec3(a)["rows"]] == [r["label"] for r in sec3(a0)["rows"]]
       and f"Правки сотрудника: снято {n_rel}, включено 0, исправлено цен 0, убрано 0, добавлено вручную 0."
       in sec3(a)["source_lines"], sec3(a)["source_lines"])
    ok("акт без блока market: market_value.available = false", a0["market_value"]["available"] is False
       and a0["market_value"]["verdict"] == "none" and not sec3(a0)["source_lines"])

    # минимум объявлений — настройка
    with db.tx() as con:
        con.execute("INSERT INTO act_settings (created_at, created_by, settings_json, calibrated, note) "
                    "VALUES (?,?,?,?,?)", (db.now(), "тест", _json.dumps({"market": {"min_listings": 6}}), 0, "тест"))
    try:
        st, a = mk_make(listings, sid)
        ok("min_listings = 6 из настроек — 5 объявлений уже «мало»", a["market_value"]["verdict"] == "few")
    finally:
        clear_settings()

    # недоверенный ввод и ПД в ручном вводе
    for bad, key in (({"price": "abc"}, "цена"), ({"currency": "EUR", "price": 1}, "currency"),
                     ({"price": 1, "url": "javascript:alert(1)"}, "адрес"), ({"price": 1, "year": 1700}, "год"),
                     ({"price": 1, "url": "http://169.254.169.254/latest/meta-data/"}, "адрес"),
                     ({"price": 1, "url": "https://example.uz/ad/1"}, "адрес"),
                     ({"price": 1, "posted_date": "31.02.2026"}, "дата")):
        st, b = mk_make([bad], sid)
        ok(f"ввод проверяется: {key} ({bad.get('url') or bad.get('posted_date') or ''})",
           st == 422 and "market" in b["errors"] and key in b["errors"]["market"], (st, b))
    st, a = mk_make([{"title": "Кран, продаёт Петров Пётр Петрович", "price": 2_700_000_000,
                      "region": "Каримов Алишер", "why_excluded": None},
                     {"title": "XCMG, тел. +998 91 555 44 33", "price": 2_600_000_000}], None)
    dump = _json.dumps(a, ensure_ascii=False)
    with db.tx() as con:
        saved = db.rows(con, "SELECT act_json FROM acts WHERE id=?", a["id"])[0]["act_json"]
    ok("ПД в ручном вводе отброшены — ни в ответе, ни в базе",
       not any(x in dump or x in saved for x in ("Петрович", "Петров", "555 44 33", "Каримов")), dump[:300])
    st, b = mk_make(listings, "не-та-сессия")
    ok("чужая или неизвестная загрузка снимков — все объявления «введено сотрудником»",
       all(r["source"] == "manual" for r in b["market_value"]["listings"]) and b["market_value"]["shots_missing"])
    return aid


# ------------------------------------------------------------------ 37. запрос филиала (30.09.2026)

BR_HEAD = ["ОСГОР бўйича белгиланган чегарадан ошиб кетиш***", "бошқа: ______________"]
BR_SUM1 = "81 250 000 000,00 (саксон бир миллиард икки юз эллик миллион сўм ва 00 тийин) сўм"
BR_SUM2 = ("47 397 852 345,04 (қирқ етти миллиард уч юз тўқсон етти миллион саккиз юз эллик икки минг уч юз "
           "қирқ беш сўм ва 04 тийин) сўм")
# два образца заказчика (сканы 24.png и 25.png) — строки как в бланке, узбекская кириллица
BR_SAMPLE1 = [
    ("1.", "Суғурта тури (буйруқ бўйича код):", "0832"),
    ("2.", "Суғурта қилдирувчи номи:", '"NAMUNA SAVDO" MCHJ'),
    ("3.", "Наф олувчи:", 'CHEKI "Namuna Bank" ATB Sinov universal BXO'),
    ("4.", "Гаровга қўювчи", '"OMAD" AJ'),
    ("5.", "Суғурта объекти:", ["«Кўчмас мулк нотурар бино қишлоқ хўжалиги махсулотларини сақлаш учун музлатгич»",
                                "Ер участкасининг умумий майдони 8 640,00 кв.м.",
                                "Умумий фойдали майдони 11 898.42 кв.м", "Умумий майдони 13 344.00 кв.м.",
                                "кадастр рақами 10:00:00:00:00:00001"]),
    ("6.", "Суғурта қиймати:", BR_SUM1),
    ("7.", "Суғурта суммаси:", BR_SUM1),
    ("8.", "Франшиза:", "Қўлланилинмайди"),
    ("9.", "Суғурта тарифи :", "0.05"),
    ("10.", "Суғурта мукофоти:", "123 322 000,00 (бир юз йигирма уч миллион уч юз йигирма икки минг сўм ва 00 тийин) сўм"),
    ("11.", "Суғурта муддати:", ["2026 йил «29» сентябрдан", "2029 йил «10» октябргача"]),
    ("12.", "Стандарт суғурта шартномаси шартларини ўзгартириш/қўшиш:*", "Стандарт"),
    ("13.", "Контрагент:**", ""),
    ("14.", "Шартнома миқдори:**", "1 дона"),
    ("15.", "Класс*** (ОСГОР бўйича)", ""),
    ("16.", "Қўшимча маълумот:", ""),
]
BR_SAMPLE2 = [
    ("1.", "Суғурта тури (буйруқ бўйича код):", "0832"),
    ("2.", "Суғурта қилдирувчи номи:", '"Namuna Bank" АТБ Намуна УБХО'),
    ("3.", "Наф олувчи:", '"SINOV" MCHJ'),
    ("4.", "Гаровга қўювчи", ""),
    ("5.", "Суғурта объекти:", "Технологик асбоб ускуна нон махсулотлари ишлаб чиқариш учун"),
    ("6.", "Суғурта қиймати:", BR_SUM2),
    ("7.", "Суғурта суммаси:", BR_SUM2),
    ("8.", "Франшиза:", "Қўлланилинмайди"),
    ("9.", "Суғурта тарифи :", "0.05"),
    ("10.", "Суғурта мукофоти:", "122 589 000,00 (бир юз йигирма икки миллион беш юз саксон тўққиз минг сўм ва 00 "
                                "тийин) сўм"),
    ("11.", "Суғурта муддати:", ["2026 йилнинг «07» сентябрдан", "2031 йилнинг «07» ноябр мобайнида"]),
    ("12.", "Стандарт суғурта шартномаси шартларини ўзгартириш/қўшиш:*", "Стандарт"),
    ("13.", "Контрагент:**", ""),
    ("14.", "Шартнома миқдори:**", "1 дона"),
    ("15.", "Класс*** (ОСГОР бўйича)", ""),
    ("16.", "Қўшимча маълумот:", ""),
]
BR_S1 = 81_250_000_000.0
BR_S2 = 47_397_852_345.04


def _x(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def docx_table(rows, head=BR_HEAD) -> bytes:
    """DOCX с абзацами-шапкой и таблицей из 16 строк (ячейка с несколькими строками — несколько абзацев)."""
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    para = lambda t: f"<w:p><w:r><w:t xml:space=\"preserve\">{_x(t)}</w:t></w:r></w:p>"   # noqa: E731
    body = "".join(para(h) for h in head) + "<w:tbl>"
    for n, lab, val in rows:
        vals = val if isinstance(val, list) else [val]
        body += "<w:tr>" + "".join(f"<w:tc>{c}</w:tc>" for c in (
            para(n), para(lab), "".join(para(v) for v in vals))) + "</w:tr>"
    body += "</w:tbl>"
    xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{ns}"><w:body>{body}</w:body></w:document>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def xlsx_table(rows) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сўров"
    for h in BR_HEAD:
        ws.append([h])
    for n, lab, val in rows:
        ws.append([n, lab, "\n".join(val) if isinstance(val, list) else val])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_table(rows) -> bytes:
    """PDF с текстовым слоем: каждая ячейка — своя строка (так pymupdf отдаёт текст таблицы)."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 40
    font = r"C:\Windows\Fonts\arial.ttf"
    for n, lab, val in rows:
        for t in [n, lab] + (val if isinstance(val, list) else [val]):
            if t:
                page.insert_text((40, y), t[:95], fontsize=7, fontname="arl", fontfile=font)
                y += 9
    return doc.tobytes()


def br_request(b):
    return (b.get("branch_request") or {}).get("request")


def nb(s) -> str:
    return str(s).replace(" ", " ")


def check_br_fields(tag, f, sample):
    """Поля бланка против расшифровки заказчика (оба образца)."""
    if sample == 1:
        ok(f"{tag}: код продукта, стороны (юрлица), залог",
           f["product_code"] == "0832" and f["policyholder"] == {"kind": "legal", "name": '"NAMUNA SAVDO" MCHJ'}
           and f["beneficiary"]["name"] == 'CHEKI "Namuna Bank" ATB Sinov universal BXO'
           and f["pledger"] == {"kind": "legal", "name": '"OMAD" AJ'} and f["has_pledger"] is True
           and f["has_beneficiary"] is True, f)
        ok(f"{tag}: объект — здание, площади, кадастр",
           f["class_hint"] == "building" and f["object_kind"] == "warehouse"
           and f["areas"] == {"land_m2": 8640.0, "useful_m2": 11898.42, "total_m2": 13344.0}
           and f["cadastre_no"] == "10:00:00:00:00:00001" and "музлатгич" in f["object_description"], f)
        ok(f"{tag}: стоимость и сумма — число без суммы прописью",
           f["object_value"] == BR_S1 and f["sum_insured"] == BR_S1, (f["object_value"], f["sum_insured"]))
        ok(f"{tag}: тариф 0,05, премия 123 322 000, франшиза не применяется",
           f["tariff_pct"] == 0.05 and f["premium"] == 123_322_000.0
           and f["franchise"] == {"applied": False, "text": "Қўлланилинмайди", "pct": None, "amount": None}, f)
        ok(f"{tag}: срок 29.09.2026–10.10.2029, 1 108 дн. включительно",
           f["term_from"] == "2026-09-29" and f["term_to"] == "2029-10-10" and f["term_days"] == 1108
           and f["term_inclusive"] is True, (f["term_from"], f["term_to"], f["term_days"]))
    else:
        ok(f"{tag}: код продукта, стороны, залога нет",
           f["product_code"] == "0832" and f["policyholder"]["name"] == '"Namuna Bank" АТБ Намуна УБХО'
           and f["beneficiary"] == {"kind": "legal", "name": '"SINOV" MCHJ'}
           and f["pledger"] == {"kind": None, "name": None} and f["has_pledger"] is False, f)
        ok(f"{tag}: объект — оборудование",
           f["class_hint"] == "equipment" and f["object_kind"] == "equipment" and f["cadastre_no"] is None
           and f["areas"] == {"land_m2": None, "useful_m2": None, "total_m2": None}, f)
        ok(f"{tag}: стоимость и сумма 47 397 852 345,04", f["object_value"] == BR_S2 and f["sum_insured"] == BR_S2,
           (f["object_value"], f["sum_insured"]))
        ok(f"{tag}: тариф, премия 122 589 000, франшиза не применяется",
           f["tariff_pct"] == 0.05 and f["premium"] == 122_589_000.0 and f["franchise"]["applied"] is False, f)
        ok(f"{tag}: срок 07.09.2026–07.11.2031 («йилнинг … мобайнида»), 1 888 дн.",
           f["term_from"] == "2026-09-07" and f["term_to"] == "2031-11-07" and f["term_days"] == 1888,
           (f["term_from"], f["term_to"], f["term_days"]))
    ok(f"{tag}: стандартные условия, 1 договор, пустые строки — null",
       f["contract_terms"] == "Стандарт" and f["contract_terms_standard"] is True and f["contracts_count"] == 1
       and f["counterparty"] is None and f["osgor_class"] is None and f["additional_info"] is None, f)


def check_branch_text():
    print("37а. Запрос филиала: файл с текстом (DOCX, XLSX, PDF) — без модели")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    CALLS.clear()
    st, b = upload([("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1))], {"lang": "ru"})
    ok("DOCX образца 1 принят, модель не вызывалась", st == 200 and b.get("ok") and not CALLS, (st, len(CALLS)))
    brq = b.get("branch_request") or {}
    ok("узнан как «запрос филиала»: 16 строк из 16, источник — документ",
       brq.get("detected") and brq.get("rows_found") == 16 and brq.get("rows_total") == 16
       and brq.get("source") == "document" and brq.get("kind_label") == "запрос филиала"
       and b["files"][0]["document_kind"] == "запрос филиала", brq)
    ok("строки бланка подписаны: «Страховой тариф», пустые отмечены",
       brq["rows"][8]["label"] == "Страховой тариф" and brq["rows"][8]["filled"]
       and not brq["rows"][12]["filled"] and all(r["found"] for r in brq["rows"]), brq["rows"])
    check_br_fields("DOCX 1", brq["fields"], 1)
    rq = br_request(b)
    ok("готовый optional.request для /act/make",
       rq["tariff_pct"] == 0.05 and rq["premium"] == 123_322_000.0 and rq["term_days"] == 1108
       and rq["franchise"]["applied"] is False and rq["source"] == "document", rq)
    pf = b.get("prefill") or {}
    ok("prefill шага 2: код продукта, сумма, стоимость, срок и даты — «из документа, проверьте»",
       pf.get("product_code", {}).get("value") == "0832" and pf["sum_insured"]["value"] == BR_S1
       and pf["object_value"]["value"] == BR_S1 and pf["term_days"]["value"] == 1108
       and pf["term_from"]["value"] == "2026-09-29" and pf["term_to"]["value"] == "2029-10-10"
       and "region" not in pf and all(v["check_label"] == "из документа, проверьте" for v in pf.values()), pf)
    rec = {(r["key"], r["value"]) for r in b["recognized"]}
    ok("распознанное: премия, тариф, срок, стороны-юрлица, площади, кадастр",
       {("premium", "123 322 000"), ("tariff_pct", "0.05"), ("term_days", "1108"), ("term_from", "2026-09-29"),
        ("pledger", '"OMAD" AJ'), ("land_area", "8 640 м²"), ("cadastre_no", "10:00:00:00:00:00001"),
        ("sum_insured", "81 250 000 000"), ("product_code", "0832")} <= rec, sorted(rec))
    ok("подписи новых полей есть", all(r["label"] != r["key"] for r in b["recognized"]),
       [r["key"] for r in b["recognized"] if r["label"] == r["key"]])
    ok("класс-подсказка из строки объекта: здание → 8 или 9", b["class_hint"] == "building"
       and b["suggest_classes"] == ["8", "9"], b["class_hint"])
    with db.tx() as con:
        audit = db.rows(con, "SELECT detail FROM audit WHERE entity=?", "act_upload:" + b["session"])
    dump = _json.dumps(audit, ensure_ascii=False)
    det = _json.loads(audit[0]["detail"]) if audit else {}
    ok("в журнале — только признак и число строк, без названий сторон и сумм",
       det.get("branch_request") is True and det.get("branch_rows") == 16 and "OMAD" not in dump and "NAMUNA" not in dump and "81250" not in dump,
       dump[:300])
    sid1 = b["session"]

    CALLS.clear()
    st, b2 = upload([("sorov2.xlsx", XLSX_MIME, xlsx_table(BR_SAMPLE2))], {"lang": "ru"})
    ok("XLSX образца 2 принят, модель не вызывалась", st == 200 and not CALLS and b2.get("branch_request"), (st, b2))
    check_br_fields("XLSX 2", b2["branch_request"]["fields"], 2)
    ok("образец 2: класс-подсказка — оборудование", b2["class_hint"] == "equipment" and b2["group"] == "equipment",
       (b2["class_hint"], b2["group"]))

    st, b3 = upload([("sorov1.pdf", "application/pdf", pdf_table(BR_SAMPLE1))], {"lang": "ru"})
    ok("PDF с текстом (ячейки построчно) — тот же разбор без модели",
       st == 200 and not CALLS and (b3.get("branch_request") or {}).get("rows_found") == 16, (st, b3))
    if b3.get("branch_request"):
        check_br_fields("PDF 1", b3["branch_request"]["fields"], 1)

    # русские подписи и русский срок
    ru_rows = [("1.", "Вид страхования (код):", "0832"), ("2.", "Страхователь:", 'ООО "Ромашка"'),
               ("3.", "Выгодоприобретатель:", 'АКБ "Капиталбанк"'), ("4.", "Залогодатель:", ""),
               ("5.", "Объект страхования:", "Нежилое здание — склад, общая площадь 1 200 кв.м"),
               ("6.", "Страховая стоимость:", "1 000 000 000,00 (один миллиард) сум"),
               ("7.", "Страховая сумма:", "1 000 000 000,00 сум"), ("8.", "Франшиза:", "не применяется"),
               ("9.", "Страховой тариф:", "0,1"), ("10.", "Страховая премия:", "3 000 000,00 сум"),
               ("11.", "Срок страхования:", "с 29.09.2026 по 10.10.2029")]
    st, b4 = upload([("ru.docx", DOCX_MIME, docx_table(ru_rows, head=[]))], {"lang": "ru"})
    f4 = (b4.get("branch_request") or {}).get("fields") or {}
    ok("русские подписи: тариф 0,1, срок «с … по …» → 1 108 дн., франшиза не применяется, общая площадь",
       f4.get("tariff_pct") == 0.1 and f4.get("term_days") == 1108 and f4.get("franchise", {}).get("applied") is False
       and f4.get("areas", {}).get("total_m2") == 1200.0 and f4.get("beneficiary", {}).get("kind") == "legal", f4)

    # обычный договор с похожими подписями, но без кода продукта — не запрос филиала
    contract = [(str(i) + ".", lab, val) for i, (lab, val) in enumerate((
        ("Вид страхования:", "страхование имущества"), ("Страхователь:", 'ООО "Ромашка"'),
        ("Выгодоприобретатель:", 'АКБ "Капиталбанк"'), ("Объект страхования:", "склад"),
        ("Страховая стоимость:", "1 000 000 000 сум"), ("Страховая сумма:", "1 000 000 000 сум"),
        ("Франшиза:", "нет"), ("Страховой тариф:", "0,1 %"), ("Страховая премия:", "1 000 000 сум"),
        ("Срок страхования:", "с 01.10.2026 по 30.09.2027")), 1)]
    st, b6 = upload([("dogovor.docx", DOCX_MIME, docx_table(contract, head=["ДОГОВОР СТРАХОВАНИЯ ИМУЩЕСТВА"]))],
                    {"lang": "ru"})
    ok("договор с похожими подписями без кода продукта — не запрос филиала",
       st == 200 and b6.get("branch_request") is None, (st, b6.get("branch_request")))
    # срок без последнего дня — настройка request_check.term_inclusive = false
    with db.tx() as con:
        con.execute("INSERT INTO act_settings (created_at, created_by, settings_json, calibrated, note) "
                    "VALUES (?,?,?,?,?)", (db.now(), "тест", _json.dumps({"request_check": {"term_inclusive": False,
                                                                                            "premium_tolerance": 1000}}),
                                           0, "тест"))
    st, b5 = upload([("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1))], {"lang": "ru"})
    ok("настройка term_inclusive = false: 1 107 дн.",
       b5["branch_request"]["fields"]["term_days"] == 1107 and b5["branch_request"]["fields"]["term_inclusive"] is False,
       b5["branch_request"]["fields"]["term_days"])
    clear_settings()
    ok("ae.check_settings: request_check проверяется",
       ae.check_settings({"request_check": {"term_inclusive": "да", "premium_tolerance": -1}}) and
       not ae.check_settings({"request_check": {"term_inclusive": True, "premium_tolerance": 500}}))
    return sid1, b


def br_model_reply(sample=1, policyholder=None):
    rows = BR_SAMPLE1 if sample == 1 else BR_SAMPLE2
    v = {br_code: (val if not isinstance(val, list) else " ".join(val)) or None
         for br_code, (_n, _l, val) in zip(
             ["product_code", "policyholder", "beneficiary", "pledger", "object", "object_value", "sum_insured",
              "franchise", "tariff", "premium", "term", "contract_terms", "counterparty", "contracts_count",
              "osgor_class", "additional_info"], rows)}
    for k in ("policyholder", "beneficiary", "pledger"):
        v[k] = {"is_legal": bool(v[k]), "name": v[k]}
    if policyholder:
        v["policyholder"] = {"is_legal": True, "name": policyholder}      # модель ошиблась: гражданин как юрлицо
    v.update(file=1, term_from="2026-09-29" if sample == 1 else "2026-09-07",
             term_to="2029-10-10" if sample == 1 else "2031-11-07",
             object_description_translated=("недвижимость: нежилое здание — холодильник для хранения "
                                            "сельхозпродукции" if sample == 1 else
                                            "технологическое оборудование для производства хлебобулочных изделий"),
             class_hint="building" if sample == 1 else "equipment")
    return "```json\n" + _json.dumps({"files": [{"n": 1, "view": "document", "document_kind": "branch_request"}],
                                      "object_kind": None, "class_hint": None, "condition": None, "fields": [],
                                      "damages": [], "branch_request": v}, ensure_ascii=False) + "\n```"


def check_branch_scan():
    print("37б. Запрос филиала: скан — ответ модели (подменён) по строгой схеме")
    fresh()
    model_on(True)
    CALLS.clear()
    REPLY["text"] = br_model_reply(1)
    st, b = upload([("scan.png", "image/png", image((250, 250, 250)))], {"lang": "ru"})
    ok("скан принят, одно обращение к модели", st == 200 and len(CALLS) == 1, (st, len(CALLS)))
    prompt = CALLS[0]["messages"][0]["content"] + CALLS[0]["messages"][1]["content"]
    ok("в инструкции модели — запрос филиала, branch_request и запрет имён граждан",
       "branch_request" in prompt and "запрос филиала" in prompt and "is_legal" in prompt
       and "суғурта мукофоти" in prompt, prompt[-400:])
    ok("инструкция не портится маскировкой ПД", llm.mask_pd(prompt) == prompt)
    brq = b.get("branch_request") or {}
    ok("вид документа — запрос филиала, источник — скан", brq.get("source") == "photo"
       and b["files"][0]["document_kind"] == "запрос филиала" and brq.get("file") == "f1", brq)
    check_br_fields("скан 1", brq["fields"], 1)
    ok("перевод описания объекта — от модели, исходный текст сохранён",
       "холодильник" in (brq["fields"]["object_description_translated"] or "")
       and "музлатгич" in brq["fields"]["object_description"], brq["fields"]["object_description_translated"])
    rec = {(r["key"], r["value"], r["source"]) for r in b["recognized"]}
    ok("распознанное со скана — источник «документ»",
       ("premium", "123 322 000", "document") in rec and ("beneficiary", 'CHEKI "Namuna Bank" ATB Sinov universal BXO',
                                                          "document") in rec, sorted(rec))
    ok("prefill со скана: код продукта и срок", (b.get("prefill") or {}).get("product_code", {}).get("value") == "0832"
       and b["prefill"]["term_days"]["value"] == 1108, b.get("prefill"))
    sid = b["session"]

    CALLS.clear()
    REPLY["text"] = br_model_reply(2)
    st, b2 = upload([("scan2.jpg", "image/jpeg", image(kind="jpg"))], {"lang": "ru"})
    check_br_fields("скан 2", (b2.get("branch_request") or {}).get("fields") or {}, 2)

    # гражданин в строке страхователя: имени нет ни в ответе, ни в базе, ни в журнале
    CALLS.clear()
    REPLY["text"] = br_model_reply(1, policyholder="Каримов Алишер Анварович")
    st, b3 = upload([("scan3.png", "image/png", image((240, 240, 240)))], {"lang": "ru"})
    f3 = b3["branch_request"]["fields"]
    dump = _json.dumps(b3, ensure_ascii=False)
    with db.tx() as con:
        saved = db.rows(con, "SELECT result_json FROM act_uploads WHERE id=?", b3["session"])[0]["result_json"]
        audit = _json.dumps(db.rows(con, "SELECT detail FROM audit WHERE entity=?", "act_upload:" + b3["session"]),
                            ensure_ascii=False)
    ok("физлицо в строке страхователя: kind = individual, имени нет",
       f3["policyholder"] == {"kind": "individual", "name": None} and "Каримов" not in dump
       and "Каримов" not in saved and "Каримов" not in audit, f3["policyholder"])
    ok("пометка «физическое лицо — данные не извлекаются»",
       any("физическое лицо" in n for n in b3["branch_request"]["notes"]), b3["branch_request"]["notes"])
    st, b4 = upload([("s.docx", DOCX_MIME, docx_table([(n, l, "Каримов Алишер" if n == "2." else v)
                                                        for n, l, v in BR_SAMPLE1]))], {"lang": "ru"})
    dump = _json.dumps(b4, ensure_ascii=False)
    ok("файл с текстом: гражданин-страхователь не извлекается",
       b4["branch_request"]["fields"]["policyholder"] == {"kind": "individual", "name": None}
       and "Каримов" not in dump, b4["branch_request"]["fields"]["policyholder"])
    return sid


def br_make(sid, sample=1, request=None, optional=None, lang="ru", must=None):
    S = BR_S1 if sample == 1 else BR_S2
    body = {"session": sid, "lang": lang,
            "must": must or {"product_code": "0832", "sum_insured": S, "object_value": S,
                             "region": "Ташкентская область"},
            "optional": dict(optional or {})}
    if request is not None:
        body["optional"]["request"] = request
    return call("POST", "/act/make", body)


def rq_items(a):
    return {i["code"]: i for i in (a.get("request_check") or {}).get("items") or []}


def check_branch_make(sid_text, b_text, sid_scan):
    print("37в. Сверка запроса филиала с расчётом акта (продукт 0832, оба образца)")
    fresh()
    with db.tx() as con:
        ref = db.load_reference(con)
        mr = min_rate(ref, "0832")
        cls = [r["class_code"] for r in db.rows(con, "SELECT class_code FROM product_classes WHERE product_code='0832'")]
    ok("0832 в справочнике копии базы: класс 8, минимальная ставка компании 0,08 %",
       cls == ["8"] and mr["company"] == 0.08 and mr["floor"] == 0.08, (cls, mr))
    req1 = br_request(b_text)
    st, a = br_make(sid_text, 1, req1)
    ok("акт по образцу 1 сформирован", st == 200 and a.get("ok"), (st, a))
    lvl = a["risk"]["level"]
    adj = ae.DEFAULT_SETTINGS["adj_pct"][lvl]
    applied = round(max(0.08 * (1 + adj / 100), 0.08), 4)
    prem = round(BR_S1 * applied / 100 * 1108 / 365)
    BR_REPORT["образец 1"] = {"level": lvl, "rate": applied, "premium": a["premium"]["amount"], "days": 1108}
    ok("многолетний срок: премия акта на 1 108 дн. по годовой ставке",
       a["premium"]["term_days"] == 1108 and a["rate"]["applied_pct"] == applied
       and a["premium"]["amount"] == prem, (a["premium"], a["rate"]["applied_pct"], prem))
    it = rq_items(a)
    rc = a["request_check"]
    ok("request_check: доступна, источник — файл запроса (сервер сверил со своей загрузкой, правок нет)",
       rc["available"] and rc["source"] == "document" and rc["source_kind"] == "document"
       and rc["source_label"] == "из документа — из файла запроса (разбор текста)" and rc["edits"]["count"] == 0,
       (rc.get("source"), rc.get("source_label")))
    ok("тариф 0,05 ниже минимума 0,08 — below_min",
       it["tariff_min"]["verdict"] == "below_min" and it["tariff_min"]["requested"] == 0.05
       and it["tariff_min"]["calculated"] == 0.08 and "ниже минимального" in it["tariff_min"]["text"], it["tariff_min"])
    ok("тариф ниже ставки акта — differs", it["tariff_act"]["verdict"] == "differs"
       and it["tariff_act"]["calculated"] == applied, it["tariff_act"])
    ok("образец 1: премия по тарифу запроса сходится (123 322 000 против 123 321 917,81)",
       it["premium_request"]["verdict"] == "ok" and it["premium_request"]["calculated"] == 123_321_917.81
       and abs(it["premium_request"]["diff"] - 82.19) < 0.01, it["premium_request"])
    ok("премия акта — справочно, в решение не идёт", it["premium_act"]["reference"]
       and it["premium_act"]["calculated"] == prem and it["premium_act"]["verdict"] == "differs", it["premium_act"])
    ok("сумма к стоимости — ссылка на раздел 3, без дубля", it["sum_value"]["reference"]
       and "раздел" in it["sum_value"]["text"] and it["sum_value"]["verdict"] == "ok", it["sum_value"])
    ok("франшиза: в запросе не применяется, акт не требует — ok",
       it["franchise"]["verdict"] == "ok", it["franchise"])
    ok("срок: 1 108 дн. взят из запроса — ok", it["term"]["verdict"] == "ok" and it["term"]["requested"] == 1108
       and "взят из запроса" in it["term"]["text"], it["term"])
    ok("итог сверки — ниже минимума", rc["summary"]["verdict"] == "below_min" and rc["summary"]["below_min"] == 1,
       rc["summary"])
    ok("как сверено: дни включительно, формула, допуск 1 000 сум",
       any("29.09.2026" in h and "1108" in h and "включены" in h for h in rc["how"])
       and any("× 0,05 %" in nb(h) for h in rc["how"]) and any("1 000 сум" in nb(h) for h in rc["how"]), rc["how"])
    ok("решение не «принять без оговорок»; в проверках — тариф ниже минимума",
       a["decision"]["code"] != "accept" and any("ниже минимального" in c for c in a["decision"]["checks"]),
       a["decision"])
    s4 = a["sections"][3]
    titles = [li["title"] for li in s4["lists"]]
    ok("раздел 4: подраздел «Сверка с запросом филиала»", "Сверка с запросом филиала" in titles, titles)
    aid1 = a["id"]

    # образец 2 со скана (запрос берётся из своей загрузки, если экран его не прислал)
    st, a2 = br_make(sid_scan, 1, None)
    ok("без optional.request — запрос из своей загрузки (source = session)",
       st == 200 and a2["request_check"]["source"] == "session", a2.get("request_check", {}).get("source"))
    req2 = {"tariff_pct": 0.05, "premium": "122 589 000,00", "franchise": {"applied": False, "text": "Қўлланилинмайди"},
            "term_from": "07.09.2026", "term_to": "2031-11-07", "source": "photo"}
    st, a2 = br_make(None, 2, req2)
    it2 = rq_items(a2)
    lvl2 = a2["risk"]["level"]
    applied2 = round(0.08 * (1 + ae.DEFAULT_SETTINGS["adj_pct"][lvl2] / 100), 4)
    BR_REPORT["образец 2"] = {"level": lvl2, "rate": applied2, "premium": a2["premium"]["amount"], "days": 1888}
    ok("образец 2: срок 1 888 дн., премия акта на весь срок",
       a2["premium"]["term_days"] == 1888 and a2["premium"]["amount"] == round(BR_S2 * applied2 / 100 * 1888 / 365),
       a2["premium"])
    ok("образец 2: премия расходится на 3 869,55 сум (differs)",
       it2["premium_request"]["verdict"] == "differs" and it2["premium_request"]["calculated"] == 122_585_130.45
       and abs(it2["premium_request"]["diff"] - 3869.55) < 0.01 and int(it2["premium_request"]["diff"]) == 3869,
       it2["premium_request"])
    ok("образец 2: в проверках андеррайтера — расхождение премии",
       any("не сходится" in c and "3 870" in nb(c) for c in a2["decision"]["checks"]), a2["decision"]["checks"])
    ok("образец 2: тариф ниже минимума", it2["tariff_min"]["verdict"] == "below_min")
    BR_REPORT["сверка 1"] = {k: (v["verdict"], v["requested"], v["calculated"], v["diff"]) for k, v in it.items()}
    BR_REPORT["сверка 2"] = {k: (v["verdict"], v["requested"], v["calculated"], v["diff"]) for k, v in it2.items()}

    # тариф не ниже минимума и не ниже ставки акта, премия сходится — расхождений сверки нет
    good = {"tariff_pct": 0.2, "premium": round(BR_S1 * 0.2 / 100 * 1108 / 365), "franchise": {"applied": False},
            "term_days": 1108}
    st, a3 = br_make(None, 1, good)
    it3 = rq_items(a3)
    ok("тариф 0,2 — не ниже минимума и ставки акта; премия сходится",
       it3["tariff_min"]["verdict"] == "ok" and it3["tariff_act"]["verdict"] == "ok"
       and it3["premium_request"]["verdict"] == "ok" and a3["request_check"]["summary"]["verdict"] == "ok",
       {k: v["verdict"] for k, v in it3.items()})
    ok("без расхождений сверка ничего не добавляет в проверки",
       not any(c.startswith(("Тариф в запросе", "Премия в запросе")) for c in a3["decision"]["checks"]),
       a3["decision"]["checks"])
    # срок сотрудника расходится с запросом
    st, a4 = br_make(None, 1, dict(good), {"term_days": 365})
    ok("срок сотрудника 365 против 1 108 в запросе — differs",
       rq_items(a4)["term"]["verdict"] == "differs" and a4["premium"]["term_days"] == 365, rq_items(a4)["term"])
    # франшиза: в запросе не применяется, а акт предлагает (клиент просит снизить премию)
    st, a5 = br_make(None, 1, dict(good), {"want_lower_premium": True})
    f5 = rq_items(a5)["franchise"]
    ok("франшиза: в запросе нет, акт предлагает — differs", f5["verdict"] == "differs"
       and "предлагает франшизу" in f5["text"], f5)
    fr5 = a5["franchise"]
    if fr5.get("premium_after") is not None and fr5.get("rate_after") is not None:
        ok("франшиза на многолетнем сроке: премия с франшизой — на все 1 108 дн.",
           fr5["premium_after"] == round(BR_S1 * fr5["rate_after"] / 100 * 1108 / 365), fr5)
    ms = a5["measures_summary"]
    ok("мероприятия считаются от премии акта на весь срок", ms["premium_before"] in (None, a5["premium"]["amount"]),
       ms)
    sc = a2["scenarios"]
    ok("сценарии PML/EML/MFL для срока 1 888 дн. (> 60 мес.) — посчитаны с пометкой о сроке",
       sc["available"] and any("1888" in x["text"] for x in sc["assumptions"]), sc.get("assumptions"))
    # проверка ввода
    for bad, key in (({"tariff_pct": 0}, "тариф 0"), ({"tariff_pct": "abc"}, "тариф не число"),
                     ({"premium": -5}, "премия < 0"), ({"term_from": "2029-10-10", "term_to": "2026-09-29"}, "даты наоборот"),
                     ({"term_from": "29.09.2026", "term_to": "10.10.2029", "term_days": 1107}, "дни не по датам"),
                     ({"franchise": {"applied": "нет"}}, "франшиза без applied"), ("строка", "не объект")):
        st, e = br_make(None, 1, bad)
        ok(f"optional.request проверяется: {key}", st == 422 and "request" in (e.get("errors") or {}), (st, e))
    st, a6 = br_make(None, 1, {"franchise": "не применяется", "term_from": "29.09.2026", "term_to": "10.10.2029"})
    ok("франшиза текстом и даты ДД.ММ.ГГГГ принимаются", st == 200 and rq_items(a6)["term"]["requested"] == 1108
       and rq_items(a6)["tariff_min"]["verdict"] == "missing", (st, a6.get("request_check")))
    return aid1, a2["id"]


BR_REPORT = {}


def check_branch_langs_files(aid):
    print("37г. Сверка: три языка, Word и PDF")
    for lang, title, word in (("ru", "Сверка с запросом филиала", "ниже минимального"),
                              ("uz", "Filial soʻrovi bilan solishtirish", "eng kam stavka"),
                              ("en", "Check against the branch request", "below the tariff-policy minimum")):
        st, a = call("GET", f"/act/{aid}", params={"lang": lang})
        s4 = a["sections"][3]
        lines = [x for li in s4["lists"] if li["title"] == title for x in li["items"]]
        ok(f"{lang}: подраздел сверки и текст строки", bool(lines) and any(word in x for x in lines), lines[:3])
        rc = a["request_check"]
        ok(f"{lang}: request_check на языке акта", rc["items"][0]["label"] and rc["summary"]["text"]
           and all(i["verdict_label"] for i in rc["items"]), rc["items"][0])
        if lang != "ru":
            txt = " ".join(lines + [c for c in a["decision"]["checks"]] + rc["how"])
            ok(f"{lang}: в сверке нет кириллицы", not re.search(r"[А-Яа-яЁё]", txt), txt[:300])
        st, blob, h = call("GET", f"/act/{aid}.docx", params={"lang": lang}, raw=True)
        xml = zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8")
        plain = re.sub(r"<[^>]+>", "", xml)
        ok(f"{lang}: сверка в DOCX", title in plain and word in plain, plain[-300:])
        st, blob, h = call("GET", f"/act/{aid}.pdf", params={"lang": lang}, raw=True)
        text = pdf_text(pymupdf.open(stream=blob, filetype="pdf"))
        ok(f"{lang}: сверка в PDF", title.replace("ʻ", "'") in text.replace("ʻ", "'") and word in text, text[-300:])
    # старый акт без сверки — блок недоступен, ничего не падает
    with db.tx() as con:
        row = db.rows(con, "SELECT act_json FROM acts WHERE id=?", aid)[0]
        stored = _json.loads(row["act_json"])
    stored["data"].pop("request_check", None)
    stored["data"].pop("request", None)
    D = stored["data"]
    out = act.render(D, "ru", stored["meta"])
    ok("старый акт без сверки: request_check.available = false", out["request_check"]["available"] is False
       and "Сверка с запросом филиала" not in [li["title"] for li in out["sections"][3]["lists"]])


# ================================================================================================
#  38. Договор страхования: чтение (файл, скан, модель по тексту) и сверка с расчётом акта (30.09.2026)
# ================================================================================================

CT_UZC = [
    "МОЛ-МУЛКНИ СУҒУРТА ҚИЛИШ ШАРТНОМАСИ № 45-ИМ/2026",
    "Тошкент шаҳри                                   2026 йил «1» октябрь",
    "\"INSON\" АЖ, бундан буён «Суғурталовчи» деб юритилади, директор Каримов Алишер Анварович номидан, "
    "бир томондан, ва \"ALFA TEXTILE\" МЧЖ, бундан буён «Суғурта қилдирувчи» деб юритилади, директор "
    "Тошпўлатов Бахтиёр Равшанович номидан, иккинчи томондан, мазкур шартномани туздилар.",
    "1. ШАРТНОМА ПРЕДМЕТИ",
    "1.1. Суғурта объекти: нотурар бино — омбор, умумий майдони 2 400 кв.м, кадастр рақами 10:09:05:01:02:0033.",
    "1.2. Объект манзили: Тошкент вилояти, Чирчиқ шаҳри, Саноат кўчаси, 7.",
    "2. СУҒУРТА СУММАСИ ВА МУКОФОТИ",
    "2.1. Суғурта қиймати: 12 000 000 000 (ўн икки миллиард) сўм.",
    "2.2. Суғурта суммаси: 12 000 000 000 (ўн икки миллиард) сўм.",
    "2.3. Суғурта тарифи: йиллик 0,1 %.",
    "2.4. Суғурта мукофоти: 36 032 877 (ўттиз олти миллион ўттиз икки минг саккиз юз етмиш етти) сўм.",
    "2.5. Суғурта мукофоти бўлиб-бўлиб тўланади: биринчи тўлов 18 016 438 сўм — 2026 йил 10 октябргача; "
    "иккинчи тўлов 18 016 439 сўм — 2027 йил 10 октябргача.",
    "3. СУҒУРТА МУДДАТИ",
    "3.1. Суғурта муддати: 2026 йил 1 октябрдан 2029 йил 30 сентябргача.",
    "4. ФРАНШИЗА",
    "4.1. Франшиза: қўлланилмайди.",
    "5. СУҒУРТА ХАВФЛАРИ",
    "5.1. Суғурта хавфлари: ёнғин, чақмоқ уриши, портлаш, сув босиши, табиий офатлар, учинчи шахсларнинг "
    "ғайриқонуний ҳаракатлари.",
    "6. ИСТИСНОЛАР",
    "6.1. Қуйидагилар суғурта ҳодисаси ҳисобланмайди: уруш ҳаракатлари, ядро портлаши, суғурта "
    "қилдирувчининг қасддан қилган ҳаракатлари.",
    "7. ЯКУНИЙ ҚОИДАЛАР",
    "7.1. Суғурта қилдирувчи суғурта ҳодисаси юз берганлиги ҳақида 3 (уч) иш куни ичида хабар беради.",
]
CT_UZL = [
    "MOL-MULKNI SUGʻURTA QILISH SHARTNOMASI № 46-IM/2026",
    "Toshkent shahri                                   2026-yil 1-oktyabr",
    "\"INSON\" AJ, bundan buyon «Sugʻurtalovchi» deb yuritiladi, direktor Karimov Alisher Anvarovich nomidan, "
    "bir tomondan, va \"ALFA TEXTILE\" MCHJ, bundan buyon «Sugʻurta qildiruvchi» deb yuritiladi, direktor "
    "Toshpoʻlatov Baxtiyor Ravshanovich nomidan, ikkinchi tomondan, mazkur shartnomani tuzdilar.",
    "1. SHARTNOMA PREDMETI",
    "1.1. Sugʻurta obyekti: noturar bino — ombor, umumiy maydoni 2 400 kv.m, kadastr raqami 10:09:05:01:02:0033.",
    "1.2. Obyekt manzili: Toshkent viloyati, Chirchiq shahri, Sanoat koʻchasi, 7.",
    "2. SUGʻURTA SUMMASI VA MUKOFOTI",
    "2.1. Sugʻurta qiymati: 12 000 000 000 (oʻn ikki milliard) soʻm.",
    "2.2. Sugʻurta summasi: 12 000 000 000 (oʻn ikki milliard) soʻm.",
    "2.3. Sugʻurta tarifi: yillik 0,1 %.",
    "2.4. Sugʻurta mukofoti: 36 032 877 (oʻttiz olti million oʻttiz ikki ming sakkiz yuz yetmish yetti) soʻm.",
    "2.5. Sugʻurta mukofoti boʻlib-boʻlib toʻlanadi: birinchi toʻlov 18 016 438 soʻm — 2026-yil 10-oktyabrgacha; "
    "ikkinchi toʻlov 18 016 439 soʻm — 2027-yil 10-oktyabrgacha.",
    "3. SUGʻURTA MUDDATI",
    "3.1. Sugʻurta muddati: 2026-yil 1-oktyabrdan 2029-yil 30-sentyabrgacha.",
    "4. FRANSHIZA",
    "4.1. Franshiza: qoʻllanilmaydi.",
    "5. SUGʻURTA XAVFLARI",
    "5.1. Sugʻurta xavflari: yongʻin, chaqmoq urishi, portlash, suv bosishi, tabiiy ofatlar, uchinchi shaxslarning "
    "gʻayriqonuniy harakatlari.",
    "6. ISTISNOLAR",
    "6.1. Quyidagilar sugʻurta hodisasi hisoblanmaydi: urush harakatlari, yadro portlashi, sugʻurta "
    "qildiruvchining qasddan qilgan harakatlari.",
    "7. YAKUNIY QOIDALAR",
    "7.1. Sugʻurta qildiruvchi sugʻurta hodisasi yuz berganligi haqida 3 (uch) ish kuni ichida xabar beradi.",
]
CT_RU = [
    "ДОГОВОР СТРАХОВАНИЯ ИМУЩЕСТВА ЮРИДИЧЕСКИХ ЛИЦ № 77/2026",
    "г. Ташкент                                              «1» октября 2026 г.",
    "Акционерное общество «INSON», именуемое в дальнейшем «Страховщик», в лице директора Иванова Ивана "
    "Ивановича, действующего на основании Устава, с одной стороны, и ООО «Ромашка Трейд», именуемое в "
    "дальнейшем «Страхователь», в лице генерального директора Петрова Петра Петровича, с другой стороны, "
    "заключили настоящий договор о нижеследующем:",
    "1. ПРЕДМЕТ ДОГОВОРА",
    "1.1. Объектом страхования являются имущественные интересы Страхователя, связанные с владением нежилым "
    "зданием склада готовой продукции, общая площадь 1 500 кв.м, кадастровый номер 10:00:00:00:00:0077.",
    "1.2. Адрес места страхования: Самаркандская область, г. Самарканд, ул. Навои, 15.",
    "2. СТРАХОВАЯ СУММА. СТРАХОВАЯ ПРЕМИЯ",
    "2.1. Страховая стоимость имущества: 5 000 000 000 (пять миллиардов) сум.",
    "2.2. Страховая сумма по настоящему договору составляет 5 000 000 000 (пять миллиардов) сум.",
    "2.3. Страховой тариф: 0,2 % годовых.",
    "2.4. Страховая премия составляет 10 000 000 (десять миллионов) сум и уплачивается единовременно "
    "до 10.10.2026.",
    "3. СРОК ДЕЙСТВИЯ ДОГОВОРА",
    "3.1. Срок страхования: с 01.10.2026 по 30.09.2027.",
    "4. ФРАНШИЗА",
    "4.1. Франшиза безусловная, 1 % от страховой суммы по каждому страховому случаю.",
    "5. СТРАХОВЫЕ РИСКИ",
    "5.1. Страховыми случаями являются гибель или повреждение имущества в результате:",
    "5.1.1. пожара, удара молнии, взрыва газа;",
    "5.1.2. стихийных бедствий: землетрясения, наводнения, бури;",
    "5.1.3. кражи со взломом, грабежа.",
    "6. ИСКЛЮЧЕНИЯ",
    "6.1. Не являются страховыми случаями события, произошедшие вследствие:",
    "6.1.1. военных действий, террористических актов;",
    "6.1.2. ядерного взрыва, радиации;",
    "6.1.3. умышленных действий Страхователя; износа и коррозии.",
    "7. ПРОЧИЕ УСЛОВИЯ",
    "7.1. Страхователь обязан уведомить Страховщика о наступлении страхового случая в течение 2 (двух) "
    "рабочих дней.",
    "7.2. Территория страхования: Республика Узбекистан.",
]
CT_EN = [
    "PROPERTY INSURANCE POLICY No. INS-2026/0045",
    "Tashkent, October 1, 2026",
    "INSON JSC, hereinafter the Insurer, and ALFA TEXTILE LLC, hereinafter the Policyholder, have agreed:",
    "1. Insured property: warehouse building, total area 2 400 sq.m, cadastral number 10:09:05:01:02:0033.",
    "2. Sum insured: UZS 12,000,000,000.",
    "3. Insured value: UZS 12,000,000,000.",
    "4. Premium rate: 0.1% per annum.",
    "5. Insurance premium: UZS 12,000,000, payable in one payment by October 10, 2026.",
    "6. Period of insurance: from October 1, 2026 to September 30, 2027.",
    "7. Deductible: not applicable.",
    "8. Insured perils: fire, lightning, explosion, earthquake, theft.",
    "9. Exclusions: war, terrorism, nuclear risks, wear and tear.",
]
CT_S = 12_000_000_000.0
CT_PREMIUM = 36_032_877.0
CT_PD = ("Каримов", "Тошпўлатов", "Бахтиёр", "Karimov", "Toshpoʻlatov", "Baxtiyor", "Иванов", "Петров", "Петра")


def pdf_lines(lines, per_page=46, width=92) -> bytes:
    """PDF с текстовым слоем: строки переносятся по ширине страницы (как в настоящем договоре)."""
    import textwrap
    doc = pymupdf.open()
    font = act._fonts()[0]
    wrapped = [w for ln in lines for w in (textwrap.wrap(ln, width) or [""])]
    for k in range(0, len(wrapped), per_page):
        page = doc.new_page()
        tw = pymupdf.TextWriter(page.rect)
        for j, ln in enumerate(wrapped[k:k + per_page]):
            if ln:
                tw.append((40, 50 + 16 * j), ln, font=font, fontsize=9)
        tw.write_text(page)
    return doc.tobytes()


def ctb(b):
    return b.get("contract") or {}


def docx_mixed(paras, tables) -> bytes:
    """DOCX: абзацы, затем таблицы (строки — списки ячеек)."""
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    para = lambda t: f"<w:p><w:r><w:t xml:space=\"preserve\">{_x(t)}</w:t></w:r></w:p>"   # noqa: E731
    body = "".join(para(x) for x in paras)
    for rows in tables:
        body += "<w:tbl>" + "".join("<w:tr>" + "".join(f"<w:tc>{para(c)}</w:tc>" for c in r) + "</w:tr>"
                                    for r in rows) + "</w:tbl>"
    xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{ns}"><w:body>{body}</w:body></w:document>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def check_ct_uz(tag, f, latin=False):
    ok(f"{tag}: номер, дата, место",
       f["contract_no"] == ("46-IM/2026" if latin else "45-ИМ/2026") and f["contract_date"] == "2026-10-01"
       and f["place"] == ("Toshkent shahri" if latin else "Тошкент шаҳри"), (f["contract_no"], f["contract_date"], f["place"]))
    ok(f"{tag}: стороны — юрлица из преамбулы",
       f["insurer"] == {"kind": "legal", "name": '"INSON" AJ' if latin else '"INSON" АЖ'}
       and f["policyholder"] == {"kind": "legal", "name": '"ALFA TEXTILE" MCHJ' if latin else '"ALFA TEXTILE" МЧЖ'},
       (f["insurer"], f["policyholder"]))
    ok(f"{tag}: объект — здание, кадастр, площадь, регион",
       f["class_hint"] == "building" and f["object_kind"] == "warehouse" and f["cadastre_no"] == "10:09:05:01:02:0033"
       and f["areas"]["total_m2"] == 2400.0 and f["region"] == "Ташкентская область"
       and "7" not in (f["address"] or "x"), (f["object_description"], f["address"], f["region"]))
    ok(f"{tag}: суммы, тариф, премия",
       f["object_value"] == CT_S and f["sum_insured"] == CT_S and f["tariff_pct"] == 0.1 and f["premium"] == CT_PREMIUM
       and f["currency"] == "UZS", (f["object_value"], f["sum_insured"], f["tariff_pct"], f["premium"]))
    ok(f"{tag}: рассрочка, график из 2 платежей",
       f["payment_mode"] == "installments" and f["payments"] == [{"date": "2026-10-10", "amount": 18016438.0},
                                                                   {"date": "2027-10-10", "amount": 18016439.0}],
       (f["payment_mode"], f["payments"]))
    ok(f"{tag}: срок 01.10.2026–30.09.2029 — 1 096 дн. включительно",
       f["term_from"] == "2026-10-01" and f["term_to"] == "2029-09-30" and f["term_days"] == 1096,
       (f["term_from"], f["term_to"], f["term_days"]))
    ok(f"{tag}: франшиза не применяется", (f["franchise"] or {}).get("applied") is False, f["franchise"])
    ok(f"{tag}: риски и исключения — короткими кодами",
       {"fire", "lightning", "explosion", "water", "natural"} <= {x["code"] for x in f["covered_risks"]}
       and [x["code"] for x in f["exclusions"]] == ["war", "nuclear", "intent"],
       (f["covered_risks"], f["exclusions"]))
    ok(f"{tag}: срок уведомления о страховом случае", "3" in (f["notice"] or ""), f["notice"])


def check_contract_text():
    print("38а. Договор страхования: файл с текстом (DOCX, PDF) — без модели, три языка")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    CALLS.clear()
    st, b = upload([("contract.docx", DOCX_MIME, CONTRACT.read_bytes())], {"lang": "ru", "product_code": "0808"})
    c = ctb(b)
    f = c.get("fields") or {}
    ok("учебный договор узнан: блок contract, источник — документ, модель не вызывалась",
       st == 200 and c.get("detected") and c["source"] == "document" and c["kind_label"] == "договор страхования"
       and b["files"][0]["document_kind"] == "договор страхования" and not CALLS, (st, c.get("source"), len(CALLS)))
    ok("учебный договор: номер, дата, место, страхователь-юрлицо",
       f.get("contract_no") == "15/2026" and f["contract_date"] == "2026-09-21" and f["place"] == "г. Ташкент"
       and f["policyholder"] == {"kind": "legal", "name": "ООО «Тестовый склад»"}, f)
    ok("учебный договор: объект, адрес без улицы, регион, конструкция, год постройки",
       f["object_description"] == "склад готовой продукции" and f["object_kind"] == "warehouse"
       and f["address"] == "г. Ташкент, Юнусабадский район" and f["region"] == "город Ташкент"
       and f["construction"] == "кирпич" and f["year_built"] == "2012", f)
    ok("учебный договор: сумма 4,2 млрд, стоимость 5 млрд, срок 12 мес. = 365 дн.",
       f["sum_insured"] == 4.2e9 and f["object_value"] == 5e9 and f["term_days"] == 365 and f["term_from"] is None, f)
    ok("учебный договор: чего нет — null (тариф, премия, франшиза, риски)",
       f["tariff_pct"] is None and f["premium"] is None and f["franchise"] is None and f["covered_risks"] == [])
    fk = [x["code"] for x in c["found"]]
    mk = [x["code"] for x in c["missing"]]
    ok("found / missing с подписями",
       fk == ["contract_no", "contract_date", "policyholder", "object", "sum_insured", "term"]
       and mk == ["tariff_pct", "premium", "franchise", "covered_risks"]
       and c["missing"][0]["label"] == "Тариф", (fk, mk))
    ess = {e["code"]: e["present"] for e in c["essentials"]}
    ok("существенные условия (ГК ст. 929): нет премии и страхового случая",
       ess == {"object": True, "insured_event": False, "sum_insured": True, "premium": False, "term": True}
       and c["legal_ref"] == "ГК РУз, ст. 929" and any("ст. 929" in n for n in c["notes"]), (ess, c["notes"]))
    rq = c["request"]
    ok("contract.request — в формате request (+ условия договора)",
       rq["sum_insured"] == 4.2e9 and rq["term_days"] == 365 and rq["tariff_pct"] is None and rq["source"] == "document"
       and rq["contract_no"] == "15/2026" and rq["contract_date"] == "2026-09-21" and rq["covered_risks"] == []
       and {"tariff_pct", "premium", "franchise", "term_from", "term_to", "term_days", "sum_insured",
            "product_code", "source"} <= set(rq), rq)
    pf = b["prefill"]
    ok("prefill: сумма, стоимость, срок, регион",
       pf["sum_insured"]["value"] == 4.2e9 and pf["object_value"]["value"] == 5e9 and pf["term_days"]["value"] == 365
       and pf["region"]["code"] == "tashkent_city", pf)
    dump = _json.dumps(b, ensure_ascii=False)
    ok("ИНН организации в блок договора не попал", "301234567" not in dump)
    with db.tx() as con:
        audit = db.rows(con, "SELECT detail FROM audit WHERE entity=?", "act_upload:" + b["session"])
    det = _json.loads(audit[0]["detail"])
    ok("журнал: только признак, источник и счётчики договора",
       det.get("contract") is True and det.get("contract_source") == "document" and det.get("contract_found") == 6
       and "Тестовый" not in audit[0]["detail"] and "15/2026" not in audit[0]["detail"], det)
    CT_REPORT["учебный договор"] = {k: f[k] for k in ("contract_no", "contract_date", "place", "policyholder",
                                                    "object_description", "address", "region", "construction",
                                                    "year_built", "sum_insured", "object_value", "term_days")}

    # узбекская кириллица и латиница: DOCX и PDF с текстом
    for tag, lines, latin in (("узб. кириллица DOCX", CT_UZC, False), ("узб. латиница DOCX", CT_UZL, True)):
        CALLS.clear()
        st, b = upload([("shartnoma.docx", DOCX_MIME, docx_bytes([_x(x) for x in lines]))], {"lang": "uz"})
        c = ctb(b)
        ok(f"{tag}: узнан, модель не вызывалась", st == 200 and c.get("detected") and not CALLS
           and c["kind_label"] == "sugʻurta shartnomasi", (st, len(CALLS)))
        check_ct_uz(tag, c["fields"], latin)
        dump = _json.dumps(b, ensure_ascii=False)
        ok(f"{tag}: ФИО директоров не извлечены", not any(x in dump for x in CT_PD), [x for x in CT_PD if x in dump])
        ok(f"{tag}: все ключевые поля найдены, существенные условия есть",
           not c["missing"] and all(e["present"] for e in c["essentials"]), c["missing"])
    for tag, lines, latin in (("узб. кириллица PDF", CT_UZC, False), ("узб. латиница PDF", CT_UZL, True)):
        CALLS.clear()
        st, b = upload([("shartnoma.pdf", "application/pdf", pdf_lines(lines))], {"lang": "uz"})
        c = ctb(b)
        ok(f"{tag}: узнан по тексту (строки перенесены), модель не вызывалась",
           st == 200 and c.get("detected") and not CALLS and b["files"][0]["parsed"], (st, len(CALLS)))
        if c:
            check_ct_uz(tag, c["fields"], latin)

    # русский и английский
    st, b = upload([("dogovor.docx", DOCX_MIME, docx_bytes([_x(x) for x in CT_RU]))], {"lang": "ru"})
    f = ctb(b).get("fields") or {}
    ok("русский: стороны из преамбулы (без представителей), объект, адрес без улицы",
       f.get("insurer", {}).get("name") == "Акционерное общество «INSON»"
       and f["policyholder"]["name"] == "ООО «Ромашка Трейд»" and f["cadastre_no"] == "10:00:00:00:00:0077"
       and f["address"] == "Самаркандская область, г. Самарканд" and f["region"] == "Самаркандская область"
       and not any(x in _json.dumps(b, ensure_ascii=False) for x in ("Иванов", "Петров", "Навои")), f)
    ok("русский: премия 10 млн единовременно до 10.10.2026, тариф 0,2 %, срок 365 дн.",
       f["premium"] == 10_000_000 and f["payment_mode"] == "single"
       and f["payments"] == [{"date": "2026-10-10", "amount": 10_000_000.0}] and f["tariff_pct"] == 0.2
       and f["term_days"] == 365, (f["premium"], f["payments"], f["tariff_pct"], f["term_days"]))
    ok("русский: франшиза безусловная 1 %, риски, исключения, территория, уведомление",
       f["franchise"] == {"applied": True, "text": "безусловная, 1 % от страховой суммы по каждому страховому случаю.",
                          "pct": 1.0, "amount": None, "type": "unconditional", "risk": None}
       and {"fire", "natural", "earthquake", "theft"} <= {x["code"] for x in f["covered_risks"]}
       and {"war", "terrorism", "nuclear", "intent", "wear"} == {x["code"] for x in f["exclusions"]}
       and f["territory"] == "Республика Узбекистан." and "2" in f["notice"], f)
    st, b = upload([("policy.docx", DOCX_MIME, docx_bytes(CT_EN))], {"lang": "en"})
    f = ctb(b).get("fields") or {}
    ok("английский: номер, дата «October 1, 2026», стороны, суммы, тариф, срок, франшиза",
       f.get("contract_no") == "INS-2026/0045" and f["contract_date"] == "2026-10-01" and f["place"] == "Tashkent"
       and f["insurer"]["name"] == "INSON JSC" and f["policyholder"]["name"] == "ALFA TEXTILE LLC"
       and f["sum_insured"] == CT_S and f["premium"] == 12e6 and f["tariff_pct"] == 0.1 and f["term_days"] == 365
       and f["franchise"]["applied"] is False and ctb(b)["kind_label"] == "insurance contract", f)

    # не договор: счёт, письмо; запрос филиала — не договор
    inv = ["СЧЁТ НА ОПЛАТУ № 45 от 01.10.2026", "Плательщик: ООО «Ромашка Трейд»",
           "Назначение платежа: страховая премия по договору страхования № 77/2026",
           "Сумма к оплате: 10 000 000 сум", "Страховая сумма по договору: 5 000 000 000 сум"]
    letter = ["Директору ООО «Ромашка Трейд»", "Просим заключить договор страхования имущества на следующих условиях.",
              "Страховая сумма 5 000 000 000 сум, страховая премия 10 000 000 сум.", "С уважением, отдел продаж"]
    for tag, lines in (("счёт на оплату", inv), ("письмо", letter)):
        st, b = upload([("x.docx", DOCX_MIME, docx_bytes(lines))], {"lang": "ru"})
        ok(f"{tag} — не договор", st == 200 and b.get("contract") is None, ctb(b).get("fields"))
    st, b = upload([("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1))], {"lang": "ru"})
    ok("запрос филиала — не договор (блок branch_request есть, contract нет)",
       b.get("branch_request") and b.get("contract") is None and b.get("cross_check") is None)

    # физлицо-страхователь: признак без имени — нигде
    lines = list(CT_RU)
    lines[2] = ("Акционерное общество «INSON», именуемое в дальнейшем «Страховщик», и гражданин Иванов Иван "
                "Иванович, именуемый в дальнейшем «Страхователь», заключили настоящий договор:")
    lines.insert(3, "Страхователь: Иванов Иван Иванович, паспорт AA1234567")
    st, b = upload([("fiz.docx", DOCX_MIME, docx_bytes(lines))], {"lang": "ru"})
    c = ctb(b)
    dump = _json.dumps(b, ensure_ascii=False)
    with db.tx() as con:
        saved = db.rows(con, "SELECT result_json FROM act_uploads WHERE id=?", b["session"])[0]["result_json"]
        journal = _json.dumps(db.rows(con, "SELECT detail FROM audit WHERE entity=?", "act_upload:" + b["session"]),
                              ensure_ascii=False)
    ok("физлицо-страхователь: kind = individual, имени и паспорта нет в ответе, базе и журнале",
       c["fields"]["policyholder"] == {"kind": "individual", "name": None}
       and not any(x in dump + saved + journal for x in ("Иванов", "AA1234567")), c["fields"]["policyholder"])
    ok("пометка «физическое лицо — данные не извлекаются»", any("физическое лицо" in n for n in c["notes"]), c["notes"])

    # транспорт: марка, модель, год, VIN, госномер; франшиза суммой по риску; срок словами на скане — в тесте скана
    kasko = ["ДОГОВОР СТРАХОВАНИЯ ТРАНСПОРТНОГО СРЕДСТВА (КАСКО) № К-12/2026", "г. Ташкент, 5 октября 2026 г.",
             "Страхователь: ООО «Автолизинг Плюс»", "Объект страхования: легковой автомобиль", "Марка: Chevrolet",
             "Модель: Cobalt", "Год выпуска: 2022", "VIN: XWBJA69V9LA123456", "Государственный номер: 01 A 123 BC",
             "Страховая сумма: 150 000 000 сум", "Страховая премия: 4 500 000 сум",
             "Срок страхования: с 05.10.2026 по 04.10.2027", "Франшиза: безусловная 500 000 сум по риску «ущерб»"]
    st, b = upload([("kasko.docx", DOCX_MIME, docx_bytes(kasko))], {"lang": "ru", "product_code": "0318"})
    f = ctb(b).get("fields") or {}
    rec = {(r["key"], r["value"]) for r in b["recognized"]}
    ok("транспорт: марка, модель, год, VIN, госномер; франшиза 500 000 сум по риску «ущерб»",
       f.get("class_hint") == "vehicle" and f["brand"] == "Chevrolet" and f["model"] == "Cobalt" and f["year"] == "2022"
       and f["vin"] == "XWBJA69V9LA123456" and f["reg_no"] == "01 A 123 BC"
       and f["franchise"] == {"applied": True, "text": "безусловная 500 000 сум по риску «ущерб»", "pct": None,
                              "amount": 500000.0, "type": "unconditional", "risk": "ущерб"}
       and {("serial_no", "XWBJA69V9LA123456"), ("reg_no", "01 A 123 BC"), ("brand", "Chevrolet")} <= rec, f)

    # перечень имущества и график платежей таблицами
    paras = [x for x in CT_RU if not x.startswith("2.4.")] + ["Страховая премия: 8 000 000 сум, уплачивается в рассрочку "
                                                              "по графику."]
    items_t = [["№", "Наименование имущества", "Страховая сумма, сум"], ["1", "Здание склада", "3 000 000 000"],
               ["2", "Стеллажи и погрузчики", "1 500 000 000"], ["", "Итого", "4 500 000 000"]]
    pays_t = [["№", "Дата платежа", "Сумма платежа, сум"], ["1", "10.10.2026", "4 000 000"], ["2", "10.04.2027", "4 000 000"]]
    st, b = upload([("tables.docx", DOCX_MIME, docx_mixed(paras, [items_t, pays_t]))], {"lang": "ru"})
    f = ctb(b).get("fields") or {}
    ok("таблицы: страховая сумма по частям (итог отдельно) и график платежей",
       f.get("items") == [{"name": "Здание склада", "sum": 3e9}, {"name": "Стеллажи и погрузчики", "sum": 1.5e9}]
       and f["items_total"] == 4.5e9 and f["payments"] == [{"date": "2026-10-10", "amount": 4e6},
                                                            {"date": "2027-04-10", "amount": 4e6}]
       and f["payment_mode"] == "installments" and f["premium"] == 8e6, (f.get("items"), f.get("payments")))
    # PDF: текст только на части страниц — разбирается правилами, в модель не уходит
    pdf = pymupdf.open(stream=pdf_lines(CT_RU), filetype="pdf")
    pdf.new_page()
    CALLS.clear()
    st, b = upload([("part.pdf", "application/pdf", pdf.tobytes())], {"lang": "ru"})
    ok("PDF с текстом не на всех страницах — разбор правилами, в модель не отправлен",
       st == 200 and ctb(b).get("detected") and not CALLS and b["files"][0]["parsed"]
       and not b["files"][0]["read_by_ai"], (st, len(CALLS)))


def check_contract_long():
    print("38б. Длинный договор (30 страниц): в пределах срока разбора, PDF с текстом до 60 страниц")
    fresh()
    model_on(True)
    CALLS.clear()
    body = [f"8.{k}. Страховщик обязан в течение 10 рабочих дней рассмотреть документы, представленные "
            f"Страхователем, и принять решение о выплате страхового возмещения либо об отказе в выплате, о чём "
            f"письменно уведомить Страхователя с указанием причин, если иное не предусмотрено правилами страхования "
            f"имущества юридических лиц, утверждёнными Страховщиком (пункт {k})." for k in range(1, 330)]
    lines = CT_RU[:26] + body + CT_RU[26:]
    text_len = sum(len(x) for x in lines)
    import time as _t
    t0 = _t.monotonic()
    st, b = upload([("long.docx", DOCX_MIME, docx_bytes([_x(x) for x in lines]))], {"lang": "ru"})
    dt_docx = _t.monotonic() - t0
    c = ctb(b)
    ok(f"DOCX ≈{text_len // 1000} тыс. знаков: разобран за {dt_docx:.1f} с, договор узнан, без пометки о сроке",
       st == 200 and c.get("detected") and dt_docx < 5 and not CALLS
       and not any("не уложился" in n or "время" in n for n in b["notes"]), (dt_docx, b["notes"]))
    ok("длинный договор: условия из начала и конца найдены",
       c["fields"]["premium"] == 10_000_000 and c["fields"]["territory"] == "Республика Узбекистан."
       and [x["code"] for x in c["fields"]["exclusions"]][:3] == ["war", "terrorism", "nuclear"], c.get("fields"))
    pdf = pdf_lines(lines, per_page=40)
    pages = pymupdf.open(stream=pdf, filetype="pdf").page_count
    t0 = _t.monotonic()
    st, b = upload([("long.pdf", "application/pdf", pdf)], {"lang": "ru"})
    dt_pdf = _t.monotonic() - t0
    c = ctb(b)
    CT_REPORT["длинный договор"] = {"знаков": text_len, "DOCX, с": round(dt_docx, 2), "PDF страниц": pages,
                                    "PDF, с": round(dt_pdf, 2)}
    ok(f"PDF {pages} стр. с текстом (больше 10) принят и разобран за {dt_pdf:.1f} с без модели",
       pages > 10 and st == 200 and not b["rejected"] and c.get("detected") and not CALLS and dt_pdf < 5
       and c["fields"]["premium"] == 10_000_000 and c.get("pages") == pages, (st, b.get("rejected"), dt_pdf))
    ok("PDF-скан больше 10 страниц по-прежнему отклоняется",
       "10 страниц" in (upload([("scan.pdf", "application/pdf", pdf_pages(12))], {"lang": "ru"})[1].get("rejected")
                        or [{}])[0].get("error", ""))
    set_limits(doc_max_text_chars=20000)
    try:
        st, b = upload([("long.docx", DOCX_MIME, docx_bytes([_x(x) for x in lines]))], {"lang": "ru"})
    finally:
        clear_settings()
    c = ctb(b)
    ok("предел текста: договор обрезан — честная пометка truncated и заметка",
       c.get("truncated") is True and any("предела разбора" in n for n in c["notes"]), (c.get("truncated"), c.get("notes")))


def ct_scan_reply(policyholder=None):
    v = {"file": 1, "contract_no": "45-ИМ/2026", "contract_date": "2026-10-01", "place": "Тошкент шаҳри",
         "product_name": "мол-мулкни суғурта қилиш", "product_code": None,
         "insurer": {"is_legal": True, "name": '"INSON" АЖ'},
         "policyholder": {"is_legal": True, "name": '"ALFA TEXTILE" МЧЖ'},
         "beneficiary": None, "pledger": None,
         "object": "нотурар бино — омбор, умумий майдони 2 400 кв.м", "class_hint": "building",
         "address": "Тошкент вилояти, Чирчиқ шаҳри", "cadastre_no": "10:09:05:01:02:0033",
         "object_value": "12 000 000 000", "sum_insured": "12 000 000 000 сўм", "currency": "UZS",
         "tariff": "0,1 %", "premium": "36 032 877", "payment_mode": "installments",
         "payments": [{"date": "2026-10-10", "amount": "18 016 438"}, {"date": "2027-10-10", "amount": 18016439}],
         "term": "2026 йил 1 октябрдан 2029 йил 30 сентябргача", "term_from": "2026-10-01", "term_to": "2029-09-30",
         "franchise": "қўлланилмайди", "covered_risks": ["ёнғин", "портлаш", "табиий офатлар", "сув босиши"],
         "exclusions": ["уруш ҳаракатлари", "ядро портлаши"], "territory": None, "special_terms": [],
         "notice": "3 иш куни"}
    if policyholder:
        v["policyholder"] = {"is_legal": True, "name": policyholder}
    return "```json\n" + _json.dumps({"files": [{"n": 1, "view": "document", "document_kind": "contract"}],
                                      "object_kind": None, "class_hint": None, "condition": None, "fields": [],
                                      "damages": [], "branch_request": None, "contract": v},
                                     ensure_ascii=False) + "\n```"


def check_contract_scan():
    print("38в. Договор: скан — ответ модели (подменён) по строгой схеме")
    fresh()
    model_on(True)
    CALLS.clear()
    REPLY["text"] = ct_scan_reply()
    st, b = upload([("scan.png", "image/png", image((250, 250, 250)))], {"lang": "ru"})
    prompt = CALLS[0]["messages"][0]["content"] + CALLS[0]["messages"][1]["content"]
    ok("в инструкции модели — договор, схема contract, «не выдумывай», запрет имён",
       '"contract": null или' in prompt and "document_kind = contract" in prompt and "не выдумывай" in prompt
       and "is_legal = false" in prompt and llm.mask_pd(prompt) == prompt, prompt[-300:])
    c = ctb(b)
    f = c.get("fields") or {}
    ok("скан: источник photo, файл f1, вид — договор страхования",
       c.get("source") == "photo" and c.get("file") == "f1" and b["files"][0]["document_kind"] == "договор страхования", c)
    ok("скан: те же поля, что у разбора текста (суммы, срок, график, риски)",
       f["sum_insured"] == CT_S and f["premium"] == CT_PREMIUM and f["tariff_pct"] == 0.1 and f["term_days"] == 1096
       and f["payments"] == [{"date": "2026-10-10", "amount": 18016438.0}, {"date": "2027-10-10", "amount": 18016439.0}]
       and [x["code"] for x in f["covered_risks"]] == ["fire", "explosion", "natural", "water"]
       and [x["code"] for x in f["exclusions"]] == ["war", "nuclear"] and f["franchise"]["applied"] is False, f)
    rec = {(r["key"], r["value"], r["source"]) for r in b["recognized"]}
    ok("распознанное со скана — источник «документ», prefill — код, сумма, срок",
       ("premium", "36 032 877", "document") in rec and ("policyholder", '"ALFA TEXTILE" МЧЖ', "document") in rec
       and b["prefill"]["term_days"]["value"] == 1096 and b["prefill"]["sum_insured"]["value"] == CT_S, sorted(rec))
    REPLY["text"] = ct_scan_reply(policyholder="Каримов Алишер Анварович")
    st, b = upload([("scan2.png", "image/png", image((240, 240, 240)))], {"lang": "ru"})
    dump = _json.dumps(b, ensure_ascii=False)
    with db.tx() as con:
        saved = db.rows(con, "SELECT result_json FROM act_uploads WHERE id=?", b["session"])[0]["result_json"]
    REPLY["text"] = ct_scan_reply().replace('"2026 йил 1 октябрдан 2029 йил 30 сентябргача"', '"36 ой"').replace(
        '"term_from": "2026-10-01", "term_to": "2029-09-30"', '"term_from": null, "term_to": null')
    st, b3 = upload([("scan3.png", "image/png", image((230, 230, 230)))], {"lang": "ru"})
    f3 = ctb(b3).get("fields") or {}
    ok("скан: срок словами без дат («36 ой») — 1 095 дн., даты не выдумываются",
       f3.get("term_days") == 1095 and f3["term_from"] is None and f3["term_to"] is None, f3.get("term_days"))
    ok("скан: гражданин вместо организации — kind = individual, имени нет",
       ctb(b)["fields"]["policyholder"] == {"kind": "individual", "name": None}
       and "Каримов" not in dump and "Каримов" not in saved, ctb(b)["fields"]["policyholder"])


CT_AI_LINES = [
    "ДОГОВОР СТРАХОВАНИЯ ИМУЩЕСТВА № 9/2026",
    "г. Ташкент, 1 октября 2026 г.",
    "Акционерное общество «INSON», именуемое в дальнейшем «Страховщик», в лице директора Иванова Ивана Ивановича, "
    "и ООО «Бета Логистик», именуемое в дальнейшем «Страхователь», в лице директора Сидорова Сидора Сидоровича, "
    "заключили настоящий договор.",
    "Страховая сумма: 1 000 000 000 сум.",
    "Страховщик принимает на себя обязательство возместить ущерб, причинённый складскому комплексу Страхователя, "
    "а Страхователь уплачивает Страховщику двенадцать миллионов сумов в течение десяти дней с даты подписания.",
    "Договор действует три года со дня, следующего за днём уплаты первого взноса.",
]


def ct_ai_reply(messages):
    # «модель» видит только текст после маскировки и возвращает условия, которых правила не нашли
    return _json.dumps({"contract": {
        "sum_insured": "2 000 000 000", "premium": "12 000 000", "tariff": "0,4 %",
        "term_from": "2026-10-11", "term_to": "2029-10-10", "object": "складской комплекс",
        "covered_risks": ["пожар", "кража"], "policyholder": {"is_legal": True, "name": "[ФИО]"}}}, ensure_ascii=False)


def check_contract_ai_assist():
    print("38г. Договор: правила нашли мало — текст (после маскировки) дочитывает модель, только пустые поля")
    fresh()
    model_on(True)
    CALLS.clear()
    REPLY["text"] = ct_ai_reply
    st, b = upload([("dogovor9.docx", DOCX_MIME, docx_bytes(CT_AI_LINES))], {"lang": "ru"})
    c = ctb(b)
    f = c.get("fields") or {}
    ok("одно обращение к модели — текстом, без файлов", len(CALLS) == 1 and not CALLS[0]["files"]
       and CALLS[0]["purpose"] == "акт: договор по тексту", [x["purpose"] for x in CALLS])
    sent = " ".join(m["content"] for m in CALLS[0]["messages"]) if CALLS else ""
    ok("в модель ушёл замаскированный текст: ФИО директоров нет, метка [ФИО] есть, сумма и слова договора есть",
       "Иванов" not in sent and "Ивана Ивановича" not in sent and "Сидоров" not in sent and "[ФИО]" in sent
       and "1 000 000 000" in sent and "двенадцать миллионов" in sent, sent[:500])
    ok("найденное правилами не заменено: страховая сумма 1 млрд, а не 2 млрд из ответа модели",
       f.get("sum_insured") == 1e9 and "sum_insured" not in c["field_sources"], (f.get("sum_insured"), c["field_sources"]))
    ok("пустые поля дополнены моделью: премия, срок, объект, риски — source = document_ai",
       f["premium"] == 12e6 and f["term_days"] == 1096 and f["object_description"] == "складской комплекс"
       and set(c["field_sources"]) >= {"premium", "term_days", "object_description", "covered_risks", "tariff_pct"}
       and set(c["field_sources"].values()) == {"document_ai"} and c["source"] == "document_ai", c.get("field_sources"))
    ok("страхователь-юрлицо из правил не заменён меткой [ФИО]",
       f["policyholder"] == {"kind": "legal", "name": "ООО «Бета Логистик»"}, f["policyholder"])
    rec = {(r["key"], r["value"]): r for r in b["recognized"]}
    ok("распознанное: премия от модели — source document_ai и пометка «прочитано моделью из текста, проверьте»",
       rec.get(("premium", "12 000 000"), {}).get("source") == "document_ai"
       and rec[("premium", "12 000 000")]["note"] == "прочитано моделью из текста, проверьте"
       and rec.get(("sum_insured", "1 000 000 000"), {}).get("source") == "document", sorted(rec))
    ok("prefill: срок от модели — с пометкой модели", b["prefill"]["term_days"]["source"] == "document_ai"
       and b["prefill"]["term_days"]["check_label"] == "прочитано моделью из текста, проверьте"
       and b["prefill"]["sum_insured"]["source"] == "document", b["prefill"])
    ok("заметка о полях, прочитанных моделью", any("прочитана моделью" in n for n in c["notes"]), c["notes"])
    with db.tx() as con:
        saved = db.rows(con, "SELECT result_json FROM act_uploads WHERE id=?", b["session"])[0]["result_json"]
    ok("текст договора в базу не сохранён", "двенадцать миллионов" not in saved and "Иванов" not in saved)
    sid_ai = b["session"]
    # настройка contract.ai_assist = false — модель не вызывается
    with db.tx() as con:
        con.execute("INSERT INTO act_settings (created_at, created_by, settings_json, calibrated, note) VALUES "
                    "(?,?,?,?,?)", (db.now(), "тест", _json.dumps({"contract": {"ai_assist": False,
                                                                              "ai_max_chars": 30000}}), 0, "тест"))
    try:
        CALLS.clear()
        st, b = upload([("dogovor9.docx", DOCX_MIME, docx_bytes(CT_AI_LINES))], {"lang": "ru"})
        ok("contract.ai_assist = false — модель не вызывается, поля остаются пустыми",
           not CALLS and ctb(b)["fields"]["premium"] is None and ctb(b)["source"] == "document", len(CALLS))
    finally:
        clear_settings()
    # модель не подключена — договор разобран правилами, без ошибок
    model_on(False)
    CALLS.clear()
    st, b = upload([("dogovor9.docx", DOCX_MIME, docx_bytes(CT_AI_LINES))], {"lang": "ru"})
    ok("модель не подключена — разбор правилами, без обращения", st == 200 and not CALLS and ctb(b)["detected"])
    model_on(True)
    # длинный текст: в модель уходит не больше ai_max_chars, кусками не длиннее 11 000 знаков
    CALLS.clear()
    many = CT_AI_LINES + [f"Пункт {k}. Стороны руководствуются законодательством Республики Узбекистан." * 3
                          for k in range(1, 700)]
    st, b = upload([("dogovor_long.docx", DOCX_MIME, docx_bytes(many))], {"lang": "ru"})
    parts = [m["content"] for m in CALLS[0]["messages"][2:]] if CALLS else []
    ok("длинный текст: в модель — не больше 30 000 знаков, куски по ≤ 11 000",
       CALLS and sum(len(p) for p in parts) <= 30000 + 200 and all(len(p) <= 11100 for p in parts)
       and len(parts) >= 2, [len(p) for p in parts])
    ok("ae.check_settings: contract проверяется",
       ae.check_settings({"contract": {"ai_assist": "да", "ai_max_chars": 10}})
       and not ae.check_settings({"contract": {"ai_assist": False, "ai_max_chars": 20000}}))
    return sid_ai


def ct_make(sid=None, contract=None, optional=None, lang="ru", must=None, request=None):
    body = {"session": sid, "lang": lang,
            "must": must or {"product_code": "0832", "sum_insured": CT_S, "object_value": CT_S,
                             "region": "Ташкентская область"},
            "optional": dict(optional or {})}
    if contract is not None:
        body["optional"]["contract"] = contract
    if request is not None:
        body["optional"]["request"] = request
    return call("POST", "/act/make", body)


def ct_items(a):
    return {i["code"]: i for i in (a.get("contract_check") or {}).get("items") or []}


def check_contract_make():
    print("38д. Сверка договора с расчётом акта (продукт 0832): премия, график, существенные условия, запрос филиала")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    st, b = upload([("shartnoma.docx", DOCX_MIME, docx_bytes([_x(x) for x in CT_UZC]))], {"lang": "ru"})
    ct = ctb(b)["request"]
    st, a = ct_make(b["session"], ct)
    it = ct_items(a)
    cc = a.get("contract_check") or {}
    lvl = a["risk"]["level"]
    applied = round(max(0.08 * (1 + ae.DEFAULT_SETTINGS["adj_pct"][lvl] / 100), 0.08), 4)
    ok("акт с договором сформирован; contract_check доступна, источник — файл договора",
       st == 200 and cc.get("available") and cc["source"] == "document"
       and cc["source_label"] == "из документа — из файла договора (разбор текста)", (st, cc.get("source_label")))
    ok("срок акта взят из договора: 1 096 дн.",
       a["premium"]["term_days"] == 1096 and it["term"]["verdict"] == "ok" and "взят из договора" in it["term"]["text"],
       it.get("term"))
    ok("тариф договора 0,1 % не ниже минимума 0,08 %",
       it["tariff_min"]["verdict"] == "ok" and it["tariff_min"]["calculated"] == 0.08, it["tariff_min"])
    ok(f"тариф договора против ставки акта {applied}", it["tariff_act"]["verdict"] == ("ok" if 0.1 >= applied else
                                                                                      "differs"), it["tariff_act"])
    ok("премия договора 36 032 877 против расчёта 12 млрд × 0,1 % × 1096/365 = 36 032 876,71 — в допуске",
       it["premium_request"]["verdict"] == "ok" and it["premium_request"]["calculated"] == 36_032_876.71
       and "по тарифу договора" in it["premium_request"]["text"], it["premium_request"])
    ok("график 2 платежей = премия", it["payments"]["verdict"] == "ok" and it["payments"]["calculated"] == CT_PREMIUM,
       it["payments"])
    ok("существенные условия — все есть (ГК РУз, ст. 929)",
       it["essentials"]["verdict"] == "ok" and all(e["present"] for e in cc["essentials"])
       and cc["legal_ref"] == "ГК РУз, ст. 929", it["essentials"])
    ok("как сверено: дословно ст. 929 и формула премии",
       any("должно быть достигнуто соглашение" in h and "ст. 929" in h for h in cc["how"])
       and any("× 0,10 %" in nb(h) and "1096" in h for h in cc["how"]), cc["how"][:2])
    s1 = {r["label"]: r for r in a["sections"][0]["rows"]}
    ok("раздел 1: номер и дата договора", s1.get("Договор страхования", {}).get("value") == "№ 45-ИМ/2026 от 01.10.2026",
       list(s1))
    s4 = {li["title"]: li["items"] for li in a["sections"][3]["lists"]}
    ok("раздел 4: «Сверка с договором», риски и исключения, «Как сверен договор»",
       "Сверка с договором" in s4 and any("Застрахованные риски по договору: пожар" in x for x in s4["Сверка с договором"])
       and any("Исключения по договору: военные действия" in x for x in s4["Сверка с договором"])
       and "Как сверен договор" in s4, list(s4))
    CT_REPORT["сверка УЗ"] = {k: (v["verdict"], v["requested"], v["calculated"]) for k, v in it.items()}
    aid = a["id"]
    st, a0 = ct_make(b["session"], None)
    ok("без optional.contract — договор из своей загрузки (source = session)",
       st == 200 and a0["contract_check"]["source"] == "session" and ct_items(a0)["premium_request"]["verdict"] == "ok",
       a0.get("contract_check", {}).get("source"))

    # график не сходится с премией
    bad = [x.replace("18 016 439", "18 000 000") for x in CT_UZC]
    st, b2 = upload([("shartnoma2.docx", DOCX_MIME, docx_bytes([_x(x) for x in bad]))], {"lang": "ru"})
    st, a2 = ct_make(b2["session"], ctb(b2)["request"])
    p2 = ct_items(a2)["payments"]
    ok("график платежей не сходится с премией — differs, разница 16 439",
       p2["verdict"] == "differs" and p2["diff"] == 16439.0 and "разница" in p2["text"], p2)
    ok("в проверках андеррайтера — график; «принять без оговорок» нельзя",
       any("График платежей" in c for c in a2["decision"]["checks"]) and a2["decision"]["code"] != "accept",
       a2["decision"])

    # договор без срока → нет существенного условия
    no_term = [x for x in CT_RU if not x.startswith(("3.", "3 "))]
    st, b3 = upload([("noterm.docx", DOCX_MIME, docx_bytes([_x(x) for x in no_term]))], {"lang": "ru"})
    c3 = ctb(b3)
    ok("договор без срока: срок не найден, существенного условия нет",
       c3["fields"]["term_days"] is None and not {e["code"]: e["present"] for e in c3["essentials"]}["term"],
       c3["fields"].get("term_days"))
    st, a3 = ct_make(b3["session"], c3["request"], must={"product_code": "0832", "sum_insured": 5e9,
                                                        "object_value": 5e9, "region": "Самаркандская область"})
    e3 = ct_items(a3)["essentials"]
    ok("сверка: essentials — no_essential, «в договоре нет: срок действия договора»",
       e3["verdict"] == "no_essential" and "срок действия договора" in e3["text"]
       and a3["contract_check"]["summary"]["verdict"] == "no_essential", e3)
    ok("решение: не «принять без оговорок», в проверках — существенное условие по ст. 929",
       a3["decision"]["code"] != "accept" and any("ст. 929" in c and "срок" in c for c in a3["decision"]["checks"]),
       a3["decision"])

    # проверка ввода optional.contract
    for badc, key in (({"payments": [{"date": "31.02.2026", "amount": 5}]}, "дата платежа"),
                      ({"payments": "раз"}, "платежи не список"), ({"items": [{"name": "x", "sum": -1}]}, "сумма части"),
                      ({"currency": "XYZ", "premium": 1}, "валюта"), ({"franchise": {"applied": True, "type": "x"}},
                                                                      "тип франшизы"),
                      ({"contract_date": "вчера", "premium": 5}, "дата договора"), ("строка", "не объект"),
                      ({"covered_risks": []}, "пусто")):
        st, e = ct_make(None, badc)
        ok(f"optional.contract проверяется: {key}", st == 422 and "contract" in (e.get("errors") or {}), (st, e))
    st, a4 = ct_make(None, {"premium": "36 032 877", "tariff_pct": "0,1", "sum_insured": CT_S,
                            "term_from": "01.10.2026", "term_to": "30.09.2029", "covered_risks": ["fire", "кража"],
                            "items": [{"name": "склад", "sum": 7e9}, {"name": "оборудование", "sum": 4e9}],
                            "object_description": "склад и оборудование", "source": "input"})
    i4 = ct_items(a4)
    ok("ввод сотрудника: суммы по частям 11 млрд ≠ 12 млрд — items_sum differs, риски из кодов и слов",
       st == 200 and i4["items_sum"]["verdict"] == "differs" and i4["items_sum"]["calculated"] == 11e9
       and [x["code"] for x in a4["contract_check"]["covered_risks"]] == ["fire", "theft"]
       and any("Суммы по объектам" in c for c in a4["decision"]["checks"]), (i4.get("items_sum"), a4.get("contract_check")))
    return aid


def check_contract_cross():
    print("38е. Запрос филиала и договор в одной загрузке: расхождения (cross_check)")
    fresh()
    model_on(True)
    CALLS.clear()
    st, b = upload([("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1)),
                    ("shartnoma.docx", DOCX_MIME, docx_bytes([_x(x) for x in CT_UZC]))], {"lang": "ru"})
    x = b.get("cross_check") or {}
    xi = {i["code"]: i for i in x.get("items") or []}
    ok("оба блока есть: branch_request и contract; модель не вызывалась",
       st == 200 and b.get("branch_request") and b.get("contract") and not CALLS, (st, len(CALLS)))
    ok("cross_check: сумма, стоимость, тариф, премия, срок, объект расходятся; франшиза и код совпадают",
       xi["sum_insured"]["verdict"] == "differs" and xi["sum_insured"]["request"] == BR_S1
       and xi["sum_insured"]["contract"] == CT_S and xi["tariff_pct"]["verdict"] == "differs"
       and xi["premium"]["verdict"] == "differs" and xi["term"]["verdict"] == "differs"
       and xi["object"]["verdict"] == "differs" and xi["franchise"]["verdict"] == "same"
       and xi["product_code"]["verdict"] == "missing", {k: v["verdict"] for k, v in xi.items()})
    ok("cross_check: итог и строки на языке экрана",
       x["summary"]["verdict"] == "differs" and "расходятся" in x["summary"]["text"]
       and any(ln.startswith("Страховая сумма — расходится: в запросе 81") for ln in x["lines"]), x.get("lines"))
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru",
                                       "must": {"product_code": "0832", "sum_insured": BR_S1, "object_value": BR_S1,
                                                "region": "Ташкентская область"},
                                       "optional": {"request": br_request(b), "contract": ctb(b)["request"]}})
    s4 = {li["title"]: li["items"] for li in a["sections"][3]["lists"]}
    ok("акт: подразделы «Сверка с запросом филиала», «Сверка с договором», «Запрос филиала и договор: расхождения»",
       {"Сверка с запросом филиала", "Сверка с договором", "Запрос филиала и договор: расхождения"} <= set(s4), list(s4))
    ok("акт: cross_check в JSON и в проверках андеррайтера",
       a["cross_check"]["available"] and a["cross_check"]["differs"] >= 5
       and any(c.startswith("Запрос филиала и договор расходятся: страховая сумма") for c in a["decision"]["checks"])
       and a["decision"]["code"] != "accept", a["decision"]["checks"])
    return a["id"]


CT_REPORT = {}


def check_contract_langs_files(aid, x_aid):
    print("38ж. Сверка с договором: три языка, Word и PDF")
    for lang, title, word, xt in (("ru", "Сверка с договором", "по тарифу договора", "Запрос филиала и договор: расхождения"),
                                  ("uz", "Shartnoma bilan solishtirish", "shartnoma tarifi", "Filial soʻrovi va shartnoma: farqlar"),
                                  ("en", "Check against the contract", "contract rate", "Branch request vs contract: differences")):
        st, a = call("GET", f"/act/{aid}", params={"lang": lang})
        lines = [x for li in a["sections"][3]["lists"] if li["title"] == title for x in li["items"]]
        ok(f"{lang}: подраздел сверки с договором и текст строки", bool(lines) and any(word in x for x in lines), lines[:3])
        cc = a["contract_check"]
        ok(f"{lang}: contract_check на языке акта (подписи, итог, существенные условия)",
           cc["items"][0]["label"] and cc["summary"]["text"] and all(i["verdict_label"] for i in cc["items"])
           and cc["essentials"][0]["label"], cc["items"][0])
        row1 = [r for r in a["sections"][0]["rows"] if r["label"] == tx.t("ct_row", lang)]
        ok(f"{lang}: раздел 1 — номер и дата договора", row1 and "45-ИМ/2026" in row1[0]["value"]
           and "01.10.2026" in row1[0]["value"], row1)
        if lang != "ru":
            txt = " ".join(lines + cc["how"] + [c for c in a["decision"]["checks"]])
            ok(f"{lang}: в сверке с договором нет кириллицы", not re.search(r"[А-Яа-яЁё]", txt),
               re.findall(r".{20}[А-Яа-яЁё].{20}", txt)[:3])
        st, blob, h = call("GET", f"/act/{aid}.docx", params={"lang": lang}, raw=True)
        plain = re.sub(r"<[^>]+>", "", zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8"))
        ok(f"{lang}: сверка с договором в DOCX", title in plain and word in plain, plain[-300:])
        st, blob, h = call("GET", f"/act/{aid}.pdf", params={"lang": lang}, raw=True)
        text = pdf_text(pymupdf.open(stream=blob, filetype="pdf")).replace("ʻ", "'")
        ok(f"{lang}: сверка с договором в PDF", title.replace("ʻ", "'") in text and word in text, text[-300:])
        st, blob, h = call("GET", f"/act/{x_aid}.docx", params={"lang": lang}, raw=True)
        plain = re.sub(r"<[^>]+>", "", zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8"))
        st, blob, h = call("GET", f"/act/{x_aid}.pdf", params={"lang": lang}, raw=True)
        text = pdf_text(pymupdf.open(stream=blob, filetype="pdf")).replace("ʻ", "'")
        ok(f"{lang}: «запрос филиала и договор» в DOCX и PDF", xt in plain and xt.replace("ʻ", "'") in text)
    with db.tx() as con:
        row = db.rows(con, "SELECT act_json FROM acts WHERE id=?", aid)[0]
        stored = _json.loads(row["act_json"])
    for k in ("contract", "contract_check", "cross_check"):
        stored["data"].pop(k, None)
    out = act.render(stored["data"], "ru", stored["meta"])
    ok("старый акт без договора: contract_check.available = false, cross_check.available = false",
       out["contract_check"]["available"] is False and out["cross_check"]["available"] is False
       and "Сверка с договором" not in [li["title"] for li in out["sections"][3]["lists"]])


# ================================================================================================
#  39. Замечания контролёра 30.09.2026 (вечер): источник условий решает сервер, правки в акте, вид документа,
#      существенные условия, сверка «запрос ↔ договор», проверка ввода, маскировка счёта, стороны-юрлица
# ================================================================================================

def _s4_lines(a, title):
    return [x for li in a["sections"][3]["lists"] if li["title"] == title for x in li["items"]]


def check_trust_edits():
    print("39а. Источник условий решает сервер: совпало с загрузкой — «из документа», иначе — правка «было → стало»")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    st, b = upload([("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1))], {"lang": "ru"})
    sid = b["session"]
    req = br_request(b)
    ok("optional.request загрузки: стоимость и объект для сверки запроса с договором",
       req.get("object_value") == BR_S1 and req.get("cadastre_no") == "10:00:00:00:00:00001"
       and req.get("class_hint") == "building" and "музлатгич" in (req.get("object_description") or ""), req)
    # 1) как прочитано — «из документа», правок нет
    st, a = br_make(sid, 1, dict(req, source="input"))          # source экрана не доверяется ни в какую сторону
    rc = a["request_check"]
    ok("условия как в загрузке: источник «из документа», правок нет (source экрана не учитывается)",
       st == 200 and rc["source_kind"] == "document" and rc["source"] == "document"
       and rc["source_label"].startswith("из документа") and rc["edits"]["count"] == 0
       and rc["edits"]["line"] == "Правки сотрудника в условиях запроса: правок нет"
       and set(rc["field_sources"].values()) == {"document"} and rc["field_sources"]["tariff_pct"] == "document",
       (rc.get("source_kind"), rc.get("field_sources")))
    lines = _s4_lines(a, "Сверка с запросом филиала")
    ok("раздел 4: источник условий и «правок нет»",
       "Источник условий: из документа — из файла запроса (разбор текста)" in lines
       and "Правки сотрудника в условиях запроса: правок нет" in lines, lines)
    ok("без правок в проверках нет пункта о правках",
       not any("правки сотрудника" in c for c in a["decision"]["checks"]), a["decision"]["checks"])
    # 2) сотрудник исправил тариф и премию, экран утверждает «из документа» — сервер видит правки
    edited = dict(req, tariff_pct=0.1, premium=246_644_000, source="document")
    st, a2 = br_make(sid, 1, edited)
    rc2 = a2["request_check"]
    codes = [e["code"] for e in rc2["edits"]["items"]]
    ok("правки тарифа и премии: «из документа с правками сотрудника (2)», поля — «введено сотрудником»",
       st == 200 and rc2["source_kind"] == "document_edited" and codes == ["tariff_pct", "premium"]
       and rc2["source_label"].startswith("из документа с правками сотрудника (2)")
       and rc2["field_sources"]["tariff_pct"] == "input" and rc2["field_sources"]["premium"] == "input"
       and rc2["field_sources"]["term_days"] == "document", (rc2.get("source_label"), codes, rc2.get("field_sources")))
    ed = {e["code"]: e for e in rc2["edits"]["items"]}
    ok("правка: что, было, стало — числа и текст на языке акта",
       ed["tariff_pct"]["was"] == 0.05 and ed["tariff_pct"]["now"] == 0.1
       and nb(ed["tariff_pct"]["text"]) == "Тариф: было 0,05 % → стало 0,10 %"
       and nb(ed["premium"]["text"]) == "Страховая премия: было 123 322 000 сум → стало 246 644 000 сум"
       and rc2["edits"]["line"] == "Правки сотрудника в условиях запроса: 2", rc2["edits"])
    lines2 = [nb(x) for x in _s4_lines(a2, "Сверка с запросом филиала")]
    ok("раздел 4: строка правок и список «было → стало»",
       "Правки сотрудника в условиях запроса: 2" in lines2 and "Тариф: было 0,05 % → стало 0,10 %" in lines2, lines2)
    ok("решение: «проверить правки сотрудника в условиях запроса», без оговорок принять нельзя",
       any(c.startswith("Проверить правки сотрудника в условиях запроса (2)") for c in a2["decision"]["checks"])
       and a2["decision"]["code"] != "accept", a2["decision"])
    st, blob, h = call("GET", f"/act/{a2['id']}.docx", params={"lang": "ru"}, raw=True)
    plain = nb(re.sub(r"<[^>]+>", "", zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8")))
    ok("Word: правки сотрудника и «было → стало»",
       "Правки сотрудника в условиях запроса: 2" in plain and "было 0,05 % → стало 0,10 %" in plain, plain[-400:])
    st, blob, h = call("GET", f"/act/{a2['id']}.pdf", params={"lang": "ru"}, raw=True)
    text = nb(pdf_text(pymupdf.open(stream=blob, filetype="pdf")))
    ok("PDF: правки сотрудника", "Правки сотрудника в условиях запроса: 2" in text and "было 0,05" in text, text[-400:])
    for lang, word in (("uz", "Soʻrov shartlaridagi xodim tuzatishlari: 2"), ("en", "Staff edits to the request terms: 2")):
        st, al = call("GET", f"/act/{a2['id']}", params={"lang": lang})
        ok(f"{lang}: строка правок на языке акта, без кириллицы",
           al["request_check"]["edits"]["line"] == word
           and not re.search(r"[А-Яа-яЁё]", " ".join(e["text"] for e in al["request_check"]["edits"]["items"])),
           al["request_check"]["edits"])
    # 3) экран говорит «из документа», а загрузки нет (истекла / сменилась сессия) — всё «введено сотрудником»
    for sess, why in ((None, "сессии нет"), ("нет-такой-загрузки", "загрузка чужая или истекла")):
        st, a3 = br_make(sess, 1, dict(req, source="document"))
        rc3 = a3["request_check"]
        ok(f"{why}: все поля «введено сотрудником», пометка «документ недоступен»",
           st == 200 and rc3["source_kind"] == "input" and rc3["source"] == "input" and rc3["document_missing"]
           and set(rc3["field_sources"].values()) == {"input"}
           and rc3["source_label"] == "введено сотрудником — документ недоступен (прошло больше 24 часов или "
                                      "сменилась сессия)", rc3.get("source_label"))
    st, a4 = br_make(None, 1, {"tariff_pct": 0.1, "term_days": 365})
    ok("ввод сотрудника без документа: «введено сотрудником», без пометки о недоступном документе",
       a4["request_check"]["source_label"] == "введено сотрудником" and not a4["request_check"]["document_missing"],
       a4["request_check"].get("source_label"))
    # старый акт (без источника по полям) рисуется как раньше
    with db.tx() as con:
        stored = _json.loads(db.rows(con, "SELECT act_json FROM acts WHERE id=?", a2["id"])[0]["act_json"])
    for k in ("source_kind", "origin", "edits", "field_sources", "doc_missing"):
        stored["data"]["request_check"].pop(k, None)
    out = act.render(stored["data"], "ru", stored["meta"])
    ok("старый акт без источника по полям: без строки правок, подпись источника прежняя",
       out["request_check"]["available"] and "edits" not in out["request_check"]
       and not any("Правки сотрудника" in x for x in _s4_lines(out, "Сверка с запросом филиала"))
       and out["request_check"]["source_label"] == "из файла запроса (разбор текста)", out["request_check"].get("source_label"))

    # договор: правка премии; риски со скана — «прочитано моделью»
    st, b = upload([("shartnoma.docx", DOCX_MIME, docx_bytes([_x(x) for x in CT_UZC]))], {"lang": "ru"})
    ct = ctb(b)["request"]
    st, a5 = ct_make(b["session"], dict(ct, premium=36_000_000.5))
    cc = a5["contract_check"]
    ok("договор: правка премии — «было → стало» с тийинами, остальное из документа",
       cc["source_kind"] == "document_edited" and [e["code"] for e in cc["edits"]["items"]] == ["premium"]
       and nb(cc["edits"]["items"][0]["text"]) == "Страховая премия: было 36 032 877 сум → стало 36 000 000,50 сум"
       and cc["field_sources"]["covered_risks"] == "document"
       and any(c.startswith("Проверить правки сотрудника в условиях договора (1)") for c in a5["decision"]["checks"]),
       cc.get("edits"))
    ok("договор: строка правок в разделе 4",
       "Правки сотрудника в условиях договора: 1" in _s4_lines(a5, "Сверка с договором"), _s4_lines(a5, "Сверка с договором"))
    REPLY["text"] = ct_scan_reply()
    st, bs = upload([("scan.png", "image/png", image((250, 250, 250)))], {"lang": "ru"})
    st, a6 = ct_make(bs["session"], ctb(bs)["request"])
    cc6 = a6["contract_check"]
    risks = [x for x in _s4_lines(a6, "Сверка с договором") if x.startswith("Застрахованные риски по договору")]
    ok("скан договора: источник «из документа — со скана», риски в акте помечены «прочитано моделью»",
       cc6["source_kind"] == "document" and cc6["field_sources"]["covered_risks"] == "photo"
       and cc6["source_label"] == "из документа — со скана договора (распознано моделью)"
       and risks and risks[0].endswith("(прочитано моделью)") and cc6["covered_risks_by_model"] is True, (risks, cc6.get("source_label")))
    ok("договор из файла: риски без пометки модели",
       all(not x.endswith("(прочитано моделью)") for x in _s4_lines(a5, "Сверка с договором")))
    REPLY["text"] = CRANE_REPLY


def check_doc_kind_title():
    print("39б. Вид документа: заголовок сильнее строк; заявление — без сверки договора")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    CALLS.clear()
    ru_rows = [("1.", "Вид страхования (код):", "0832"), ("2.", "Страхователь:", 'ООО "Ромашка"'),
               ("3.", "Выгодоприобретатель:", 'АКБ "Намунабанк"'), ("4.", "Залогодатель:", ""),
               ("5.", "Объект страхования:", "Нежилое здание — склад, общая площадь 1 200 кв.м"),
               ("6.", "Страховая стоимость:", "1 000 000 000,00 сум"),
               ("7.", "Страховая сумма:", "1 000 000 000,00 сум"), ("8.", "Франшиза:", "не применяется"),
               ("9.", "Страховой тариф:", "0,1"), ("10.", "Страховая премия:", "1 000 000,00 сум"),
               ("11.", "Срок страхования:", "с 01.10.2026 по 30.09.2027")]
    st, b0 = upload([("zapros.docx", DOCX_MIME, docx_table(ru_rows, head=[]))], {"lang": "ru"})
    ok("те же строки без заголовка — запрос филиала", bool(b0.get("branch_request")) and not b0.get("contract"))
    st, b = upload([("polis.docx", DOCX_MIME, docx_table(ru_rows, head=["СТРАХОВОЙ ПОЛИС № 7/2026"]))], {"lang": "ru"})
    ok("«СТРАХОВОЙ ПОЛИС» со строками «подпись: значение» — договор/полис, не запрос филиала",
       st == 200 and b.get("branch_request") is None and ctb(b).get("detected")
       and b["documents"][0]["kind"] == "contract", (b.get("branch_request"), b.get("documents")))
    for head in ("SUGʻURTA POLISI № 12", "СУҒУРТА ПОЛИСИ", "ПОЛИС", "ДОГОВОР СТРАХОВАНИЯ ИМУЩЕСТВА № 9"):
        st, b = upload([("p.docx", DOCX_MIME, docx_table(BR_SAMPLE1, head=[head]))], {"lang": "ru"})
        ok(f"«{head}» + строки бланка — договор/полис", b.get("branch_request") is None and ctb(b).get("detected"),
           (head, b.get("documents")))
    app_lines = ["ЗАЯВЛЕНИЕ НА СТРАХОВАНИЕ ИМУЩЕСТВА", "Страхователь: ООО «Ромашка»",
                 "Объект страхования: нежилое здание — склад, общая площадь 1 200 кв.м",
                 "Страховая стоимость: 1 200 000 000 сум", "Страховая сумма: 1 000 000 000 сум",
                 "Страховой тариф: 0,2 %", "Срок страхования: с 01.10.2026 по 30.09.2027"]
    for name, lines in (("ЗАЯВЛЕНИЕ НА СТРАХОВАНИЕ", app_lines),
                        ("АРИЗА", ["АРИЗА"] + app_lines[1:]), ("ARIZA", ["Sugʻurta qilish uchun ARIZA"] + app_lines[1:])):
        st, b = upload([("z.docx", DOCX_MIME, docx_bytes([_x(x) for x in lines]))], {"lang": "ru"})
        rec = {(r["key"], r["value"]) for r in b["recognized"]}
        ok(f"{name}: вид «заявление», данные в распознанном и подсказке, сверки договора нет",
           st == 200 and b["documents"][0]["kind"] == "application"
           and b["files"][0]["document_kind"] == "заявление на страхование"
           and b.get("contract") is None and b.get("branch_request") is None
           and ("sum_insured", "1 000 000 000") in rec and (b.get("prefill") or {}).get("sum_insured", {}).get("value") == 1e9
           and any("это заявление, а не договор" in n.lower() for n in b["notes"]), (b.get("documents"), b.get("notes")))
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru",
                                       "must": {"product_code": "0808", "sum_insured": 1e9, "object_value": 1.2e9,
                                                "region": "Ташкентская область"}})
    ok("акт по заявлению: сверки с договором и существенных условий нет",
       st == 200 and a["contract_check"]["available"] is False
       and not any("929" in c for c in a["decision"]["checks"]), a.get("contract_check"))


def check_essentials_wording():
    print("39в. Существенные условия: «в тексте договора не найдено … проверьте договор» (три языка)")
    fresh()
    no_term = [x for x in CT_RU if not x.startswith(("3.", "3 "))]
    st, b = upload([("noterm.docx", DOCX_MIME, docx_bytes([_x(x) for x in no_term]))], {"lang": "ru"})
    ok("экран: заметка не утверждает, что условия нет",
       any(n.startswith("В тексте договора не найдено условие: срок действия договора (ГК РУз, ст. 929). "
                        "Проверьте договор") for n in ctb(b)["notes"]), ctb(b)["notes"])
    st, a = ct_make(b["session"], ctb(b)["request"], must={"product_code": "0832", "sum_insured": 5e9,
                                                          "object_value": 5e9, "region": "Самаркандская область"})
    e = ct_items(a)["essentials"]
    ok("акт: строка сверки и проверка андеррайтера — «не найдено … если условия действительно нет — дополнить»",
       e["verdict"] == "no_essential" and e["verdict_label"] == "условие не найдено"
       and "в тексте договора не найдено условие: срок действия договора" in e["text"]
       and "если условия действительно нет — договор нужно дополнить" in e["text"]
       and any(c.startswith("В тексте договора не найдено условие: срок действия договора") for c in a["decision"]["checks"])
       and "не найдено существенное условие" in a["contract_check"]["summary"]["text"], (e, a["decision"]["checks"]))
    for lang, word in (("uz", "Shartnoma matnida shart topilmadi"), ("en", "The contract text does not contain the term")):
        st, al = call("GET", f"/act/{a['id']}", params={"lang": lang})
        ok(f"{lang}: формулировка без категоричности",
           any(c.startswith(word) for c in al["decision"]["checks"]), al["decision"]["checks"])
    for lg, v in (("ru", "нет в одном из документов"), ("uz", "hujjatlardan birida yoʻq"),
                  ("en", "missing in one of the documents")):
        ok(f"x_v_missing ({lg}): «{v}»", tx.t("x_v_missing", lg) == v)


def check_cross_same():
    print("39г. Сверка «запрос ↔ договор» в акте совпадает со сверкой при загрузке; объект — сопоставимое")
    fresh()
    model_on(True)
    REPLY["text"] = CRANE_REPLY
    st, b = upload([("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1)),
                    ("shartnoma.docx", DOCX_MIME, docx_bytes([_x(x) for x in CT_UZC]))], {"lang": "ru"})
    up = {i["code"]: (i["verdict"], i["text"]) for i in b["cross_check"]["items"]}
    st, a = call("POST", "/act/make", {"session": b["session"], "lang": "ru",
                                       "must": {"product_code": "0832", "sum_insured": BR_S1, "object_value": BR_S1,
                                                "region": "Ташкентская область"},
                                       "optional": {"request": br_request(b), "contract": ctb(b)["request"]}})
    got = {i["code"]: (i["verdict"], i["text"]) for i in a["cross_check"]["items"]}
    ok("все строки сверки акта = строки сверки загрузки (вывод и текст)", got == up,
       {k: (up.get(k), got.get(k)) for k in set(up) | set(got) if up.get(k) != got.get(k)})
    ok("стоимость сверяется (раньше в акте её не было в запросе), объект — по кадастру",
       got["object_value"][0] == "differs" and got["object"][0] == "differs"
       and a["cross_check"]["items"][-1]["compared"] == "cadastre", got)
    # при живой загрузке объект — сохранённый, присланный вид объекта не подменяет его
    st, a1 = call("POST", "/act/make", {"session": b["session"], "lang": "ru",
                                        "must": {"product_code": "0832", "sum_insured": BR_S1, "object_value": BR_S1,
                                                 "region": "Ташкентская область"},
                                        "optional": {"request": dict(br_request(b), cadastre_no=None, class_hint="cargo"),
                                                     "contract": ctb(b)["request"]}})
    xo = [i for i in a1["cross_check"]["items"] if i["code"] == "object"][0]
    ok("живая загрузка: объект из загрузки (кадастр сохранён), правка кадастра видна в правках",
       xo["compared"] == "cadastre" and xo["request"]["cadastre_no"] == "10:00:00:00:00:00001"
       and "cadastre_no" in [e["code"] for e in a1["request_check"]["edits"]["items"]], (xo, a1["request_check"]["edits"]))
    # сравнить нечем: кадастр только в договоре, вида объекта в запросе нет
    st, a2 = ct_make(None, {"premium": CT_PREMIUM, "sum_insured": CT_S, "cadastre_no": "10:00:00:00:00:00002",
                            "tariff_pct": 0.1, "term_days": 1096},
                     request={"tariff_pct": 0.1, "premium": CT_PREMIUM, "term_days": 1096})
    x2 = [i for i in a2["cross_check"]["items"] if i["code"] == "object"][0]
    lines = _s4_lines(a2, "Запрос филиала и договор: расхождения")
    ok("объект: сравнить нечем — вердикт missing, строка «сравнить нечем: в запросе …, в договоре …»",
       x2["verdict"] == "missing" and x2["text"] == "Объект — сравнить нечем: в запросе данные недоступны, "
                                                     "в договоре 10:00:00:00:00:00002"
       and x2["text"] in lines and not any("объект" in c.lower() for c in a2["decision"]["checks"]
                                           if c.startswith("Запрос филиала и договор")), (x2, lines))
    # вид с видом: кадастра нет в запросе, вид — в обоих
    st, a3 = ct_make(None, {"premium": CT_PREMIUM, "sum_insured": CT_S, "cadastre_no": "10:00:00:00:00:00002",
                            "object_description": "нежилое здание склада", "tariff_pct": 0.1, "term_days": 1096},
                     request={"tariff_pct": 0.1, "term_days": 1096, "object_description": "технологическое оборудование"})
    x3 = [i for i in a3["cross_check"]["items"] if i["code"] == "object"][0]
    ok("объект: кадастр только в одном — сравнивается вид объекта с видом (здание ≠ оборудование)",
       x3["compared"] == "kind" and x3["verdict"] == "differs"
       and x3["text"] == "Объект — расходится: в запросе оборудование, в договоре здание, помещение", x3)


def check_input_bounds():
    print("39д. Проверка ввода: даты 2000–2100, срок — только целое, разумный предел сумм")
    fresh()
    for bad, key in (({"term_from": "1999-12-31", "term_to": "2000-12-30"}, "срок с 1999 года"),
                     ({"term_from": "2100-06-01", "term_to": "2101-01-01"}, "срок по 2101 год"),
                     ({"term_days": 365.5}, "term_days 365.5"), ({"term_days": "365,5"}, "term_days «365,5»"),
                     ({"term_days": "365 дней"}, "term_days текстом"),
                     ({"premium": 2e14}, "премия больше 10^14"),
                     ({"premium": 5e9, "sum_insured": 1e9}, "премия больше страховой суммы")):
        st, e = br_make(None, 1, bad)
        ok(f"optional.request: {key} — 422", st == 422 and "request" in (e.get("errors") or {}), (st, e.get("errors")))
    for bad, key in (({"premium": 5, "contract_date": "1990-01-01"}, "дата договора 1990"),
                     ({"payments": [{"date": "2150-01-01", "amount": 5}]}, "платёж в 2150 году"),
                     ({"premium": 5, "term_days": 365.5}, "срок договора 365.5"),
                     ({"items": [{"name": "склад", "sum": 5e14}]}, "сумма части больше 10^14")):
        st, e = ct_make(None, badc := bad)
        ok(f"optional.contract: {key} — 422", st == 422 and "contract" in (e.get("errors") or {}), (st, badc, e.get("errors")))
    st, e = br_make(None, 1, None, {"term_days": 365.5})
    ok("optional.term_days 365.5 — 422", st == 422 and "term_days" in (e.get("errors") or {}), (st, e.get("errors")))
    st, a = br_make(None, 1, {"term_days": 365.0, "premium": "1 000 000"})
    ok("целое в записи 365.0 принимается", st == 200 and rq_items(a)["term"]["requested"] == 365, st)


def check_ct_ai_mask_labels():
    print("39е. Договор с р/с — счёт не уходит в модель; заметка «прочитано моделью» — подписями, без служебных полей")
    fresh()
    model_on(True)
    CALLS.clear()
    REPLY["text"] = ct_ai_reply
    lines = CT_AI_LINES + ["Реквизиты Страхователя: р/с 20208000900123456001, МФО 00873, х/р 2020 8000 9051 2345 6001, "
                           "карта 8600 1234 5678 9012."]
    st, b = upload([("dogovor_rs.docx", DOCX_MIME, docx_bytes(lines))], {"lang": "ru"})
    sent = " ".join(m["content"] for m in CALLS[0]["messages"]) if CALLS else ""
    ok("в модель ушёл текст без счёта, МФО и карты (метки [СЧЁТ], [МФО])",
       CALLS and not any(x in sent for x in ("20208000900123456001", "2020 8000 9051", "00873", "8600 1234"))
       and "[СЧЁТ]" in sent and "[МФО]" in sent and "1 000 000 000" in sent, sent[-300:])
    notes = [n for n in ctb(b)["notes"] if "прочитана моделью" in n]
    ok("заметка: поля подписями по-русски, служебных имён нет",
       notes and "страховая премия" in notes[0] and "срок страхования" in notes[0]
       and not re.search(r"[a-z]+_[a-z]+|class_hint|object_kind|term_text", notes[0]), notes)
    st, b2 = upload([("dogovor_rs.docx", DOCX_MIME, docx_bytes(lines))], {"lang": "uz"})
    n2 = [n for n in ctb(b2)["notes"] if "model" in n.lower()]
    ok("uz: подписи на узбекском", n2 and "sugʻurta mukofoti" in n2[0] and not re.search(r"[А-Яа-я]|_", n2[0]), n2)
    REPLY["text"] = CRANE_REPLY


def check_parties_amounts():
    print("39ж. Стороны-юрлица (банк слитно, маркеры, ЧП, «в лице …»), отрицательные суммы, множители, тийины")
    from app import branch_request as brm, contract_read as crm
    cases = (('"NAMUNA SAVDO" MCHJ', "legal"), ("Namunabank", "legal"), ("Намунабанк", "legal"),
             ("Sinov banki", "legal"), ('"OMAD" AJ', "legal"), ("АТБ «Намуна»", "legal"), ('"SINOV" XK', "legal"),
             ('"SINOV" QK', "legal"), ("DUK «Namuna»", "legal"), ('"Namuna" OK', "legal"),
             ('"Bahor" fermer xoʻjaligi', "legal"), ("фермер хўжалиги «Баҳор»", "legal"), ("ООО «Намуна»", "legal"),
             ("АО «Намуна»", "legal"), ("АКБ «Намуна»", "legal"), ("ЧП Каримов", "legal"),
             ("ИП Каримов", "individual"), ("ЯТТ Каримов", "individual"), ("YaTT Karimov", "individual"),
             ("Каримов Алишер Анварович", "individual"))
    bad = [(n, brm.party(n)) for n, kind in cases if brm.party(n)["kind"] != kind]
    ok("маркеры организаций и физлиц", not bad, bad)
    ok("«ЧП Каримов» — частное предприятие: юрлицо, название как есть",
       brm.party("ЧП Каримов") == {"kind": "legal", "name": "ЧП Каримов"})
    ok("«ООО «Ромашка», в лице директора Иванова И.И.» → «ООО «Ромашка»»",
       brm.party("ООО «Ромашка», в лице директора Иванова И.И.") == {"kind": "legal", "name": "ООО «Ромашка»"})
    ok("отрицательное число не превращается в положительное",
       brm.amount("-123 322 000,00 сўм") is None and brm.amount_ex("−5 000")[1] == "negative"
       and brm.amount("5 000") == 5000.0, brm.amount("-123 322 000,00 сўм"))
    ok("множители «млн/млрд/mln/mlrd» — как в договоре (общая функция)",
       brm.amount("1,5 млрд сум") == 1.5e9 and brm.amount("250 mln soʻm") == 2.5e8 and brm.amount("2 mlrd") == 2e9
       and crm.money("1,5 млрд сум")["value"] == 1.5e9 and brm.SCALE is crm._SCALE)
    rows = [(n, l, "-123 322 000,00 сўм" if n == "10." else v) for n, l, v in BR_SAMPLE1]
    st, b = upload([("neg.docx", DOCX_MIME, docx_table(rows))], {"lang": "ru"})
    f = b["branch_request"]["fields"]
    ok("бланк с отрицательной премией: премии нет, пометка для сотрудника",
       f["premium"] is None and f["amount_errors"] == ["premium"]
       and any("отрицательное число (страховая премия)" in n for n in b["branch_request"]["notes"]), b["branch_request"]["notes"])
    req2 = {"tariff_pct": 0.05, "premium": "122 589 000,00", "term_from": "07.09.2026", "term_to": "2031-11-07"}
    st, a = br_make(None, 2, req2)
    ok("«как считали»: сумма с тийинами (47 397 852 345,04)",
       any("47 397 852 345,04 сум" in nb(h) for h in a["request_check"]["how"]), a["request_check"]["how"])


def check_pdf_time_limit():
    print("39з. Запас времени для длинного договора: PDF с текстом — 8 с на файл, общий срок — 12 с")
    from app import act_extras as axm
    ok("настройки по умолчанию: doc_file_sec_pdf = 8, doc_parse_total_sec = 12, doc_parse_sec = 5",
       ae.DEFAULT_SETTINGS["limits"]["doc_file_sec_pdf"] == 8 and ae.DEFAULT_SETTINGS["limits"]["doc_parse_total_sec"] == 12
       and ae.DEFAULT_SETTINGS["limits"]["doc_parse_sec"] == 5 and axm.DOC_LIMITS["doc_file_sec_pdf"] == 8
       and "doc_file_sec_pdf" in ae.LIMIT_BOUNDS)
    seen = []
    orig = axm.parse_document_limited

    def spy(con, path, cls, dl, sec, inclusive=True):
        seen.append((Path(path).suffix, round(sec, 1)))
        return orig(con, path, cls, dl, sec, inclusive)
    axm.parse_document_limited = spy
    try:
        fresh()
        st, b = upload([("dog.pdf", "application/pdf", pdf_lines(CT_RU)),
                        ("dog.docx", DOCX_MIME, docx_bytes([_x(x) for x in CT_RU]))], {"lang": "ru"})
    finally:
        axm.parse_document_limited = orig
    ok("PDF — срок 8 с, DOCX — 5 с", st == 200 and seen and seen[0] == (".pdf", 8.0) and seen[1] == (".docx", 5.0), seen)


# ------------------------------------------------------------------ 40. аналитика раздела 4 (30.09.2026)

AN_TITLES_RU = ["Разбор по рискам: доля в нетто-ставке и уровень",
                "Учтённые факторы: значение, источник, вклад в техническую ставку",
                "Что изменит ставку (посчитано расчётным модулем)", "Состав тарифа", "Сценарии убытка подробно",
                "Сценарии «что если» (расчёт при других данных объекта)",
                "Лимит удержания и перестрахование", "Балл риска 0–100 (справочно)", "Рынок и статистика",
                "Франшиза: варианты (справочно)", "Мероприятия: эффект на ставку и премию"]
AN_SOURCES = {"input", "document", "photo", "plate", "marking", "text", "kind", "default", "not_set", "act_terms"}
WH8_MUST = {"product_code": "0807", "sum_insured": 4_200_000_000, "object_value": 4_200_000_000,
            "region": "Ташкентская область"}
WH8_OPT = {"object_kind": "warehouse", "protection": "alarm", "seismic_zone": 8, "construction": "reinforced",
           "losses_3y": {"count": 0, "small_count": 0}}
AN_REPORT = {}


def s4_titles(a):
    return [li["title"] for li in a["sections"][3]["lists"]]


def an_common(tag, a, S, term):
    """Проверки, общие для всех объектов: блоки на месте, факторы с источниками, вклады сходятся, балл сходится."""
    an = a.get("analytics") or {}
    ok(f"{tag}: analytics есть, calibrated = 0", an.get("available") and an.get("calibrated") == 0, an.get("reason"))
    titles = s4_titles(a)
    ok(f"{tag}: в разделе 4 все блоки аналитики", all(x in titles for x in AN_TITLES_RU),
       [x for x in AN_TITLES_RU if x not in titles])
    items = an["factors"]["items"]
    ok(f"{tag}: у каждого фактора источник и множитель", items and all(
        f["source"] in AN_SOURCES and f["source_label"] and f["multiplier"] > 0 for f in items), items[:2])
    total = an["factors"]["base_pct"] + sum(f["rate_pp"] or 0 for f in items)
    ok(f"{tag}: вклады факторов складываются в техническую ставку",
       abs(total - an["factors"]["technical_pct"]) < 2e-3, (total, an["factors"]["technical_pct"]))
    sc = an["score"]
    ok(f"{tag}: балл 0–100 = сумма вкладов составляющих", sc["available"] and 0 <= sc["score"] <= 100 and abs(
        sum(c["contribution"] or 0 for c in sc["components"] if c["applicable"]) - sc["score"]) < 0.3, sc.get("score"))
    sm = an["summary"]["sentences"]
    ok(f"{tag}: резюме 5–7 предложений и первым абзацем раздела 4", 5 <= len(sm) <= 7
       and a["sections"][3]["paragraphs"][0].startswith("Кратко: "), sm)
    ok(f"{tag}: у внешних данных — ссылка на источник", all(s.get("url", "").startswith("http")
                                                          for s in an["sources"]) and bool(an["sources"])
       or not an["market"]["available"], an.get("sources"))
    mk = an["market"]
    ok(f"{tag}: рынок — ставка, дата среза и источник НАПП (или честно «нет данных»)",
       (mk["available"] and mk["rate_pct"] and mk["rate_date"] and "napp.uz" in (mk["source"] or {}).get("url", ""))
       or (not mk["available"] and any("нет данных" in x for li in a["sections"][3]["lists"]
                                       if li["title"] == "Рынок и статистика" for x in li["items"])), mk)
    li = next(x for x in a["sections"][3]["lists"] if x["title"] == "Рынок и статистика")
    ok(f"{tag}: под рынком и статистикой — строки «Источник: …»",
       any(x.startswith("Источник: ") for x in li["items"]), li["items"][-3:])
    for it in an["stats"]["indicators"]:
        if it["status"] == "ok":
            ok(f"{tag}: показатель «{it['name']}» — с источником", it["sources"] and all(
                (s.get("url") or "").startswith("http") for s in it["sources"]), it["sources"])
            break
    return an


def check_analytics():
    print("40. Аналитика раздела 4: риски, факторы, чувствительность, состав тарифа, сценарии, балл, рынок, франшиза")
    from app import act_analytics as aa, act_extras as ax, market_picture as mp, risk_analytics as ra
    from app.engine import calculate, PURPOSE_ANALYSIS
    fresh()
    model_on(False)

    # --- 40а. автокран, класс 3 ---
    t0 = monotonic()
    st, a = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                       "optional": dict(CRANE_OPT, object_kind="truck_crane", year=2026)})
    sec = monotonic() - t0
    ok("40а автокран: акт сформирован быстрее 2 с", st == 200 and sec < 2, (st, sec))
    an = an_common("40а автокран", a, CRANE_MUST["sum_insured"], 365)
    rk = an["risks"]
    ok("40а: класс 3 не разбит на риски — одна строка 100 % и честная пометка",
       rk["whole_class"] and len(rk["items"]) == 1 and rk["items"][0]["share_of_net_pct"] == 100.0
       and any("не разбит на отдельные риски" in n for n in rk["notes"]), rk)
    fx = {f["code"]: f for f in an["factors"]["items"]}
    ok("40а: источники факторов — тип по виду объекта, убытки введены, год введён",
       fx["veh_type"]["source"] == "kind" and fx["loss_history"]["source"] == "input"
       and fx["veh_age"]["source"] == "input" and fx["antitheft"]["source"] == "not_set", fx)
    with db.tx() as con:
        ref = db.load_reference(con)
        inp = Input(product_code="0318", class_code="3", object_type=CRANE_TYPE,
                    value_amount=CRANE_MUST["object_value"], sum_insured=CRANE_MUST["sum_insured"], term_days=365,
                    factors={"veh_type": "special", "veh_age": "a3", "loss_history": "clean"})
        calc = calculate(ref, inp, purpose=PURPOSE_ANALYSIS)
    tb = an["tariff"]
    ok("40а: состав тарифа сходится с engine.calculate (нетто, техническая, минимум)",
       tb["available"] and abs(tb["technical_pct"] - calc["rates"]["technical_pct"]) < 1e-9
       and abs(tb["net_pct"] - calc["rates"]["net_pct"]) < 1e-9 and tb["min_pct"] == calc["rates"]["min_pct"],
       (tb.get("technical_pct"), calc["rates"]))
    ok("40а: тариф акта по-прежнему — ставка политики × поправка (техническая — справка)",
       tb["act_rate_pct"] == a["rate"]["applied_pct"] and tb["policy_rate_pct"] == a["rate"]["base_pct"]
       and a["premium"]["amount"] == round(a["rate"]["applied_pct"] / 100 * CRANE_MUST["sum_insured"])
       and "считается по тарифной политике" in tb["conclusion"], tb.get("conclusion"))
    ok("40а: вывод «Тариф акта … при технической ставке расчётного модуля … и рыночной …»",
       tb["conclusion"].startswith("Тариф акта " + act.pct(a["rate"]["applied_pct"], "ru") + " при технической "
                                   "ставке расчётного модуля " + act.pct(calc["rates"]["technical_pct"], "ru")),
       tb["conclusion"])
    sens = an["sensitivity"]["items"]
    good = bool(sens)
    for s in sens:
        alt = Input(**{**inp.__dict__})
        alt.factors = {**inp.factors, s["factor"]: next(o for (f, o), v in ref.coefficients.items()
                                                         if f == s["factor"] and s["to"] ==
                                                         tx.label(tx.OPTION_LABELS, f"{f}:{o}", "ru"))}
        g1 = rate_for(ref, alt)["gross_pct"]
        good = good and abs(round(g1, 4) - s["tech_after"]) < 1e-9
    ok("40а: чувствительность посчитана движком (сверка с прямым rate_for), 3–5 вариантов",
       good and 3 <= len(sens) <= 5, sens)
    ok("40а: мера страхователя — эффект и на премию акта (не ниже минимума)",
       any(s["kind"] == "measure" and s["act_premium_after"] is not None
           and s["act_premium_after"] >= round(0.35 / 100 * CRANE_MUST["sum_insured"]) for s in sens), sens)
    wi = {w["value"]: w for w in an["scenarios"]["whatif"]}
    with db.tx() as con:
        an2 = ra_same(con, {"class_code": "3", "product_code": "0318", "object_type": CRANE_TYPE,
                            "sum_insured": CRANE_MUST["sum_insured"], "object_value": CRANE_MUST["object_value"],
                            "region": CRANE_MUST["region"], "vehicle_type": "special", "year": 2026},
                      {"losses_3y": {"count": 0}, "protection": "immo"})
    ok("40а: «что если» с иммобилайзером — EML акта = PML модуля с тем же входом",
       "immo" in wi and wi["immo"]["eml"] == round(an2["scenarios"]["PML"]["amount"])
       and wi["immo"]["eml"] < an["scenarios"]["items"][1]["amount"], (wi.get("immo"), an2["scenarios"]["PML"]))
    fr = an["franchise"]
    ok("40а: таблица франшиз 0,5/1/2/5 % (потолок класса 3 — 5 %), вывод акта прежний",
       fr["available"] and [r["pct"] for r in fr["rows"]] == [0.5, 1.0, 2.0, 5.0]
       and a["franchise"]["text"] == "Франшиза не требуется" and fr["verdict"] == "Франшиза не требуется"
       and all(r["premium"] >= round(0.35 / 100 * CRANE_MUST["sum_insured"]) for r in fr["rows"]), fr)
    AN_REPORT["автокран"] = {"техническая": tb["technical_pct"], "тариф акта": tb["act_rate_pct"],
                             "рынок": tb["market_rate_pct"], "балл": an["score"]["score"],
                             "чувствительность": [(s["factor"], s["to"], s["delta_pct"]) for s in sens]}

    # --- 40б. склад, класс 8 (4,2 млрд, сейсмозона 8, сигнализация, железобетон) ---
    st, a = call("POST", "/act/make", {"lang": "ru", "must": WH8_MUST, "optional": WH8_OPT})
    an = an_common("40б склад кл. 8", a, WH8_MUST["sum_insured"], 365)
    items = an["risks"]["items"]
    with db.tx() as con:
        an_ra = ra_same(con, {"class_code": "8", "product_code": "0807", "object_type": "Склад",
                              "sum_insured": WH8_MUST["sum_insured"], "object_value": WH8_MUST["object_value"],
                              "region": WH8_MUST["region"], "construction": "reinforced", "activity": "warehouse"},
                        {"protection": "alarm", "seismic_zone": 8, "losses_3y": {"count": 0}})
    ok("40б: доли рисков — из risk_analytics, сумма 100 %",
       abs(sum(i["share_of_net_pct"] for i in items) - 100) <= 0.3
       and {i["code"]: i["share_of_net_pct"] for i in items} == {r["code"]: r["share_of_net_pct"] for r in an_ra["risks"]},
       [(i["code"], i["share_of_net_pct"]) for i in items])
    by = {i["code"]: i for i in items}
    ok("40б: уровни по рискам с причиной — землетрясение от сейсмозоны, пожар ниже среднего",
       by["earthquake"]["reason"]["code"] == "by_factor" and by["earthquake"]["reason"].get("factor") == "seismic"
       and "Сейсмическая зона" in by["earthquake"]["why"] and by["fire"]["level"] in ("low", "moderate")
       and all(i["level"] in ("low", "moderate", "high") and i["why"] for i in items), by["earthquake"])
    sc = an["scenarios"]
    ok("40б: сценарии — формула с числами, пожар по объекту и землетрясение по площадке 8 баллов",
       sc["available"] and any("землетрясение: 4 200 000 000 сум (вся площадка, 8 баллов)" in x["formula"].replace(" ", " ")
                               for x in sc["items"]), [x["formula"] for x in sc["items"]])
    ok("40б: «что если» — спринклеры снижают EML (посчитано модулем)",
       any(w["change"] == "protection" and w["value"] == "sprinkler" and w["eml"] < sc["items"][1]["amount"]
           for w in sc["whatif"]), sc["whatif"])
    ret = an["retention"]
    ok("40б: удержание — 20 % × (средства + резервы) по Положению 1806, EML в пределах, «временно», оценка",
       ret["known"] and ret["verdict"] == "within" and ret["status"] == "temporary"
       and "20 % × (собственные средства" in ret["text"] and "Положению № 1806, п. 15" in ret["text"]
       and ret["estimate"] is True and "оценка, не факт" in (ret["estimate_note"] or ""), ret)
    s1 = {r["label"]: r for r in a["sections"][0]["rows"]}
    ok("40б: раздел 1 здания — вид, адрес, кадастр, конструкция (введена), год постройки",
       s1.get("Вид объекта", {}).get("value") == "склад" and "Адрес / место нахождения" in s1
       and s1.get("Конструкция, материал стен", {}).get("value") == "железобетон, кирпич"
       and ("Кадастровый номер" in s1 or "кадастровый номер" in s1.get("Не указано", {}).get("value", "")), s1)
    AN_REPORT["склад кл. 8"] = {"PML/EML/MFL": [x["amount"] for x in sc["items"]], "балл": an["score"]["score"],
                                "уровни": {k: v["level"] for k, v in by.items()}}

    # --- 40в. склад, класс 9 ---
    st, a = call("POST", "/act/make", {"lang": "ru", "must": WH_MUST, "optional": WH_OPT})
    an = an_common("40в склад кл. 9", a, WH_MUST["sum_insured"], 365)
    items = an["risks"]["items"]
    ok("40в: риски класса 9 (кража со взломом, град…) — сумма 100 %",
       abs(sum(i["share_of_net_pct"] for i in items) - 100) <= 0.3 and items[0]["code"] == "burglary", items[:2])
    ok("40в: сейсмозона в балл класса 9 не входит", not next(c for c in an["score"]["components"]
                                                          if c["code"] == "seismic")["applicable"])
    ok("40в: сценарии класса 9 — кража, залив по помещению",
       an["scenarios"]["available"] and all(x["parts"][0]["peril"] == "damage9" for x in an["scenarios"]["items"]))

    # --- 40г. оборудование класса 8 по запросу филиала (продукт 0832, 47 397 852 345,04, 1 888 дн.) ---
    st, b = upload([("sorov2.docx", DOCX_MIME, docx_table(BR_SAMPLE2))], {"lang": "ru"})
    t0 = monotonic()
    st, a = br_make(b["session"], 2, br_request(b))
    sec = monotonic() - t0
    ok("40г оборудование: акт быстрее 2 с", st == 200 and sec < 2, sec)
    an = an_common("40г оборудование", a, BR_S2, 1888)
    fx = {f["code"]: f for f in an["factors"]["items"]}
    ok("40г: деятельность — по описанию документа («нон махсулотлари» → пищевое производство), не «склад»",
       fx["activity"]["option"] == "food" and fx["activity"]["source"] == "text"
       and "по описанию объекта" in fx["activity"]["source_label"], fx["activity"])
    lvl = a["risk"]["level"]
    applied = round(max(0.08 * (1 + ae.DEFAULT_SETTINGS["adj_pct"][lvl] / 100), 0.08), 4)
    ok("40г: правила тарифа акта не изменились (0,08 % × поправка, премия на 1 888 дн.)",
       a["rate"]["applied_pct"] == applied and a["premium"]["amount"] == round(BR_S2 * applied / 100 * 1888 / 365),
       (a["rate"]["applied_pct"], a["premium"]["amount"]))
    tb = an["tariff"]
    ok("40г: такафул — нагрузка 25 % без прибыли компании, техническая ставка справочно",
       tb["takaful"] and abs(tb["load_share"] - 0.25) < 1e-6 and tb["technical_pct"] > tb["act_rate_pct"], tb)
    s1rows = a["sections"][0]["rows"]
    s1 = {r["label"]: r for r in s1rows}
    ok("40г: раздел 1 оборудования — наименование на языке акта, текст документа в примечании",
       s1.get("Наименование", {}).get("value") == "машины и оборудование — пищевое производство"
       and "Технологик асбоб ускуна нон" in (s1["Наименование"].get("note") or "")
       and "по словарю" in s1["Наименование"]["note"], s1.get("Наименование"))
    ok("40г: строки оборудования (производитель, модель, заводской номер, год, место установки), "
       "«данные недоступны» ≤ 3, остальное — «Не указано»",
       sum(1 for r in s1rows if r["value"] == "данные недоступны") <= 3 and "Не указано" in s1
       and "Габариты" not in s1 and "Мощность двигателя" not in s1
       and all(x in " ".join(r["label"] + " " + str(r["value"]) for r in s1rows).lower()
               for x in ("модель", "заводской", "год выпуска", "место установки", "производитель")),
       [(r["label"], r["value"]) for r in s1rows])
    ok("40г: удержание — EML выше (защита не указана) → перестрахование или решение андеррайтера",
       an["retention"]["verdict"] in ("eml_excess", "mfl_excess", "within") and an["retention"]["known"]
       and "Вывод:" in an["retention"]["text"], an["retention"])
    fr = an["franchise"]
    with db.tx() as con:
        ctx = ax.ra_context(con, cls="8", product_code="0832", otype="Машины и оборудование", group="equipment",
                            kind="equipment", S=BR_S2, V=BR_S2, region="Ташкентская область", term_days=1888,
                            year=None, o={}, recognized=[], text="нон махсулотлари ишлаб чиқариш")
        em = ax._engine_multiplier(con, ctx, 1.0)
    rr = {"applied_pct": a["rate"]["applied_pct"], "min_pct": a["rate"]["min_pct"], "term_days": 1888}
    rate1, prem1, _fl = ax.apply_multiplier(rr, em["mult"], BR_S2)
    row1 = next(r for r in fr["rows"] if r["pct"] == 1.0)
    ok("40г: франшиза 1 % = ставка акта × множитель what_if (сверка с прямым вызовом), экономия = разница",
       row1["premium"] == prem1 and row1["rate_pct"] == rate1
       and row1["saving"] == a["premium"]["amount"] - prem1, (row1, prem1))
    ms = an["measures"]
    ok("40г: мероприятия — эффект на техническую ставку и премию акта", "items" in ms and all(
        "техническая ставка" in m["text"] or "не влияет" in m["text"] for m in ms["items"]), ms)
    ok("40г: вывод про франшизу прежний — «не требуется»", a["franchise"]["text"] == "Франшиза не требуется")
    aid_eq = a["id"]
    AN_REPORT["оборудование 0832"] = {"уровень акта": lvl, "тариф акта": a["rate"]["applied_pct"],
                                      "премия": a["premium"]["amount"], "техническая": tb["technical_pct"],
                                      "рынок": tb["market_rate_pct"], "балл": an["score"]["score"],
                                      "PML/EML/MFL": [x["amount"] for x in an["scenarios"]["items"]],
                                      "удержание": an["retention"].get("limit"),
                                      "франшизы": [(r["pct"], r["premium"], r["saving"]) for r in fr["rows"]]}

    # скан запроса: перевод описания — от модели (подменена), на языке акта
    model_on(True)
    REPLY["text"] = br_model_reply(2)
    st, b3 = upload([("scan2.png", "image/png", image((250, 250, 250)))], {"lang": "ru"})
    st, a3 = br_make(b3["session"], 2, br_request(b3))
    s13 = {r["label"]: r for r in a3["sections"][0]["rows"]}
    ok("40г: скан — наименование переводом модели, исходный текст в примечании",
       "хлебобулочных" in str(s13.get("Наименование", {}).get("value")) and "перевод модели" in (
           s13["Наименование"].get("note") or ""), s13.get("Наименование"))
    model_on(False)

    # --- 40д. три языка ---
    for lang in ("uz", "en"):
        st, x = call("GET", f"/act/{aid_eq}", params={"lang": lang})
        s4 = x["sections"][3]
        texts = [li["title"] for li in s4["lists"]] + s4["paragraphs"] + \
                [c for li in s4["lists"] if li.get("table") for c in li["table"]["columns"]] + \
                [str(c) for li in s4["lists"] if li.get("table") for r in li["table"]["rows"] for c in r] + \
                [n for li in s4["lists"] for n in li.get("notes") or []] + x["analytics"]["summary"]["sentences"]
        cyr = [t_ for t_ in texts if re.search(r"[А-Яа-яЁё]", t_ or "")]
        ok(f"40д {lang}: аналитика раздела 4 без кириллицы (заголовки, таблицы, резюме)", not cyr, cyr[:4])
        nm = x["sections"][0]["rows"][2]
        ok(f"40д {lang}: наименование объекта — на языке акта, текст документа — в примечании",
           not re.search(r"[А-Яа-яЁё]", str(nm["value"])) and "Технологик" in (nm.get("note") or ""), nm)

    # --- 40е. Word и PDF ---
    st, blob, h = call("GET", f"/act/{aid_eq}.docx", raw=True)
    plain = re.sub(r"<[^>]+>", "", zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf-8"))
    need = ["Разбор по рискам", "Учтённые факторы", "Что изменит ставку", "Состав тарифа", "Сценарии убытка подробно",
            "Балл риска 0–100", "Рынок и статистика", "Франшиза: варианты", "Источник: НАПП", "Кратко:"]
    ok("40е: DOCX — раздел 4 с таблицами аналитики и источниками", all(x in plain for x in need),
       [x for x in need if x not in plain])
    st, blob, h = call("GET", f"/act/{aid_eq}.pdf", raw=True)
    text = pdf_text(pymupdf.open(stream=blob, filetype="pdf"))
    ok("40е: PDF — те же блоки", all(x in text for x in need), [x for x in need if x not in text])

    # --- 40ж. старый акт без аналитики показывается ---
    with db.tx() as con:
        row = db.rows(con, "SELECT act_json FROM acts WHERE id=?", aid_eq)[0]
    stored = _json.loads(row["act_json"])
    D = stored["data"]
    D.pop("analytics", None)
    D.pop("object_doc", None)
    old = act.render(D, "ru", stored["meta"])
    ok("40ж: старый акт без блока analytics показывается (available = false)",
       old["analytics"]["available"] is False and old["analytics"]["reason"] == "old_act"
       and len(old["sections"]) == 5)

    # --- 40з. словарь деятельности и вида объекта (три языка) ---
    ok("40з: деятельность по описанию — ru, uz кириллица и латиница, en; «анонс» — не хлеб",
       ax.activity_from_text("Технологик асбоб ускуна нон махсулотлари ишлаб чиқариш учун") == "food"
       and ax.activity_from_text("хлебопекарное производство") == "food"
       and ax.activity_from_text("non mahsulotlari ishlab chiqarish uskunasi") == "food"
       and ax.activity_from_text("bakery equipment") == "food"
       and ax.activity_from_text("холодильник для хранения сельхозпродукции") == "warehouse"
       and ax.activity_from_text("АЗС и склад ГСМ") == "flammable"
       and ax.activity_from_text("анонс оборудования") is None and ax.activity_from_text("") is None)
    ok("40з: склад-холодильник узнаётся по «музлатгич» и «холодильник»",
       ax.cold_store("қишлоқ хўжалиги махсулотларини сақлаш учун музлатгич")
       and ax.cold_store("холодильник") and not ax.cold_store("склад"))
    ok("40з: пороги уровня риска помечены экспертными", aa.CALIBRATED == 0 and aa.PERIL_LEVEL["low_max"] < 1)
    del mp, ra


# ------------------------------------------------------------------ 41. замечания контролёра по аналитике (30.09.2026)

ROOT_DIR = Path(__file__).resolve().parent.parent
EQ_MUST = {"product_code": "0832", "sum_insured": BR_S2, "object_value": BR_S2, "region": "tashkent_region"}
EQ_OPT = {"term_days": 1888, "object_kind": "equipment", "activity": "food",
          "object_type": "Технологическое оборудование для производства хлебобулочных изделий"}
# «движок» как жаргон (двигатель техники — «Dvigatelni bloklash», «engine immobilisation» — не жаргон)
JARGON = re.compile(r"движ(?!ени)|dvigatel(?!ni)|\bengine\b(?! hours| lock| immobil)|Балл старого|what_if", re.I)


def s4_texts(a):
    """Все строки раздела 4: абзацы, строки, списки, таблицы, примечания."""
    s4 = a["sections"][3]
    out = list(s4["paragraphs"]) + [r["label"] + " " + str(r["value"]) + " " + str(r.get("note") or "") for r in s4["rows"]]
    for li in s4["lists"]:
        out += [li["title"]] + list(li["items"]) + list(li.get("notes") or [])
        if li.get("table"):
            out += list(li["table"]["columns"]) + [str(c) for r in li["table"]["rows"] for c in r]
    return out


def s4_list(a, title):
    return next((li for li in a["sections"][3]["lists"] if li["title"] == title), {"items": []})


def check_review_fixes():
    print("41. Замечания контролёра по аналитике: регион, удержание-оценка, тип по умолчанию, доли, рынок, балл, "
          "жаргон, уровни, склонения, вид документа, франшиза от неокруглённой ставки")
    from app import act_analytics as aa, act_extras as ax, market_picture as mp
    fresh()
    model_on(False)

    # --- 41.2. регион кодом экрана и названием — одна и та же статистика и один балл ---
    st1, a1 = call("POST", "/act/make", {"lang": "ru", "must": EQ_MUST, "optional": EQ_OPT})
    st2, a2 = call("POST", "/act/make", {"lang": "ru", "must": dict(EQ_MUST, region="Ташкентская область"),
                                         "optional": EQ_OPT})
    an1, an2 = a1["analytics"], a2["analytics"]
    strip = lambda st: [{k: v for k, v in i.items() if k != "text"} for i in st["indicators"]]   # noqa: E731
    ok("41.2: tashkent_region и «Ташкентская область» — одинаковые показатели региона и балл",
       st1 == st2 == 200 and strip(an1["stats"]) == strip(an2["stats"])
       and an1["score"]["score"] == an2["score"]["score"] == 50.9, (an1["score"]["score"], an2["score"]["score"]))
    vh = next(i for i in an1["stats"]["indicators"] if i["id"] == "vulnerable_housing")
    ok("41.2: показатель региона — Ташкентская область, 49,95 % (регион, а не республика)",
       vh["scope"] == "region" and abs(vh["value"] - 49.95) < 1e-9 and "49,95" in vh["value_text"], vh)
    with db.tx() as con:
        ok("41.2: все 14 кодов регионов экрана узнаются модулями по названию",
           all(mp.resolve_region(act.region_for_modules({"region": c, "region_code": c}))[0]
               for c in act._region_names()) and len(act._region_names()) == 14)
    del con

    # --- 41.3. удержание — оценка: норма на временных цифрах, таблица линий — экспертная ---
    ret = an1["retention"]
    lines = " ".join(ret["lines"])
    ok("41.3: verdict сохранён, добавлен estimate_note", ret["verdict"] == "eml_excess" and ret["estimate"] is True
       and ret["estimate_note"] and ret["estimate_note"].startswith("Это оценка, не факт"), ret)
    ok("41.3: лимит по Положению 1806 п. 15 — «цифры временные, до данных бухгалтерии» (company_financials)",
       "по Положению № 1806, п. 15 = 20 % × (собственные средства" in nb(lines)
       and "цифры временные, до данных бухгалтерии (источник: company_financials)" in lines, ret["lines"])
    ok("41.3: «страховая сумма 47,4 млрд в лимит 84 млрд укладывается»",
       "Страховая сумма 47 397 852 345 сум в лимит 84 000 000 000 сум укладывается" in nb(lines), ret["lines"])
    ok("41.3: таблица линий класса 8 — внутреннее экспертное правило, не норма, не калибровано",
       "Лимит по таблице линий класса 8 — 43 409 395 973 сум: внутреннее экспертное правило "
       "(capacity.retention_table), не норма, не калибровано" in nb(lines), ret["lines"])
    phrase = ("EML выше расчётного удержания по экспертной таблице — рекомендуем рассмотреть перестрахование или "
              "решение андеррайтера (оценочно, цифры временные)")
    ok("41.3: в «Кратко» — рекомендация, а не факт", any(phrase in s for s in an1["summary"]["sentences"]),
       an1["summary"]["sentences"])
    row = next(r for r in a1["sections"][3]["rows"] if r["label"] == "Лимит собственного удержания")
    ok("41.3: строка раздела 4 — «рекомендуем рассмотреть…», без «нужно перестрахование»",
       "рекомендуем рассмотреть перестрахование или решение андеррайтера (оценочно, цифры временные)" in row["note"]
       and "до данных бухгалтерии" in row["note"] and "нужно перестрахование (оценочно)" not in row["note"], row)
    how = s4_list(a1, "Как посчитаны сценарии убытка")["items"]
    ok("41.3: «как посчитаны сценарии» — таблица линий не выдана за норму",
       any("внутреннее экспертное правило, не норма" in h and "расчётное удержание" in h for h in how), how)
    for lang in ("uz", "en"):
        st, x = call("GET", f"/act/{a1['id']}", params={"lang": lang})
        r = x["analytics"]["retention"]
        ok(f"41.3 {lang}: удержание — оценка на языке акта (estimate_note, без кириллицы)",
           r["estimate_note"] and "company_financials" in r["estimate_note"]
           and not re.search(r"[А-Яа-яЁё]", " ".join(r["lines"])), r["lines"])

    # --- 41.4. пример оборудования для снимков: запрос филиала даёт вид, деятельность и описание ---
    br_file = ROOT_DIR / "sandbox" / "br30" / "sorov_equipment.docx"
    ok("41.4: образец запроса филиала для снимков оборудования лежит в sandbox/br30", br_file.exists())
    if br_file.exists():
        st, b = upload([("sorov_equipment.docx", DOCX_MIME, br_file.read_bytes())], {"lang": "ru"})
        st, ae_ = call("POST", "/act/make", {"session": b["session"], "lang": "ru", "recognized": b["recognized"],
                                             "must": EQ_MUST, "optional": {"term_days": 1888,
                                                                           "request": br_request(b)}})
        base = next(r for r in ae_["analytics"]["tariff"]["rows"] if r["code"] == "base_net")
        ok("41.4: по запросу филиала — оборудование 0,20 % и пищевое производство (не «производственное здание»)",
           st == 200 and "машины и оборудование" in base["label"] and base["value"] == act.pct(0.2, "ru")
           and next(f for f in ae_["analytics"]["factors"]["items"] if f["code"] == "activity")["option"] == "food",
           base)

    # --- 41.5. оборудование без документа: тип по умолчанию назван, деятельность согласована с типом ---
    st, a0 = call("POST", "/act/make", {"lang": "ru", "must": EQ_MUST, "optional": {"term_days": 1888}})
    s1 = {r["label"]: r for r in a0["sections"][0]["rows"]}
    ok("41.5: раздел 1 — «принят по умолчанию: производственное здание», а не «данные недоступны»",
       s1["Вид объекта"]["value"] == "принят по умолчанию: производственное здание"
       and "уточните" in s1["Вид объекта"]["note"], s1.get("Вид объекта"))
    base = next(r for r in a0["analytics"]["tariff"]["rows"] if r["code"] == "base_net")
    ok("41.5: состав тарифа — пометка «вид объекта принят по умолчанию»",
       "вид объекта принят по умолчанию: производственное здание" in base["note"], base)
    fx = {f["code"]: f for f in a0["analytics"]["factors"]["items"]}
    asm = s4_list(a0, "Принято по умолчанию (уточните)")["items"]
    ok("41.5: деятельность по умолчанию согласована с типом (производство → пищевое производство, не склад)",
       fx["activity"]["option"] == "food" and fx["activity"]["source"] == "default"
       and any("по типу объекта «производственное здание»" in x for x in asm)
       and not any("склад общего назначения" in x for x in asm), (fx["activity"], asm))
    with db.tx() as con:
        ok("41.5: правило согласования — таблица модуля (производство → food, склад → warehouse)",
           ax._default_activity(con, "8", "0832", None) == ("Производство", "food")
           and ax._default_activity(con, "8", "0807", "Склад") == ("Склад", "warehouse"))
    ok("41.5: фактор по умолчанию помечен в «Кратко» и в разборе рисков",
       any("(×1,2, принято по умолчанию)" in s for s in a0["analytics"]["summary"]["sentences"])
       and "принято по умолчанию" in a0["analytics"]["risks"]["items"][0]["why"])

    # --- 41.6. доли рисков: округление; автокран — без «Доли из справочника» ---
    notes8 = a0["analytics"]["risks"]["notes"]
    ok("41.6: класс 8 — «сумма 99,8 % — из-за округления долей»",
       any("сумма 99,8 % — из-за округления долей" in nb(n) for n in notes8), notes8)
    st, ac = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST, "optional": CRANE_OPT})
    nc = ac["analytics"]["risks"]["notes"]
    ok("41.6: автокран — «не разбит на отдельные риски», без «Доли — из справочника рисков»",
       any("не разбит на отдельные риски" in n for n in nc) and not any("Доли — из справочника" in n for n in nc), nc)

    # --- 41.7. строка рынка: полный год отдельно, какая строка отчёта взята — по market_stats ---
    mk_items = s4_list(a1, "Рынок и статистика")["items"]
    with db.tx() as con:
        last = db.rows(con, "SELECT MAX(report_date) d FROM market_stats WHERE row_key='cls8_9'")[0]["d"]
        pk = db.rows(con, "SELECT premiums_ytd p FROM market_stats WHERE row_key='cls8_9' AND report_date=?", last)[0]["p"]
        p8 = db.rows(con, "SELECT premiums_ytd p FROM market_stats WHERE row_key='cls8' AND report_date=?", last)[0]["p"]
    fy = [x for x in mk_items if x.startswith("За ") and "год: ставка" in x]
    ok("41.7: «За 2025 год: ставка …, убыточность …» — отдельной строкой, не хвостом убыточности среза",
       len(fy) == 1 and not any("за 20" in x for x in mk_items if x.startswith("Убыточность класса")), mk_items[:4])
    mli = s4_list(a1, "Рынок и статистика")
    trows = {r[0]: r for r in mli["table"]["rows"]}
    ok("41.7: в таблице документа полный год — две строки: ставка и убыточность",
       "Рыночная ставка за 2025 год" in trows and "Убыточность рынка за 2025 год" in trows
       and ";" not in trows["Рыночная ставка за 2025 год"][1], list(trows))
    ok("41.7: пояснение о строке отчёта — под таблицей документа",
       any(n.startswith("Взята строка классов 8 и 9 (cls8_9)") for n in mli["notes"]), mli.get("notes"))
    rown = [x for x in mk_items if x.startswith("Взята строка классов 8 и 9 (cls8_9)")]
    ok("41.7: «взята строка классов 8 и 9 (cls8_9); отдельная строка класса 8 мала» — числа из market_stats",
       len(rown) == 1 and act.tx._num(round(pk), "ru") + " млн сум" in rown[0]
       and act.tx._num(round(p8), "ru") + " млн сум" in rown[0] and "мала по объёму" in rown[0]
       and "одной строкой" not in " ".join(mk_items), rown)

    # --- 41.8. доля глинобитного жилья — только для зданий и складов ---
    comp1 = next(c for c in an1["score"]["components"] if c["code"] == "external_stats")
    vh1 = next(i for i in an1["stats"]["indicators"] if i["id"] == "vulnerable_housing")
    ok("41.8: оборудование — жилой фонд не в балле, честное «нет показателей региона для этого вида объекта»",
       not comp1["applicable"] and comp1["why"] == "нет показателей региона для этого вида объекта"
       and vh1["used_in_score"] is False and vh1["excluded_for_kind"] and "в балл не входит" in vh1["text"], comp1)
    comp0 = next(c for c in a0["analytics"]["score"]["components"] if c["code"] == "external_stats")
    ok("41.8: здание (тип по умолчанию «производство») — показатель жилого фонда в балле",
       comp0["applicable"] and comp0["points"] > 0, comp0)
    ok("41.8: балл без показателя — та же формула модуля (сумма вкладов = балл)",
       abs(sum(c["contribution"] or 0 for c in an1["score"]["components"] if c["applicable"])
           - an1["score"]["score"]) < 0.3)
    st, am_ = call("POST", "/act/make", {"lang": "ru", "must": dict(EQ_MUST, region="Марс"), "optional": EQ_OPT})
    cm = next(c for c in am_["analytics"]["score"]["components"] if c["code"] == "external_stats")
    ok("41.8: регион не распознан — так и написано, а не «нет показателей для класса»",
       not cm["applicable"] and cm["why"].startswith("регион не распознан"), cm)

    # --- 41.9. без жаргона на трёх языках ---
    for tag, aid in (("оборудование", a1["id"]), ("автокран", ac["id"])):
        for lang in ("ru", "uz", "en"):
            st, x = call("GET", f"/act/{aid}", params={"lang": lang})
            bad = [s for s in s4_texts(x) + x["analytics"]["summary"]["sentences"] + x["franchise"]["how"]
                   if JARGON.search(s or "")]
            ok(f"41.9 {tag} {lang}: в разделе 4 нет «движок», «Балл старого движка», what_if", not bad, bad[:3])
    st, x = call("GET", f"/act/{a1['id']}", params={"lang": "ru"})
    ok("41.9: «Балл риска (справочно)», «расчёт при других данных объекта», «посчитано расчётным модулем»",
       x["analytics"]["score"]["text"].startswith("Балл риска (справочно) — ")
       and "Сценарии «что если» (расчёт при других данных объекта)" in s4_titles(x)
       and "Что изменит ставку (посчитано расчётным модулем)" in s4_titles(x))

    # --- 41.10. одни и те же слова уровня у рисков и у акта ---
    ok("41.10: уровни рисков и акта — низкий / умеренный / высокий на трёх языках",
       all(tx.PERIL_LEVEL_LABELS[k] == tx.LEVEL_LABELS[k] for k in ("low", "moderate", "high")))
    lv = {i["level_label"] for i in an1["risks"]["items"]}
    ok("41.10: в разборе рисков нет «средний»", lv <= {"низкий", "умеренный", "высокий"}
       and not any("иначе средний" in n or "уровень средний" in n for n in an1["risks"]["notes"])
       and not any("уровень средний" in i["text"] for i in an1["risks"]["items"]), lv)

    # --- мелочи: склонения, вид документа, франшиза от неокруглённой ставки ---
    ok("склонение: 1 балл, 2 балла, 5 баллов, 11 баллов, 21 балл, 24 балла, 112 баллов, 2,5 балла",
       [tx.count_text(n, "points", "ru", 1 if n == 2.5 else 0) for n in (1, 2, 5, 11, 21, 24, 112, 2.5)]
       == ["1 балл", "2 балла", "5 баллов", "11 баллов", "21 балл", "24 балла", "112 баллов", "2,5 балла"]
       and tx.count_text(24, "points", "uz") == "24 ball" and tx.count_text(1, "points", "en") == "1 point")
    sc_items = s4_list(a1, "Балл риска 0–100 (справочно)")["items"]
    ok("склонение в акте: «24 балла», «32 случая»",
       any("MFL к лимиту удержания: 24 балла ×" in x for x in sc_items)
       and any("32 случая" in nb(x) for x in mk_items) and not any("24 баллов" in x for x in sc_items),
       [x for x in sc_items if "MFL" in x])
    if br_file.exists():
        for lang, want in (("ru", "запрос филиала"), ("uz", "filial soʻrovi"), ("en", "branch request")):
            st, x = call("GET", f"/act/{ae_['id']}", params={"lang": lang})
            r2 = [r for r in x["sections"][1]["rows"] if want in str(r["value"])]
            ok(f"вид документа в разделе 2 на языке акта ({lang}): «{want}»", bool(r2)
               and (lang == "ru" or "запрос филиала" not in " ".join(str(r["value"]) for r in x["sections"][1]["rows"])),
               x["sections"][1]["rows"])
    f05 = next(r for r in a0["analytics"]["franchise"]["rows"] if r["pct"] == 0.5)
    rr = a0["rate"]
    raw = rr["applied_pct"] * f05["multiplier"]
    ok("франшиза 0,5 %: премия от неокруглённой ставки (216 534 374), ставка показана округлённой",
       f05["premium"] == round(raw / 100 * BR_S2 * 1888 / 365) == 216_534_374
       and f05["rate_pct"] == round(raw, 4), (f05, raw))
    ok("франшиза: apply_multiplier — премия от неокруглённой ставки, ставка до 4 знаков",
       ax.apply_multiplier({"applied_pct": 0.096, "min_pct": 0.08, "term_days": 365}, 0.9197, 1e9)
       == (0.0883, round(0.096 * 0.9197 / 100 * 1e9), False))


# ================================================================================================
#  40. Бланк договора компании (шаблон с подчёркиваниями) и тот же шаблон заполненный — 30.09.2026
#  Копия структуры шаблона «Договор №____ / Страхования спортсменов от несчастных случаев» на выдуманных
#  данных: заголовок в две строки, подчёркивания, таблица приложения 1, пункты 2.6 и 5.4. Word-разметка —
#  как у настоящего файла: слова разрезаны на прогоны w:r с rsid, закладки, проверка правописания, табуляции.
# ================================================================================================

TPL_DIRECTOR = "Т.Т. Тестов"                    # выдуманный руководитель в шапке («УТВЕРЖДАЮ»)
TPL_SIGNER = "Сидоров Сидор Сидорович"          # выдуманный представитель страхователя (заполненный договор)
TPL_ATHLETE = "Спортов Тест Тестович"           # выдуманный спортсмен в списке приложения 1
TPL_LICENSE = "00099"
TPL_REQ_INSURER = ("СТРАХОВЩИК: АО СО «INSON» Адрес: ______________________________ тел: "
                   "_______________________________ факс: ______________________________ р/с: "
                   "_______________________________ в __________________________________ МФО: "
                   "_______________________________ ИНН: _______________________________ ОКОНХ: "
                   "_____________________________")
TPL_REQ_HOLDER = ("СТРАХОВАТЕЛЬ: _______________________ Адрес: ______________________________ тел: "
                  "_______________________________ факс: ______________________________ р/с: "
                  "_______________________________ в __________________________________ МФО: "
                  "_______________________________ ИНН: _______________________________ ОКОНХ: конец реквизитов")


def tpl_blocks(filled: bool = False) -> list:
    """Блоки шаблона: ("p", текст) | ("sdt", текст) | ("tbl", строки). filled — те же места заполнены."""
    u = lambda n: "_" * n                                                      # noqa: E731
    no = "Договор № 17-НС/2026" if filled else "Договор №" + u(12)
    place_date = ("г. Ташкент\t\t\t\t\t \t\t   «1» октября 2026 г." if filled
                  else "г. " + u(15) + "\t\t\t\t\t \t\t   «____» ________ 20___г.")
    holder = ("ООО «Спорт Клуб Тест», именуемое в дальнейшем «Страхователь», в лице директора " + TPL_SIGNER
              if filled else u(48) + ", именуемая в дальнейшем «Страхователь», в лице " + u(44))
    s_ins = "600 000 000 (шестьсот миллионов)" if filled else u(28) + " (" + u(60) + ")"
    prem = "9 000 000 (девять миллионов)" if filled else u(33) + " (" + u(60) + ")"
    term = ("с «1» октября 2026 года по «30» сентября 2027 года" if filled
            else "с «_____» ___________ 20___ года по «_____» _________ 20____ года")
    sched = [["Профессия (род занятия)", "Количество застрахованных \nлиц", "Персональная страховая \nсумма (сум)",
              "Процентная ставка (%)", "Страховой \nплатеж за одного застрахованного лица (сум)",
              "Страховая \nсумма ВСЕГО (гр2 х гр3)", "Страховая \nпремия ВСЕГО (гр2 х гр5)"],
             ["1", "2", "3", "4", "5", "6", "7"]]
    if filled:
        sched.append(["Футболист", "20", "30 000 000", "1,5", "450 000", "600 000 000", "9 000 000"])
    else:
        sched.append(["", "", "", "", "", "", ""])
    names = [["№ п/п", "Фамилия, Имя и Отчество", "Персональная страховая сумма, сум", "Выгодоприобретатель"],
             ["1", "2", "3", "4"]]
    for k in range(1, 16):
        who = TPL_ATHLETE if filled and k == 1 else ""
        names.append([f"{k}.", who, "30 000 000" if who else "", ""])
    names.append(["", "Итого:", "", ""])
    sign = [["ПОДПИСИ СТОРОН:", ""],
            ["От имени Страховщика: Генеральный директор/ Директор " + u(17) + " филиала/ Иное уполномоченное лицо "
             + u(18) + " (Ф.И.О.)", "От имени Страхователя: " + u(25) + " должность Ф.И.О."],
            [u(25) + " подпись\t\t м.п.", u(25) + " подпись\t\t м.п."]]
    filler = [("p", f"7.{k}. Стороны обязуются добросовестно исполнять условия настоящего Договора, своевременно "
                    f"информировать друг друга об изменении реквизитов, адресов и банковских счетов, соблюдать "
                    f"конфиденциальность сведений, полученных при исполнении настоящего Договора, и не передавать их "
                    f"третьим лицам без письменного согласия другой стороны, за исключением случаев, прямо "
                    f"предусмотренных законодательством Республики Узбекистан (пункт {k}).") for k in range(1, 61)]
    return [
        ("p", "«УТВЕРЖДАЮ»"), ("p", "Генеральный директор"), ("p", "АО СО «INSON»"), ("p", TPL_DIRECTOR),
        ("p", "Приложение №___"), ("p", "К приказу №___ от «___» _________ 20___г."),
        ("p", no), ("p", "Страхования спортсменов от несчастных случаев"), ("p", place_date),
        ("p", "Акционерное Общество Страховая организация «INSON», действующее на основании Лицензии на осуществление "
              "страховой деятельности серия ТС № " + TPL_LICENSE + " от «5» мая 2021 года, выданной уполномоченным "
              "государственным органом по регулированию страхового рынка Республики Узбекистан, именуемое в "
              "дальнейшем «Страховщик», в лице " + u(47) + ", действующего на основании " + u(21) + ", с одной "
              "стороны, и " + holder + ", действующего на основании " + u(24) + ", с другой стороны заключили "
              "настоящий Договор о нижеследующем:"),
        ("p", "1. ПРЕДМЕТ ДОГОВОРА"),
        ("p", "Страховщик обязуется в соответствии с условиями настоящего Договора выплатить при наступлении "
              "страхового случая Застрахованным лицам, указанным в Приложении 1 к настоящему Договору, обусловленную "
              "сумму, при условии, что Страхователь обязуется оплатить страховую премию в размере и сроки, указанные "
              "в настоящем Договоре."),
        ("p", "2. ОПРЕДЕЛЕНИЯ"),
        ("p", "2.1. Страховой Полис – документ, удостоверяющий факт заключения настоящего Договора."),
        ("p", "2.2. Страховая сумма – сумма денежных средств, представляющая собой предельный объем обязательств "
              "Страховщика."),
        ("p", "2.3. Страховая премия – плата за страхование, уплачиваемая Страхователем Страховщику."),
        ("p", "2.4.\tЗастрахованное лицо (спортсмен) – физическое лицо, профессионально занимающееся спортом, чьи "
              "имущественные интересы, связанные с жизнью и здоровьем, являются объектом страхования."),
        ("p", "2.5. Выгодоприобретатель – физическое лицо, названное в Приложении 1 к настоящему Договору, в качестве "
              "получателя страховой выплаты."),
        ("p", "2.6.\tСтраховой случай – получение травматического повреждения или смерть Застрахованного лица в "
              "результате несчастного случая, произошедшего во время спортивного соревнования в течение периода "
              "страхования, с наступлением которого возникает обязанность Страховщика произвести страховую выплату."),
        ("p", "2.7.\tНесчастный случай – внезапное, кратковременное событие, которое извне воздействует на организм "
              "человека."),
        ("p", "3. страховое ПОКРЫТИЕ"),
        ("p", "3.1. Страховая защита предоставляется Застрахованным лицам от несчастных случаев, произошедших во время "
              "участия в спортивных соревнованиях и приведших к:"),
        ("p", "3.1.1. травматическим повреждениям Застрахованного лица;"),
        ("p", "3.1.2. смерти Застрахованного лица."),
        ("p", "4. ОБЩИЕ ИСКЛЮЧЕНИЯ"),
        ("p", "4.1. По настоящему Договору не признаются страховым случаем события, произошедшие вследствие:"),
        ("p", "а) военных действий и их последствий, народных волнений и забастовок;"),
        ("p", "б) ядерного взрыва, радиации и радиоактивного заражения;"),
        ("p", "в) умышленных действий Страхователя или Застрахованного лица, направленных на наступление страхового "
              "случая;"),
        ("p", "г) доказанного факта применения допинга;"),
        ("p", "д) нахождения Застрахованного лица в состоянии алкогольного, наркотического или токсического "
              "опьянения;"),
        ("p", "е) самоубийства или покушения на самоубийство Застрахованного лица;"),
        ("p", "ж) совершения Застрахованным лицом умышленного преступления."),
        ("p", "5. CТРАХОВАЯ СУММА И СТРАХОВАЯ ПРЕМИЯ"),
        ("p", "5.1. Общая страховая сумма по настоящему Договору составляет " + s_ins + " сум."),
        ("p", "5.2. Персональная страховая сумма, установленная для каждого Застрахованного лица, указана в "
              "Приложении 1 к настоящему Договору."),
        ("p", "5.3. Страховая премия по настоящему Договору составляет " + prem + " сум."),
        ("tbl", sched),
        ("p", "5.4. Страховая премия оплачивается единовременно в течение 5 (пяти) банковских дней после подписания "
              "настоящего Договора сторонами."),
        ("sdt", "5.5. Все взаиморасчеты по настоящему Договору производятся в сумах Республики Узбекистан."),
        ("p", "6. ВСТУПЛЕНИЕ В СИЛУ И СРОК ДЕЙСТВИЯ ДОГОВОРА"),
        ("p", "6.1. Настоящий Договор вступает в силу с момента подписания сторонами. Обязательства Страховщика по "
              "страховой выплате вступают в силу " + term + "."),
        ("p", "7. ПРАВА И ОБЯЗАННОСТИ СТОРОН"),
    ] + filler + [
        ("p", "8. РАССМОТРЕНИЕ СТРАХОВОЙ ПРЕТЕНЗИИ"),
        ("p", "8.1. При наступлении события, которое могло бы обосновать требование к Страховщику, Застрахованное лицо "
              "обязано:"),
        ("p", "– в течение 30 (тридцати) календарных дней после наступления события, направить Страховщику "
              "письменное заявление с указанием причин и обстоятельств наступившего события."),
        ("p", "9. ПОРЯДОК ОСУЩЕСТВЛЕНИЯ СТРАХОВОЙ ВЫПЛАТЫ"),
        ("p", "9.1. При временной потере трудоспособности Застрахованным лицом страховая выплата производится по "
              "таблице выплат, но не более 50% от персональной страховой суммы.\nПереносы строки внутри пункта "
              "сохраняются."),
        ("p", "9.2. При установлении Застрахованному лицу группы инвалидности страховая выплата производится в "
              "размере от 60% до 100% персональной страховой суммы."),
        ("p", "14. ЮРИДИЧЕСКИЕ АДРЕСА И РЕКВИЗИТЫ СТОРОН:"),
        ("tbl", [[TPL_REQ_INSURER, TPL_REQ_HOLDER]] + sign),
        ("p", "Приложение 1"), ("p", "к Договору страхования спортсменов"),
        ("p", "от несчастных случаев №" + u(16)), ("p", "от «____» _________ 20___г."),
        ("p", "СПИСОК"), ("p", "ЗАСТРАХОВАННЫХ СПОРТСМЕНОВ"),
        ("tbl", names),
        ("tbl", sign),
    ]


def _w_runs(text: str, k: int) -> str:
    """Текст абзаца так, как его пишет Word: куски по нескольку знаков в разных w:r с rsid, между ними —
    закладка и отметка правописания; «\\t» — w:tab, «\\n» — w:br."""
    out = []
    for i, part in enumerate(re.split(r"(\t|\n)", text)):
        if part == "\t":
            out.append("<w:r><w:tab/></w:r>")
            continue
        if part == "\n":
            out.append("<w:r><w:br/></w:r>")
            continue
        for j in range(0, len(part), 7):
            piece = part[j:j + 7]
            if j and j % 21 == 0:
                out.append(f'<w:proofErr w:type="spellStart"/><w:bookmarkStart w:id="{k}{j}" w:name="_x{k}{j}"/>'
                           f'<w:bookmarkEnd w:id="{k}{j}"/>')
            out.append(f'<w:r w:rsidR="00A1{k:04d}" w:rsidRPr="00B2{j:04d}"><w:rPr><w:rFonts w:ascii="Times New Roman"/>'
                       f'<w:sz w:val="24"/></w:rPr><w:t xml:space="preserve">{_x(piece)}</w:t></w:r>')
    return "".join(out)


def _w_p(text: str, k: int) -> str:
    # позиции табуляции в свойствах абзаца — не знак табуляции в тексте
    return (f'<w:p w:rsidR="00C3{k:04d}"><w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs>'
            f'<w:jc w:val="both"/></w:pPr>{_w_runs(text, k)}</w:p>')


def docx_word(blocks: list) -> bytes:
    """DOCX с разметкой, как у Word: прогоны, rsid, закладки, блок w:sdt, таблицы с w:tcPr."""
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"')
    body, k = [], 0
    for kind, val in blocks:
        k += 1
        if kind == "p":
            body.append(_w_p(val, k))
        elif kind == "sdt":
            body.append(f"<w:sdt><w:sdtPr><w:alias w:val=\"поле\"/></w:sdtPr><w:sdtContent>{_w_p(val, k)}"
                        f"</w:sdtContent></w:sdt>")
        else:
            rows = "".join("<w:tr>" + "".join(
                '<w:tc><w:tcPr><w:tcW w:w="1400" w:type="dxa"/></w:tcPr>'
                + "".join(_w_p(line, k) for line in (c.split("\n") if c else [""])) + "</w:tc>" for c in r) + "</w:tr>"
                for r in val)
            body.append(f"<w:tbl><w:tblPr><w:tblW w:w=\"0\" w:type=\"auto\"/></w:tblPr>{rows}</w:tbl>")
    xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>{"".join(body)}'
           f'<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def check_docx_reader():
    print("40а. Чтение DOCX: текст из дерева XML (прогоны, табуляции, переносы, w:sdt), пределы 200 000 знаков / "
          "3 000 абзацев / 10 000 ячеек")
    from app import act_extras as ax, ingest
    folder = Path(tempfile.mkdtemp(prefix="act-docx-"))
    try:
        p = folder / "tpl.docx"
        p.write_bytes(docx_word(tpl_blocks()))
        got = ax.read_limited(p, {})
        text = got["text"]
        ok(f"шаблон ≈{len(text) // 1000} тыс. знаков (больше 20 774) прочитан целиком, без пометки «часть»",
           len(text) > 21000 and got["truncated"] is False and got["status"] is None, (len(text), got["truncated"]))
        ok("в тексте нет разметки Word (ни «<w:», ни rsid, ни «</w:r>»)",
           "<w:" not in text and "rsid" not in text and "</w:" not in text and "w:r" not in text, text[:200])
        ok("слово, разрезанное на прогоны, склеено: «Страхования спортсменов от несчастных случаев»",
           "\nСтрахования спортсменов от несчастных случаев\n" in text)
        ok("w:tab — табуляция, w:br — перенос строки; позиции табуляции абзаца в текст не попали",
           "г. " + "_" * 15 + "\t\t\t\t\t \t\t   «____»" in text and "50% от персональной страховой суммы.\nПереносы" in text
           and "\n\t" not in text and not text.startswith("\t"))
        ok("абзац в блоке w:sdt прочитан", "5.5. Все взаиморасчеты по настоящему Договору производятся в сумах" in text)
        rq = next((t for t in got["tables"] if t["rows"] and t["rows"][0][0].startswith("СТРАХОВЩИК")), None)
        ok("строка реквизитов длиннее 500 знаков — целиком (правая ячейка не обрезана)",
           rq and rq["rows"][0][1].endswith("ОКОНХ: конец реквизитов") and len(" | ".join(rq["rows"][0])) > 500,
           rq and len(" | ".join(rq["rows"][0])))
        ok("таблица приложения 1 и список застрахованных прочитаны целиком (4 таблицы)",
           len(got["tables"]) == 4 and got["tables"][0]["rows"][0][0] == "Профессия (род занятия)",
           [t["rows"][0][:2] for t in got["tables"]])
        r_old = ingest.read_file(p)
        sdt = "5.5. Все взаиморасчеты по настоящему Договору производятся в сумах Республики Узбекистан."
        ok("потоковое чтение = прежнее чтение без пределов, плюс абзац w:sdt (прежнее его теряло)",
           r_old["text"].split() == text.replace(sdt + "\n", "").split() and sdt not in r_old["text"])
        # пределы: 2 000 абзацев и ≈190 000 знаков — целиком; 3 001 абзац — «часть»; строка таблицы > 4 000 знаков
        big = folder / "big.docx"
        big.write_bytes(docx_bytes([_x(f"{k}. " + "Страховщик обязан рассмотреть заявление в срок. " * 2)[:94]
                                    for k in range(2000)]))
        g = ax.read_limited(big, {})
        ok(f"2 000 абзацев, {len(g['text']) // 1000} тыс. знаков — прочитано целиком",
           g["truncated"] is False and len(g["text"].splitlines()) == 2000 and len(g["text"]) > 180000,
           (g["truncated"], len(g["text"])))
        many = folder / "many.docx"
        many.write_bytes(docx_bytes([f"пункт {k}" for k in range(3001)]))
        g = ax.read_limited(many, {})
        ok("3 001 абзац — предел абзацев (3 000): честная пометка «часть»",
           g["truncated"] is True and len(g["text"].splitlines()) == 3000, len(g["text"].splitlines()))
        wide = folder / "wide.docx"
        wide.write_bytes(docx_mixed(["Договор"], [[["а" * 2500, "б" * 2500]]]))
        g = ax.read_limited(wide, {})
        ok("строка таблицы длиннее 4 000 знаков — обрезана с пометкой",
           g["truncated"] is True and len(" | ".join(g["tables"][0]["rows"][0])) <= 4000)
        cells = folder / "cells.docx"
        cells.write_bytes(docx_mixed(["Договор"], [[[f"{r}-{c}" for c in range(10)] for r in range(150)]]))
        g = ax.read_limited(cells, {})
        ok("1 500 ячеек (меньше 10 000) — целиком", g["truncated"] is False
           and len(g["tables"][0]["rows"]) == 150, (g["truncated"], len(g["tables"][0]["rows"])))
        ok("предел ячеек из настроек по-прежнему работает (100)",
           ax.read_limited(cells, {"doc_max_cells": 100})["truncated"] is True)
        ok("пределы по умолчанию: 10 000 ячеек, 3 000 абзацев, 200 000 знаков, 4 000 знаков в строке DOCX, 5 с / 8 с",
           ax.DOC_LIMITS["doc_max_cells"] == 10000 and ax.DOC_LIMITS["doc_max_paras"] == 3000
           and ax.DOC_LIMITS["doc_max_text_chars"] == 200000 and ax.DOC_LIMITS["doc_max_row_chars"] == 4000
           and ax.DOC_LIMITS["doc_parse_sec"] == 5 and ax.DOC_LIMITS["doc_file_sec_pdf"] == 8
           and not ae.check_settings({"limits": dict(ae.DEFAULT_SETTINGS["limits"])})
           and ae.check_settings({"limits": dict(ae.DEFAULT_SETTINGS["limits"], doc_max_paras=1)}))
        # защита от «zip-бомбы» — прежняя: огромный распакованный объём отклоняется до разбора
        st, b = upload([("bomb.docx", DOCX_MIME, docx_bomb(60))], {"lang": "ru"})
        ok("zip-бомба в DOCX по-прежнему отклоняется до разбора", st == 422 and b["rejected"]
           and "слишком большой" in b["rejected"][0]["error"], (st, b.get("rejected")))
    finally:
        shutil.rmtree(folder, ignore_errors=True)


def check_contract_template():
    print("40б. Бланк договора компании: заголовок в две строки, пустые поля, личное страхование (класс 1)")
    from app import contract_read as cr
    fresh()
    model_on(True)
    CALLS.clear()
    REPLY["text"] = ct_ai_reply
    blob = docx_word(tpl_blocks())
    st, b = upload([("0102_contract.docx", DOCX_MIME, blob)], {"lang": "ru"})
    c = ctb(b)
    f = c.get("fields") or {}
    dump = _json.dumps(b, ensure_ascii=False)
    ok("бланк узнан как договор, прочитан целиком (без «только часть»)",
       st == 200 and c.get("detected") and not c.get("truncated")
       and not any("только часть" in n for n in b["notes"] + c.get("notes", [])), (st, b.get("notes")))
    ok("заголовок в две строки: вид — договор, название продукта — «Страхование спортсменов от несчастных случаев»",
       f.get("product_name") == "Страхование спортсменов от несчастных случаев"
       and cr.title_kind("Договор №____\nСтрахования спортсменов от несчастных случаев\nг. ____") == "contract",
       f.get("product_name"))
    ok("код продукта из имени файла не взят (0102 нет ни в полях, ни в подсказке)",
       f.get("product_code") is None and "product_code" not in (b.get("prefill") or {}), f.get("product_code"))
    ok("«Приложение № к приказу» — пометка шаблона компании", f.get("template_hint") is True
       and c.get("template_hint") is True and any("приложения к приказу" in n for n in c["notes"]), c.get("notes"))
    ok("подсказка класса: несчастные случаи → класс 1, объект — люди",
       f["class_hint"] == "accident" and b["class_hint"] == "accident" and b["suggest_classes"] == ["1"]
       and f["object_kind"] == "people", (f["class_hint"], b.get("suggest_classes"), f["object_kind"]))
    ok("бланк: is_template и незаполненные поля (номер, дата, место, страхователь, сумма, премия, срок)",
       c.get("is_template") is True and f["is_template"] is True
       and {"contract_no", "contract_date", "place", "policyholder", "sum_insured", "premium", "term"}
       <= set(f["blank"]) and [x["code"] for x in c["blank"]][:3] == ["contract_no", "contract_date", "place"]
       and c["blank_label"] == "не заполнено", (f.get("blank"), c.get("is_template")))
    ok("пустые поля — не значения: номер не «00099» из лицензии, дата не дата лицензии, место не «г. ____»",
       f["contract_no"] is None and f["contract_date"] is None and f["place"] is None
       and TPL_LICENSE not in dump and "2021-05-05" not in dump, (f["contract_no"], f["contract_date"], f["place"]))
    ok("страхователь не заполнен — не «физическое лицо»; страховщик — юрлицо «INSON»",
       f["policyholder"] == {"kind": None, "name": None}
       and f["insurer"] == {"kind": "legal", "name": "Акционерное Общество Страховая организация «INSON»"}
       and f["beneficiary"]["kind"] is None and not any("физическое лицо" in n for n in c["notes"]),
       (f["policyholder"], f["insurer"], f["beneficiary"]))
    ok("честный текст: «Это бланк договора: поля … не заполнены. Существенные условия проверяются по заполненному»",
       c["notes"][0].startswith("Это бланк договора: поля номер договора, дата договора")
       and "не заполнены. Существенные условия проверяются по заполненному договору" in c["notes"][0], c["notes"])
    ok("ст. 929 и сверка не выполняются: essentials пуст, request = null, missing — незаполненные поля",
       c["essentials"] == [] and c["request"] is None and not any("ст. 929" in n for n in c["notes"])
       and [x["code"] for x in c["missing"]] == ["contract_no", "contract_date", "policyholder", "sum_insured",
                                                   "tariff_pct", "premium", "term"], (c["essentials"], c["missing"]))
    ok("объект: жизнь и здоровье застрахованных лиц — спортсменов",
       f["object_description"] == "жизнь и здоровье застрахованных лиц — спортсменов", f["object_description"])
    ok("страховой случай (п. 2.6): травма или смерть от несчастного случая во время соревнования",
       f["insured_event"].startswith("получение травматического повреждения или смерть Застрахованного лица")
       and "с наступлением которого" not in f["insured_event"]
       and f["cover_period"] == "во время спортивного соревнования", (f["insured_event"], f["cover_period"]))
    ok("порядок оплаты (п. 5.4): единовременно в течение 5 банковских дней после подписания",
       f["payment_mode"] == "single" and f["payment_text"] == "единовременно в течение 5 (пяти) банковских дней "
       "после подписания настоящего Договора сторонами" and c["payment_mode_label"] == "единовременно",
       f.get("payment_text"))
    ok("срок уведомления: 30 календарных дней (письменное заявление страховщику)",
       f["notice"] == "30 (тридцати) календарных дней", f["notice"])
    ok("риски личного страхования: травма и смерть — подписями",
       [x["code"] for x in f["covered_risks"]] == ["injury", "death"]
       and [x["label"] for x in c["covered_risks"]] == ["травма", "смерть"], f["covered_risks"])
    ok("исключения: военные действия, ядерные, умысел, допинг, опьянение, самоубийство, преступление",
       {"war", "riots", "nuclear", "intent", "doping", "intoxication", "suicide", "crime"}
       == {x["code"] for x in f["exclusions"]} and "применение допинга" in [x["label"] for x in c["exclusions"]],
       f["exclusions"])
    sch = f.get("schedule") or {}
    ok("приложение 1: структура таблицы (7 колонок) распознана, строк нет — пусто",
       sch.get("columns") == ["profession", "count", "personal_sum", "rate", "premium_one", "sum_total",
                              "premium_total"] and sch["items"] == [] and sch["blank"] is True
       and c["schedule"]["columns"][0]["label"] == "Профессия (род занятий)" and f["persons_listed"] == 0, sch)
    ok("франшизы в шаблоне нет — null (не выдумана)", f["franchise"] is None)
    ok("бланк: модель не вызывалась (дочитывать нечего)", not CALLS, len(CALLS))
    with db.tx() as con:
        saved = db.rows(con, "SELECT result_json FROM act_uploads WHERE id=?", b["session"])[0]["result_json"]
        journal = _json.dumps(db.rows(con, "SELECT detail FROM audit WHERE entity=?", "act_upload:" + b["session"]),
                              ensure_ascii=False)
    ok("ФИО руководителя из шапки нигде: ни в ответе, ни в базе, ни в журнале",
       "Тестов" not in dump + saved + journal, [x for x in ("Тестов",) if x in dump + saved + journal])
    # бланк + запрос филиала: сверки «запрос ↔ договор» нет
    st, b2 = upload([("tpl.docx", DOCX_MIME, blob), ("sorov1.docx", DOCX_MIME, docx_table(BR_SAMPLE1))],
                    {"lang": "ru"})
    ok("бланк и запрос филиала: сверки «запрос ↔ договор» нет", st == 200 and b2.get("branch_request")
       and ctb(b2).get("is_template") and b2.get("cross_check") is None, b2.get("cross_check"))
    # акт по бланку: contract_check не строится (как для заявления)
    st, a = ct_make(b["session"], None, must={"product_code": "0102", "sum_insured": 600_000_000,
                                              "object_value": 600_000_000, "region": "Ташкентская область"})
    ok("акт по бланку: сверки договора и ст. 929 нет", st == 200 and not (a.get("contract_check") or {}).get("available")
       and not ct_items(a), (st, (a.get("contract_check") or {}) if isinstance(a, dict) else a))
    st, a2 = ct_make(b["session"], {"premium": 9_000_000, "term_days": 365},
                     must={"product_code": "0102", "sum_insured": 600_000_000, "object_value": 600_000_000,
                           "region": "Ташкентская область"})
    ok("условия, присланные к бланку, в сверку договора не идут", st == 200
       and not (a2.get("contract_check") or {}).get("available"), (a2.get("contract_check") or {}).get("available"))
    # uz и en: тот же текст бланка на языке экрана
    for lang, head in (("uz", "Bu shartnoma blankasi"), ("en", "This is a blank contract form")):
        st, b3 = upload([("tpl.docx", DOCX_MIME, blob)], {"lang": lang})
        ok(f"{lang}: пометка бланка и «не заполнено» на языке экрана",
           ctb(b3)["notes"][0].startswith(head) and ctb(b3)["blank_label"] == ("toʻldirilmagan" if lang == "uz"
                                                                                else "not filled in")
           and [x["label"] for x in ctb(b3)["covered_risks"]] == (["jarohat", "vafot etish"] if lang == "uz"
                                                                   else ["injury", "death"]), ctb(b3)["notes"][:1])
    CT_REPORT["бланк 0102 (копия структуры)"] = {"blank": f["blank"], "product_name": f["product_name"],
                                                 "found": [x["code"] for x in c["found"]]}
    return b["session"]


def check_contract_template_filled():
    print("40в. Тот же шаблон заполнен: номер, дата, страхователь-юрлицо, сумма, премия, срок — полноценный договор")
    fresh()
    model_on(True)
    CALLS.clear()
    REPLY["text"] = ct_ai_reply
    st, b = upload([("dogovor_nс.docx", DOCX_MIME, docx_word(tpl_blocks(filled=True)))], {"lang": "ru"})
    c = ctb(b)
    f = c.get("fields") or {}
    dump = _json.dumps(b, ensure_ascii=False)
    ok("заполненный шаблон — не бланк: is_template = false, пустых ключевых полей нет",
       st == 200 and c.get("detected") and c["is_template"] is False and f["is_template"] is False
       and not (set(f["blank"]) & {"contract_no", "contract_date", "policyholder", "sum_insured", "premium", "term"}),
       f.get("blank"))
    ok("номер, дата, место — из шапки (не номер и дата лицензии)",
       f["contract_no"] == "17-НС/2026" and f["contract_date"] == "2026-10-01" and f["place"] == "г. Ташкент",
       (f["contract_no"], f["contract_date"], f["place"]))
    ok("страхователь — юрлицо без представителя; страховщик — «INSON»",
       f["policyholder"] == {"kind": "legal", "name": "ООО «Спорт Клуб Тест»"}
       and f["insurer"]["kind"] == "legal", f["policyholder"])
    ok("сумма 600 млн, премия 9 млн, срок 01.10.2026–30.09.2027 = 365 дн.",
       f["sum_insured"] == 6e8 and f["premium"] == 9e6 and f["term_from"] == "2026-10-01"
       and f["term_to"] == "2027-09-30" and f["term_days"] == 365, (f["sum_insured"], f["premium"], f["term_days"]))
    sch = f.get("schedule") or {}
    ok("приложение 1: строка «Футболист · 20 · 30 млн · 1,5 % · 450 000 · 600 млн · 9 млн»",
       sch.get("blank") is False and sch["items"] == [{"profession": "Футболист", "count": 20, "personal_sum": 3e7,
                                                        "rate": 1.5, "premium_one": 450000.0, "sum_total": 6e8,
                                                        "premium_total": 9e6}], sch)
    ok("список застрахованных: одна строка — только число, без фамилии",
       f["persons_listed"] == 1 and TPL_ATHLETE.split()[0] not in dump, f.get("persons_listed"))
    ess = {e["code"]: e["present"] for e in c["essentials"]}
    ok("существенные условия (ст. 929) проверяются — все есть; request готов",
       ess == {"object": True, "insured_event": True, "sum_insured": True, "premium": True, "term": True}
       and c["request"] and c["request"]["premium"] == 9e6 and c["request"]["term_days"] == 365, ess)
    ok("ФИО руководителя и представителя нигде нет; модель не вызывалась (правила нашли главное)",
       "Тестов" not in dump and "Сидоров" not in dump and not CALLS, len(CALLS))
    st, a = ct_make(b["session"], c["request"], must={"product_code": "0102", "sum_insured": 600_000_000,
                                                      "object_value": 600_000_000, "region": "Ташкентская область"})
    it = ct_items(a)
    ok("акт: сверка договора есть (contract_check), существенные условия — все есть",
       st == 200 and (a.get("contract_check") or {}).get("available") and it.get("essentials", {}).get("verdict") == "ok"
       and a["premium"]["term_days"] == 365, (st, list(it)))
    # договор, где правила нашли мало: текст уходит в модель только замаскированным (ФИО руководителя — метка)
    CALLS.clear()
    lines = ["«УТВЕРЖДАЮ»", "Генеральный директор", "АО СО «INSON»", TPL_DIRECTOR, "Договор № 18-НС/2026",
             "Страхования спортсменов от несчастных случаев", "г. Ташкент «2» октября 2026 г.",
             "Акционерное Общество Страховая организация «INSON», именуемое в дальнейшем «Страховщик», и ООО «Спорт "
             "Клуб Тест», именуемое в дальнейшем «Страхователь», заключили настоящий Договор.",
             "Страховая сумма и премия определяются по приложению к договору.",
             "2.6. Страховой случай – травма застрахованного лица во время соревнования."]
    st, b2 = upload([("d18.docx", DOCX_MIME, docx_word([("p", x) for x in lines]))], {"lang": "ru"})
    sent = " ".join(m["content"] for x in CALLS for m in x["messages"])
    ok("правила нашли мало — модель дочитывает; ФИО руководителя в модель не ушло (метка [ФИО])",
       len(CALLS) == 1 and "Тестов" not in sent and "[ФИО]" in sent and ctb(b2)["is_template"] is False,
       (len(CALLS), sent[:300]))


# ------------------------------------------------------------------ 41. шаблоны анализа по классам (30.09.2026)

TPL_REPORT = {}


def _admin_header(login: str) -> tuple:
    """Сессия администратора в копии базы: заголовок Authorization."""
    now = datetime.now().isoformat(timespec="seconds")
    token = secrets.token_urlsafe(32)
    with db.tx() as con:
        cur = con.execute("INSERT INTO users (login, full_name, role, branch, password_hash, salt, status, created_at,"
                          " approved_by, approved_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                          (login, "Test Admin", "админ", "тест", secrets.token_hex(32), secrets.token_hex(16),
                           "активен", now, "test", now))
        con.execute("INSERT INTO sessions (token, user_id, created_at, expires_at, ip, user_agent) VALUES (?,?,?,?,?,?)",
                    (token, cur.lastrowid, now, (datetime.now() + timedelta(hours=2)).isoformat(timespec="seconds"),
                     "127.0.0.1", "test_act"))
    return (b"authorization", f"Bearer {token}".encode())


def check_templates_ref():
    print("41а. Шаблоны 17 классов: файл, таблица class_templates, структура, доли, оговорки, мероприятия, ракурсы")
    from app import class_templates as ctm
    data = ctm.load_file()
    ok("файл шаблонов: версия 1.0.1 от 30.09.2026 (замечания контролёра), 17 классов 1–17, 13з → 14, 16у → 16",
       data["version"] == "1.0.1" and data["date"] == "2026-09-30" and sorted(data["classes"], key=int) ==
       [str(i) for i in range(1, 18)] and data["aliases"] == {"13з": "14", "16у": "16"}, list(data["classes"]))
    with db.tx() as con:
        ctm.ensure(con)
        rows = ctm.all_current(con)
        n_db = con.execute("SELECT COUNT(DISTINCT class_code) FROM class_templates").fetchone()[0]
        errs = {r["class_code"]: ctm.validate(r["template"], r["class_code"], con) for r in rows}
        mcodes = ctm.measure_codes(con)
        perils = {r[0] for r in con.execute("SELECT DISTINCT class_code FROM perils")}
    ok("17 шаблонов загружены в таблицу class_templates (calibrated = 0)",
       len(rows) == 17 and n_db == 17 and all(r["calibrated"] == 0 for r in rows), (len(rows), n_db))
    ok("структура всех 17 шаблонов проходит проверку", not any(errs.values()), {k: v for k, v in errs.items() if v})
    clauses_all = ctm.clause_codes()
    for r in rows:
        c, tp = r["class_code"], r["template"]
        rk = tp["risks"]
        if c in ("8", "9"):
            ok(f"класс {c}: риски — ссылка на справочник perils (доли не дублируются)",
               rk["source"] == "perils" and rk["items"] == [] and c in perils, rk["source"])
        else:
            total = sum(x["share_pct"] for x in rk["items"])
            ok(f"класс {c}: экспертные доли рисков, сумма 100, calibrated = 0",
               rk["source"] == "template" and abs(total - 100) < 1e-9 and rk["calibrated"] == 0
               and all(set(x["label"]) == {"ru", "uz", "en"} for x in rk["items"]), total)
        cl = {x for codes in tp["clauses"].values() for x in codes}
        ms = {x for codes in tp["measures"].values() for x in codes}
        ok(f"класс {c}: оговорок ≥ 3 и мероприятий ≥ 3 в каждой группе, коды существуют",
           all(len(v) >= 3 for v in tp["clauses"].values()) and all(len(v) >= 3 for v in tp["measures"].values())
           and cl <= clauses_all and ms <= mcodes, (sorted(cl - clauses_all), sorted(ms - mcodes)))
        ok(f"класс {c}: ракурсы есть (каждая группа — непустой список из известных кодов), must ≤ 4",
           tp["required_views"] and all(v and set(v) <= set(ae.VIEWS) for v in tp["required_views"].values())
           and len(tp["must"]) <= 4 and tp["scenario_rule"]["code"] in ctm.SCENARIO_RULES, tp["required_views"])
    # классы 3, 8, 9: шаблон повторяет прежний выбор по группе объекта — акт не меняется
    cat = act.clause_catalog()
    mcat = __import__("app.act_extras", fromlist=["x"]).measures_catalog()
    same = True
    for c, groups in (("3", ("vehicle", "special")), ("8", ("property", "equipment")), ("9", ("property", "equipment"))):
        tp = next(r["template"] for r in rows if r["class_code"] == c)
        for g in groups:
            same = same and tp["clauses"][g] == [x["code"] for x in cat["groups"][g]]
            same = same and tp["required_views"][g] == ae.REQUIRED_VIEWS[g]
            by_group = [x["code"] for x in mcat["catalog"] if g in (x.get("groups") or [])]
            listed = [x for x in tp["measures"][g] if x in {m["code"] for m in mcat["catalog"]}]
            same = same and sorted(listed) == sorted(by_group)
    ok("классы 3, 8, 9: оговорки, ракурсы и мероприятия шаблона = прежний выбор по группе объекта", same)
    ok("класс 3: граница транспорта и спецтехники — по регистрационному документу (особое правило)",
       any(n["code"] == "vehicle_boundary" and "регистрационному документу" in n["text"]["ru"]
           for n in data["classes"]["3"]["notes"]))
    ok("класс 14: только необеспеченная часть и не более 50 %, страхователь — банк",
       any("не более 50 %" in n["text"]["ru"] and "банк" in n["text"]["ru"] for n in data["classes"]["14"]["notes"]))
    ok("перестрахование почти всегда — у 5, 6, 11, 12",
       [c for c, t_ in data["classes"].items() if t_["reinsurance_usually"]] == ["5", "6", "11", "12"])
    ci = ctm.credit_insurable(100_000_000, 60_000_000)
    ok("кредит 100 млн при залоге 60 млн: страхуется min(40; 50) = 40 млн (необеспеченная часть)",
       ci["insurable"] == 40_000_000 and ci["by"] == "unsecured", ci)
    ci2 = ctm.credit_insurable(100_000_000, 20_000_000)
    ok("кредит 100 млн при залоге 20 млн: min(80; 50) = 50 млн (предел 50 %)",
       ci2["insurable"] == 50_000_000 and ci2["by"] == "cap", ci2)
    TPL_REPORT["классы"] = {r["class_code"]: [(x["code"], x["share_pct"]) for x in r["template"]["risks"]["items"]]
                            or "perils" for r in rows}


def check_templates_api():
    print("41б. API шаблонов: список, класс на трёх языках, история; PUT — проверка структуры и новая версия")
    fresh()
    st, lst = call("GET", "/act/templates", params={"lang": "ru"})
    ok("GET /act/templates — 17 шаблонов кратко, версия файла 1.0.1",
       st == 200 and len(lst["templates"]) == 17 and lst["file_version"] == "1.0.1"
       and [x["class_code"] for x in lst["templates"]] == [str(i) for i in range(1, 18)], (st, str(lst)[:300]))
    one = {x["class_code"]: x for x in lst["templates"]}
    ok("кратко: класс 3 — ДТП 45 %, угон 20 %; классы 8/9 — источник perils",
       one["3"]["risks"][:2] == [{"code": "mv_accident", "label": "ДТП", "share_pct": 45, "catastrophic": False},
                                 {"code": "mv_theft", "label": "Угон", "share_pct": 20, "catastrophic": False}]
       and one["8"]["risks_source"] == "perils", one["3"]["risks"][:2])
    names = {}
    for lg in ("ru", "uz", "en"):
        st, t1 = call("GET", "/act/templates/1", params={"lang": lg})
        names[lg] = (st, t1.get("template", {}).get("name"), t1["template"]["risks"]["items"][0]["label"]
                     if st == 200 else None, t1["template"]["must"][0]["label"] if st == 200 else None)
    ok("GET /act/templates/1 на трёх языках: название, риск, поле",
       names["ru"][1:] == ("Несчастные случаи", "Смерть", "Число застрахованных")
       and names["uz"][1:] == ("Baxtsiz hodisalar", "Oʻlim", "Sugʻurtalanganlar soni")
       and names["en"][1:] == ("Accident", "Death", "Number of insured persons"), names)
    st, t8 = call("GET", "/act/templates/8", params={"lang": "ru"})
    ok("класс 8: риски из perils базы (пожар 40 %, землетрясение 20 %) и документы из checklists",
       st == 200 and {p["code"]: p["share_pct"] for p in t8["perils"]}.get("fire") == 40.0
       and {p["code"]: p["share_pct"] for p in t8["perils"]}.get("earthquake") == 20.0
       and any(d["doc"] == "Сведения о пожарной сигнализации и охране" for d in t8["checklists"]), str(t8)[:300])
    st, t13z = call("GET", "/act/templates/13з", params={"lang": "ru"})
    ok("13з — по шаблону класса 14", st == 200 and t13z["class_code"] == "14" and t13z["alias_of"] == "14", st)
    st, _x = call("GET", "/act/templates/99")
    ok("неизвестный класс — 404", st == 404, st)
    st, raw = call("GET", "/act/templates/1", params={"raw": 1})
    good = _json.loads(_json.dumps(raw["raw"]))
    # гость не правит
    st, _x = call("PUT", "/act/templates/1", {"template": good})
    ok("гость шаблон не меняет (401/403)", st in (401, 403), st)
    HEADERS.append(_admin_header("tpl_test_admin"))
    try:
        bad = _json.loads(_json.dumps(good))
        bad["risks"]["items"][0]["share_pct"] = 15            # сумма 90
        bad["clauses"]["default"].append("no_such_clause")
        bad["measures"]["default"].append("no_such_measure")
        bad["must"] = bad["must"] + [dict(bad["must"][0], code="x1"), dict(bad["must"][0], code="x2")]
        st, e = call("PUT", "/act/templates/1", {"template": bad})
        errs = " | ".join(e.get("errors") or [])
        ok("PUT с плохой структурой — 422: доли, оговорка, мероприятие, must > 4",
           st == 422 and "сумма долей 90" in errs and "no_such_clause" in errs and "no_such_measure" in errs
           and "не больше 4" in errs, (st, errs))
        st, e = call("PUT", "/act/templates/1", {"template": {"name": {"ru": "x"}}})
        ok("PUT без обязательных полей шаблона — 422", st == 422 and any("нет поля" in x for x in e["errors"]), e)
        new = _json.loads(_json.dumps(good))
        new["risks"]["items"][0]["share_pct"] = 20.2          # 20,2 + 30 + 30 + 15 = 95,2 → поправим травму
        new["risks"]["items"][3]["share_pct"] = 19.6          # сумма 99,8 — в пределах ± 0,5
        st, r = call("PUT", "/act/templates/1", {"template": new, "note": "тест: доли НС"})
        ok("PUT с хорошей структурой — новая версия 1.1 (правка администратора)",
           st == 200 and r["version"] == "1.1" and r["source"] == "admin"
           and r["template"]["risks"]["items"][0]["share_pct"] == 20.2, (st, str(r)[:300]))
        st, h = call("GET", "/act/templates/1/history")
        ok("история: версия файла 1.0.1 и правка 1.1 — обе сохранены",
           st == 200 and [(x["version"], x["source"]) for x in h["history"]] == [("1.0.1", "file"), ("1.1", "admin")],
           h)
        # файл той же версии правку не затирает
        from app import class_templates as ctm
        ctm.reset_cache()
        with db.tx() as con:
            ctm.ensure(con)
            cur = ctm.current(con, "1")
        ok("ensure с файлом 1.0.1 не затирает правку 1.1", cur["version"] == "1.1" and cur["source"] == "admin",
           cur["version"])
        st, a = call("POST", "/act/make", {"lang": "ru", "must": {"class_code": "1", "sum_insured": 1_000_000_000,
                                                                "object_value": 1_000_000_000, "region": "Ташкент"}})
        shares = {i["code"]: i["share_of_net_pct"] for i in a["analytics"]["risks"]["items"]}
        ok("акт класса 1 берёт действующую версию шаблона (смерть 20,2 %)",
           st == 200 and shares.get("pa_death") == 20.2 and a["template"]["version"] == "1.1", shares)
        # вернуть доли файла: ещё одна версия (история не удаляется)
        st, r = call("PUT", "/act/templates/1", {"template": good})
        ok("возврат долей — версия 1.2, история из трёх строк", st == 200 and r["version"] == "1.2", r.get("version"))
    finally:
        HEADERS.clear()


def _cls_act(cls, fields=None, S=1_000_000_000, V=None, lang="ru", optional=None):
    o = dict(optional or {})
    if fields is not None:
        o["class_fields"] = fields
    return call("POST", "/act/make", {"lang": lang, "must": {"class_code": cls, "sum_insured": S,
                                                             "object_value": V or S, "region": "Ташкент"},
                                      "optional": o})


def check_templates_act():
    print("41в. Акт по классу без perils (1, 7, 13, 14): риски и сценарии из шаблона; классы 3/8/9 — цифры прежние")
    fresh()
    model_on(False)
    # класс 1: 50 человек × 20 млн, в одном месте 10
    st, a = _cls_act("1", {"insured_count": 50, "occupation": "строители", "sum_per_person": 20_000_000,
                           "people_in_one_place": 10})
    rk = a["analytics"]["risks"]
    sc = a["scenarios"]
    ok("класс 1: риски из шаблона (смерть, инвалидность, ВУТ, травма), сумма 100, пометка «экспертные доли шаблона»",
       st == 200 and rk["source"] == "template" and [i["code"] for i in rk["items"]] ==
       ["pa_death", "pa_disability", "pa_temp_disability", "pa_injury"] and rk["total_pct"] == 100.0
       and rk["label"] == "экспертные доли шаблона" and any("шаблона класса" in n for n in rk["notes"])
       and not rk["whole_class"], rk.get("notes"))
    ok("класс 1: PML = EML = 20 млн (сумма на человека), MFL = 20 млн × 10 = 200 млн",
       sc["available"] and sc["source"] == "template" and
       [sc[k]["amount"] for k in ("pml", "eml", "mfl")] == [20_000_000, 20_000_000, 200_000_000]
       and "катастрофа" in sc["mfl"]["what"], [sc[k]["amount"] for k in ("pml", "eml", "mfl")])
    ok("класс 1: «как посчитано» — простое правило шаблона, экспертное; формула с числами",
       any("простым правилом шаблона класса 1" in h for h in sc["how"])
       and sc["mfl"]["formula"].replace(" ", " ") == "MFL = 20 000 000 сум × 10 чел. в одном месте = 200 000 000 сум",
       (sc["how"], sc["mfl"]["formula"]))
    ok("класс 1: оговорки и мероприятия — из шаблона",
       [c["code"] for c in a["clauses"]] == ["pa_list", "pa_hazard", "pa_cover_time", "ot_declared"]
       and [m["code"] for m in a["measures"]] == ["pa_briefing", "pa_ppe", "pa_medical"], [c["code"] for c in a["clauses"]])
    ok("класс 1: блок template — версия, поля класса, документы, статистика",
       a["template"]["class_code"] == "1" and a["template"]["class_fields"]["insured_count"] == 50
       and a["template"]["documents"]["items"] and "mortality_rate" in a["template"]["stats"]
       and a["template"]["required_views"] == ["document"], str(a.get("template"))[:200])
    st, u = call("GET", f"/act/{a['id']}", params={"lang": "uz"})
    ok("класс 1 по-узбекски: риск «Oʻlim», подпись сценария из шаблона",
       u["analytics"]["risks"]["items"][0]["name"] == "Oʻlim" and u["scenarios"]["mfl"]["what"].startswith("falokat"),
       (u["analytics"]["risks"]["items"][0]["name"], u["scenarios"]["mfl"]["what"]))
    st, e = call("GET", f"/act/{a['id']}", params={"lang": "en"})
    ok("класс 1 по-английски: риск «Death»", e["analytics"]["risks"]["items"][0]["name"] == "Death")
    # без полей класса — всё по страховой сумме, с пометками «по умолчанию»
    st, a0 = _cls_act("1")
    ok("класс 1 без полей: сценарии по страховой сумме, в «принято по умолчанию» — что не указано",
       a0["scenarios"]["available"] and a0["scenarios"]["pml"]["amount"] == 1_000_000_000
       and {x["code"] for x in a0["scenarios"]["assumptions"]} >= {"as_tpl_per_person_sum", "as_tpl_place"},
       a0["scenarios"]["assumptions"])
    # класс 7: отправка 200 млн, накопление 600 млн
    st, a = _cls_act("7", {"cargo_kind": "бытовая техника", "transport_mode": "auto", "route": "Ташкент — Самарканд",
                           "limit_per_shipment": 200_000_000, "accumulation_value": 600_000_000})
    sc = a["scenarios"]
    ok("класс 7: PML = EML = одна отправка 200 млн, MFL = накопление 600 млн",
       [sc[k]["amount"] for k in ("pml", "eml", "mfl")] == [200_000_000, 200_000_000, 600_000_000], sc)
    ok("класс 7: риски шаблона (повреждение 30 %, кража 20 % …)",
       {i["code"]: i["share_of_net_pct"] for i in a["analytics"]["risks"]["items"]}.get("cg_damage") == 30.0
       and a["analytics"]["risks"]["source"] == "template")
    ok("класс 7: неверный вид транспорта в полях класса — 422",
       _cls_act("7", {"transport_mode": "teleport"})[0] == 422)
    # класс 13: лимит на случай 300 млн, годовой 1 млрд
    st, a = _cls_act("13", {"activity_kind": "trade", "turnover_or_payroll": 5_000_000_000,
                            "limit_per_case": 300_000_000, "limit_aggregate": 1_000_000_000})
    sc = a["scenarios"]
    ok("класс 13: PML = EML = лимит на случай 300 млн, MFL = годовой лимит 1 млрд",
       [sc[k]["amount"] for k in ("pml", "eml", "mfl")] == [300_000_000, 300_000_000, 1_000_000_000], sc)
    ok("класс 13: удержание сравнивается с EML (Положение 1806, п. 15)",
       sc["retention"]["compared_with"] == "eml" and (not sc["retention"]["known"] or
                                                      sc["retention"]["eml_excess"] == max(300_000_000 - sc["retention"]["limit"], 0)),
       sc["retention"])
    ok("класс 13: раздел 4 — аналитика: сценарии шаблона формулой с суммой",
       a["analytics"]["scenarios"]["available"] and a["analytics"]["scenarios"]["items"][0]["formula"].startswith(
           "один случай: лимит на случай"), a["analytics"]["scenarios"])
    # класс 14: кредит 100 млн, залог 60 млн
    st, a = _cls_act("14", {"credit_amount": 100_000_000, "collateral_value": 60_000_000, "credit_term_months": 24},
                     S=40_000_000)
    ok("класс 14: кредит 100 млн, залог 60 млн, сумма 40 млн — в пределах (проверка в «как посчитано»)",
       any("допустимая страховая сумма 40" in h.replace(" ", " ") and "в пределах" in h for h in a["scenarios"]["how"])
       and not any("Кредит: страховая сумма" in c for c in a["decision"]["checks"]), a["scenarios"]["how"])
    st, a = _cls_act("14", {"credit_amount": 100_000_000, "collateral_value": 60_000_000}, S=50_000_000)
    ok("класс 14: сумма 50 млн выше допустимой 40 млн — проверка андеррайтеру, превышение 10 млн",
       any("Кредит: страховая сумма 50" in c.replace(" ", " ") and "10 000 000" in c.replace(" ", " ")
           for c in a["decision"]["checks"]) and a["decision"]["code"] != "accept", a["decision"])
    # классы 3, 8, 9: контрольные цифры прежние (как до шаблонов)
    st, c3 = call("POST", "/act/make", {"lang": "ru", "must": CRANE_MUST,
                                        "optional": dict(CRANE_OPT, object_kind="truck_crane")})
    ok("класс 3 (автокран 2,945 млрд): ставка 0,42 %, премия 12 369 000, PML/EML/MFL 1 472 500 000 / 2 945 000 000 × 2",
       c3["rate"]["applied_pct"] == 0.42 and c3["premium"]["amount"] == 12_369_000
       and [c3["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")] == [1_472_500_000, 2_945_000_000, 2_945_000_000]
       and c3["analytics"]["risks"]["whole_class"] and c3["scenarios"].get("source") is None,
       (c3["rate"]["applied_pct"], c3["premium"]["amount"], [c3["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")]))
    ok("класс 3: оговорки спецтехники прежние (5 шт.), ракурсы — 7",
       [c["code"] for c in c3["clauses"]] == ["sp_attachments_storage", "sp_territory", "sp_reinspection", "sp_operator",
                                              "sp_rated_load"] and len(c3["template"]["required_views"]) == 7)
    st, c8 = call("POST", "/act/make", {"lang": "ru", "must": WH8_MUST, "optional": WH8_OPT})
    ok("класс 8 (склад 4,2 млрд): PML/EML/MFL 2 100 000 000 / 3 360 000 000 / 4 200 000 000, риски — perils",
       [c8["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")] == [2_100_000_000, 3_360_000_000, 4_200_000_000]
       and c8["analytics"]["risks"]["source"] == "perils"
       and c8["analytics"]["risks"]["items"][0]["code"] == "fire", [c8["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")])
    st, c9 = call("POST", "/act/make", {"lang": "ru", "must": WH_MUST, "optional": WH_OPT})
    ok("класс 9: риски из perils (кража со взломом первой), правило property9",
       c9["analytics"]["risks"]["items"][0]["code"] == "burglary" and c9["scenarios"]["rule"] == "property9")
    TPL_REPORT["класс 3"] = (c3["rate"]["applied_pct"], c3["premium"]["amount"],
                             [c3["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")])
    TPL_REPORT["класс 8"] = [c8["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")]
    TPL_REPORT["класс 9"] = [c9["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")]


PT_REPORT = {}
PT_CAR = {"class_code": "3", "sum_insured": 60_000_000, "object_value": 60_000_000,
          "object_description": "легковой автомобиль в залоге"}
PT_CREDIT = {"class_code": "14", "sum_insured": 40_000_000,
             "fields": {"class_fields": {"credit_amount": 100_000_000, "collateral_value": 60_000_000,
                                         "credit_term_months": 24}}}
PT_MUST = {"product_code": "0312", "sum_insured": 100_000_000, "object_value": 60_000_000, "region": "Ташкент"}
PT_OPT = {"losses_3y": {"count": 0, "small_count": 0}, "documents_provided": True}


def _pt_make(must, optional, lang="ru"):
    return call("POST", "/act/make", {"lang": lang, "must": must, "optional": optional})


def check_parts_engine():
    print("42а. Комплексный продукт: чистые функции разбора по частям (act_engine, раздел 9)")
    ok("ставки частей из текста тарифа: 0305 «ТС 1,1% · НС 0,5% · ОТВ 1%» → 3: 1,1; 1: 0,5; 13: 1",
       ae.part_rates_from_text("ТС 1,1% · НС 0,5% · ОТВ 1%", ["3", "1", "13"]) == {"3": 1.1, "1": 0.5, "13": 1.0})
    ok("0312 «фин. риск 0,5% · залог 0,5%» → 14: 0,5; 3: 0,5; 1415 «ГБО 0,3% · ГО 0,5% · кредит 1,7%» → 8/13/14",
       ae.part_rates_from_text("фин. риск 0,5% · залог 0,5%", ["3", "14"]) == {"14": 0.5, "3": 0.5}
       and ae.part_rates_from_text("ГБО 0,3% · ГО 0,5% · кредит 1,7%", ["8", "13", "14"]) ==
       {"8": 0.3, "13": 0.5, "14": 1.7})
    ok("одна ставка без подписи «0,25% фикс.» — общая для частей; «по согласованию» — ставок нет",
       ae.part_rates_from_text("0,25% фикс.", ["13", "1"]) == {"*": 0.25}
       and ae.part_rates_from_text("по согласованию с ЦО", ["8", "9"]) == {})
    sp = ae.split_sum(1_000_000_001, ["3", "14"])
    ok("поровну: 1 000 000 001 → 500 000 000 + 500 000 001 (остаток — последней части), сумма точно S",
       [x["sum_insured"] for x in sp] == [500_000_000, 500_000_001] and ae.check_parts_sum(sp, 1_000_000_001) is None)
    sp2 = ae.split_sum(1_000_000_000, ["8", "9"], {"8": 70, "9": 30})
    ok("по долям тарифной политики 70/30: 700 млн + 300 млн",
       [x["sum_insured"] for x in sp2] == [700_000_000, 300_000_000] and [x["share_pct"] for x in sp2] == [70, 30])
    bad = ae.check_parts_sum([{"sum_insured": 60e6}, {"sum_insured": 39_999_998}], 100e6, 1)
    ok("сумма частей на 2 сума меньше при допуске 1 сум — расхождение; на 1 сум — сходится",
       bad == {"total": 99_999_998, "sum_insured": 100e6, "diff": 2}
       and ae.check_parts_sum([{"sum_insured": 60e6}, {"sum_insured": 39_999_999}], 100e6, 1) is None, bad)
    rows = ae.parts_from_items([{"name": "Легковой автомобиль в залоге", "sum": 60e6},
                                {"name": "Финансовый риск непогашения кредита", "sum": 40e6}], ["3", "14"])
    ok("объекты договора → части: автомобиль → класс 3, непогашение кредита → класс 14",
       [(r["class_code"], r["sum_insured"], r["class_guess"]) for r in rows] == [("3", 60e6, False), ("14", 40e6, False)],
       rows)
    rows = ae.parts_from_items([{"name": "Склад", "sum": 1e9}, {"name": "Ущерб от залива и кражи", "sum": 2e8},
                                {"name": "Прочее", "sum": 1e8}], ["8", "9"], {"8": ["склад"], "9": ["склад"]})
    ok("8/9: «склад» → 8, «ущерб от залива и кражи» → 9, неузнанное «прочее» — к первому классу с пометкой",
       [(r["class_code"], r["sum_insured"], r["class_guess"]) for r in rows] == [("8", 1.1e9, True), ("9", 2e8, False)],
       rows)
    ok("один объект по умолчанию — только 8/9/16", ae.default_same_object(["8", "9", "16"])
       and not ae.default_same_object(["3", "14"]) and not ae.default_same_object(["8", "1"]))
    sc = lambda p, e, m: {"available": True, "items": {"PML": {"amount": p}, "EML": {"amount": e}, "MFL": {"amount": m}}}
    same = ae.aggregate_scenarios([{"index": 1, "class_code": "8", "main": True, "scenarios": sc(70, 100, 100)},
                                   {"index": 2, "class_code": "9", "main": True, "scenarios": sc(15, 37, 75)}])
    diff = ae.aggregate_scenarios([{"index": 1, "class_code": "3", "main": True, "scenarios": sc(30, 60, 60)},
                                   {"index": 2, "class_code": "14", "main": False, "scenarios": sc(40, 40, 40)},
                                   {"index": 3, "class_code": "13", "main": False, "scenarios": {"available": False,
                                                                                                 "reason": "sc_na_class"}}])
    ok("сценарии: один объект — большее (70/100/100), разные объекты — сумма (70/100/100), часть без правила — "
       "не входит", same["rule"] == "max" and [same["items"][s]["amount"] for s in ("PML", "EML", "MFL")] == [70, 100, 100]
       and diff["rule"] == "sum" and [diff["items"][s]["amount"] for s in ("PML", "EML", "MFL")] == [70, 100, 100]
       and diff["excluded"] == [{"index": 3, "class_code": "13", "reason": "sc_na_class"}], (same, diff))
    ret = ae.contract_retention([{"known": True, "limit": 11e9}, {"known": True, "limit": 7e9}], 8e9, 9e9)
    ok("удержание договора: наименьший лимит частей 7 млрд против EML договора 8 млрд — превышение 1 млрд",
       ret["limit"] == 7e9 and ret["eml_excess"] == 1_000_000_000 and ret["within"] is False
       and ret["mfl_excess"] == 2_000_000_000, ret)
    tot = ae.contract_totals([{"premium": 3_720_000, "level": "moderate", "sum_insured": 146.1e6, "term_days": 365},
                              {"premium": 4_480_000, "level": "high", "sum_insured": 226.1e6, "term_days": 365}], 372.2e6)
    ok("итоги: премия = сумма премий частей, уровень — самый высокий, ставка договора только справочно",
       tot["premium"] == 8_200_000 and tot["level"] == "high" and tot["reference_only"]
       and tot["reference_rate_pct"] == 2.2031, tot)
    with db.tx() as con:
        ref = db.load_reference(con)
        prod = db.rows(con, "SELECT code, name, pricing_mode, rate_text FROM products WHERE code='0312'")[0]
        r14 = ae.part_rate(ref, prod, "14", "moderate", 40e6, 365, None, None, act.load_settings(con))
        r3 = ae.part_rate(ref, prod, "3", "moderate", 60e6, 365, None, None, act.load_settings(con))
    ok("0312: минимум класса 3 — справочник min_rates 0,5 %; класса 14 — из текста тарифа 0,5 % (в min_rates его нет)",
       r3["class_min"] == {"pct": 0.5, "source": "min_rates"} and r14["class_min"] == {"pct": 0.5, "source": "rate_text"}
       and r14["min_pct"] == 0.5 and r14["applied_pct"] == 0.6 and r14["premium"] == 240_000,
       (r3["class_min"], r14["class_min"], r14["applied_pct"], r14["premium"]))


def check_parts_make():
    print("42б. Комплексный продукт в /act/make: части по умолчанию, подтверждённые, из договора, 8/9, обязательная часть")
    fresh()
    model_on(False)
    with db.tx() as con:
        ref = db.load_reference(con)
        st_ = act.load_settings(con)
        name = db.rows(con, "SELECT name, pricing_mode FROM products WHERE code='0312'")[0]
        cls0312 = [r["class_code"] for r in db.rows(con, "SELECT class_code FROM product_classes "
                                                         "WHERE product_code='0312' ORDER BY part_no")]
    ok("0312 — «Автокредит» (Хамкорбанк): классы 3 (залоговый автомобиль) и 14 (невозврат кредита) — разные объекты, "
       "делить осмысленно", cls0312 == ["3", "14"] and "Автокредит" in name["name"] and name["pricing_mode"] == "ставка",
       (cls0312, name))
    # 1. части по умолчанию: поровну, предложение с пометкой
    st, a = _pt_make(PT_MUST, PT_OPT)
    P = a.get("parts") or {}
    text = all_text(a) if st == 200 else str(a)
    ok("0312 без частей: mode multi, source default, не подтверждено, suggested_parts 50 + 50 млн",
       st == 200 and P["mode"] == "multi" and P["source"] == "default" and not P["confirmed"]
       and [(x["class_code"], x["sum_insured"]) for x in a["suggested_parts"]] == [("3", 50e6), ("14", 50e6)]
       and a["suggested_parts"] == P["suggested_parts"], (st, P.get("source"), a.get("suggested_parts")))
    ok("пометка «распределение по умолчанию — подтвердите» в акте и в решении, «разбор отложен» больше нет",
       "разделена поровну на 2 части по умолчанию" in text and "принято поровну по умолчанию — подтвердить" in text
       and any("Подтвердить распределение" in c for c in a["decision"]["checks"])
       and "разбор по частям отложен" not in text and a["rate"]["multi_class"], text[:300])
    ok("переключатель по умолчанию: автомобиль и кредит — разные объекты (object_mode different)",
       P["object_mode"] == "different" and [x["same_object"] for x in P["items"]] == [True, False])
    # 2. подтверждённые части: автомобиль 60 млн (залог), кредит 40 млн = min(100 − 60; 50)
    st, a = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, PT_CREDIT]))
    P = a["parts"]
    it = P["items"]
    exp = []
    for p in it:
        adj = st_["adj_pct"][p["level"]]
        applied = round(0.5 * (1 + adj / 100), 4)
        exp.append((applied, round(premium_of(applied, p["sum_insured"], 365))))
    ok("части сотрудника: source employee, подтверждено, suggested_parts на верхнем уровне нет",
       st == 200 and P["source"] == "employee" and P["confirmed"] and "suggested_parts" not in a, (st, P.get("source")))
    ok("ставка каждой части — по своему классу: 0,5 % × поправка уровня, не ниже 0,5 % класса",
       [(p["rate"]["applied_pct"], p["premium"]) for p in it] == exp
       and [p["rate"]["min_pct"] for p in it] == [0.5, 0.5]
       and [p["rate"]["min_source"] for p in it] == ["min_rates", "rate_text"], ([(p["rate"], p["premium"]) for p in it], exp))
    ok("премия договора = сумма премий частей; верхнее поле premium — итог договора",
       a["premium"]["amount"] == sum(p["premium"] for p in it) == P["totals"]["premium"],
       (a["premium"]["amount"], [p["premium"] for p in it]))
    lv = {"low": 0, "moderate": 1, "high": 2}
    ok("уровень договора — самый высокий среди частей (верхнее поле risk)",
       a["risk"]["level"] == max((p["level"] for p in it), key=lv.get) == P["totals"]["level"])
    ok("средняя ставка договора — только справочно: rate.applied_pct и min_pct пустые, reference_pct = премия / сумма",
       a["rate"]["mode"] == "multi" and a["rate"]["applied_pct"] is None and a["rate"]["min_pct"] is None
       and a["rate"]["reference_pct"] == round(a["premium"]["amount"] / 100e6 * 100, 4)
       and "средняя не используется для проверки минимума" in all_text(a), a["rate"])
    sc_parts = [[p["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")] for p in it]
    ok("разные объекты — сценарии складываются: EML договора = EML автомобиля + EML кредита",
       a["scenarios"]["available"] and [a["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")] ==
       [sc_parts[0][i] + sc_parts[1][i] for i in range(3)] and P["totals"]["scenarios"]["rule"] == "sum", sc_parts)
    ok("кредит 40 млн при залоге 60 млн из 100 млн — в пределах: проверки превышения нет",
       not any("выше допустимой" in c for c in a["decision"]["checks"])
       and sc_parts[1] == [40_000_000] * 3)
    ok("удержание договора сравнивается с EML договора",
       a["scenarios"]["retention"]["compared_with"] == "eml" and (not a["scenarios"]["retention"]["known"] or
       a["scenarios"]["retention"]["eml_excess"] == max(a["scenarios"]["eml"]["amount"] - a["scenarios"]["retention"]["limit"], 0)))
    ok("сумма к стоимости по частям: автомобиль 100 %, у кредита — «не применяется»; по договору 100 %",
       it[0]["value"]["ratio_pct"] == 100.0 and it[1]["value"]["verdict"] == "na" and not it[1]["value"]["applicable"]
       and a["value"]["ratio_pct"] == 100.0, [p["value"] for p in it])
    ok("аналитика — по каждой части (parts.items[].analytics), сводка договора — общая",
       all(p["analytics"].get("available") for p in it) and a["analytics"]["reason"] == "by_parts"
       and a["analytics"]["summary"]["text"].startswith("Договор из 2 частей"), a["analytics"].get("summary"))
    t4 = a["sections"][3]
    tbl = next((li for li in t4["lists"] if li["title"] == "Части договора"), None)
    ok("раздел 4: таблица частей (класс, сумма, уровень, ставка, премия, франшиза) с итогом договора",
       tbl and tbl["table"]["columns"] == ["№", "Класс", "Страховая сумма", "Уровень", "Ставка", "Премия", "Франшиза"]
       and len(tbl["table"]["rows"]) == 3 and tbl["table"]["rows"][-1][1] == "Итого по договору", tbl)
    ok("раздел 1 — перечень частей, раздел 3 — сумма к стоимости по частям, раздел 5 — подтверждено сотрудником",
       sum(1 for r in a["sections"][0]["rows"] if r["label"].startswith("Часть ")) == 2
       and sum(1 for r in a["sections"][2]["rows"] if r["label"].startswith("Часть ")) == 2
       and "Распределение страховой суммы по классам подтверждено сотрудником." in a["sections"][4]["paragraphs"])
    PT_REPORT["0312 подтверждено"] = {"части": [(p["class_code"], p["sum_insured"], p["level"], p["rate"]["applied_pct"],
                                                 p["premium"]) for p in it], "премия": a["premium"]["amount"],
                                       "PML/EML/MFL": [a["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")]}
    aid = a["id"]
    # кредит 60 млн — выше допустимых 40 млн: проверка по части
    st, a2 = _pt_make(dict(PT_MUST, sum_insured=120_000_000),
                      dict(PT_OPT, parts=[PT_CAR, dict(PT_CREDIT, sum_insured=60_000_000)]))
    ok("кредит 60 млн при допустимых 40 млн — проверка по части 2 с превышением 20 млн",
       any("Часть 2 (класс 14)" in c and "20 000 000" in flat(c) for c in a2["decision"]["checks"]),
       a2["decision"]["checks"])
    # франшиза по частям: своя франшиза только у автомобиля
    st, a3 = _pt_make(PT_MUST, dict(PT_OPT, parts=[dict(PT_CAR, deductible={"pct": 1}), PT_CREDIT]))
    f = [p["franchise"] for p in a3["parts"]["items"]]
    ok("франшиза по каждой части отдельно: у части 1 применена 1 %, у части 2 — нет; премия договора — сумма частей",
       f[0]["status"] == "applied" and f[0]["size_pct"] == 1.0 and f[1]["status"] == "none"
       and a3["premium"]["amount"] == sum(p["premium"] for p in a3["parts"]["items"])
       and "часть 1 — применена 1" in flat(all_text(a3)), f)
    ok("часть 1 на минимуме класса (0,5 %): со франшизой ставка не опускается ниже минимума — премия та же, пометка",
       f[0]["floor_applied"] and a3["parts"]["items"][0]["premium"] == a3["parts"]["items"][0]["premium_before_franchise"]
       and "упёрлась в минимальную ставку" in f[0]["text"], (f[0]["floor_applied"], f[0]["text"]))
    # 3. ошибки ввода
    st, e = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, dict(PT_CREDIT, sum_insured=39_999_998)]))
    ok("сумма частей ≠ страховой сумме (разница 2 сума) — 422 с объяснением",
       st == 422 and "parts" in e["errors"] and "не равна страховой сумме договора" in e["errors"]["parts"], (st, e))
    st, e1 = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, dict(PT_CREDIT, sum_insured=39_999_999)]))
    ok("разница 1 сум — в допуске, акт формируется", st == 200, (st, e1 if st != 200 else ""))
    st, e = _pt_make(PT_MUST, dict(PT_OPT, parts=[dict(PT_CAR, share_pct=150), PT_CREDIT]))
    ok("доля больше 100 — 422", st == 422 and "share_pct" in e["errors"]["parts"], e)
    st, e = _pt_make(PT_MUST, dict(PT_OPT, parts=[dict(PT_CAR, class_code="99"), PT_CREDIT]))
    ok("класс не из справочника — 422", st == 422 and "не найден" in e["errors"]["parts"], e)
    st, e = _pt_make({"product_code": "0807", "sum_insured": 1e9, "object_value": 1e9, "region": "Ташкент"},
                     {"parts": [{"class_code": "8", "sum_insured": 1e9}]})
    ok("у продукта с одним классом одна часть — 422 (это обычный акт)", st == 422, e)
    st, e = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, dict(PT_CREDIT, deductible={"pct": 80})]))
    ok("франшиза части больше 50 % — 422", st == 422 and "deductible" in e["errors"]["parts"], e)
    # 4. из договора: перечень объектов
    ct = {"sum_insured": 100_000_000, "premium": 500_000, "term_from": "2026-10-01", "term_to": "2027-09-30",
          "items": [{"name": "Легковой автомобиль в залоге", "sum": 60_000_000},
                    {"name": "Финансовый риск непогашения кредита", "sum": 40_000_000}]}
    st, a4 = _pt_make(PT_MUST, dict(PT_OPT, contract=ct))
    P4 = a4.get("parts") or {}
    ok("из договора: части по перечню объектов (автомобиль → 3, кредит → 14), source contract, на подтверждение",
       st == 200 and P4["source"] == "contract" and not P4["confirmed"]
       and [(x["class_code"], x["sum_insured"]) for x in a4["suggested_parts"]] == [("3", 60e6), ("14", 40e6)],
       (st, P4.get("source"), a4.get("suggested_parts")))
    ok("сверка с договором: премия договора против премии акта (суммы частей), суммы объектов = общей сумме",
       a4["contract_check"]["available"]
       and next(i for i in a4["contract_check"]["items"] if i["code"] == "premium_act")["calculated"] == a4["premium"]["amount"]
       and next(i for i in a4["contract_check"]["items"] if i["code"] == "items_sum")["verdict"] == "ok",
       a4["contract_check"]["items"])
    # 5. КАСКО-пример заметок: средняя ниже минимума первого класса, но каждая часть — не ниже своего минимума
    st, k = _pt_make({"product_code": "0305", "sum_insured": 400_000_000, "object_value": 100_000_000,
                      "region": "Ташкент"},
                     {"losses_3y": {"count": 0}, "documents_provided": True,
                      "parts": [{"class_code": "3", "sum_insured": 100_000_000},
                                {"class_code": "1", "sum_insured": 200_000_000},
                                {"class_code": "13", "sum_insured": 100_000_000}]})
    ki = k["parts"]["items"]
    ok("0305 (ТС + НС + ОТВ): минимумы по классам 1,1 / 0,5 / 1 %; каждая часть не ниже своего минимума",
       [p["rate"]["min_pct"] for p in ki] == [1.1, 0.5, 1.0]
       and all(p["rate"]["applied_pct"] >= p["rate"]["min_pct"] for p in ki), [p["rate"] for p in ki])
    ok("0305: средняя по договору ниже минимума класса 3 (1,1 %) — но это справочно, остановки нет (правило № 5)",
       k["rate"]["reference_pct"] < 1.1 and not any("ниже минимал" in c for c in k["decision"]["checks"])
       and k["decision"]["code"] != "decline", (k["rate"]["reference_pct"], k["decision"]))
    PT_REPORT["0305 по классам"] = {"ставки": [p["rate"]["applied_pct"] for p in ki],
                                    "премии": [p["premium"] for p in ki], "премия": k["premium"]["amount"],
                                    "справочная ставка": k["rate"]["reference_pct"]}
    # 6. 8/9 один объект: 0824 гостиница (8, 9, 16), по согласованию — ставка по каждой части не определена
    st, h = _pt_make({"product_code": "0824", "sum_insured": 3_000_000_000, "object_value": 3_000_000_000,
                      "region": "Ташкент"}, {"object_kind": "hotel", "losses_3y": {"count": 0}})
    hi = h["parts"]["items"]
    ok("0824 (8 + 9 + 16): по умолчанию один объект, сценарии договора — большее из частей",
       st == 200 and h["parts"]["object_mode"] == "one" and all(p["same_object"] for p in hi)
       and [h["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")] ==
       [max(p["scenarios"][k]["amount"] for p in hi if p["scenarios"]["available"]) for k in ("pml", "eml", "mfl")],
       [[p["scenarios"][k]["amount"] for k in ("pml", "eml", "mfl")] for p in hi])
    ok("0824 по согласованию: у частей ставки нет, премия договора не определена, проверка по каждой части",
       h["premium"]["amount"] is None and all(p["rate"]["mode"] == "undefined" for p in hi)
       and sum(1 for c in h["decision"]["checks"] if "определить ставку" in c) == 3, h["decision"]["checks"])
    # 8/9 с явными частями + обязательная часть (0820, ПКМ № 532) на другом объекте
    wh = {"product_code": "0807", "sum_insured": 5_000_000_000, "object_value": 5_000_000_000, "region": "Ташкент"}
    parts = [{"class_code": "8", "sum_insured": 3_000_000_000},
             {"class_code": "9", "sum_insured": 1_500_000_000, "same_object": True,
              "fields": {"class_fields": {"largest_room_value": 600_000_000}}},
             {"class_code": "8", "product_code": "0820", "sum_insured": 500_000_000, "same_object": False,
              "deductible": {"pct": 1}, "object_description": "строящийся склад (СМР)"}]
    st, w = _pt_make(wh, {"object_kind": "warehouse", "losses_3y": {"count": 0}, "parts": parts})
    wi = w["parts"]["items"]
    ok("склад: 8 и 9 — один объект (большее), СМР 0820 — другой объект (прибавляется): правило mixed",
       st == 200 and w["parts"]["totals"]["scenarios"]["rule"] == "mixed"
       and w["scenarios"]["eml"]["amount"] == max(wi[0]["scenarios"]["eml"]["amount"], wi[1]["scenarios"]["eml"]["amount"])
       + wi[2]["scenarios"]["eml"]["amount"], [p["scenarios"]["eml"]["amount"] for p in wi])
    ok("обязательная часть 0820: ставка 0,4 % по ПКМ № 532 без поправок, франшиза сотрудника не применена",
       wi[2]["rate"]["mode"] == "statutory" and wi[2]["rate"]["applied_pct"] == 0.4 and wi[2]["rate"]["adj_pct"] == 0
       and wi[2]["franchise"]["status"] == "statutory" and wi[2]["premium"] == 2_000_000
       and wi[2]["rate"]["min_source"] == "act", (wi[2]["rate"], wi[2]["franchise"]["status"]))
    ok("класс 9 не из состава продукта 0807 — принят с пометкой и проверкой, база — техническая ставка класса",
       wi[1]["class_outside"] and wi[1]["rate"]["base_source"] == "technical"
       and any("Часть 2 (класс 9): проверить класс" in c for c in w["decision"]["checks"]), wi[1]["rate"])
    ok("премия склада = сумма трёх частей", w["premium"]["amount"] == sum(p["premium"] for p in wi))
    PT_REPORT["склад 8/9 + СМР 0820"] = {"премии": [p["premium"] for p in wi], "премия": w["premium"]["amount"],
                                         "EML частей": [p["scenarios"]["eml"]["amount"] for p in wi],
                                         "EML договора": w["scenarios"]["eml"]["amount"]}
    # переключатель «разные объекты» для 0824
    st, h2 = _pt_make({"product_code": "0824", "sum_insured": 3_000_000_000, "object_value": 3_000_000_000,
                       "region": "Ташкент"}, {"object_kind": "hotel", "losses_3y": {"count": 0}, "same_object": False})
    ok("переключатель «разные объекты» (optional.same_object = false): сценарии складываются",
       h2["parts"]["object_mode"] == "different" and h2["scenarios"]["eml"]["amount"] ==
       sum(p["scenarios"]["eml"]["amount"] for p in h2["parts"]["items"] if p["scenarios"]["available"]))
    return aid


def check_parts_langs_files(aid):
    print("42в. Комплексный продукт: три языка, Word и PDF; однопродуктовые акты без изменений")
    fresh()
    st, u = call("GET", f"/act/{aid}", params={"lang": "uz"})
    st2, e = call("GET", f"/act/{aid}", params={"lang": "en"})
    tu, te = all_text(u), all_text(e)
    ok("по-узбекски: таблица «Shartnoma qismlari», «Jami», части «1-qism»",
       "Shartnoma qismlari" in tu and "Shartnoma boʻyicha jami" in tu and "1-qism" in tu
       and u["parts"]["totals"]["premium"] == e["parts"]["totals"]["premium"], tu[:200])
    ok("по-английски: «Contract parts», «Part 2: class 14», средняя «not used to check the minimum»",
       "Contract parts" in te and "Part 2: class 14" in te and "not used to check the minimum" in te)
    ok("ни одного незаполненного шаблона {…} ни на одном языке",
       not any(re.search(r"\{[a-z_]+\}", x) for x in (all_text(u), te, all_text(call("GET", f"/act/{aid}")[1]))))
    ok("русских слов в английском акте частей нет (кроме наименований объектов)",
       not re.search(r"[А-Яа-я]{4,}", "\n".join(li["title"] for s in e["sections"] for li in s.get("lists") or [])),
       [li["title"] for s in e["sections"] for li in s.get("lists") or [] if re.search(r"[А-Яа-я]{4,}", li["title"])][:5])
    for lang, word in (("ru", "Части договора"), ("uz", "Shartnoma qismlari"), ("en", "Contract parts")):
        dx, pd = docx_plain(aid, lang), pdf_plain(aid, lang)
        a = call("GET", f"/act/{aid}", params={"lang": lang})[1]
        prem = flat(act.money(a["premium"]["amount"], lang))
        ok(f"Word и PDF ({lang}): таблица частей и премия договора {prem}",
           word in dx and prem in dx and word in pd and prem in pd, (word in dx, prem in dx, word in pd, prem in pd))
    # однопродуктовые акты: блок parts — single, остальное как прежде
    for code, must, opt in (("0318", CRANE_MUST, CRANE_OPT), ("0807", WH8_MUST, WH8_OPT),
                            ("0832", {"product_code": "0832", "sum_insured": 5e9, "object_value": 5e9,
                                      "region": "Ташкент"}, {"losses_3y": {"count": 0}})):
        st, a = call("POST", "/act/make", {"lang": "ru", "must": must, "optional": opt})
        r = a["rate"]
        ok(f"{code}: один класс — parts.mode single, ставка и премия по-прежнему, без частей в тексте",
           st == 200 and a["parts"]["mode"] == "single" and not a["parts"]["items"] and "suggested_parts" not in a
           and r["mode"] == "tariff" and not r["multi_class"] and "reference_pct" not in r
           and a["premium"]["amount"] == round(premium_of(r["applied_pct"], must["sum_insured"], 365))
           and a["analytics"]["available"] and "Части договора" not in all_text(a)
           and not any(x["label"].startswith("Часть ") for s in a["sections"] for x in s["rows"]),
           (st, a.get("parts"), r.get("mode")))


def check_review_0930():
    print("42г. Замечания контролёра 30.09.2026: кредит, лимиты сценариев, проверка шаблона, части, минимум класса")
    from app import class_templates as ctm
    from app import act_extras as axm
    fresh()
    model_on(False)
    data = ctm.load_file()
    good14 = _json.loads(_json.dumps(data["classes"]["14"]))
    good2 = _json.loads(_json.dumps(data["classes"]["2"]))
    # ---------- 1. кредит: параметры правила ----------
    bad = _json.loads(_json.dumps(good14))
    bad["scenario_rule"]["params"]["max_share_of_loan"] = 0.6
    e1 = ctm.validate(bad, "14")
    bad["scenario_rule"]["params"]["max_share_of_loan"] = "0.5"
    e2 = ctm.validate(bad, "14")
    ok("шаблон 14: доля кредита 0,6 — ошибка (правило № 6, не более 50 %); «0.5» строкой — «нужно число»",
       any("max_share_of_loan" in x and "не больше 0,5" in x for x in e1)
       and any("max_share_of_loan" in x and "нужно число" in x for x in e2), (e1, e2))
    b2 = _json.loads(_json.dumps(good2))
    b2["scenario_rule"]["params"] = {"epidemic_share": 1.5, "months": -3, "x": None}
    e3 = " | ".join(ctm.validate(b2, "2"))
    ok("шаблон 2: доля эпидемии 1,5 — вне (0; 1], отрицательный параметр и null — ошибки",
       "epidemic_share: доля больше 0 и не больше 1" in e3 and "months: число больше нуля" in e3
       and "x: нужно число" in e3, e3)
    ok("проверка параметров — чистая функция: 0,5 у кредита и 0,3 у эпидемии проходят",
       ctm.check_params("credit", {"max_share_of_loan": 0.5}) == [] and ctm.check_params("frequency", {"epidemic_share": 0.3}) == []
       and ctm.check_params("credit", {"max_share_of_loan": 0}) != [])
    # ---------- 1. кредит: чистая функция проверки ----------
    ok("банк по названию: «АКБ Хамкорбанк», «Xalq banki», «Kapitalbank ATB» — банк; «ООО Ромашка» — нет",
       ae.is_bank("АКБ «Хамкорбанк»") and ae.is_bank("Xalq banki") and ae.is_bank("Kapitalbank ATB")
       and ae.is_bank("ООО «Ромашка»") is False and ae.is_bank(None) is None)
    c0 = ae.credit_check(40e6, {}, 0.5, None)
    c1 = ae.credit_check(50e6, {"credit_amount": 100e6, "collateral_value": 60e6}, 0.5,
                         {"kind": "legal", "name": "АКБ «Хамкорбанк»", "source": "contract"})
    c2 = ae.credit_check(40e6, {"credit_amount": 100e6, "collateral_value": 60e6}, 0.5,
                         {"kind": "legal", "name": "ООО «Ромашка»", "source": "request"})
    ok("кредит без суммы и залога: «введите сумму кредита и обеспечение» + страхователь неизвестен",
       [c["code"] for c in c0] == ["credit_need_data", "credit_holder_unknown"], c0)
    ok("кредит 100 млн, залог 60 млн, сумма 50 млн, страхователь банк: превышение 10 млн, уменьшить до 40 млн",
       c1 == [{"code": "credit_over", "params": {"credit": 100_000_000, "collateral": 60_000_000,
                                                 "insurable": 40_000_000, "by": "unsecured", "sum": 50_000_000,
                                                 "excess": 10_000_000, "share_pct": 50}}], c1)
    ok("сумма 40 млн в пределах, страхователь ООО — только проверка страхователя",
       [c["code"] for c in c2] == ["credit_holder_not_bank"] and c2[0]["params"]["holder"] == "ООО «Ромашка»", c2)
    # ---------- 1. кредит в акте ----------
    st, a = _cls_act("14", None, S=40_000_000)
    chk = [flat(c) for c in a["decision"]["checks"]]
    ok("акт класса 14 без суммы кредита и залога: проверка андеррайтеру, рекомендация не выше «с оговорками»",
       st == 200 and any(c.startswith("Кредит: проверить, что страховая сумма не превышает необеспеченную часть и "
                                      "50 % суммы кредита — введите сумму кредита и обеспечение") for c in chk)
       and a["decision"]["code"] in ("accept_with_clauses", "decline"), (st, chk, a.get("decision", {}).get("code")))
    ok("страхователь неизвестен — пункт проверки «страхователь и плательщик премии — банк-кредитор»",
       any("страхователь и плательщик премии — банк-кредитор" in c and "не найден" in c for c in chk), chk)
    st, a = _cls_act("14", {"credit_amount": 100_000_000, "collateral_value": 60_000_000}, S=50_000_000)
    chk = [flat(c) for c in a["decision"]["checks"]]
    ok("кредит 100 млн, залог 60 млн, сумма 50 млн: «страховая сумма … выше допустимой … — уменьшить до 40 000 000 сум»",
       any("выше допустимой" in c and "уменьшить до 40 000 000 сум" in c and "превышение 10 000 000 сум" in c
           for c in chk) and a["decision"]["code"] != "accept", chk)
    ct_bad = {"sum_insured": 40_000_000, "premium": 240_000, "policyholder": "ООО «Ромашка»"}
    st, a = _cls_act("14", {"credit_amount": 100_000_000, "collateral_value": 60_000_000}, S=40_000_000,
                     optional={"contract": ct_bad})
    chk = [flat(c) for c in a["decision"]["checks"]]
    ok("страхователь в договоре — ООО, не банк: «в документе (договор страхования) страхователь — ООО «Ромашка»»",
       st == 200 and any("страхователь и плательщик премии — банк-кредитор" in c and "ООО «Ромашка»" in c
                         and "договор страхования" in c for c in chk) and a["decision"]["code"] != "accept", (st, chk))
    st, a = _cls_act("14", {"credit_amount": 100_000_000, "collateral_value": 60_000_000}, S=40_000_000,
                     optional={"contract": dict(ct_bad, policyholder="АКБ «Хамкорбанк»")})
    chk = [flat(c) for c in a["decision"]["checks"]]
    ok("страхователь — банк, сумма в пределах: кредитных проверок нет",
       st == 200 and not any("банк-кредитор" in c or "Кредит:" in c for c in chk), chk)
    st, a = _cls_act("14", {"credit_amount": 100_000_000, "collateral_value": 60_000_000}, S=50_000_000, lang="en",
                     optional={"contract": ct_bad})
    ok("кредитные проверки на английском: «reduce to 40,000,000 UZS», «the lending bank»",
       any("reduce to" in flat(c) for c in a["decision"]["checks"])
       and any("lending bank" in c and "ООО" in c for c in a["decision"]["checks"]), a["decision"]["checks"])
    # ---------- 2. лимиты ответственности не выше страховой суммы ----------
    r = axm.simple_scenarios("limit", 100, 100, {"limit_per_case": 300, "limit_aggregate": 500})
    r2 = axm.simple_scenarios("full_limit", 100, 100, {"limit_per_case": 300})
    ok("limit: лимит на случай 300 и годовой 500 при сумме 100 → PML = EML = MFL = 100, две пометки",
       [r[s]["amount"] for s in ("PML", "EML", "MFL")] == [100, 100, 100]
       and [x["code"] for x in r["assumptions"]] == ["as_tpl_limit_case_over", "as_tpl_limit_aggregate_over"], r)
    ok("full_limit: лимит 300 при сумме 100 → всё 100, пометка «лимит выше страховой суммы»",
       [r2[s]["amount"] for s in ("PML", "EML", "MFL")] == [100, 100, 100]
       and r2["assumptions"] == [{"code": "as_tpl_limit_case_over", "params": {"limit": 300, "sum": 100}}], r2)
    st, a = _cls_act("13", {"activity_kind": "trade", "limit_per_case": 300_000_000, "limit_aggregate": 1_000_000_000},
                     S=200_000_000)
    sc = a["scenarios"]
    asm = " ".join(flat(x["text"]) for x in sc["assumptions"])
    ok("класс 13, сумма 200 млн, лимиты 300 млн и 1 млрд: PML/EML/MFL = 200 млн, в допущениях — «взята страховая сумма»",
       [sc[k]["amount"] for k in ("pml", "eml", "mfl")] == [200_000_000] * 3
       and "лимит на один случай 300 000 000 сум выше страховой суммы 200 000 000 сум — взята страховая сумма" in asm
       and "годовой лимит 1 000 000 000 сум выше страховой суммы" in asm, (sc, asm))
    # ---------- 4. проверка шаблона: размер, подписи, переводы ----------
    errs_file = {c: ctm.validate(t_, c) for c, t_ in data["classes"].items()}
    ok("поставленный docs/act_class_templates.json проходит новую проверку (все 17 классов, переводы uz/en)",
       not any(errs_file.values()), {k: v for k, v in errs_file.items() if v})
    b = _json.loads(_json.dumps(good14))
    b["must"][0]["label"].pop("uz")
    b["optional"][0]["label"]["en"] = ""
    b["risks"]["items"][0]["label"].pop("en")
    b["name"].pop("uz")
    b["notes"][0]["text"].pop("ru")
    b["documents"]["items"][0]["ru"] = "х" * 501
    e = " | ".join(ctm.validate(b, "14"))
    ok("шаблон без переводов: 422-перечень — name.uz, must.credit_amount.uz, optional.borrower_industry.en, "
       "risks.cr_insolvency.en; подпись без ru; подпись длиннее 500",
       "нет перевода uz/en" in e and "name.uz" in e and "must.credit_amount.uz" in e
       and "optional.borrower_industry.en" in e and "risks.cr_insolvency.en" in e
       and "подпись без ru: notes.credit_rule.text" in e and "длиннее 500 знаков" in e, e)
    big = _json.loads(_json.dumps(good14))
    big["notes"] = [{"code": f"n{i}", "text": {"ru": "я" * 400, "uz": "a" * 400, "en": "a" * 400}} for i in range(200)]
    e = ctm.validate(big, "14")
    ok("шаблон больше 200 КБ — ошибка размера", len(e) == 1 and "больше 200 КБ" in e[0], e)
    HEADERS.append(_admin_header("tpl_review_admin"))
    try:
        bad = _json.loads(_json.dumps(good14))
        bad["scenario_rule"]["params"]["max_share_of_loan"] = 0.7
        bad["must"][1]["label"].pop("en")
        st, e = call("PUT", "/act/templates/14", {"template": bad})
        errs = " | ".join(e.get("errors") or [])
        ok("PUT шаблона 14 с долей 0,7 и без en у поля — 422 с перечнем",
           st == 422 and "max_share_of_loan" in errs and "must.collateral_value.en" in errs, (st, errs))
        st, e = call("PUT", "/act/templates/14", {"template": _json.loads(_json.dumps(good14))})
        ok("PUT хорошего шаблона 14 проходит (новая версия)", st == 200, (st, e))
    finally:
        HEADERS.clear()
    # ---------- мелочи: пометки «сверх приложения А», убытки за 5 лет, единственный вид объекта ----------
    C = data["classes"]
    fld = lambda c, code: next(f for f in C[c]["must"] + C[c]["optional"] if f["code"] == code)
    exp_ok = lambda x: x.get("expert") is True and x["note"]["ru"] == "сверх приложения А, экспертно"
    ok("сверх приложения А — expert: true: факторы 13 liab_limit/liab_turnover, MFL класса 10, поля сценариев 1, 2, 7, 16, 17",
       all(exp_ok(f) for f in C["13"]["factors"] if f["code"] in ("liab_limit", "liab_turnover"))
       and exp_ok(C["10"]["scenario_rule"]["expert_scenarios"]["MFL"])
       and all(exp_ok(fld(c, k)) for c, k in (("1", "people_in_one_place"), ("2", "avg_visits"), ("2", "avg_bill"),
                                              ("7", "limit_per_shipment"), ("7", "accumulation_value"),
                                              ("16", "monthly_loss"), ("17", "limit_per_dispute")))
       and not fld("1", "insured_count").get("expert"))
    ok("классы 5, 6, 11, 12: убытки за 5 лет — поле losses_3y (так понимает акт), подпись «за 5 лет (… история убытков)»",
       all(fld(c, "losses_3y")["type"] == "losses" and "за 5 лет (в расчёте используется как история убытков)"
           in fld(c, "losses_3y")["label"]["ru"] and fld(c, "losses_3y")["period_years"] == 5
           and not any(f["code"] == "losses_5y" for f in C[c]["optional"]) for c in ("5", "6", "11", "12")))
    st, t14 = call("GET", "/act/templates/14", params={"lang": "ru"})
    st3, t3 = call("GET", "/act/templates/3", params={"lang": "ru"})
    st_, lst = call("GET", "/act/templates", params={"lang": "ru"})
    one = {x["class_code"]: x for x in lst["templates"]}
    ok("GET /act/templates/14: один вид объекта «кредит» — single_kind: true, default_kind: loan; у класса 3 — false",
       t14["template"]["object"]["single_kind"] is True and t14["template"]["object"]["default_kind"] == "loan"
       and t3["template"]["object"]["single_kind"] is False and one["14"]["single_kind"] and not one["3"]["single_kind"],
       (t14["template"]["object"], t3["template"]["object"].get("single_kind")))
    # ---------- 5, 6. части: недостающие поля частей 2+, минимум класса из текста тарифа ----------
    cr_no_term = dict(PT_CREDIT, fields={"class_fields": {"credit_amount": 100_000_000, "collateral_value": 60_000_000}})
    st, p = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, cr_no_term]))
    chk = [flat(c) for c in p["decision"]["checks"]]
    ok("часть 2 (класс 14) без срока кредита: «Часть 2 (класс 14): уточнить срок кредита, месяцев» в решении",
       st == 200 and "Часть 2 (класс 14): уточнить Срок кредита, месяцев" in chk, chk)
    ok("у части 2 класса 14 — вид объекта по умолчанию «loan» (единственный), в расчёт не идёт",
       p["parts"]["items"][1]["object_kind_default"] == "loan" and p["parts"]["items"][1]["object_kind"] is None)
    ok("часть 2 кредита: страхователь не найден — проверка по части, сумма 40 млн в пределах — превышения нет",
       any(c.startswith("Часть 2 (класс 14): проверить, что страхователь и плательщик премии — банк-кредитор")
           for c in chk) and not any("выше допустимой" in c for c in chk), chk)
    st, p = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, {"class_code": "14", "sum_insured": 40_000_000}]))
    chk = [flat(c) for c in p["decision"]["checks"]]
    ok("часть 2 кредита без полей: «Часть 2 (класс 14). Кредит: … введите сумму кредита и обеспечение» и перечень полей",
       any(c.startswith("Часть 2 (класс 14). Кредит: проверить, что страховая сумма не превышает") for c in chk)
       and any(c.startswith("Часть 2 (класс 14): уточнить Сумма кредита, Стоимость обеспечения") for c in chk), chk)
    how = {}
    for lg in ("ru", "uz", "en"):
        st, x = _pt_make(PT_MUST, dict(PT_OPT, parts=[PT_CAR, PT_CREDIT]), lang=lg)
        how[lg] = (x["parts"]["items"][1]["rate"]["how"], x["parts"]["items"][0]["rate"]["how"], all_text(x))
    ok("часть 2 (класс 14): «минимум класса 14 по тарифной политике (из текста тарифа продукта 0312)», не «ставка продукта»",
       any("минимум класса 14 по тарифной политике (из текста тарифа продукта 0312)" in flat(h) for h in how["ru"][0])
       and not any("ставка продукта" in h or "минимальной ставки продукта" in h for h in how["ru"][0])
       and any("ставка продукта 0312" in h for h in how["ru"][1]), how["ru"][:2])
    ok("то же на узбекском и английском",
       any("14-klass minimumi (0312 mahsuloti tarif matnidan)" in flat(h) for h in how["uz"][0])
       and any("class 14 minimum under the tariff policy" in h and "product 0312" in h for h in how["en"][0]),
       (how["uz"][0], how["en"][0]))
    ok("аналитика части 2: строка «Минимум класса 14 по тарифной политике (из текста тарифа продукта 0312)»",
       "Минимум класса 14 по тарифной политике (из текста тарифа продукта 0312)" in how["ru"][2]
       and "Class 14 minimum under the tariff policy (from the tariff text of product 0312)" in how["en"][2])
    ok("ни одного незаполненного шаблона {…} в актах с кредитными проверками",
       not any(re.search(r"\{[a-z_]+\}", v[2]) for v in how.values()))


def check_templates_sync():
    print("41г. Сервер: база без таблицы class_templates — refsync доводит; новая версия файла; db_build на копии")
    import importlib
    from app import class_templates as ctm, refsync
    folder = Path(tempfile.mkdtemp(prefix="surveyor-tpl-"))
    try:
        disk = folder / "disk.db"
        db.snapshot(db.DB_PATH, disk)
        import sqlite3
        con = sqlite3.connect(str(disk))
        con.execute("DROP TABLE IF EXISTS class_templates")
        con.commit()
        con.close()
        ctm.reset_cache()
        res = refsync.sync_templates(disk)
        con = sqlite3.connect(str(disk))
        n = con.execute("SELECT COUNT(*), COUNT(DISTINCT class_code) FROM class_templates").fetchone()
        ok("база без таблицы: refsync.sync_templates создал таблицу и довёл 17 шаблонов версии 1.0.1",
           res["status"] == "обновлено" and n == (17, 17), (res.get("status"), n))
        res2 = refsync.sync_templates(disk)
        ok("повторный запуск — «актуально», строк не прибавилось",
           res2["status"] == "актуально" and con.execute("SELECT COUNT(*) FROM class_templates").fetchone()[0] == 17)
        # правка администратора 1.1 на «сервере», потом образ приносит файл 2.0
        con.execute("INSERT INTO class_templates (class_code, version, json, source, file_version, updated_at, updated_by,"
                    " calibrated) SELECT class_code, '1.1', json, 'admin', '1.0', '2026-09-30T12:00:00', 'админ', 0 "
                    "FROM class_templates WHERE class_code='13'")
        con.commit()
        src = _json.loads(ctm.TEMPLATES_FILE.read_text(encoding="utf-8"))
        src["version"] = "2.0"
        newer = folder / "tpl.json"
        newer.write_text(_json.dumps(src, ensure_ascii=False), encoding="utf-8")
        orig = ctm.TEMPLATES_FILE
        ctm.TEMPLATES_FILE = newer
        try:
            res3 = refsync.sync_templates(disk)
        finally:
            ctm.TEMPLATES_FILE = orig
            ctm.reset_cache()
            ctm._file_cache.update(mtime=None, data=None)
        hist = [tuple(r) for r in con.execute("SELECT version, source FROM class_templates WHERE class_code='13' ORDER BY id")]
        ok("файл 2.0 новее — добавлен всем 17 классам; история класса 13: 1.0.1 файл, 1.1 админ, 2.0 файл",
           len(res3["added"]) == 17 and hist == [("1.0.1", "file"), ("1.1", "admin"), ("2.0", "file")], (res3, hist))
        con.close()
        # tools/db_build.py — заполнение из JSON на копии (рабочая база не открывается)
        build_copy = folder / "build.db"
        db.snapshot(db.DB_PATH, build_copy)
        con = sqlite3.connect(str(build_copy))
        con.execute("DROP TABLE IF EXISTS class_templates")
        con.commit()
        con.close()
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
        dbb = importlib.import_module("db_build")
        saved = dbb.DB
        dbb.DB = build_copy
        buf = io.StringIO()
        try:
            import contextlib
            with contextlib.redirect_stdout(buf):
                dbb.main()
        finally:
            dbb.DB = saved
        con = sqlite3.connect(str(build_copy))
        nb = con.execute("SELECT COUNT(*) FROM class_templates").fetchone()[0]
        con.close()
        ok("tools/db_build.py на копии: таблица class_templates заполнена из JSON (17 строк)",
           nb == 17 and "шаблонов классов добавлено: 17" in buf.getvalue(), (nb, buf.getvalue()[-300:]))
    finally:
        ctm.reset_cache()
        shutil.rmtree(folder, ignore_errors=True)


def main():
    ORIG.update(chat_raw=llm.chat_raw, enabled=llm.enabled, supports_files=llm.supports_files, post=llm._post)
    llm.chat_raw = fake_chat_raw
    llm._post = no_network
    folder = Path(tempfile.mkdtemp(prefix="surveyor-act-"))
    try:
        with temp_db("surveyor-act-test.db"):
            db.ensure_schema()
            act.DIR = folder
            guest.reset()
            check_engine()
            a, sid = check_crane()
            aid = a["id"]
            check_value()
            check_franchise()
            check_no_rate()
            check_no_model()
            check_pd()
            check_files(aid)
            check_langs(aid)
            check_polish()
            check_access(aid)
            check_bomb()
            check_threadpool()
            check_body_limit()
            check_limits()
            check_budget()
            check_sources()
            check_misc()
            check_scenarios()
            check_documents()
            fr_aid, fr_hand = check_franchise_apply()
            check_measures()
            check_new_langs(fr_aid)
            check_new_files(fr_aid, fr_hand)
            check_speed()
            check_doc_limits()
            check_dtd_prolog()
            check_fr_proposed()
            check_alt_base()
            check_protection_other()
            check_lang_fields()
            check_minor()
            check_screen()
            check_market_engine()
            check_market_links()
            shots_sid = check_market_shots()
            check_market_make(shots_sid)
            br_sid, br_b = check_branch_text()
            br_scan = check_branch_scan()
            br_aid, _ = check_branch_make(br_sid, br_b, br_scan)
            check_branch_langs_files(br_aid)
            check_contract_text()
            check_contract_long()
            check_contract_scan()
            check_contract_ai_assist()
            ct_aid = check_contract_make()
            x_aid = check_contract_cross()
            check_contract_langs_files(ct_aid, x_aid)
            check_trust_edits()
            check_doc_kind_title()
            check_essentials_wording()
            check_cross_same()
            check_input_bounds()
            check_ct_ai_mask_labels()
            check_parties_amounts()
            check_pdf_time_limit()
            check_analytics()
            check_review_fixes()
            check_docx_reader()
            check_contract_template()
            check_contract_template_filled()
            check_templates_ref()
            check_templates_api()
            check_templates_act()
            check_templates_sync()
            check_parts_engine()
            pt_aid = check_parts_make()
            check_parts_langs_files(pt_aid)
            check_review_0930()
            check_send(aid)
            check_cleanup(sid, aid)
    finally:
        llm.chat_raw, llm.enabled, llm.supports_files, llm._post = (ORIG["chat_raw"], ORIG["enabled"],
                                                                    ORIG["supports_files"], ORIG["post"])
        if "cbu" in ORIG:
            from app import valuation_sources as vs
            vs.cbu_usd_rate = ORIG["cbu"]
        shutil.rmtree(folder, ignore_errors=True)
    if BR_REPORT:
        print("\nзапрос филиала, продукт 0832:")
        for k, v in BR_REPORT.items():
            print("  ", k, v)
    if CT_REPORT:
        print("\nдоговор страхования:")
        for k, v in CT_REPORT.items():
            print("  ", k, v)
    if AN_REPORT:
        print("\nаналитика раздела 4:")
        for k, v in AN_REPORT.items():
            print("  ", k, v)
    if TPL_REPORT:
        print("\nшаблоны классов:")
        for k, v in TPL_REPORT.items():
            print("  ", k, v)
    if PT_REPORT:
        print("\nкомплексные продукты по частям:")
        for k, v in PT_REPORT.items():
            print("  ", k, v)
    if SCEN_REPORT:
        print("\nсценарии (сумма, % страховой суммы):")
        for k, v in SCEN_REPORT.items():
            print("  ", k, v)
    print(f"\nитог: ок {passed}, плохо {failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
