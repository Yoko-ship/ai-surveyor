"""
Финансы компании: отчёты о страховых резервах по периодам (app/finance.py).

Запуск из корня проекта:
    set PYTHONIOENCODING=utf-8
    sandbox\\.venv\\Scripts\\python.exe tests\\test_finance.py

Рабочая база не открывается на запись: STORAGE_DIR указывает на копию в sandbox/test_finance,
а поверх неё tests/tmpdb.temp_db даёт ещё одну временную копию. Живой сервер не трогаем.
"""
import asyncio
import io
import json as _json
import os
import shutil
import sqlite3
import sys
from pathlib import Path
from urllib.parse import urlencode

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

# копия рабочей базы в sandbox/ — до импорта app.db, который читает STORAGE_DIR при импорте
STORE = ROOT / "sandbox" / "test_finance"
shutil.rmtree(STORE, ignore_errors=True)
STORE.mkdir(parents=True)
_src = sqlite3.connect(f"file:{ROOT / 'data' / 'surveyor.db'}?mode=ro", uri=True)
_dst = sqlite3.connect(STORE / "surveyor.db")
_src.backup(_dst)
_src.close()
_dst.close()
os.environ["STORAGE_DIR"] = str(STORE)
os.environ.setdefault("SURVEYOR_DEV", "1")

import openpyxl  # noqa: E402

from tmpdb import temp_db  # noqa: E402
from app import auth, db, capacity, finance  # noqa: E402
from app.main import app  # noqa: E402

assert db.DB_PATH == STORE / "surveyor.db", db.DB_PATH
TOKENS = {}
USERS = {"тест-фин-админ": "админ", "тест-фин-сотр": "андеррайтер"}
ADMIN, STAFF = "тест-фин-админ", "тест-фин-сотр"


def call(method, path, body=None, who=ADMIN, raw=None, ctype="application/json"):
    """Вызов приложения по ASGI от имени вошедшего who. Возвращает (статус, тело, content-type)."""
    if "?" in path:
        path, query = path.split("?", 1)
    else:
        query = ""
    payload = raw if raw is not None else (_json.dumps(body, ensure_ascii=False).encode() if body is not None else b"")
    headers = [(b"host", b"test"), (b"content-type", ctype.encode()), (b"content-length", str(len(payload)).encode())]
    if who:
        headers.append((b"cookie", f"sid={TOKENS[who]}".encode()))
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method, "scheme": "http",
             "path": path, "raw_path": path.encode(), "root_path": "", "query_string": query.encode(),
             "headers": headers, "client": ("127.0.0.1", 0), "server": ("test", 80)}
    out = {"status": None, "chunks": [], "ctype": ""}

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"] = msg["status"]
            out["ctype"] = dict(msg.get("headers") or {}).get(b"content-type", b"").decode()
        elif msg["type"] == "http.response.body":
            out["chunks"].append(msg.get("body") or b"")

    asyncio.run(app(scope, receive, send))
    data = b"".join(out["chunks"])
    if "json" in out["ctype"]:
        return out["status"], _json.loads(data.decode("utf-8")), out["ctype"]
    return out["status"], data, out["ctype"]


def upload(data: bytes, name="резервы.xlsx", who=ADMIN):
    b = "----finance-test-boundary"
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() + \
        data + f"\r\n--{b}--\r\n".encode()
    return call("POST", "/finance/reserves/parse", raw=body, who=who, ctype=f"multipart/form-data; boundary={b}")


def xlsx(rows, title="Резервы") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def setup():
    db.ensure_schema()
    with db.tx() as con:
        con.execute("DELETE FROM reserve_reports")          # только во временной копии: мешают выбору периода
        for login, role in USERS.items():
            con.execute("DELETE FROM sessions WHERE user_id IN (SELECT id FROM users WHERE login=?)", (login,))
            con.execute("DELETE FROM users WHERE login=?", (login,))
            cur = con.execute("INSERT INTO users (login, full_name, role, branch, password_hash, salt, status, created_at)"
                              " VALUES (?,?,?,?,?,?,?,?)", (login, "Тест Финансы", role, "тест", "x", "y", "активен", db.now()))
            u = db.rows(con, "SELECT * FROM users WHERE id=?", cur.lastrowid)[0]
            TOKENS[login], _ = auth.create_session(con, u, "127.0.0.1", "test")


def ok(cond, what):
    if not cond:
        raise AssertionError(what)
    print("  ok —", what)


MANUAL = [{"reserve_code": "rnp", "group_code": "1", "amount": 1_000_000_000},
          {"reserve_code": "rnp", "group_code": "2", "amount": 500_000_000.5},
          {"reserve_code": "rzu", "group_code": "", "amount": 200_000_000},
          {"reserve_code": "rpnu", "group_code": "", "amount": 100_000_000},
          {"reserve_code": "stab", "group_code": "ОСГОР", "amount": 30_000_000}]


def test_manual():
    print("ручной ввод")
    s, st, _ = call("GET", "/finance/reserves/structure")
    ok(s == 200 and [k["code"] for k in st["kinds"]] == ["rnp", "rzu", "rpnu", "stab", "cat_reserve", "other"],
       "состав резервов по Положению 1882 — 6 видов, как колонки reserve_reports")
    ok([p["code"] for p in st["periods"]] == ["Q1", "Q2", "Q3", "Q4", "Y"], "периоды: 4 квартала и годовая")
    s, r, _ = call("POST", "/finance/reserves", {"period_year": 2025, "period_type": "Q2", "lines": MANUAL,
                                                 "base_premium_12m": 2_000_000_000, "source": "вручную"})
    ok(s == 200 and r["report_date"] == "2025-06-30", "сохранено, отчётная дата — конец квартала")
    ok(abs(r["total"] - 1_830_000_000.5) < 1e-6, "итог считается сам (1 830 000 000,5)")
    s, d, _ = call("GET", f"/finance/reserves/{r['id']}")
    ok(s == 200 and len(d["lines"]) == 5 and d["created_by"] == ADMIN and d["source"] == "вручную",
       "просмотр: строки, кто внёс (логин, не ФИО), источник")
    return r["id"]


def test_validation():
    print("проверка входных данных")
    base = {"period_year": 2025, "period_type": "Q1", "source": "вручную"}
    cases = [
        (dict(base, lines=[{"reserve_code": "rnp", "group_code": "1", "amount": -5}]), "меньше нуля"),
        (dict(base, lines=[{"reserve_code": "rnp", "group_code": "", "amount": 5}]), "не подходит"),
        (dict(base, lines=[{"reserve_code": "rpm", "group_code": "", "amount": 5}]), "неизвестный вид"),
        (dict(base, lines=[{"reserve_code": "rzu", "group_code": "", "amount": 5},
                           {"reserve_code": "rzu", "group_code": "", "amount": 6}]), "дважды"),
        (dict(base, lines=[]), "ни одной суммы"),
    ]
    for body, needle in cases:
        s, r, _ = call("POST", "/finance/reserves", body)
        ok(s == 400 and any(needle in e for e in r["detail"]["errors"]), f"400: «{needle}»")
    line = [{"reserve_code": "rzu", "group_code": "", "amount": 5}]
    s, r, _ = call("POST", "/finance/reserves", dict(base, period_year=1999, lines=line))
    ok(s == 400 and "2000" in r["detail"], "год вне 2000–2100 — 400")
    s, r, _ = call("POST", "/finance/reserves", dict(base, period_type="Q5", lines=line))
    ok(s == 400, "период Q5 — 400")
    s, r, _ = call("POST", "/finance/reserves", dict(base, period_year=2100, period_type="Y", lines=line))
    ok(s == 400 and "не наступила" in r["detail"], "будущий период — 400")
    s, r, _ = call("POST", "/finance/reserves", dict(base, lines=[{"reserve_code": "rzu", "amount": "abc"}]))
    ok(s in (400, 422), "не число в сумме — отказ")
    s, r, _ = call("POST", "/finance/reserves", dict(base, source="робот", lines=line))
    ok(s == 400, "источник только «вручную» или «файл»")


def test_replace(first_id):
    print("повторный период — замена с историей")
    body = {"period_year": 2025, "period_type": "Q2", "source": "вручную",
            "lines": [{"reserve_code": "rnp", "group_code": "3", "amount": 777}]}
    s, r, _ = call("POST", "/finance/reserves", body)
    ok(s == 409 and r["detail"]["existing"]["id"] == first_id and "Заменить" in r["detail"]["message"],
       "без согласия — 409 «заменить?» и сведения о прежнем отчёте")
    s, r, _ = call("POST", "/finance/reserves", dict(body, replace=True))
    ok(s == 200 and r["replaced"] == first_id and r["total"] == 777, "с согласием — заменено")
    s, d, _ = call("GET", f"/finance/reserves/{first_id}")
    ok(d["status"] == "заменён" and d["closed_by"] == ADMIN and len(d["lines"]) == 5,
       "старая версия не удалена: статус «заменён», строки на месте")
    ok([h["id"] for h in d["history"]] == [r["id"], first_id], "история версий периода")
    s, lst, _ = call("GET", "/finance/reserves")
    ok(sum(1 for p in lst["items"] if p["period_year"] == 2025 and p["period_type"] == "Q2") == 1,
       "в списке один действующий отчёт на период")
    s, lst, _ = call("GET", "/finance/reserves?history=1")
    ok(any(p["id"] == first_id for p in lst["items"]), "с историей — видна и прежняя версия")
    with db.tx() as con:
        a = db.rows(con, "SELECT * FROM audit WHERE action=? AND entity=?", "резервы: отчёт заменён", "резервы:2025-Q2")
    ok(a and _json.loads(a[-1]["detail"])["old_id"] == first_id and a[-1]["who"] == ADMIN, "замена записана в audit")
    return r["id"]


def test_template_roundtrip():
    print("шаблон: скачать, заполнить, загрузить обратно")
    s, data, ctype = call("GET", "/finance/reserves/template.xlsx")
    ok(s == 200 and data[:2] == b"PK" and "spreadsheetml" in ctype, "шаблон отдаётся как xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ok(wb.sheetnames == ["Резервы", "Инструкция"], "листы «Резервы» и «Инструкция»")
    ws = wb["Резервы"]
    ok([c.value for c in ws[1]] == ["Вид резерва", "Учётная группа", "Сумма, сум"], "колонки шаблона")
    s, p, _ = upload(data, "пустой.xlsx")
    ok(s == 200 and any("ни одной суммы" in e for e in p["errors"]), "пустой шаблон — ошибка «ни одной суммы»")
    expected, n = {}, 0
    for row in ws.iter_rows(min_row=2):
        n += 1
        kind = finance.detect_kind(row[0].value)
        if kind == "base":
            row[2].value = "3 000 000 000"
            continue
        amt = 1_000_000 * n + 0.25 if n % 2 else f"{n} 234 567,89"      # числа и текст с пробелами
        row[2].value = amt
        grp = finance.detect_group(row[1].value)
        expected[(kind, grp)] = finance.to_amount(amt)
    buf = io.BytesIO()
    wb.save(buf)
    s, p, _ = upload(buf.getvalue(), "резервы 2025.xlsx")
    ok(s == 200 and not p["errors"] and not p["warnings"], "заполненный шаблон разобран без ошибок и предупреждений")
    got = {(x["reserve_code"], x["group_code"]): x["amount"] for x in p["lines"]}
    ok(got == expected and len(got) == 11, "все 11 строк распознаны, суммы совпали до копейки")
    ok(p["base_premium_12m"] == 3e9 and abs(p["total"] - sum(expected.values())) < 1e-6, "итог и базовая премия")
    body = {"period_year": 2025, "period_type": "Y", "source": "файл", "file_name": p["file_name"],
            "base_premium_12m": p["base_premium_12m"],
            "lines": [{k: x[k] for k in ("reserve_code", "group_code", "amount")} for x in p["lines"]]}
    s, r, _ = call("POST", "/finance/reserves", body)
    ok(s == 200 and r["report_date"] == "2025-12-31", "сохранено из предпросмотра; годовая — на 31.12")
    s, d, _ = call("GET", f"/finance/reserves/{r['id']}")
    saved = {(x["reserve_code"], x["group_code"]): x["amount"] for x in d["lines"]}
    ok(saved == expected and d["source"] == "файл" and d["file_name"] == "резервы 2025.xlsx", "в базе — без потерь")
    return r["id"]


def test_bad_files():
    print("кривые файлы — понятные ошибки")
    s, r, _ = upload(b"hello, it is not excel", "резервы.xlsx")
    ok(s == 400 and "PK" in r["detail"], "не xlsx — 400 с объяснением")
    s, r, _ = upload(xlsx([["a"]]), "резервы.xls")
    ok(s == 400 and ".xlsx" in r["detail"], "расширение .xls — 400")
    s, r, _ = upload(b"PK\x03\x04broken", "резервы.xlsx")
    ok(s == 400 and "не читается" in r["detail"], "битый архив — 400")
    s, r, _ = upload(xlsx([["Показатель", "Значение"], ["РНП", 5]]))
    ok(s == 400 and "заголовок" in r["detail"], "нет нужных колонок — 400 «не найден заголовок»")
    # заголовки по-узбекски и с другим регистром, лист называется иначе
    rows = [["Отчёт о резервах"], [],
            ["ZAXIRA TURI", "Hisob guruhi", "Summa, so'm"],
            ["РНП", "I", "1 000 000"],                       # римская цифра группы
            ["резерв незаработанной премии", "2-я", -5],       # минус
            ["РЗУ", "", "abc"],                               # не число
            ["Стабилизационный резерв ОСГО", None, 10],        # вид ОСГО — из названия
            ["Прибыль прошлых лет", "", 999],                 # неизвестная строка
            ["РПМ", "", 50],                                  # РПМ компания не ведёт
            ["РНП", "", 7],                                   # РНП без группы
            ["РПНУ", "1", 3], ["РПНУ", "2", 4],               # РПНУ по группам — сложить с предупреждением
            ["РНП", "1", 1],                                  # повтор
            ["Итого", "", 123]]
    s, p, _ = upload(xlsx(rows, "Лист1"))
    ok(s == 200, "файл с ошибками всё равно показывается в предпросмотре")
    errs, warns = " | ".join(p["errors"]), " | ".join(p["warnings"])
    ok("строка 5" in errs and "меньше нуля" in errs, "минус — ошибка с номером строки")
    ok("строка 6" in errs and "не число" in errs, "текст вместо суммы — ошибка с номером строки")
    ok("строка 10" in errs and "учётную группу" in errs, "РНП без группы — ошибка")
    ok("строка 13" in errs and "уже была" in errs, "повтор строки — ошибка")
    ok("Прибыль прошлых лет" in warns and "не распознан" in warns, "неизвестная строка — в предупреждениях, не молча")
    ok("РПМ" in warns, "РПМ — предупреждение, в итог не входит")
    ok("группа не нужна" in warns, "РПНУ с группами — предупреждение, суммы сложены")
    ok("Итог в файле" in warns, "итог файла сверяется с суммой строк")
    got = {(x["reserve_code"], x["group_code"]): x["amount"] for x in p["lines"]}
    ok(got == {("rnp", "1"): 1_000_000, ("stab", "ОСГО"): 10, ("rpnu", ""): 7}, "распознанное — ровно то, что верно")


def test_rights(pid):
    print("права: сотрудник — 403, без входа — 401")
    for method, path in (("GET", "/finance/reserves"), ("GET", "/finance/reserves/structure"),
                         ("GET", "/finance/reserves/template.xlsx"), ("GET", f"/finance/reserves/{pid}"),
                         ("DELETE", f"/finance/reserves/{pid}")):
        s, _, _ = call(method, path, who=STAFF)
        ok(s == 403, f"сотрудник {method} {path} — 403")
    s, _, _ = call("POST", "/finance/reserves", {"period_year": 2024, "period_type": "Y", "source": "вручную",
                                                 "lines": MANUAL}, who=STAFF)
    ok(s == 403, "сотрудник не может внести отчёт")
    s, _, _ = upload(xlsx([["Вид резерва", "Сумма"]]), who=STAFF)
    ok(s == 403, "сотрудник не может загрузить файл")
    s, _, _ = call("GET", "/finance/reserves", who=None)
    ok(s == 401, "без входа — 401")


def test_capacity(q2_id, y_id):
    print("ёмкость берёт последний период")
    with db.tx() as con:
        c = capacity.capacity(con)
    ok(c["reserves_period"]["id"] == y_id, "годовая 2025 новее 2 квартала 2025")
    s, q4, _ = call("POST", "/finance/reserves", {"period_year": 2025, "period_type": "Q4", "source": "вручную",
                                                  "lines": MANUAL})
    with db.tx() as con:
        c = capacity.capacity(con)
        _, rows = finance.latest_reserve_rows(con)
    ok(c["reserves_period"]["id"] == y_id, "при одной дате 31.12 годовая важнее 4 квартала")
    ok(abs(c["reserves"] - c["reserves_period"]["total"]) < 1e-6, "резервы в ёмкости = итог периода")
    ok(any(r["scope_type"] == "группа" for r in rows) and rows[-1]["scope_type"] == "итого"
       and rows[-1]["base_premium_12m"] == 3e9, "строки в форме reserve_reports: группы, виды, итого с базовой премией")
    s, _, _ = call("DELETE", f"/finance/reserves/{y_id}")
    ok(s == 200, "админ удалил годовую")
    s, r, _ = call("DELETE", f"/finance/reserves/{y_id}")
    ok(s == 409, "повторное удаление — 409")
    with db.tx() as con:
        c = capacity.capacity(con)
        a = db.rows(con, "SELECT * FROM audit WHERE action=? AND entity=?", "резервы: отчёт удалён", "резервы:2025-Y")
        left = db.rows(con, "SELECT status FROM reserve_periods WHERE id=?", y_id)
    ok(c["reserves_period"]["id"] == q4["id"] and abs(c["reserves"] - 1_830_000_000.5) < 1e-6,
       "после удаления годовой — берётся 4 квартал")
    ok(a and a[-1]["who"] == ADMIN and left[0]["status"] == "удалён", "удаление — в audit, строка осталась в истории")
    with db.tx() as con:   # отчёт /admin/reserves на более позднюю дату — берётся он, как раньше
        con.execute("INSERT INTO reserve_reports (report_date, scope_type, scope_code, rnp, source) VALUES (?,?,?,?,?)",
                    ("2026-06-30", "итого", "итого", 42, "тест"))
        c = capacity.capacity(con)
    ok(c["reserves_period"] is None and c["reserves"] == 42, "старые строки reserve_reports с более поздней датой не потеряны")


if __name__ == "__main__":
    try:
        with temp_db("surveyor-finance.db"):
            setup()
            first = test_manual()
            test_validation()
            q2 = test_replace(first)
            y = test_template_roundtrip()
            test_bad_files()
            test_rights(q2)
            test_capacity(q2, y)
            print("\nВсе проверки финансов пройдены.")
    finally:
        db.close_pool()
        shutil.rmtree(STORE, ignore_errors=True)
