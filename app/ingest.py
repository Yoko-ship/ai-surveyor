"""
Приём документов: PDF, DOCX и XLSX → единое внутреннее представление.

Задача заказчика от 21.09.2026: «когда загружают документы в PDF, Word или Excel, агенту нужно
перевести их в свой LLM-язык для лучшей эффективности и точности анализа. Не забудь добавить
узбекский и английский язык».

Смысл: любой загруженный файл один раз превращается в структуру (текст, таблицы, поля, факты)
и складывается в таблицу document_extracts. Все остальные модули читают ЕЁ, а не сырой файл:
  ingest.fields_for_request(request_id)   — сведённые поля по всем документам запроса;
  ingest.summary_for_request(request_id)  — готовый материал для раздела выгрузки.

Что умеем и чего не умеем — честно:
  * PDF с текстовым слоем — читаем (pymupdf). Скан и фотография без текстового слоя —
    статус «нужно распознавание», поля пустые: локального распознавания текста в проекте нет
    и значения мы не угадываем;
  * DOCX — читаем сами (zipfile + xml.etree): абзацы w:p/w:t и таблицы w:tbl/w:tr/w:tc.
    Пакеты не ставим;
  * XLSX — читаем openpyxl (read_only, data_only): формулы приходят значениями, до 2000 строк
    на лист, обрезка помечается честно;
  * .doc, .xls, .rtf, .odt — статус «не поддерживается» и понятное сообщение агенту, что сделать;
  * языки — ru, uz-latn, uz-cyrl, en, mixed или «не определён». Словари — docs/ingest_dicts.json,
    свёртка написаний (Sugʻurta / Sug'urta / Суғурта / sugurta) — docs/i18n_terms.json;
  * виды документов — техпаспорт, кадастр, отчёт оценщика, договор, выписка, штатное расписание,
    прочее. Маркеры и пороги — тоже docs/ingest_dicts.json. Образцов бланков у компании мы не
    видели, поэтому у каждого вида есть уверенность, а не «да/нет».

Персональные данные (ЗРУ-547, правило PD-01): строки с подписью человека (ФИО, паспорт,
ПИНФЛ, адрес проживания, страхователь и т. п.) пропускаются целиком — значения из них
в fields и facts не попадают; сверх того каждое значение поля перед записью проходит
маскировку (см. _mask_field_value). Сохраняемый текст документа проходит через llm.mask_pd():
в базе лежит текст с метками [ФИО], [ПАСПОРТ], [ПИНФЛ] вместо самих данных. В журнал и в логи
пишем только id, вид, язык, статус и число полей — ни текста, ни имени файла клиента.
"""
import json
import re
import zipfile
from datetime import date as _date
from pathlib import Path
from typing import List, Optional, Tuple
from xml.etree import ElementTree as ET

from fastapi import APIRouter, HTTPException

from . import db
from . import docparse as D

router = APIRouter()
ROOT = Path(__file__).resolve().parent.parent
DICTS_FILE = ROOT / "docs" / "ingest_dicts.json"
TERMS_FILE = ROOT / "docs" / "i18n_terms.json"

# --------------------------------------------------------------------------- #
# Статусы и виды
# --------------------------------------------------------------------------- #
ST_OK = D.ST_OK                      # «разобран»
ST_PARTIAL = D.ST_PARTIAL            # «частично»
ST_OCR = D.ST_OCR                    # «нужно распознавание»
ST_ERROR = D.ST_ERROR                # «ошибка»
ST_UNSUPPORTED = "не поддерживается"

METHOD_REGEX = "regex"
METHOD_LLM = "llm"
METHOD_MANUAL = "manual"

KIND_OTHER = D.KIND_OTHER
KINDS = (D.KIND_PASSPORT, D.KIND_CADASTRE, D.KIND_VALUATION, D.KIND_CONTRACT,
         D.KIND_STATEMENT, D.KIND_STAFF, KIND_OTHER)

MAX_TEXT_CHARS = 200_000             # больше в базу не кладём: остальное не помогает разбору
MAX_SHEET_ROWS = 2000                # разумный предел на лист XLSX
MAX_TABLE_COLS = 60

OLD_FORMAT_HINT = {
    "doc": "Формат .doc не поддерживается. Откройте файл в Word и сохраните как .docx.",
    "xls": "Формат .xls не поддерживается. Откройте файл в Excel и сохраните как .xlsx.",
    "rtf": "Формат .rtf не поддерживается. Откройте файл в Word и сохраните как .docx.",
    "odt": "Формат .odt не поддерживается. Откройте файл в Word и сохраните как .docx.",
}
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"      # старые .doc и .xls (OLE2)


# --------------------------------------------------------------------------- #
# Словари: язык, виды документов, регулярки, поля и факты
# --------------------------------------------------------------------------- #
_cache = {}


def dicts() -> dict:
    """Словари аналитика (docs/ingest_dicts.json). Файла нет — работаем без него, но честно."""
    if "dicts" not in _cache:
        try:
            _cache["dicts"] = json.loads(DICTS_FILE.read_text(encoding="utf-8"))
        except Exception as e:
            print("словари приёма документов не прочитаны:", type(e).__name__)
            _cache["dicts"] = {"lang": {}, "doc_kinds": [], "patterns": {},
                               "fields": {"extra": []}, "facts": {"items": []}}
    return _cache["dicts"]


def terms() -> dict:
    """Термины юриста (docs/i18n_terms.json): подписи полей и маркеры на четырёх написаниях."""
    if "terms" not in _cache:
        try:
            _cache["terms"] = json.loads(TERMS_FILE.read_text(encoding="utf-8"))
        except Exception as e:
            print("словарь терминов не прочитан:", type(e).__name__)
            _cache["terms"] = {"doc_kinds": [], "fields": [], "normalization": {}}
    return _cache["terms"]


def dict_sources() -> dict:
    """Что из словарей реально доступно — показываем в ответе, чтобы не гадать."""
    return {"ingest_dicts.json": DICTS_FILE.exists(), "i18n_terms.json": TERMS_FILE.exists()}


def patterns() -> dict:
    """Скомпилированные регулярки из словаря. Флаги — по patterns_meta."""
    if "patterns" not in _cache:
        raw = dicts().get("patterns") or {}
        strict = {"vin", "reg_no_uz", "reg_no_uz_new", "account_no"}   # только верхний регистр
        out = {}
        for name, rx in raw.items():
            try:
                flags = 0 if name in strict else (re.IGNORECASE | re.UNICODE)
                out[name] = re.compile(rx, flags)
            except re.error as e:                 # битую регулярку не глотаем молча
                print("регулярка %s не скомпилирована: %s" % (name, e))
        _cache["patterns"] = out
    return _cache["patterns"]


# --------------------------------------------------------------------------- #
# Определение языка
# --------------------------------------------------------------------------- #
TOKEN_RE = re.compile(r"[a-zа-яўқғҳ']+")
CYR_RE = re.compile(r"[а-яўқғҳ]")
LAT_RE = re.compile(r"[a-z]")
LANG_CODES = ("ru", "uz-latn", "uz-cyrl", "en")


def _lang_rules() -> dict:
    r = (dicts().get("lang") or {}).get("rules") or {}
    return {"min_tokens": r.get("min_tokens", 8), "min_score": r.get("min_score", 0.06),
            "mixed": r.get("mixed_threshold", 0.6),
            "gate_ok": (r.get("script_gate") or {}).get("match", 1.0),
            "gate_bad": (r.get("script_gate") or {}).get("mismatch", 0.2),
            "bonus_cap": r.get("letter_bonus_cap", 0.2),
            "bonus_k": r.get("letter_bonus_k") or {},
            "segment": r.get("segment") or {}}


def _lang_words() -> dict:
    if "lang_words" not in _cache:
        lang = dicts().get("lang") or {}
        out = {}
        for code in LANG_CODES:
            d = lang.get(code) or {}
            out[code] = {"words": {D.norm(w) for w in d.get("words") or []},
                         "script": d.get("script") or "latn",
                         "letters": [D.norm(x) for x in d.get("marker_letters") or []],
                         "suffixes": [D.norm(x) for x in d.get("marker_suffixes") or []]}
        _cache["lang_words"] = out
    return _cache["lang_words"]


def _scores(tokens: List[str], t: str) -> dict:
    """Оценка каждого языка по алгоритму из ingest_dicts.json → lang.rules.algorithm."""
    rules, words = _lang_rules(), _lang_words()
    cyr = len(CYR_RE.findall(t)) or 0
    lat = len(LAT_RE.findall(t)) or 0
    letters = cyr + lat
    scores = {}
    for code in LANG_CODES:
        d = words[code]
        if not d["words"]:
            continue
        share = sum(1 for tk in tokens if tk in d["words"]) / len(tokens)
        own = cyr if d["script"] == "cyrl" else lat
        gate = rules["gate_ok"] if letters and own / letters >= 0.35 else rules["gate_bad"]
        k = rules["bonus_k"].get(code, 0.0)
        bonus = 0.0
        if k:
            if d["script"] == "cyrl" and cyr:
                hits = sum(t.count(ch) for ch in d["letters"])
                bonus = k * hits / cyr
            elif d["script"] == "latn" and tokens:
                hits = sum(t.count(ch) for ch in d["letters"])
                hits += sum(1 for tk in tokens if any(tk.endswith(s) for s in d["suffixes"]))
                lat_tokens = sum(1 for tk in tokens if LAT_RE.search(tk)) or 1
                bonus = k * hits / lat_tokens
        scores[code] = round(share * gate + min(bonus, rules["bonus_cap"]), 4)
    return scores


def _decide(scores: dict) -> Tuple[Optional[str], float, list]:
    rules = _lang_rules()
    if not scores:
        return None, 0.0, []
    order = sorted(scores.items(), key=lambda kv: kv[1], reverse=True)
    best, best_v = order[0]
    if best_v < rules["min_score"]:
        return None, 0.0, []
    second, second_v = order[1] if len(order) > 1 else (None, 0.0)
    if second and best_v and second_v / best_v >= rules["mixed"] and second_v >= rules["min_score"]:
        return "mixed", round(min(1.0, best_v), 3), [best, second]
    return best, round(min(1.0, best_v), 3), [best]


def detect_language(text: str) -> dict:
    """
    Язык документа: ru | uz-latn | uz-cyrl | en | mixed | None (не определён).

    Словари и алгоритм — docs/ingest_dicts.json → lang. Сторонние пакеты не используются.
    Возвращает {"language", "confidence", "scores", "languages", "reason"}.
    Короткий текст и мусор после плохого распознавания честно дают None.
    """
    rules = _lang_rules()
    t = D.norm(text or "")
    tokens = TOKEN_RE.findall(t)
    base = {"language": None, "confidence": 0.0, "scores": {}, "languages": [], "reason": None}
    if not _lang_words().get("ru", {}).get("words"):
        return base | {"reason": "нет словаря языков (docs/ingest_dicts.json)"}
    if len(tokens) < rules["min_tokens"]:
        return base | {"reason": "слишком мало слов (%d), язык не определяем" % len(tokens)}
    short = sum(1 for tk in tokens if len(tk) <= 2) / len(tokens)
    if short > 0.5:
        return base | {"reason": "текст похож на мусор после распознавания — язык не определяем"}

    scores = _scores(tokens, t)
    lang, conf, langs = _decide(scores)

    # длинный или табличный документ считаем ещё и по строкам: так видно двуязычные бланки
    seg = rules["segment"]
    lines = [ln for ln in t.splitlines() if ln.strip()]
    if lang and lang != "mixed" and len(lines) >= 20:
        per_line, decided = {}, 0
        for ln in lines:
            tks = TOKEN_RE.findall(ln)
            if len(tks) < seg.get("min_tokens", 4):
                continue
            L, _, _ = _decide(_scores(tks, ln))
            if L and L != "mixed":
                per_line[L] = per_line.get(L, 0) + 1
                decided += 1
        if decided >= 5:
            top, n = max(per_line.items(), key=lambda kv: kv[1])
            if n / decided < seg.get("dominant", 0.7):
                lang, langs = "mixed", sorted(per_line, key=per_line.get, reverse=True)[:2]

    return {"language": lang, "confidence": conf, "scores": scores, "languages": langs,
            "reason": None if lang else "ни один язык не набрал порога"}


# --------------------------------------------------------------------------- #
# Определение вида документа
# --------------------------------------------------------------------------- #
HEAD_LINES = 40                     # «шапка» документа: маркер в ней весит в 1.5 раза больше
HEADER_ROWS = 5                     # шапка листа XLSX

# Настоящая шапка БЛАНКА — это строка-заголовок, а не «первые 40 строк». Маркер с флагом
# "head" (заголовок бланка) считается сильным только здесь: первые BLANK_HEAD_LINES непустых
# строк, и в самой строке кроме маркера почти ничего нет (номер бланка). Иначе оглавление
# учебника, служебная записка «прошу выдать страховой полис…» и акт сверки со ссылкой
# на полис получали вид «договор» с уверенностью 1.000.
BLANK_HEAD_LINES = 3
BLANK_HEAD_REST = 4                 # сколько слов сверх маркера допускаем в строке-заголовке

# Одного заголовка мало: «ЗАЯВЛЕНИЕ / Прошу заключить договор страхования склада»,
# «ПРАЙС-ЛИСТ / Страховой полис (бланк) — 5 000 сум», «СЧЁТ-ФАКТУРА / Страховой полис
# (услуга)», «СОДЕРЖАНИЕ / 1. Страховой полис ... 5» тоже кладут маркер в первые строки.
# Поэтому маркер с флагом "head" становится сильным только когда выполняются ОБА условия:
#   1) строка — действительно заголовок бланка: кроме маркера в ней нет значащих слов
#      (допустимы только номер, цифры и короткие коды вроде «AB»);
#   2) у документа есть ВТОРОЙ признак бланка: номер бланка в той же строке
#      («№ AB 0001234», «No. 12», «seriya AB») либо не меньше BLANK_REQUISITES_MIN
#      заполненных реквизитов (страховая сумма, премия, срок, стороны — метка со значением).
# Правило про класс документов, а не про три разобранных примера: заявление, прайс-лист,
# счёт-фактура, акт, оглавление и служебная записка не проходят ни по одному из условий.
BLANK_REQUISITES_MIN = 2

# Что, кроме названия бланка, допустимо в строке-заголовке. Раньше здесь стояло
# «любое слово из трёх букв» (_BLANK_SERVICE_RE с веткой [a-z]{1,3}), и после свёртки
# кириллицы в латиницу служебными становились «сум» → sum, «шт» → sht, «акт» → akt:
# строка сметы «Страховой полис № 5 | 5 000 сум» проходила как заголовок бланка и
# закрывала пункт чек-листа «Документ о праве на объект». Теперь перечислены конкретные
# вещи: номер и серия, уточнение вида страхования, пометка копии. Единицы измерения
# и валюты служебными не считаются вовсе — единица рядом с номером означает смету
# или прайс, а не бланк.
_BLANK_CODE_RE = re.compile(r"^(?:\d+|[a-z]*\d[a-z0-9]*)$")      # 12, AB0001234, 5a
_BLANK_NUM_WORDS = {"seriya", "seriyasi", "serii", "serij", "series", "raqam", "raqami",
                    "nomer", "nomeri", "nomera", "number", "vypisan", "berilgan",
                    "no", "nr", "bl", "ser"}

# Единицы измерения, валюты и слова прайса: в заголовке бланка их не бывает.
BLANK_UNIT_WORDS = [
    "сум", "сўм", "сумов", "сума", "so'm", "so'mi", "sum", "sums", "uzs", "usd", "eur",
    "rub", "руб", "рубл", "рублей", "доллар", "долларов", "евро",
    "шт", "штук", "штука", "штуки", "штуке", "ед", "единиц", "единица", "компл", "уп",
    "dona", "donasi", "birlik", "so'mda",
    "дона", "донаси", "бирлик",
    "pcs", "pc", "piece", "pieces", "item", "items", "unit", "units", "each", "per",
    "qty", "price", "cena", "цена", "стоимость", "сумма", "итого", "всего", "total",
    "кг", "kg", "тонн", "ton", "литр", "litr", "м2", "m2", "кв", "куб",
]

# Уточнение ВИДА страхования: бланк так и печатают — «СТРАХОВОЙ ПОЛИС (КАСКО)»,
# «ПОЛИС ДОБРОВОЛЬНОГО СТРАХОВАНИЯ ИМУЩЕСТВА», «INSURANCE POLICY (PROPERTY)»,
# «SUG'URTA POLISI (MULK)». Эти слова и допустимы в заголовке, и выбрасываются
# при сборке «скелета» строки (_title_skeleton), чтобы название бланка снова стало
# сплошным: «полис добровольного страхования имущества» → «полис страхования».
BLANK_TYPE_WORDS = [
    "каско", "осго", "осаго", "огпо", "осгор", "автокаско",
    "добровольного", "добровольное", "добровольной", "добровольном",
    "обязательного", "обязательное", "обязательной", "обязательном",
    "имущества", "имущество", "имущественного", "имущественное",
    "гражданской", "ответственности", "транспортных", "транспортного", "средств",
    "средства", "грузов", "груза", "недвижимости", "жилья", "строений", "здания",
    "от", "несчастных", "случаев", "работников",
    "ixtiyoriy", "majburiy", "mulk", "mulkni", "mulkiy", "mol-mulk", "transport",
    "transportni", "yuk", "yuklarni", "ko'chmas", "kasko",
    "ихтиёрий", "мажбурий", "мулк", "мулкни", "мулкий", "транспорт", "юк", "каско",
    # «мулкни ихтиёрий суғурта қилиш полиси» — узбекское название бланка: без этих слов
    # маркер «суғурта полиси» снова становится сплошным
    "qilish", "qildirish", "qildiruvchi", "қилиш", "қилдириш",
    "property", "properties", "voluntary", "compulsory", "mandatory", "motor", "casco",
    "cargo", "goods", "real", "estate", "buildings", "liability",
]

# Пометка на копии бланка: сам бланк от неё бланком быть не перестаёт.
BLANK_MARK_WORDS = [
    "образец", "образцы", "копия", "копии", "дубликат", "экземпляр", "форма", "бланк",
    "namuna", "nusxa", "dublikat", "shakl",
    "намуна", "нусха", "дубликат", "шакл",
    "sample", "specimen", "copy", "duplicate", "form", "blank",
]

# Слова самого названия бланка в любых падежах: они допустимы в заголовке рядом
# с маркером («СТРАХОВОЙ ПОЛИС добровольного страхования имущества»).
BLANK_NAME_WORDS = [
    "страхование", "страхования", "страховой", "страховая", "страхового", "страховое",
    "полис", "полиса", "полису", "полисы", "договор", "договора", "договору",
    "sug'urta", "sug'urtasi", "polis", "polisi", "polisning", "shartnoma", "shartnomasi",
    "суғурта", "суғуртаси", "полис", "полиси", "шартнома", "шартномаси",
    "insurance", "policy", "contract", "certificate",
]

# Падежные формы, которые «скелет» приводит к именительному: иначе маркер
# «страховой полис» не находится в строке «ОБРАЗЕЦ СТРАХОВОГО ПОЛИСА № AB 1».
BLANK_TITLE_FORMS = {"страхового": "страховой", "страховому": "страховой",
                     "полиса": "полис", "полису": "полис", "полисе": "полис",
                     "договора": "договор", "договору": "договор",
                     "polisining": "polisi", "polisiga": "polisi",
                     "полисининг": "полиси", "полисига": "полиси"}
# Номер бланка в строке: «№ 12», «# 5», «No. AB 0001234», «seriya AB», «раками AB»
_BLANK_NUM_SIGN = re.compile(r"[№#]\s*\S")
_BLANK_NUM_WORD = re.compile(r"\b(?:no|nr|seriya|seriyasi|serii|series|raqam|raqami|"
                             r"nomer|nomeri|number)\.?\s*[a-z0-9]")

# Значение реквизита: сразу после метки разделитель (двоеточие, тире, ячейка таблицы)
# и число неподалёку — «Страховая сумма: 100 000 000 сум», «Срок страхования — с 01.06.2026».
_BLANK_VALUE_RE = re.compile(r"^\s*[:：|\t\-–—]+\s*.{0,12}\d")

# Маркеры «номер бланка»: название бланка + слово «номер/серия» («страховой полис серия»,
# «номер полиса», insurance policy no, polis raqami). Вес у них >= 4, и в шапке они делали
# сильным любой ЧУЖОЙ документ, где бланк лишь упомянут фразой: «MEMORANDUM / Please issue
# an insurance policy No. 5 for the branch» сворачивается в «insurance policy no 5» и давало
# «договор» с уверенностью 1.000. Поэтому такой маркер, как и маркер с флагом "head",
# считается сильным только в строке-заголовке бланка (_is_blank_title).
_NUM_MARKER_NUM_WORDS = {"no", "nr", "nomer", "nomera", "nomeri", "number",
                         "seriya", "seriyasi", "serii", "serij", "series", "raqam", "raqami"}
_NUM_MARKER_NAME_WORDS = {"polis", "polisa", "polisi", "polisning", "shartnoma", "shartnomasi",
                          "shartnomaning", "contract", "policy", "dogovor", "dogovora"}

# Реквизиты бланка: метка + значение в той же строке. Без значения это текст о страховании
# («страховая сумма определяется соглашением сторон»), а не заполненный бланк.
BLANK_REQUISITE_LABELS = [
    "страховая сумма", "страховая премия", "страховой тариф", "срок страхования",
    "срок действия", "период страхования", "объект страхования", "страхователь",
    "страховщик", "выгодоприобретатель",
    "sug'urta summasi", "sug'urta mukofoti", "sug'urta muddati", "sug'urta obyekti",
    "sug'urta qildiruvchi", "sug'urtalovchi",
    "суғурта суммаси", "суғурта мукофоти", "суғурта муддати",
    "sum insured", "insurance premium", "period of insurance", "policyholder",
    "the insurer", "insured property",
]


def _kind_defs() -> list:
    """Маркеры видов: словарь аналитика + маркеры юриста из i18n_terms.json (4 написания)."""
    if "kinds" not in _cache:
        by_kind = {}
        for k in dicts().get("doc_kinds") or []:
            # 4-й элемент — исходное написание маркера: его и показываем в ответе (matched),
            # свёрнутое написание нужно только для сравнения
            # 5-й элемент — флаг head: такой маркер (заголовок бланка) считается сильным
            # только если стоит в шапке; в глубине текста он остаётся слабым
            markers = [(D.fold(m.get("t") or ""), float(m.get("w") or 1), m.get("lang"),
                        (m.get("t") or "").strip(), bool(m.get("head")))
                       for m in k.get("markers") or []]
            by_kind[k.get("kind")] = {
                "kind": k.get("kind"),
                "norm_weight": float(k.get("norm_weight") or 1),
                "threshold": float(k.get("threshold") or 0.4),
                "markers": [m for m in markers if m[0]],
                # header_markers из словаря — справочные данные: в подсчёте уверенности вида
                # они НЕ участвуют (вес шапки задаётся HEAD_LINES). Держим их здесь, чтобы
                # аналитик видел, что словарь прочитан целиком; см. документацию, раздел 3.
                "headers": [D.fold(h) for h in k.get("header_markers") or [] if D.fold(h)],
            }
        # маркеры юриста: тот же вид, вес 2 — они из нормативных актов, но без весов
        alias = {"кадастровый документ": D.KIND_CADASTRE}
        for k in terms().get("doc_kinds") or []:
            name = alias.get(k.get("kind"), k.get("kind"))
            slot = by_kind.get(name)
            if not slot:
                continue
            have = {m[0] for m in slot["markers"]}
            for lang, lst in (k.get("markers") or {}).items():
                for m in lst:
                    f = D.fold(m)
                    if f and f not in have:
                        slot["markers"].append((f, 2.0, lang, m.strip(), False))
                        have.add(f)
        _cache["kinds"] = list(by_kind.values())
    return _cache["kinds"]


def _marker_hit(marker: str, folded: str) -> bool:
    """Короткий маркер («vin», «мфо») ищем как отдельное слово, длинный — как подстроку."""
    if len(marker) < 4:
        return re.search(r"(?<![a-z0-9])%s(?![a-z0-9])" % re.escape(marker), folded) is not None
    return marker in folded


_PUNCT = "()[]{}«»\"'.,;:!?/\|—–-№#*_"


def _title_words() -> dict:
    """Словари слов строки-заголовка в свёрнутом написании (считаем один раз)."""
    if "title_words" not in _cache:
        fold = lambda lst: {D.fold(w) for w in lst if D.fold(w)}
        drop = fold(BLANK_TYPE_WORDS) | fold(BLANK_MARK_WORDS)
        units = fold(BLANK_UNIT_WORDS)
        _cache["title_words"] = {
            "units": units,
            "drop": drop - units,                       # единица измерения не «уточнение»
            "ok": (drop | fold(BLANK_NAME_WORDS) | set(_BLANK_NUM_WORDS)) - units,
            "forms": {D.fold(k): D.fold(v) for k, v in BLANK_TITLE_FORMS.items()
                      if D.fold(k) and D.fold(v)},
        }
    return _cache["title_words"]


def _title_skeleton(folded: str) -> str:
    """
    «Скелет» строки-заголовка: без уточнения вида страхования и без пометки копии,
    падежные формы приведены к именительному. Нужен потому, что маркеры ищутся сплошной
    подстрокой: в «ПОЛИС ДОБРОВОЛЬНОГО СТРАХОВАНИЯ ИМУЩЕСТВА» маркер «полис страхования»
    разорван уточнением, а в «ОБРАЗЕЦ СТРАХОВОГО ПОЛИСА» стоит в родительном падеже.
    """
    w = _title_words()
    out = []
    for token in (folded or "").split():
        t = token.strip(_PUNCT)
        if not t or t in w["drop"]:
            continue
        out.append(w["forms"].get(t, t))
    return " ".join(out)


def _blank_head_lines(text: str, tables: list = None) -> List[Tuple[str, str]]:
    """
    Строки-кандидаты на заголовок бланка: первые BLANK_HEAD_LINES непустых строк текста
    и шапки листов XLSX (название листа и первые строки). Пара (исходная строка, свёрнутая):
    свёртка выбрасывает «№», а он нужен как признак номера бланка.
    Для каждой строки добавляем ещё и её «скелет» (_title_skeleton) с тем же исходным
    текстом: по нему находится название бланка, разорванное уточнением вида страхования.
    """
    out = []
    for ln in (text or "").splitlines():
        if not ln.strip():
            continue
        out.append((ln.strip(), D.fold(ln.strip())))
        if len(out) >= BLANK_HEAD_LINES:
            break
    for t in tables or []:
        name = str(t.get("name") or "")
        out.append((name, D.fold(name)))
        for row in (t.get("rows") or [])[:BLANK_HEAD_LINES]:
            row_s = " ".join(str(c) for c in row if c)
            out.append((row_s, D.fold(row_s)))
    lines = [(raw, f) for raw, f in out if f]
    for raw, f in list(lines):
        skeleton = _title_skeleton(f)
        if skeleton and skeleton != f:
            lines.append((raw, skeleton))
    return lines


def _in_blank_head(marker: str, lines: List[Tuple[str, str]]) -> bool:
    """
    Маркер стоит в строке-заголовке: кроме него в строке не больше BLANK_HEAD_REST слов
    и все они допустимы в заголовке бланка (номер, серия, вид страхования, пометка копии).
    Одного счёта слов мало: «Прошу выдать страховой полис … № 5» укладывается в четыре
    слова, но заголовком бланка не является и удвоенного веса не заслуживает.
    """
    for raw, ln in lines:
        if not _marker_hit(marker, ln):
            continue
        rest = ln.replace(marker, " ").split()
        if len(rest) <= BLANK_HEAD_REST and all(_blank_title_word(w, raw) for w in rest):
            return True
    return False


def _has_blank_number(raw: str, folded: str) -> bool:
    """Номер бланка в строке: «№ AB 0001234», «No. 12», «seriya AB», «# 5»."""
    if _BLANK_NUM_SIGN.search(raw or ""):
        return True
    return _BLANK_NUM_WORD.search(folded or "") is not None


def _blank_requisites(text: str) -> int:
    """
    Сколько разных реквизитов бланка заполнено: метка, за ней разделитель и число
    («Страховая сумма: 100 000 000 сум»). Без разделителя это фраза о страховании
    («страховая сумма в примере 100 000 000 сум рассчитана условно») — не бланк.
    Смотрим только шапку документа: в бланке реквизиты стоят сразу под заголовком.
    """
    found = set()
    for ln in (text or "").splitlines()[:HEAD_LINES]:
        if not any(c.isdigit() for c in ln):
            continue
        f, idx = D.fold_map(ln)          # свёртка выбрасывает двоеточие — значение ищем в исходной
        for lab in _blank_requisite_folded():
            at = f.find(lab)
            if at < 0:
                continue
            end = at + len(lab) - 1
            pos = (idx[end] + 1) if end < len(idx) else len(ln)
            if _BLANK_VALUE_RE.match(ln[pos:]):
                found.add(lab)
    return len(found)


def _blank_requisite_folded() -> List[str]:
    if "blank_req" not in _cache:
        _cache["blank_req"] = sorted({D.fold(x) for x in BLANK_REQUISITE_LABELS if D.fold(x)})
    return _cache["blank_req"]


def _is_number_marker(marker: str) -> bool:
    """
    Маркер «номер бланка»: в нём есть и слово «номер/серия», и название бланка (полис,
    договор). «Номер кузова», «кадастровый номер», «номер счёта» под правило не попадают:
    это реквизит внутри документа, а не заголовок бланка.
    """
    cache = _cache.setdefault("num_marker", {})
    if marker not in cache:
        words = set(marker.split())
        cache[marker] = bool(words & _NUM_MARKER_NUM_WORDS) and bool(words & _NUM_MARKER_NAME_WORDS)
    return cache[marker]


def _blank_series_code(word: str, raw: str) -> bool:
    """
    Слово — серия бланка («AB», «ABC»): короткий код ЗАГЛАВНОЙ латиницей в исходной строке.
    Проверяем именно исходную строку: после свёртки кириллица тоже становится латиницей,
    и «сум», «шт», «акт» выглядели бы кодом серии.
    """
    if len(word) > 4 or not word.isalpha():
        return False
    return re.search(r"(?<![A-Za-z])%s(?![A-Za-z])" % re.escape(word.upper()),
                     raw or "") is not None


def _blank_title_word(word: str, raw: str) -> bool:
    """
    Слово допустимо в строке-заголовке бланка: номер, серия, уточнение вида страхования,
    пометка копии, слово из названия бланка. Единица измерения, валюта и слово прайса
    («сум», «шт», pcs, usd, «цена») недопустимы: рядом с номером это смета, а не бланк.
    Связная речь («прошу», «оказаны», «глава») сюда не попадает — её здесь просто нет.
    """
    w = (word or "").strip(_PUNCT)
    if not w:
        return True                        # «—», «№», «...» — разделители, не слова
    words = _title_words()
    if w in words["units"]:
        return False
    if _BLANK_CODE_RE.match(w) or w in words["ok"]:
        return True
    return _blank_series_code(w, raw)


def _is_blank_title(marker: str, lines: List[Tuple[str, str]], requisites: int) -> bool:
    """
    Строка — заголовок бланка, а не первая строка чужого документа. Условия описаны
    у BLANK_REQUISITES_MIN: кроме маркера в строке допустимы только номер, серия,
    уточнение вида страхования и пометка копии (_blank_title_word) И есть второй
    признак бланка (номер в этой же строке или заполненные реквизиты в шапке).
    """
    for raw, ln in lines:
        if not _marker_hit(marker, ln):
            continue
        rest = ln.replace(marker, " ").split()
        if any(not _blank_title_word(w, raw) for w in rest):
            continue                       # «прошу выдать…», «(услуга) — 5 000 сум», «2 шт»
        if _has_blank_number(raw, ln) or requisites >= BLANK_REQUISITES_MIN:
            return True
    return False


def _header_text(tables: list) -> str:
    """Шапки листов XLSX и первых строк таблиц — в них стоит название вида документа."""
    out = []
    for t in tables or []:
        out.append(str(t.get("name") or ""))
        for row in (t.get("rows") or [])[:HEADER_ROWS]:
            out.append(" ".join(str(c) for c in row if c))
    return D.fold("\n".join(out))


def detect_kind(text: str, tables: list = None, filename: str = None) -> dict:
    """
    Вид документа по маркерам: техпаспорт | кадастр | отчёт оценщика | договор | выписка |
    штатное расписание | прочее.

    Правила счёта — docs/ingest_dicts.json → doc_kinds_scoring: вес маркера, строка-заголовок
    бланка ×2.0, шапка (первые HEAD_LINES строк) ×1.5, имя файла ×0.5, маркеры чужого языка
    ×0.7 (кроме строки-заголовка: двуязычный бланк для РУз — норма), без сильного маркера
    уверенность вдвое ниже.
    Сильным считается маркер с весом ≥ 4 вне глубины текста (шапка, заголовок, имя файла)
    и маркер со словарным флагом "head" (заголовок бланка: «страховой полис», «договор
    страхования», sug'urta polisi) — только в строке-заголовке. Иначе учебник с оглавлением,
    служебная записка «прошу выдать страховой полис…» и тарифная политика становятся
    «договором» и «техпаспортом» с уверенностью 1.000.
    Возвращает {"kind", "confidence", "candidates", "matched"}.
    """
    folded = D.fold(text or "")
    head = D.fold("\n".join((text or "").splitlines()[:HEAD_LINES]))
    header = _header_text(tables)
    blank_head = _blank_head_lines(text, tables)
    requisites = _blank_requisites(text)
    fname = D.fold(filename or "")
    lang = detect_language(text or "")["language"]

    results, matched = {}, {}
    for k in _kind_defs():
        total, strong, hits = 0.0, False, []
        for marker, w, mlang, shown, head_only in k["markers"]:
            where = None
            if _in_blank_head(marker, blank_head):
                where = "шапка бланка"
            elif _marker_hit(marker, head) or (header and _marker_hit(marker, header)):
                where = "шапка"
            elif _marker_hit(marker, folded):
                where = "текст"
            elif fname and len(marker) >= 5 and _marker_hit(marker, fname):
                where = "имя файла"
            if not where:
                continue
            weight = w
            if where == "шапка бланка":
                weight *= 2.0          # заголовок бланка — самый надёжный признак вида
            elif where == "шапка":
                weight *= 1.5
            elif where == "имя файла":
                weight *= 0.5          # имя файла даёт клиент — это подсказка, а не документ
            if (lang and mlang and mlang != lang and lang != "mixed"
                    and where != "шапка бланка"):
                # двуязычный бланк для РУз — норма (ОСГО по ПКМ № 141 печатают сразу
                # на латинице и по-русски), поэтому заголовок бланка чужим языком не штрафуем
                weight *= 0.7
            if head_only or _is_number_marker(marker):
                # заголовок бланка и номер бланка сильны только в строке-заголовке:
                # «прошу выдать страховой полис № 5», «please issue an insurance policy
                # No. 5 for the branch» и «СОДЕРЖАНИЕ / 1. Страховой полис ... 5» —
                # это упоминание бланка в чужом документе, а не бланк
                strong_hit = _is_blank_title(marker, blank_head, requisites)
            else:
                strong_hit = w >= 4 and where != "текст"
            if strong_hit:
                # сильный маркер (бланковая формулировка) засчитываем только в шапке,
                # в шапке листа XLSX или в имени файла: те же слова в глубине длинного
                # текста — это упоминание, а не бланк (пример: docs/source/Тарифная
                # политика.txt, где «свидетельство о регистрации транспортного средства»
                # стоит на 554-й строке и делало документ техпаспортом с уверенностью 1.000).
                # Заголовок бланка (флаг head) — только в строке-заголовке
                strong = True
            total += weight
            hits.append(shown or marker)
        if not total:
            continue
        conf = min(1.0, total / (k["norm_weight"] or 1))
        if not strong:
            conf *= 0.5                # иначе длинный нормативный текст становится «договором»
        results[k["kind"]] = round(conf, 3)
        matched[k["kind"]] = hits[:12]

    if not results:
        return {"kind": KIND_OTHER, "confidence": 0.0, "candidates": [], "matched": []}
    order = sorted(results.items(), key=lambda kv: kv[1], reverse=True)
    kind, conf = order[0]
    thresholds = {k["kind"]: k["threshold"] for k in _kind_defs()}
    candidates = [{"вид": k, "уверенность": v} for k, v in order[:3]]

    # разбор спорных пар — doc_kinds_disambiguation
    pair = {k for k, _ in order[:2]}
    if pair == {D.KIND_CADASTRE, D.KIND_STATEMENT} and any(
            _marker_hit(m, folded) for m in ("kadastr raqami", "kadastroviy nomer",
                                             "kadastr raqami belgisi")):
        kind, conf = D.KIND_CADASTRE, results[D.KIND_CADASTRE]
    if pair == {D.KIND_CONTRACT, D.KIND_VALUATION} and any(
            _marker_hit(m, folded) for m in ("otchet ob osenke", "baholash togrisida hisobot",
                                             "baholash togrisidagi hisobot")):
        kind, conf = D.KIND_VALUATION, results[D.KIND_VALUATION]

    if conf < thresholds.get(kind, 0.4):
        return {"kind": KIND_OTHER, "confidence": 0.0, "candidates": candidates,
                "matched": matched.get(kind, [])}
    return {"kind": kind, "confidence": round(conf, 3), "candidates": candidates,
            "matched": matched.get(kind, [])}


# --------------------------------------------------------------------------- #
# Чтение файлов
# --------------------------------------------------------------------------- #

def read_pdf(path: Path) -> Tuple[str, list]:
    """Текст PDF (pymupdf). Таблицы берём, если их находит сам pymupdf, иначе список пуст."""
    text = D.pdf_text(Path(path))
    tables = []
    try:
        import pymupdf
    except ImportError:                                   # pragma: no cover
        try:
            import fitz as pymupdf
        except ImportError:
            return text, tables
    try:
        with pymupdf.open(path) as doc:
            for n, page in enumerate(doc, 1):
                finder = getattr(page, "find_tables", None)
                if not finder:
                    break
                for t in (finder() or []):
                    rows = _clean_rows(t.extract())
                    if rows:
                        tables.append({"name": "стр. %d" % n, "rows": rows})
    except Exception as e:      # таблицы — приятное дополнение, из-за них разбор не роняем
        print("таблицы PDF не прочитаны:", type(e).__name__)
    return text, tables


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_para(p) -> str:
    """Абзац DOCX: текст прогонов w:t плюс переносы w:br и табуляции w:tab."""
    out = []
    for node in p.iter():
        tag = node.tag
        if tag == W_NS + "t":
            out.append(node.text or "")
        elif tag == W_NS + "tab":
            out.append("\t")
        elif tag == W_NS + "br":
            out.append("\n")
    return "".join(out).strip()


def read_docx(path: Path) -> Tuple[str, list]:
    """
    DOCX своими силами: zipfile + xml.etree (пакеты не ставим).
    Абзацы — w:p/w:t, таблицы — w:tbl → w:tr → w:tc. Текст таблиц идёт и в общий текст,
    чтобы подписи полей внутри таблиц тоже находились.
    """
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml")
    except KeyError:
        raise ValueError("в файле нет word/document.xml — это не документ Word")
    root = ET.fromstring(xml)
    body = root.find(W_NS + "body")
    if body is None:
        return "", []
    lines, tables = [], []
    for child in body:
        if child.tag == W_NS + "p":
            txt = _docx_para(child)
            if txt:
                lines.append(txt)
        elif child.tag == W_NS + "tbl":
            rows = []
            for tr in child.findall(W_NS + "tr"):
                cells = []
                for tc in tr.findall(W_NS + "tc"):
                    parts = [_docx_para(p) for p in tc.findall(W_NS + "p")]
                    cells.append(" ".join(x for x in parts if x).strip())
                if any(cells):
                    rows.append(cells[:MAX_TABLE_COLS])
            rows = _clean_rows(rows)
            if rows:
                tables.append({"name": "таблица %d" % (len(tables) + 1), "rows": rows})
                lines.extend(" | ".join(r) for r in rows)
    return "\n".join(lines), tables


def _clean_rows(rows: list) -> list:
    """Пустые строки и полностью пустые колонки выбрасываем: они мешают искать шапку."""
    grid = [[("" if c is None else str(c)).strip() for c in (r or [])] for r in (rows or [])]
    grid = [r for r in grid if any(r)]
    if not grid:
        return []
    width = max(len(r) for r in grid)
    keep = [i for i in range(width) if any(i < len(r) and r[i] for r in grid)]
    return [[r[i] if i < len(r) else "" for i in keep] for r in grid]


def read_xlsx(path: Path) -> Tuple[str, list]:
    """
    XLSX через openpyxl (read_only, data_only — формулы приходят значениями).
    Каждый лист — список списков строк. Пустые строки и колонки отбрасываем,
    больше MAX_SHEET_ROWS строк не берём и помечаем обрезку в названии листа.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:                                   # pragma: no cover
        raise RuntimeError("openpyxl не установлен — XLSX прочитать нечем")
    wb = load_workbook(path, read_only=True, data_only=True)
    lines, tables = [], []
    try:
        for ws in wb.worksheets:
            rows, cut = [], False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    cut = True
                    break
                rows.append(["" if c is None else str(c).strip() for c in row[:MAX_TABLE_COLS]])
            rows = _clean_rows(rows)
            if not rows:
                continue
            name = str(ws.title)
            if cut:
                name += " (показаны первые %d строк)" % MAX_SHEET_ROWS
            tables.append({"name": name, "rows": rows, "обрезан": cut})
            lines.append(name)
            lines.extend(" | ".join(r) for r in rows)
    finally:
        wb.close()
    return "\n".join(lines), tables


def sniff_format(path: Path, mime: str = None) -> str:
    """
    Формат по содержимому файла, а не по расширению и не по тому, что сказал клиент.
    Возвращает: 'pdf' | 'docx' | 'xlsx' | 'image' | 'doc' | 'xls' | 'rtf' | 'odt' | 'unknown'.
    """
    p = Path(path)
    try:
        head = p.open("rb").read(8)
    except Exception:
        return "unknown"
    if head[:5] == b"%PDF-":
        return "pdf"
    if head.startswith(OLE_SIGNATURE):
        # старый OLE-контейнер: и .doc, и .xls. Различаем по расширению, иначе считаем .doc
        return "xls" if p.suffix.lower() in (".xls", ".xlt") else "doc"
    if head[:5] == b"{\\rtf":
        return "rtf"
    if head[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(p) as z:
                names = set(z.namelist())
            if "word/document.xml" in names:
                return "docx"
            if "xl/workbook.xml" in names:
                return "xlsx"
            if any(n.startswith("xl/") for n in names):
                return "xlsx"
            if "mimetype" in names:
                return "odt"
        except Exception:
            return "unknown"
        return "unknown"
    if (mime or "").startswith("image/") or head[:3] == b"\xff\xd8\xff" or head[:4] == b"\x89PNG":
        return "image"
    return "unknown"


def read_file(path: Path, mime: str = None) -> dict:
    """
    Чтение любого поддерживаемого файла.
    {"format", "text", "tables", "status", "note"} — status заполняется только у отказов.
    """
    fmt = sniff_format(path, mime)
    if fmt in OLD_FORMAT_HINT:
        return {"format": fmt, "text": "", "tables": [], "status": ST_UNSUPPORTED,
                "note": OLD_FORMAT_HINT[fmt]}
    if fmt == "image":
        return {"format": fmt, "text": "", "tables": [], "status": ST_OCR, "note": D.OCR_TEXT}
    if fmt == "unknown":
        return {"format": fmt, "text": "", "tables": [], "status": ST_UNSUPPORTED,
                "note": "Формат файла не распознан. Принимаются PDF, DOCX и XLSX."}
    try:
        if fmt == "pdf":
            text, tables = read_pdf(path)
        elif fmt == "docx":
            text, tables = read_docx(path)
        else:
            text, tables = read_xlsx(path)
    except Exception as e:
        return {"format": fmt, "text": "", "tables": [], "status": ST_ERROR,
                "note": "Файл прочитать не удалось: %s" % type(e).__name__}
    if len(D.norm(text)) < D.MIN_TEXT_CHARS:
        note = D.OCR_TEXT if fmt == "pdf" else "В файле нет текста — проверьте, тот ли это файл."
        return {"format": fmt, "text": text, "tables": tables,
                "status": ST_OCR if fmt == "pdf" else ST_PARTIAL, "note": note}
    return {"format": fmt, "text": text, "tables": tables, "status": None, "note": None}


# --------------------------------------------------------------------------- #
# Извлечение полей
# --------------------------------------------------------------------------- #

def _lines(text: str, tables: list = None) -> List[str]:
    """Строки документа для поиска подписей: абзацы плюс строки таблиц ячейка | ячейка."""
    out = [ln.strip() for ln in (text or "").splitlines()]
    for t in tables or []:
        for row in t.get("rows") or []:
            out.append(" | ".join(str(c) for c in row))
    return [ln for ln in out if ln]


# Поля из словаря юриста (docs/i18n_terms.json → fields), которые имеет смысл извлекать:
# ключ → (название, вид документа, регулярка). Персональные данные (person_name, tin,
# passport_id, phone) сюда НЕ включены — их не извлекаем никогда (PD-01).
I18N_EXTRA = {
    "headcount": ("Численность работников", D.KIND_STAFF, "number"),
    "payroll_fund": ("Фонд оплаты труда", D.KIND_STAFF, "money"),
    "market_value": ("Рыночная стоимость", D.KIND_VALUATION, "money"),
    "appraised_value": ("Оценочная стоимость", D.KIND_VALUATION, "money"),
    "book_value": ("Балансовая (остаточная) стоимость", "*", "money"),
    "valuation_date": ("Дата оценки", D.KIND_VALUATION, "date_dmy"),
    "bank_account": ("Номер счёта организации", D.KIND_STATEMENT, "account_no"),
    "org_name": ("Наименование организации", "*", None),
    "activity_type": ("Вид деятельности", "*", None),
    "policy_series_no": ("Серия и номер полиса", D.KIND_CONTRACT, None),
}

# Порядок попыток: у суммы без валюты и у даты написаний много, перебираем от точного к общему.
# Цепочки берём из словаря аналитика (patterns_meta.chains), здесь — запасной вариант,
# если словаря нет. «year» в цепочку даты не входит: одиночный год датой не считаем.
PATTERN_CHAIN_DEFAULT = {
    "money": ["money", "money_prefix", "money_scaled_bare", "number"],
    "date_dmy": ["date_dmy", "date_iso", "date_ru_words", "date_uz_latn_words",
                 "date_uz_latn_ymd", "date_uz_cyrl_words", "date_uz_cyrl_ymd"],
}
# имя поля-типа в наших словарях → имя цепочки в patterns_meta.chains
CHAIN_ALIAS = {"money": "money", "date_dmy": "date"}


def pattern_chain(name: str) -> List[str]:
    """Порядок шаблонов для типа значения: сначала словарь аналитика, затем запасной список."""
    if "chains" not in _cache:
        _cache["chains"] = (dicts().get("patterns_meta") or {}).get("chains") or {}
    chain = _cache["chains"].get(CHAIN_ALIAS.get(name, name))
    if isinstance(chain, list) and chain:
        return chain
    return PATTERN_CHAIN_DEFAULT.get(name, [name])
MONTHS = {"январ": 1, "феврал": 2, "март": 3, "апрел": 4, "май": 5, "ма": 5, "июн": 6, "июл": 7,
          "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12,
          "yanvar": 1, "fevral": 2, "mart": 3, "aprel": 4, "may": 5, "iyun": 6, "iyul": 7,
          "avgust": 8, "sentabr": 9, "sentyabr": 9, "oktabr": 10, "oktyabr": 10,
          "noyabr": 11, "dekabr": 12}


MONTH_KEYS = sorted(MONTHS, key=len, reverse=True)      # длинные раньше: «март» до «ма»


def _month(value: str) -> Optional[int]:
    """Месяц словом → номер. Шаблон отдаёт основу («январ», «мая», «iyun»), окончания разные."""
    v = str(value).lower().replace("ё", "е").strip(". ")
    if v.isdigit():
        return int(v)
    for key in MONTH_KEYS:
        if v.startswith(key):
            return MONTHS[key]
    return None


def _as_date(groups: dict) -> Optional[str]:
    """
    День, месяц и год из шаблона даты → «ДД.ММ.ГГГГ», одинаково для всех шести написаний.
    Год лежит в группе «yyyy» (так в docs/ingest_dicts.json с 20.09.2026); имя «y» читаем тоже,
    чтобы разбор не зависел от переименования групп в словаре.
    """
    d, m = groups.get("d"), groups.get("m")
    y = groups.get("yyyy") or groups.get("y")
    if not (d and m and y):
        return None                        # одиночный год датой не считаем — это год, а не дата
    mm = _month(m)
    if not mm or not (1 <= mm <= 12):
        return None
    if not str(d).isdigit() or not str(y).isdigit():
        return None
    try:                                   # честный календарь: «31.02.2026» — не дата
        _date(int(y), mm, int(d))
    except ValueError:
        return None
    return "%02d.%02d.%s" % (int(d), mm, y)


def _extra_fields() -> list:
    """
    Дополнительные поля: словарь аналитика (ingest_dicts.json → fields.extra) плюс подписи
    юриста на четырёх написаниях (i18n_terms.json → fields). Подписи объединяются по ключу.
    """
    if "extra" not in _cache:
        by_key = {}
        for f in (dicts().get("fields") or {}).get("extra") or []:
            labels = []
            for key in ("labels_ru", "labels_uz_latn", "labels_uz_cyrl", "labels_en"):
                labels += f.get(key) or []
            by_key[f.get("key")] = {"key": f.get("key"), "name": f.get("name"),
                                    "kind": f.get("kind") or "*", "pattern": f.get("pattern"),
                                    "raw": labels}
        for f in terms().get("fields") or []:
            spec = I18N_EXTRA.get(f.get("key"))
            if not spec:
                continue
            name, kind, pattern = spec
            labels = []
            for lst in (f.get("labels") or {}).values():
                labels += lst
            slot = by_key.setdefault(f["key"], {"key": f["key"], "name": name, "kind": kind,
                                                "pattern": pattern, "raw": []})
            slot["raw"] += labels
        out = []
        for f in by_key.values():
            labels = [x for x in f["raw"] if x and len(x) > 2]
            out.append({"key": f["key"], "name": f["name"], "kind": f["kind"],
                        "pattern": f["pattern"],
                        "labels": sorted({D.norm(x) for x in labels}, key=len, reverse=True),
                        "folded": sorted({D.fold(x) for x in labels if D.fold(x)},
                                         key=len, reverse=True)})
        _cache["extra"] = out
    return _cache["extra"]


def _blank_dates(value: str) -> str:
    """
    Затирает даты пробелами. Нужно для денежных величин: иначе из «01.06.2026» запасной
    шаблон number вытаскивает «2026» и год становится страховой суммой (ГК ст. 936/938).
    """
    out = value
    for rx_name, rx in patterns().items():
        if rx_name.startswith("date"):
            out = rx.sub(lambda m: " " * len(m.group(0)), out)
    return out


def _apply_pattern(name: str, value: str) -> Optional[str]:
    """Значение по регулярке из словаря. Не подошло — None: лучше пусто, чем мусор."""
    v = (value or "").strip(" \t:.,;|—–-")
    if CHAIN_ALIAS.get(name, name) == "money":
        v = _blank_dates(v).strip(" \t:.,;|—–-")   # год из даты деньгами не считаем
    if not v:
        return None
    if not name:
        return D.clean_value(v) or None
    for rx_name in pattern_chain(name):
        rx = patterns().get(rx_name)
        if not rx:
            continue
        m = rx.search(v if rx_name not in ("vin", "account_no") else v.upper())
        if not m:
            continue
        groups = m.groupdict()
        got = _as_date(groups)             # у даты забираем всю дату, а не один год (он идёт ниже)
        if got:
            return got
        if rx_name.startswith("date") and groups.get("d") and groups.get("m"):
            continue                       # «31.02.2026» — не дата: дальше по цепочке
        for key in ("amount", "val", "no", "acc", "inn", "y", "vin", "oked", "mfo"):
            if groups.get(key):
                return groups[key].strip()
        return m.group(0).strip()
    return None


def extract_fields(text: str, tables: list, kind: str, language: str = None, *,
                   use_llm: bool = True, with_reg_no: bool = False) -> List[dict]:
    """
    Поля документа регулярками и подписями. Персональные данные не извлекаются (PD-01).
    Каждое поле: ключ, название, значение, откуда взято, уверенность и метод.
    use_llm=False — техпаспорт и кадастр без ИИ-дозаполнения; with_reg_no — госномер ТС (по запросу).
    """
    lines = _lines(text, tables)
    out, seen = [], set()

    # техпаспорт и кадастр разбирает docparse — там подписи, нормы и предупреждения
    if kind in D.FIELDS:
        parsed = D.parse_text(text if not tables else "\n".join(lines), kind, with_reg_no=with_reg_no,
                              use_llm=use_llm)
        for it in parsed.get("поля") or []:
            if it.get("значение") is None or it["поле"] in seen:
                continue
            seen.add(it["поле"])
            out.append({"ключ": it["поле"], "название": it["название"], "значение": it["значение"],
                        "найдено_по": it.get("найдено_по"), "уверенность": it.get("уверенность"),
                        "метод": METHOD_REGEX, "требует_проверки": False,
                        "норма": it.get("норма")})

    # дополнительные поля по словарю аналитика
    for f in _extra_fields():
        if f["key"] in seen or (f["kind"] != "*" and f["kind"] != kind):
            continue
        hit = _find_labelled(lines, f)
        if not hit:
            continue
        seen.add(f["key"])
        out.append({"ключ": f["key"], "название": f["name"], "значение": hit["value"],
                    "найдено_по": hit["label"], "уверенность": hit["confidence"],
                    "метод": METHOD_REGEX, "требует_проверки": False, "норма": None})
    _split_brand_model(out)
    return out


# Марки из двух слов: по первому слову их делить нельзя — «Land Rover Discovery» давал
# марку «Land». Список закрытый; марки вне списка делятся по первому слову — это записано
# в ограничениях (docs/Приём документов — как работает.md, раздел 2).
COMPOUND_BRANDS = (
    "land rover", "range rover", "alfa romeo", "great wall", "great wall motor",
    "aston martin", "rolls royce", "mercedes benz", "general motors", "dong feng",
    "chery automobile", "hyundai motor", "kia motors", "isuzu motors", "byd auto",
)


def _brand_split(value: str) -> Optional[Tuple[str, str]]:
    """Марка и модель из одного значения; вторая часть пустая — модели в документе нет."""
    parts = str(value or "").split()
    if len(parts) < 2:
        return None
    # дефис в марке пишут по-разному («Mercedes-Benz» и «Mercedes Benz»), поэтому сравниваем
    # по словам, а режем по исходным частям значения
    acc = []
    for i, part in enumerate(parts):
        acc += part.lower().replace("-", " ").split()
        cur = " ".join(acc)
        if cur in COMPOUND_BRANDS:
            return " ".join(parts[:i + 1]), " ".join(parts[i + 1:])
        if not any(cb.startswith(cur + " ") for cb in COMPOUND_BRANDS):
            break
    return parts[0], " ".join(parts[1:])


def _split_brand_model(fields: List[dict]) -> None:
    """
    «Марка и модель: Chevrolet Malibu» — одна подпись на два поля, и модель дублировала марку.
    Делим: марка — первое слово, модель — остальное; составную марку («Land Rover») берём
    из списка COMPOUND_BRANDS целиком. Раздельные подписи не трогаем.
    """
    by_key = {f["ключ"]: f for f in fields}
    brand, model = by_key.get("brand"), by_key.get("model")
    if not (brand and model) or brand["значение"] != model["значение"]:
        return
    got = _brand_split(brand["значение"])
    if not got:
        return
    brand["значение"], rest = got
    if rest:
        model["значение"] = rest
    else:
        fields.remove(model)               # модели в значении нет — пустое поле не выдумываем


def _in_heading(line: str, tail: str) -> bool:
    """
    Подпись стоит внутри длинной фразы без разделителя — это заголовок, а не подпись значения:
    «VALUATION REPORT ON THE MARKET VALUE OF THE PROPERTY». Значения у неё нет, и следующая
    строка к ней не относится (там оказывалась дата, и год шёл в страховую сумму).
    По-русски так почти не бывает из-за падежа, а хвост подписи («…стоимость объекта»)
    короткий — поэтому порогом служит число слов после подписи.
    """
    t = D.cut_tail(tail or "").strip(" \t:.,;|—–-")
    if not t or ":" in (line or "") or " | " in (line or ""):
        return False
    return len(t.split()) >= 2


def _find_labelled(lines: List[str], field: dict) -> Optional[dict]:
    """Значение справа от подписи; два прохода — обычный и по свёрнутому написанию."""
    for use_fold in (False, True):
        labels = field["folded"] if use_fold else sorted(field["labels"], key=len, reverse=True)
        for i, line in enumerate(lines):
            D.tick()                               # срок разбора акта (без срока — ничего не делает)
            if D.is_personal_label(line):          # строка про человека — пропускаем (PD-01)
                continue
            got = D._tail_after_label(line, labels, use_fold)
            if not got:
                continue
            lab, tail = got
            val = _apply_pattern(field.get("pattern"), D.cut_tail(tail))
            if val:
                return {"value": val, "label": lab, "confidence": "высокая"}
            if " | " in line:
                # строка таблицы: подпись в шапке колонки, а не подпись значения.
                # Брать число из следующей строки нельзя — это была бы первая строка данных
                continue
            if _in_heading(line, tail):
                continue
            for j in range(i + 1, min(i + 3, len(lines))):
                if D.is_personal_label(lines[j]):
                    continue
                val = _apply_pattern(field.get("pattern"), lines[j])
                if val:
                    return {"value": val, "label": lab, "confidence": "средняя"}
                break
    return None


def extract_facts(text: str, kind: str, status: str = None) -> List[dict]:
    """Короткие признаки для движка: ипотека, ограничения, полис и т. д. (словарь → facts)."""
    folded = D.fold(text or "")
    out = []
    for item in (dicts().get("facts") or {}).get("items") or []:
        markers = [D.fold(m) for m in item.get("markers") or [] if D.fold(m)]
        if not markers:
            continue
        if item.get("kind") and item["kind"] != kind:
            continue
        if any(_marker_hit(m, folded) for m in markers):
            out.append({"ключ": item["key"], "значение": True, "источник": "маркер в тексте",
                        "правило": item.get("rule")})
    if status == ST_OCR:
        out.append({"ключ": "needs_ocr", "значение": True,
                    "источник": "в файле нет текстового слоя", "правило": None})
    if any(D.is_personal_label(ln) for ln in (text or "").splitlines()):
        out.append({"ключ": "has_personal_data", "значение": True,
                    "источник": "в документе есть подписи персональных данных — "
                                "их значения не сохраняются (ЗРУ-547, правило PD-01)",
                    "правило": "PD-01"})
    return out


CONF_BY_WORD = {"высокая": 0.9, "средняя": 0.7, D.ST_AI: 0.5}


def _confidence(fields: List[dict], kind_conf: float) -> float:
    """Уверенность разбора: среднее по полям; полей нет — половина уверенности в виде."""
    vals = [CONF_BY_WORD.get(f.get("уверенность"), 0.6) for f in fields]
    if not vals:
        return round((kind_conf or 0.0) * 0.5, 3)
    return round(sum(vals) / len(vals), 3)


# --------------------------------------------------------------------------- #
# ИИ-дораскрытие
# --------------------------------------------------------------------------- #

def ai_extend(text: str, kind: str, language: str, fields: List[dict]) -> dict:
    """
    Дораскрытие полей ИИ. Без ключа ИИ ничего не делает и честно об этом говорит.
    Всё, что предложил ИИ, помечается method = 'llm' и «требует проверки».
    """
    try:
        from . import llm
    except Exception:
        return {"ok": False, "reason": "модуль ИИ не подключён", "added": []}
    if not llm.enabled():
        return {"ok": False, "reason": llm.status()["reason"], "added": []}
    try:
        got = llm.extract_document(text, kind, language)
    except Exception as e:                 # ошибка уже в журнале llm_calls
        return {"ok": False, "reason": "ошибка обращения к ИИ: %s" % type(e).__name__, "added": []}
    if not got.get("ok"):
        return {"ok": False, "reason": got.get("reason"), "added": []}
    have = {f["ключ"] for f in fields}
    added = []
    for key, value in (got.get("fields") or {}).items():
        if key in have or not str(value).strip():
            continue
        fields.append({"ключ": key, "название": key, "значение": str(value).strip()[:200],
                       "найдено_по": "ИИ", "уверенность": D.ST_AI, "метод": METHOD_LLM,
                       "требует_проверки": True, "норма": None})
        added.append(key)
    return {"ok": True, "reason": got.get("reason"), "added": added,
            "facts": got.get("facts") or {}, "summary": got.get("summary")}


# --------------------------------------------------------------------------- #
# Конвейер: файл → document_extracts
# --------------------------------------------------------------------------- #

def _photo(con, photo_id: int) -> dict:
    r = db.rows(con, "SELECT * FROM photos WHERE id=?", photo_id)
    if not r:
        raise HTTPException(404, "Файл не найден")
    return r[0]


def _mask(text: str) -> str:
    """Текст в базу кладём с замаскированными персональными данными (правило проекта № 8)."""
    return _mask_str(text)[:MAX_TEXT_CHARS]


def _mask_str(text: str) -> str:
    try:
        from . import llm
        return llm.mask_pd(text or "")
    except Exception:
        return text or ""


def _mask_deep(value):
    """
    Та же маскировка для вложенных структур: таблицы и факты пишутся в базу как JSON,
    и ячейка таблицы — такой же текст документа, как и всё остальное (ЗРУ-547, правило PD-01).
    """
    if isinstance(value, str):
        return _mask_str(value)
    if isinstance(value, dict):
        return {k: _mask_deep(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_mask_deep(v) for v in value]
    return value


# Значения полей маскировать целиком нельзя: кадастровый номер, ИНН организации и госномер —
# это данные объекта, их движок использует. Поэтому к полям применяем только те правила,
# которые не могут совпасть с данными объекта: контакты и документы человека.
FIELD_MASK_KEEP = {"[КАДАСТР]", "[ИНН]", "[ГОСНОМЕР]"}

# Реквизиты документа, похожие на паспорт человека, но им не являющиеся: серия и номер полиса
# («AB 0001234») — это реквизит договора, без него поле бесполезно.
FIELD_MASK_KEEP_BY_KEY = {"policy_series_no": {"[ПАСПОРТ]"}}

# Поля-характеристики объекта и организации. Правило ФИО к ним не применяем: «Шота Руставели»
# в адресе объекта, «Chevrolet Cobalt» в марке и название фирмы — это не имя человека.
# Строки с подписью человека до значений полей не доходят (docparse.is_personal_label),
# поэтому здесь остаётся только имущество. Инициалы («Каримов А. А.») маскируем и тут.
OBJECT_FIELD_KEYS = {"address", "region", "object_kind", "activity_type",
                     "brand", "model", "vehicle_type", "color", "walls", "right_kind"}


def _mask_field_value(value: str, key: str = None) -> str:
    """
    Маскировка значения поля. Уровень тот же, что и у текста (llm.mask_pd), за двумя
    исключениями: данные объекта (кадастровый номер, ИНН организации, госномер, серия полиса)
    сохраняются, а правило ФИО не применяется к полям-характеристикам объекта.
    """
    try:
        from . import llm
    except Exception:
        return value or ""
    out = value or ""
    keep = FIELD_MASK_KEEP | FIELD_MASK_KEEP_BY_KEY.get(key or "", set())
    for rx, placeholder in llm.PD_RULES:
        if placeholder in keep:
            continue
        out = rx.sub(placeholder, out)
    if key in OBJECT_FIELD_KEYS:
        return llm.mask_name_initials(out)
    return llm.mask_names(out)


def _mask_fields(fields: List[dict]) -> List[dict]:
    out = []
    for f in fields or []:
        g = dict(f)
        if isinstance(g.get("значение"), str):
            g["значение"] = _mask_field_value(g["значение"], g.get("ключ"))
        out.append(g)
    return out


def _save(con, photo: dict, data: dict) -> int:
    """
    Запись в document_extracts. Идемпотентно: у одного файла одна строка, она обновляется.
    Всё, что уходит в базу текстом (text, tables, facts, значения полей), проходит маскировку
    персональных данных — это единственное место записи, поэтому маскируем здесь (ЗРУ-547).
    """
    row = db.rows(con, "SELECT id FROM document_extracts WHERE photo_id=?", photo["id"])
    values = (photo["request_id"], photo["id"], photo["path"], photo.get("filename"),
              photo.get("mime"), data["kind"], data["kind_confidence"], data["language"],
              data["language_confidence"], data["status"], _mask(data["text"]),
              json.dumps(_mask_deep(data["tables"]), ensure_ascii=False),
              json.dumps(_mask_fields(data["fields"]), ensure_ascii=False),
              json.dumps(_mask_deep(data["facts"]), ensure_ascii=False),
              data["confidence"], data["method"])
    if row:
        con.execute("""UPDATE document_extracts SET request_id=?, photo_id=?, file=?, filename=?,
                       mime=?, kind=?, kind_confidence=?, language=?, language_confidence=?,
                       status=?, text=?, tables=?, fields=?, facts=?, confidence=?, method=?,
                       updated_at=? WHERE id=?""", values + (db.now(), row[0]["id"]))
        return row[0]["id"]
    cur = con.execute("""INSERT INTO document_extracts (request_id, photo_id, file, filename, mime,
                         kind, kind_confidence, language, language_confidence, status, text,
                         tables, fields, facts, confidence, method, created_at, updated_at)
                         VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                      values + (db.now(), db.now()))
    return cur.lastrowid


def parse_path(full: Path, mime: str = None, filename: str = None, kind_hint: str = None,
               use_llm: bool = True) -> dict:
    """
    Ядро разбора одного файла без базы: прочитать → язык и вид → поля и факты (поля уже
    замаскированы). Его используют ingest_file (документы запроса) и app/analysis_docs.py
    (договор для вкладки «Аналитика», без запроса).
    """
    read = read_file(full, mime)
    text, tables = read["text"], read["tables"]
    lang = detect_language(text)
    kind = detect_kind(text, tables, filename)

    # вид, указанный агентом при загрузке, сильнее автоопределения, если оно ничего не нашло
    if kind["kind"] == KIND_OTHER and kind_hint in KINDS and kind_hint != KIND_OTHER:
        kind = {"kind": kind_hint, "confidence": 0.0, "candidates": kind.get("candidates") or [],
                "matched": ["вид указан агентом при загрузке"]}

    ai = {"ok": False, "reason": "не запрашивалось", "added": []}
    if read["status"] in (ST_UNSUPPORTED, ST_OCR, ST_ERROR):
        fields, facts = [], extract_facts(text, kind["kind"], read["status"])
        status, method = read["status"], METHOD_REGEX
    else:
        fields = extract_fields(text, tables, kind["kind"], lang["language"])
        facts = extract_facts(text, kind["kind"], read["status"])
        if use_llm:
            ai = ai_extend(text, kind["kind"], lang["language"], fields)
        method = METHOD_LLM if ai.get("added") else METHOD_REGEX
        status = ST_OK if fields else ST_PARTIAL
        if fields and kind["kind"] == KIND_OTHER:
            status = ST_PARTIAL           # поля есть, а вид не опознан — это не полный разбор

    fields = _mask_fields(fields)     # в ответ и в базу уходят одинаково очищенные поля
    data = {"kind": kind["kind"], "kind_confidence": kind["confidence"],
            "language": lang["language"], "language_confidence": lang["confidence"],
            "status": status, "text": text, "tables": tables,
            "fields": fields, "facts": facts,
            "confidence": _confidence(fields, kind["confidence"]), "method": method}
    return {"read": read, "kind": kind, "lang": lang, "ai": ai, "fields": fields, "facts": facts,
            "status": status, "method": method, "data": data}


# --------------------------------------------------------------------------- #
# Условия договора страхования: суммы, стоимость, место, год — для вкладки «Аналитика»
# --------------------------------------------------------------------------- #

# Подписи в договоре страхования. Длинные подписи проверяются раньше (сортировка ниже).
CONTRACT_TERMS = [
    ("sum_insured", "Страховая сумма", "money",
     ["общая страховая сумма", "страховая сумма", "лимит ответственности",
      "sug'urta summasi", "umumiy sug'urta summasi", "суғурта суммаси",
      "sum insured", "insured amount", "limit of liability"]),
    ("object_value", "Стоимость имущества", "money",
     ["действительная стоимость имущества", "действительная стоимость", "страховая стоимость имущества",
      "страховая стоимость", "стоимость имущества", "стоимость объекта",
      "стоимость застрахованного имущества", "рыночная стоимость", "балансовая стоимость",
      "mulk qiymati", "sug'urta qiymati", "haqiqiy qiymati", "мулк қиймати", "суғурта қиймати",
      "insured value", "property value", "actual value"]),
    ("address", "Местонахождение объекта", None,
     ["адрес места страхования", "место страхования", "территория страхования", "адрес объекта",
      "местонахождение имущества", "местонахождение объекта", "адрес", "место нахождения",
      "sug'urta hududi", "obyekt manzili", "manzil", "манзил", "location", "address"]),
    ("year", "Год постройки или выпуска", "year",
     ["год постройки", "год ввода в эксплуатацию", "год выпуска", "qurilgan yili",
      "ishlab chiqarilgan yili", "қурилган йили", "ишлаб чиқарилган йили",
      "year of construction", "year of manufacture"]),
    ("walls", "Материал стен и конструкция", None,
     ["материал стен", "конструкция здания", "конструктивные элементы", "devor materiali",
      "девор материали"]),
    ("object_kind", "Объект страхования", None,
     ["объект страхования", "застрахованное имущество", "sug'urta obyekti", "суғурта объекти",
      "insured property", "subject of insurance"]),
    ("activity", "Деятельность на объекте", None,
     ["вид деятельности на объекте", "назначение объекта", "использование объекта",
      "faoliyat turi", "фаолият тури", "occupancy"]),
]

MONEY_SCALE = (("трлн", 1e12), ("триллион", 1e12), ("trln", 1e12), ("млрд", 1e9), ("миллиард", 1e9),
               ("mlrd", 1e9), ("млн", 1e6), ("миллион", 1e6), ("mln", 1e6), ("тыс", 1e3),
               ("тысяч", 1e3), ("ming", 1e3), ("минг", 1e3))
_MONEY_SPACES = re.compile(r"[    ']")


def money_value(raw: str) -> Optional[dict]:
    """
    Сумма из строки: {"value": число, "currency": "UZS"|"USD"|"EUR"|"RUB"|None}. Множитель
    (млн, млрд) учитывается — _apply_pattern отдаёт только число, для подстановки в форму мало.
    """
    v = _blank_dates(raw or "")
    for rx_name in ("money", "money_prefix", "money_scaled_bare"):
        rx = patterns().get(rx_name)
        m = rx.search(v) if rx else None
        if not m:
            continue
        g = m.groupdict()
        try:
            num = float(_MONEY_SPACES.sub("", g.get("amount") or "").replace(",", "."))
        except ValueError:
            continue
        scale = (g.get("scale") or "").lower()
        for word, mult in MONEY_SCALE:
            if scale.startswith(word):
                num *= mult
                break
        cur = (g.get("cur") or "").lower()
        currency = None
        if cur:
            currency = ("USD" if ("usd" in cur or "долл" in cur or "$" in cur) else
                        "EUR" if ("eur" in cur or "€" in cur) else
                        "RUB" if "rub" in cur else "UZS")
        return {"value": num, "currency": currency}
    return None


def extract_contract_terms(text: str, tables: list = None) -> List[dict]:
    """
    Условия договора страхования по подписям: страховая сумма, стоимость, место, год и т. д.
    Строки с подписью человека пропускает _find_labelled (PD-01); значения маскируются.
    У денег дополнительно value_num и currency — с учётом «млн»/«млрд».
    """
    lines = _lines(text, tables)
    out = []
    for key, name, pattern, labels in CONTRACT_TERMS:
        spec = {"key": key, "pattern": pattern,
                "labels": sorted({D.norm(x) for x in labels}, key=len, reverse=True),
                "folded": sorted({D.fold(x) for x in labels if D.fold(x)}, key=len, reverse=True)}
        hit = _find_labelled(lines, spec)
        if not hit:
            continue
        item = {"ключ": key, "название": name, "значение": _mask_field_value(hit["value"], key),
                "найдено_по": hit["label"], "уверенность": hit["confidence"], "метод": METHOD_REGEX}
        if pattern == "money":
            mv = None
            for ln in lines:                   # строка с подписью целиком: там и множитель, и валюта
                D.tick()
                if D.is_personal_label(ln):
                    continue
                got = D._tail_after_label(ln, [hit["label"]], False) or \
                    D._tail_after_label(ln, [hit["label"]], True)
                if got:
                    mv = money_value(D.cut_tail(got[1]))
                    if mv:
                        break
            mv = mv or money_value(hit["value"])
            if not mv:
                continue
            item.update({"value_num": mv["value"], "currency": mv["currency"]})
        out.append(item)
    return out


def ingest_file(con, photo_id: int, who: str = "api", use_llm: bool = True) -> dict:
    """
    Весь конвейер по одному файлу: прочитать → определить язык и вид → извлечь поля и факты →
    записать в document_extracts → отметить пункт чек-листа.

    Идемпотентно: повторный вызов обновляет ту же строку.
    В журнал пишем только id, вид, язык, статус и число полей — ни текста, ни имени файла.
    """
    photo = _photo(con, photo_id)
    full = ROOT / photo["path"]
    if not full.exists():
        raise HTTPException(404, "Файл не найден на диске")

    kind_hint = photo.get("doc_kind")
    core = parse_path(full, photo.get("mime"), photo.get("filename"), kind_hint, use_llm)
    read, kind, lang, ai = core["read"], core["kind"], core["lang"], core["ai"]
    tables, fields, facts = read["tables"], core["fields"], core["facts"]
    status, method = core["status"], core["method"]
    data = core["data"]
    extract_id = _save(con, photo, data)

    # пункт чек-листа: отмечаем только если вид документа опознан
    checklist = {"отмечен": False, "причина": "вид документа не определён"}
    if kind["kind"] != KIND_OTHER:
        checklist = D.mark_received(con, photo["request_id"], kind["kind"], photo["path"],
                                    None, who)
        if photo.get("doc_kind") != kind["kind"] and kind["confidence"] >= 0.4:
            con.execute("UPDATE photos SET doc_kind=? WHERE id=?", (kind["kind"], photo["id"]))

    db.audit(con, who, "документ разобран", "request:%s" % photo["request_id"],
             {"document_extract_id": extract_id, "photo_id": photo["id"], "вид": kind["kind"],
              "язык": lang["language"], "статус": status, "полей": len(fields),
              "метод": method})

    return {"id": extract_id, "photo_id": photo["id"], "request_id": photo["request_id"],
            "формат": read["format"], "вид_документа": kind["kind"],
            "уверенность_вида": kind["confidence"], "кандидаты": kind.get("candidates"),
            "язык": lang["language"], "уверенность_языка": lang["confidence"],
            "языки": lang.get("languages"), "статус": status, "метод": method,
            "поля": fields, "факты": facts, "таблиц": len(tables),
            "уверенность": data["confidence"], "сообщение": read["note"],
            "ии": {"подключён": ai.get("ok", False), "пояснение": ai.get("reason"),
                   "дозаполнил": ai.get("added") or []},
            "чек_лист": checklist, "словари": dict_sources(),
            "сверено_с_бланком": False}


def ingest_request(con, rid: int, who: str = "api", force: bool = False,
                   use_llm: bool = True) -> dict:
    """Разбор всех файлов запроса. force=False — только те, которых ещё нет в document_extracts."""
    if not db.rows(con, "SELECT id FROM requests WHERE id=?", rid):
        raise HTTPException(404, "Запрос не найден")
    done = {r["photo_id"] for r in
            db.rows(con, "SELECT photo_id FROM document_extracts WHERE request_id=?", rid)}
    files = db.rows(con, "SELECT id, doc_kind FROM photos WHERE request_id=? ORDER BY id", rid)
    out, skipped = [], 0
    for f in files:
        if f["doc_kind"] == D.KIND_PHOTO:          # фотографии объекта здесь не разбираем
            skipped += 1
            continue
        if f["id"] in done and not force:
            skipped += 1
            continue
        try:
            out.append(ingest_file(con, f["id"], who, use_llm))
        except HTTPException as e:                 # ошибку не глотаем — она видна в ответе
            out.append({"photo_id": f["id"], "статус": ST_ERROR, "сообщение": e.detail})
    return {"request_id": rid, "разобрано": len(out), "пропущено": skipped, "документы": out}


# --------------------------------------------------------------------------- #
# Чтение результата смежными командами
# --------------------------------------------------------------------------- #

def _extracts(con, rid: int) -> List[dict]:
    return db.rows(con, "SELECT * FROM document_extracts WHERE request_id=? ORDER BY id", rid)


def _load(row: dict, key: str, default):
    try:
        return json.loads(row.get(key) or "null") or default
    except Exception:
        return default


def fields_for_request(request_id: int) -> dict:
    """
    Сведённые поля по всем документам запроса. Более свежий разбор перекрывает старый.
    У каждого поля видно, из какого документа оно взято и каким методом.
    """
    with db.tx() as con:
        rows = _extracts(con, request_id)
    merged, docs = {}, []
    for r in sorted(rows, key=lambda x: (x.get("updated_at") or x.get("created_at") or "", x["id"])):
        docs.append({"id": r["id"], "photo_id": r["photo_id"], "вид": r["kind"],
                     "язык": r["language"], "статус": r["status"]})
        for f in _load(r, "fields", []):
            merged[f["ключ"]] = {
                "ключ": f["ключ"], "название": f.get("название"), "значение": f.get("значение"),
                "метод": f.get("метод"), "уверенность": f.get("уверенность"),
                "требует_проверки": bool(f.get("требует_проверки")),
                "документ": {"id": r["id"], "photo_id": r["photo_id"], "вид": r["kind"],
                             "имя_файла": r["filename"]},
            }
    return {"request_id": request_id, "поля": merged, "документы": docs,
            "всего_полей": len(merged)}


def summary_for_request(request_id: int) -> dict:
    """
    Материал для раздела «Документы и извлечённые данные» в PDF и XLSX выгрузках.
    Структура плоская и самодостаточная: списки словарей, ничего досчитывать не надо.
    """
    with db.tx() as con:
        rows = _extracts(con, request_id)
    documents, fields, facts, warnings = [], [], [], []
    for r in rows:
        documents.append({
            "id": r["id"], "photo_id": r["photo_id"],
            "имя_файла": r["filename"] or "без имени", "вид": r["kind"] or KIND_OTHER,
            "уверенность_вида": r["kind_confidence"],
            "язык": r["language"] or "не определён", "статус": r["status"],
            "метод": r["method"], "дата": r.get("updated_at") or r.get("created_at"),
            "полей": len(_load(r, "fields", [])),
        })
        for f in _load(r, "fields", []):
            fields.append({"документ": r["filename"] or ("id %s" % r["id"]),
                           "вид_документа": r["kind"], "поле": f.get("название") or f.get("ключ"),
                           "ключ": f.get("ключ"), "значение": f.get("значение"),
                           "метод": f.get("метод"), "уверенность": f.get("уверенность"),
                           "требует_проверки": bool(f.get("требует_проверки"))})
            if f.get("требует_проверки"):
                warnings.append({"вид": "требует проверки",
                                 "текст": "Поле «%s» предложено ИИ — проверьте по оригиналу"
                                          % (f.get("название") or f.get("ключ")),
                                 "документ": r["filename"] or ("id %s" % r["id"])})
        for f in _load(r, "facts", []):
            facts.append({"документ": r["filename"] or ("id %s" % r["id"]),
                          "признак": f.get("ключ"), "значение": f.get("значение"),
                          "источник": f.get("источник"), "правило": f.get("правило")})
        if r["status"] == ST_OCR:
            warnings.append({"вид": "нужно распознавание",
                             "текст": D.OCR_TEXT, "документ": r["filename"] or ("id %s" % r["id"])})
        if r["status"] == ST_UNSUPPORTED:
            warnings.append({"вид": "не поддерживается",
                             "текст": "Формат файла не поддерживается — пересохраните документ "
                                      "в DOCX, XLSX или PDF",
                             "документ": r["filename"] or ("id %s" % r["id"])})
        if (r["kind"] or KIND_OTHER) == KIND_OTHER:
            warnings.append({"вид": "требует проверки",
                             "текст": "Вид документа не определён — проверьте вручную",
                             "документ": r["filename"] or ("id %s" % r["id"])})
    return {"request_id": request_id, "документы": documents, "поля": fields, "факты": facts,
            "пометки": warnings, "сверено_с_бланком": False,
            "примечание": "Данные извлечены автоматически; персональные данные физических лиц "
                          "не извлекаются (ЗРУ-547, правило PD-01)"}


# --------------------------------------------------------------------------- #
# Эндпоинты
# --------------------------------------------------------------------------- #

@router.post("/requests/{rid}/ingest")
def ingest_request_endpoint(rid: int, who: str = "api", force: bool = False, use_llm: bool = True):
    """Разобрать документы запроса. По умолчанию — только ещё не разобранные."""
    with db.tx() as con:
        return ingest_request(con, rid, who, force, use_llm)


@router.post("/documents/{photo_id}/ingest")
def ingest_document_endpoint(photo_id: int, who: str = "api", use_llm: bool = True):
    """Разобрать один загруженный файл заново."""
    with db.tx() as con:
        return ingest_file(con, photo_id, who, use_llm)


@router.get("/requests/{rid}/extracts")
def request_extracts(rid: int):
    """Разобранные документы запроса — материал для выгрузок и экрана агента."""
    with db.tx() as con:
        if not db.rows(con, "SELECT id FROM requests WHERE id=?", rid):
            raise HTTPException(404, "Запрос не найден")
    return summary_for_request(rid)


@router.get("/requests/{rid}/extracted-fields")
def request_extracted_fields(rid: int):
    """Сведённые поля по всем документам запроса."""
    with db.tx() as con:
        if not db.rows(con, "SELECT id FROM requests WHERE id=?", rid):
            raise HTTPException(404, "Запрос не найден")
    return fields_for_request(rid)
