"""
Портфельный аудит: выгрузка договоров из учётной системы (xlsx) → массовая проверка движком.
С 22.09.2026 — ещё и шаблон отчётности портфеля (листы «Портфель», «Убытки», «Инструкция», «Справочники»).

Что делает:
  1. читает лист, сам находит строку заголовков и сопоставляет колонки по синонимам (рус/узб/лат);
  2. проверяет каждую строку: типы, суммы, даты, значения из списков — ошибки по строкам и колонкам;
  3. по каждой строке строит engine.Input и вызывает engine.calculate;
  4. складывает результат в portfolio_batches / portfolio_reviews / portfolio_contracts / portfolio_claims
     и отдаёт сводку.

Два режима файла:
  * «шаблон» — в книге есть лист «Портфель» (или «Portfel»): ошибки в строках блокируют загрузку;
  * «выгрузка» — любой другой xlsx учётной системы: как раньше, ошибки становятся пометками строк.

Точки подключения:
  GET  /portfolio/template.xlsx          — пустой шаблон для заполнения (?lang=uz — заголовки по-узбекски)
  POST /portfolio/preview                — разобрать файл без сохранения: ошибки по строкам и сводка
  POST /portfolio/import                 — загрузить xlsx (параметр sheet — имя листа, необязателен)
  GET  /portfolio/batches                — список загрузок
  GET  /portfolio/batches/{id}?verdict=  — строки загрузки
  GET  /portfolio/batches/{id}/summary   — сводка проверки движком
  GET  /portfolio/batches/{id}/overview  — сводка портфеля: договоры, суммы, премии, убытки по классам и регионам
  GET  /portfolio/overview               — то же по последней загрузке
  GET  /portfolio                        — страница
"""
import io
import json
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

import openpyxl
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, Response

from . import db, web
from .engine import Input, calculate, premium_of

ROOT = Path(__file__).resolve().parent.parent
UPLOAD_DIR = db.DATA_DIR / "uploads" / "portfolio"

router = APIRouter()

# Порядок важен: поле, которое идёт раньше, забирает колонку первым
# (например, «страховая стоимость» — до «страховой суммы», «дата начала» — до «срока»).
# Поля шаблона стоят первыми: их заголовки содержат слова старых полей
# («Статус договора» — «договор», «Премия оплаченная» — «премия», «Тип объекта» — «тип»).
# Шаблон «=слово» — только точное совпадение заголовка целиком.
SYNONYMS = [
    ("claims_count", ["число убытков", "количество убытков", "убытков, шт", "кол-во убытков", "zararlar soni"]),
    ("claims_claimed", ["заявлено по убыткам", "убытки заявлено", "заявлено", "da'vo qilingan", "zarar da'vo"]),
    ("claims_paid", ["выплачено по убыткам", "убытки выплачено", "выплачено", "to'langan zarar", "zarar to'langan"]),
    ("last_loss_date", ["дата последнего убытка", "дата убытка", "oxirgi zarar", "zarar sanasi"]),
    ("premium_paid", ["премия оплач", "оплаченная премия", "премия (оплач", "оплачено", "to'langan mukofot",
                      "mukofot to'langan"]),
    ("date_signed", ["дата заключения", "заключ", "tuzilgan"]),
    ("holder_type", ["тип страхователя", "юр/физ", "юр./физ", "sug'urtalanuvchi turi", "shaxs turi"]),
    ("holder_inn", ["=инн", "инн ", "инн,", "инн(", "stir"]),
    ("value_amount", ["страховая стоимость", "страх. стоимость", "стоимость", "qiymat", "haqiqiy qiymat"]),
    ("object_type", ["тип объекта", "вид объекта", "obyekt turi"]),
    ("class_code", ["класс страхования", "класс (", "sug'urta klassi", "klassi"]),
    ("seismic_zone", ["сейсм", "seysm", "zilzila"]),
    ("currency", ["валюта", "valyuta", "currency"]),
    ("status", ["статус договора", "статус", "shartnoma holati", "holati"]),
    ("sum_insured", ["страховая сумма", "страх. сумма", "страховая сум", "ответственность", "sug'urta summasi",
                     "sugurta summasi", "summa", "javobgarlik"]),
    ("franchise", ["франшиза", "franshiza", "franchise"]),
    ("rate", ["ставка", "тариф", "tarif", "stavka", "%"]),
    ("premium", ["премия", "страховой взнос", "взнос", "mukofot", "premiya", "badal"]),
    ("product", ["вид страхования", "вид страх", "продукт", "вид", "класс", "sug'urta turi",
                 "turi", "mahsulot", "product"]),
    ("external_no", ["договор", "полис", "shartnoma", "polis", "contract", "№"]),
    ("policyholder", ["страхователь", "клиент", "sug'urtalovchi", "sugurtalovchi", "mijoz"]),
    ("branch", ["филиал", "подразделение", "отделение", "filial", "bo'lim"]),
    ("region", ["регион", "область", "город", "viloyat", "hudud", "shahar"]),
    ("date_from", ["дата начала", "начало", "срок с", "действует с", "boshlanish", "boshlan"]),
    ("date_to", ["дата окончания", "окончание", "срок по", "действует по", "tugash", "tugallan"]),
    ("term_days", ["срок", "дней", "muddat", "kun"]),
]
# лист «Убытки» шаблона
CLAIM_SYNONYMS = [
    ("claim_no", ["номер убытка", "№ убытка", "zarar raqami"]),
    ("contract_no", ["номер договора", "договор", "полис", "shartnoma", "polis"]),
    ("event_date", ["дата события", "дата убытка", "дата происшествия", "voqea sanasi", "hodisa sanasi"]),
    ("reported_date", ["дата заявления", "ariza sanasi", "murojaat sanasi"]),
    ("claimed", ["заявлено", "сумма заявлен", "da'vo summasi", "da'vo qilingan"]),
    ("paid", ["выплачено", "выплата", "to'langan"]),
    ("status", ["статус", "holat"]),
    ("cause", ["причина", "риск", "sabab", "xavf"]),
]
# заголовки, которые никогда не считаем полезными
IGNORE = ["п/п", "порядк", "t/r", "№ п"]

LABELS = {"external_no": "Договор №", "product": "Вид страхования / продукт", "policyholder": "Страхователь",
          "branch": "Филиал", "sum_insured": "Страховая сумма", "value_amount": "Страховая стоимость",
          "premium": "Премия", "rate": "Ставка", "franchise": "Франшиза", "region": "Регион",
          "date_from": "Дата начала", "date_to": "Дата окончания", "term_days": "Срок, дней"}

REQUIRED = ["sum_insured"]          # без этого проверять нечего
OBJECT_BY_CLASS = {"8": "Склад", "9": "Склад", "3": "Легковой", "14": "Кредит"}
DEFAULT_FACTORS = {"construction": "mixed", "activity": "warehouse", "protection": "alarm",
                   "seismic": "z8", "wear": "mid", "loss_history": "clean"}

# ---------- шаблон отчётности портфеля ----------
# (поле, заголовок ru, заголовок uz, обязательная, вид значения, ширина, что вписать, пример)
PORTFOLIO_COLUMNS = [
    ("external_no", "Номер договора", "Shartnoma raqami", True, "text", 20,
     "Номер как в учётной системе. Повторяться на листе не должен.", "0275/0807/26/0001"),
    ("date_signed", "Дата заключения", "Tuzilgan sanasi", False, "date", 14,
     "Дата подписания договора, ДД.ММ.ГГГГ.", "25.08.2026"),
    ("date_from", "Дата начала", "Boshlanish sanasi", True, "date", 14,
     "Начало ответственности, ДД.ММ.ГГГГ.", "01.09.2026"),
    ("date_to", "Дата окончания", "Tugash sanasi", True, "date", 14,
     "Окончание ответственности; позже даты начала.", "31.08.2027"),
    ("product", "Код продукта", "Mahsulot kodi", True, "product", 12,
     "Четыре цифры из справочника продуктов (лист «Справочники»).", "0807"),
    ("class_code", "Класс страхования", "Sug'urta klassi", False, "class", 11,
     "Если пусто — возьмём первый класс продукта по справочнику.", "8"),
    ("branch", "Филиал", "Filial", False, "text", 18, "Название филиала.", "Ташкентский"),
    ("region", "Регион", "Hudud (viloyat)", True, "region", 26,
     "Где находится объект (или страхователь, если объекта нет). Из списка.", "город Ташкент"),
    ("holder_type", "Тип страхователя", "Sug'urtalanuvchi turi", True, "holder", 15,
     "юрлицо или физлицо. ФИО не вносить.", "юрлицо"),
    ("holder_inn", "ИНН юрлица", "Yuridik shaxs STIR", False, "inn", 13,
     "9 цифр, только для юрлица. У физлица оставить пустым — это персональные данные.", "301234567"),
    ("object_type", "Тип объекта", "Obyekt turi", False, "object", 22,
     "Из списка, если подходит; иначе своими словами.", "Склад"),
    ("sum_insured", "Страховая сумма", "Sug'urta summasi", True, "money", 18,
     "В валюте договора, число без пробелов.", 4200000000),
    ("value_amount", "Стоимость объекта", "Obyekt qiymati", False, "money", 18,
     "Действительная стоимость; если пусто — примем равной страховой сумме.", 4200000000),
    ("premium", "Премия начисленная", "Hisoblangan mukofot", True, "money", 16,
     "Премия по договору за весь срок.", 12600000),
    ("premium_paid", "Премия оплаченная", "To'langan mukofot", False, "money", 16,
     "Сколько фактически поступило на дату отчёта.", 12600000),
    ("franchise", "Франшиза", "Franshiza", False, "money", 14, "Сумма франшизы; 0 — без франшизы.", 42000000),
    ("rate", "Ставка, % годовых", "Tarif stavkasi, % yillik", False, "rate", 12,
     "В процентах: 0,3 — это 0,3 %. Если пусто — посчитаем из премии.", 0.3),
    ("seismic_zone", "Сейсмозона, баллов", "Seysmik zona, ball", False, "seismic", 11,
     "Целое от 6 до 10 (для имущества).", 8),
    ("currency", "Валюта", "Valyuta", False, "currency", 9, "UZS, USD, EUR или RUB. Пусто — UZS.", "UZS"),
    ("status", "Статус договора", "Shartnoma holati", True, "status", 14,
     "действует, расторгнут или истёк.", "действует"),
    ("claims_count", "Число убытков", "Zararlar soni", False, "int", 10,
     "Сколько убытков по договору. Подробности — на листе «Убытки».", 0),
    ("claims_claimed", "Заявлено по убыткам", "Da'vo qilingan zararlar", False, "money", 16,
     "Сумма всех заявленных убытков.", 0),
    ("claims_paid", "Выплачено по убыткам", "To'langan zararlar", False, "money", 16,
     "Сумма всех выплат.", 0),
    ("last_loss_date", "Дата последнего убытка", "Oxirgi zarar sanasi", False, "date", 14,
     "Если убытков не было — пусто.", None),
]
CLAIM_COLUMNS = [
    ("contract_no", "Номер договора", "Shartnoma raqami", True, "text", 20,
     "Номер договора с листа «Портфель».", "0275/0807/26/0001"),
    ("claim_no", "Номер убытка", "Zarar raqami", False, "text", 14, "Номер дела в учётной системе.", "У-2026-015"),
    ("event_date", "Дата события", "Voqea sanasi", True, "date", 14, "Когда случилось, ДД.ММ.ГГГГ.", "12.03.2027"),
    ("reported_date", "Дата заявления", "Ariza sanasi", False, "date", 14, "Не раньше даты события.", "14.03.2027"),
    ("claimed", "Заявлено", "Da'vo summasi", True, "money", 16, "Заявленная сумма.", 85000000),
    ("paid", "Выплачено", "To'langan", False, "money", 16, "Выплачено на дату отчёта.", 80000000),
    ("status", "Статус убытка", "Zarar holati", True, "claim_status", 13, "заявлен, оплачен или отказ.", "оплачен"),
    ("cause", "Причина / риск", "Sabab / xavf", False, "cause", 24, "Из списка, если подходит.", "Пожар"),
]
SHEETS = {"ru": {"portfolio": "Портфель", "claims": "Убытки", "help": "Инструкция", "ref": "Справочники"},
          "uz": {"portfolio": "Portfel", "claims": "Zararlar", "help": "Yo'riqnoma", "ref": "Ma'lumotnomalar"}}
PORTFOLIO_SHEETS = {"портфель", "portfel"}
CLAIM_SHEETS = {"убытки", "zararlar"}

HOLDER_TYPES = {"юрлицо": ["юрлицо", "юр", "юл", "юридическое лицо", "юр. лицо", "юр.лицо", "yuridik shaxs", "yuridik"],
                "физлицо": ["физлицо", "физ", "фл", "физическое лицо", "физ. лицо", "физ.лицо", "ип",
                            "jismoniy shaxs", "jismoniy", "yatt"]}
CURRENCIES = {"UZS": ["uzs", "сум", "сўм", "so'm", "sum"], "USD": ["usd", "доллар", "$"],
              "EUR": ["eur", "евро", "€"], "RUB": ["rub", "руб", "рубль"]}
STATUSES = {"действует": ["действует", "действующий", "amalda", "faol"],
            "расторгнут": ["расторгнут", "расторгнут досрочно", "досрочно расторгнут", "bekor qilingan", "bekor"],
            "истёк": ["истёк", "истек", "закончился", "muddati tugagan", "tugagan"]}
CLAIM_STATUSES = {"заявлен": ["заявлен", "в работе", "на рассмотрении", "arizada", "ko'rib chiqilmoqda"],
                  "оплачен": ["оплачен", "выплачен", "to'langan"],
                  "отказ": ["отказ", "отказано", "rad etilgan", "rad"]}
CAUSES_EXTRA = ["Залив водой", "ДТП", "Угон", "Противоправные действия третьих лиц", "Поломка оборудования",
                "Несчастный случай", "Болезнь", "Невозврат кредита", "Прочее"]

FILL_REQ = "FCE4D6"     # обязательные колонки — тёплый цвет
FILL_OPT = "DDE3F5"     # необязательные
MAX_ROWS = 5000         # на сколько строк ставим выпадающие списки


# ---------- чтение файла ----------

def _norm(s) -> str:
    s = str(s).replace("\n", " ")
    s = re.sub(r"[‘’ʻʼ`´]", "'", s)          # узбекские апострофы пишут по-разному
    return re.sub(r"\s+", " ", s).strip().lower()


def read_sheet(path, sheet: Optional[str] = None) -> dict:
    """Лист → {'headers': [...], 'rows': [[...]], 'header_row': n, 'sheet': имя, 'formats': [...]}.
    Строка заголовков — первая, где не меньше трёх непустых текстовых ячеек."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if sheet:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Листа «{sheet}» нет в файле. Есть: {', '.join(wb.sheetnames)}")
        ws = wb[sheet]
    else:
        ws = wb[wb.sheetnames[0]]
        # первый лист может быть пустым — берём первый, где есть данные
        for name in wb.sheetnames:
            if wb[name].max_row and wb[name].max_row > 1:
                ws = wb[name]
                break
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)), start=1):
        texts = [c.value for c in row if isinstance(c.value, str) and c.value.strip()]
        if len(texts) >= 3:
            header_row = i
            break
    if header_row is None:
        raise ValueError(f"На листе «{ws.title}» не нашёл строку заголовков: нужна строка, где не меньше трёх текстовых ячеек.")
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row]]
    rows, formats, row_nos = [], None, []
    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1), start=header_row + 1):
        vals = [c.value for c in row]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        if formats is None:   # форматы ячеек первой строки данных: чтобы понять, процент ли «ставка»
            formats = [c.number_format or "" for c in row]
        rows.append(vals)
        row_nos.append(i)
    wb.close()
    return {"headers": headers, "rows": rows, "header_row": header_row, "sheet": ws.title,
            "formats": formats or [], "row_nos": row_nos}


def _match(h: str, pats: list) -> bool:
    return any((h == p[1:]) if p.startswith("=") else (p in h) for p in pats)


def map_columns(headers: list, synonyms: list = None, required: list = None) -> dict:
    """Сопоставление поле -> индекс колонки. Возвращает {'index': {...}, 'names': {...}, 'missing': [...]}."""
    synonyms = SYNONYMS if synonyms is None else synonyms
    required = REQUIRED if required is None else required
    norm = [_norm(h) for h in headers]
    taken, index = set(), {}
    for field, pats in synonyms:
        pats = [_norm(p) if not p.startswith("=") else "=" + _norm(p[1:]) for p in pats]
        for i, h in enumerate(norm):
            if not h or i in taken or any(x in h for x in IGNORE):
                continue
            if _match(h, pats):
                index[field] = i
                taken.add(i)
                break
    names = {f: headers[i] for f, i in index.items()}
    missing = [f for f in required if f not in index]
    return {"index": index, "names": names, "missing": missing}


def explain_missing(missing: list, labels: dict = None) -> str:
    labels = labels or LABELS
    return "Не найдены колонки: " + "; ".join(f"«{labels.get(m, m)}»" for m in missing) + \
        ". Проверьте заголовки в файле или переименуйте колонки."


# ---------- разбор значений ----------

def to_number(v) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(" ", "").replace("%", "")
    if not s or s in ("-", "—"):
        return None
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def to_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    # дата, которую Excel отдал числом (ячейка без формата даты): 45000 ≈ 2023 год
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 <= v <= 80000:
        return date(1899, 12, 30) + timedelta(days=int(v))
    if isinstance(v, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
            try:
                return datetime.strptime(v.strip()[:10], fmt).date()
            except ValueError:
                pass
    return None


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _pick(v, variants: dict) -> Optional[str]:
    """Значение из списка: 'Юр. лицо' → 'юрлицо'. Не узнали — None."""
    s = _norm(v)
    for canon, alts in variants.items():
        if s == _norm(canon) or s in [_norm(a) for a in alts]:
            return canon
    return None


def product_code_of(product_cell, external_no, known: dict) -> Optional[str]:
    """Код продукта: из колонки вида («0807.Страхование…», 807, «0807»), иначе из номера договора 0275/0311/26/0030."""
    if product_cell is not None:
        if isinstance(product_cell, (int, float)) and not isinstance(product_cell, bool):
            code = f"{int(product_cell):04d}"
            if code in known:
                return code
        s = str(product_cell).strip()
        m = re.match(r"^\s*(\d{4})\b", s) or re.search(r"\b(\d{4})\b", s)
        if m and m.group(1) in known:
            return m.group(1)
        if m and not known:
            return m.group(1)
    if external_no:
        parts = re.split(r"[/\-\s]+", str(external_no).strip())
        if len(parts) >= 2 and re.fullmatch(r"\d{4}", parts[1]) and (not known or parts[1] in known):
            return parts[1]
        for p in parts:
            if re.fullmatch(r"\d{4}", p) and p in known:
                return p
    return None


def _cell(row, idx: dict, field: str):
    i = idx.get(field)
    return row[i] if i is not None and i < len(row) else None


def _text(v) -> Optional[str]:
    if _blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_row(row: list, mapping: dict, formats: list, known: dict, issues: list = None) -> dict:
    """Строка выгрузки → плоский словарь полей договора.
    issues — куда сложить «не число» / «не дата» по непустым ячейкам (поле, сообщение)."""
    idx = mapping["index"]
    g = lambda f: _cell(row, idx, f)
    bad = (lambda f, msg: issues.append((f, msg))) if issues is not None else (lambda f, msg: None)

    def num(f):
        v = g(f)
        n = to_number(v)
        if n is None and not _blank(v) and str(v).strip() not in ("-", "—"):
            bad(f, "не число")
        return n

    def dt(f):
        v = g(f)
        d = to_date(v)
        if d is None and not _blank(v):
            bad(f, "дата не распознана — нужна ДД.ММ.ГГГГ")
        return d

    external_no = _text(g("external_no"))
    rate = num("rate")
    # ставка в процентном формате Excel хранится долей: 0,0005 → 0,05 %
    if rate is not None and "rate" in idx and idx["rate"] < len(formats) and "%" in formats[idx["rate"]]:
        rate *= 100
    d_from, d_to = dt("date_from"), dt("date_to")
    d_signed, d_loss = dt("date_signed"), dt("last_loss_date")
    term = num("term_days")
    if term is None and d_from and d_to and d_to > d_from:
        term = (d_to - d_from).days
    fr = num("franchise")
    cnt = num("claims_count")
    sz = num("seismic_zone")
    return {
        "external_no": external_no,
        "product_code": product_code_of(g("product"), external_no, known),
        "product_text": _text(g("product")),
        "policyholder": _text(g("policyholder")),
        "branch": _text(g("branch")),
        "region": _text(g("region")),
        "sum_insured": num("sum_insured"),
        "value_amount": num("value_amount"),
        "premium": num("premium"),
        "rate": rate,
        "franchise": fr or 0.0,
        "term_days": int(term) if term and term > 0 else 365,
        "date_from": d_from.isoformat() if d_from else None,
        "date_to": d_to.isoformat() if d_to else None,
        # поля шаблона отчётности
        "date_signed": d_signed.isoformat() if d_signed else None,
        "class_code": _text(g("class_code")),
        "holder_type": _text(g("holder_type")),
        "holder_inn": _text(g("holder_inn")),
        "object_type": _text(g("object_type")),
        "premium_paid": num("premium_paid"),
        "seismic_zone": sz,
        "currency": _text(g("currency")),
        "status": _text(g("status")),
        "claims_count": cnt,
        "claims_claimed": num("claims_claimed"),
        "claims_paid": num("claims_paid"),
        "last_loss_date": d_loss.isoformat() if d_loss else None,
    }


# ---------- справочники для шаблона и проверки ----------

def reference_lists(con) -> dict:
    """Продукты, классы, регионы, типы объектов и причины убытков — из базы."""
    from .stat_sources import REGION_NAMES_RU
    pcs = {}
    for r in db.rows(con, "SELECT product_code, class_code FROM product_classes ORDER BY product_code, part_no"):
        pcs.setdefault(r["product_code"], []).append(r["class_code"])
    products = [dict(p, classes=", ".join(pcs.get(p["code"], [])))
                for p in db.rows(con, "SELECT code, name FROM products ORDER BY code")]
    # «13з», «16у» — рядом со своими номерами; сортируем в Python, без приведения типов в SQL
    num = lambda c: (int(re.match(r"\d+", c).group()) if re.match(r"\d+", c) else 999, c)
    # как на экранах и в /reference/classes (db.CLASSES_SHOWN): строк L* (жизнь) в шаблоне и проверке нет
    classes = sorted(db.rows(con, f"SELECT code, name FROM classes WHERE {db.CLASSES_SHOWN}"),
                     key=lambda c: num(c["code"]))
    regions = db.rows(con, """SELECT row_key, MAX(row_name) row_name FROM market_stats
                              WHERE row_key LIKE 'region:%' GROUP BY row_key ORDER BY 2""")
    regions = [{"key": r["row_key"], "name": r["row_name"]} for r in regions] or \
        [{"key": k, "name": v} for k, v in REGION_NAMES_RU.items() if k != "total"]
    objects = sorted({r["object_type"] for r in db.rows(con, "SELECT DISTINCT object_type FROM base_rates")})
    perils = [r["name"] for r in db.rows(con, "SELECT name FROM perils ORDER BY class_code, code")]
    return {"products": products, "classes": classes, "regions": regions, "objects": objects,
            "causes": perils + [c for c in CAUSES_EXTRA if c not in perils]}


def region_of(v, regions: list) -> Optional[str]:
    """Регион из свободного текста → русское название из справочника. Не узнали — None."""
    s = _norm(v)
    if not s:
        return None
    for r in regions:
        if s in (_norm(r["name"]), _norm(r["key"]), _norm(r["key"].split(":", 1)[-1])):
            return r["name"]
    # «г. Ташкент» / «Toshkent shahri» — город; «Ташкентская обл.» — область
    if "ташкент" in s or "toshkent" in s:
        city = not any(x in s for x in ("обл", "viloyat", "ская"))
        want = "region:TOSHKENT SHAHRI" if city else "region:TOSHKENT"
        return next((r["name"] for r in regions if r["key"] == want), None)
    if "каракалп" in s or "qoraqalp" in s:
        return next((r["name"] for r in regions if "QORAQALP" in r["key"]), None)
    for r in regions:
        stem_ru = _norm(r["name"])[:5]
        stem_uz = _norm(r["key"].split(":", 1)[-1])[:5]
        if stem_ru in s or stem_uz in s:
            return r["name"]
    return None


# ---------- проверка строк ----------

def _issue(sheet, row_no, column, message, level="ошибка") -> dict:
    return {"sheet": sheet, "row": row_no, "column": column, "level": level, "message": message}


def _col_title(field: str, cols: list, names: dict) -> str:
    """Название колонки для сообщения: как в файле, иначе как в шаблоне."""
    if names.get(field):
        return names[field]
    return next((c[1] for c in cols if c[0] == field), LABELS.get(field, field))


def check_contract(p: dict, row_no: int, sheet: str, names: dict, lists: dict, known: dict,
                   template: bool, parse_issues: list) -> list:
    """Ошибки и предупреждения одной строки листа «Портфель». Нормализует значения в p на месте."""
    out = []
    col = lambda f: _col_title(f, PORTFOLIO_COLUMNS, names)
    err = lambda f, m: out.append(_issue(sheet, row_no, col(f), m))
    warn = lambda f, m: out.append(_issue(sheet, row_no, col(f), m, "предупреждение"))
    for f, m in parse_issues:
        err(f, m)
    bad_fields = {f for f, _ in parse_issues}
    if template:
        for f, _ru, _uz, req, *_ in PORTFOLIO_COLUMNS:
            key = {"product": "product_text", "rate": "rate"}.get(f, f)
            if req and f not in bad_fields and p.get(key) in (None, ""):
                err(f, "обязательная колонка не заполнена")
    # продукт и класс
    if p.get("product_text") and not p.get("product_code"):
        (err if template else warn)("product", "код продукта не найден в справочнике")
    classes = known.get(p.get("product_code") or "", [])
    if p.get("class_code"):
        cc = p["class_code"]
        valid = {c["code"] for c in lists["classes"]}
        if cc not in valid:
            (err if template else warn)("class_code", "нет такого класса (1–18, 13з, 16у)")
        elif classes and cc not in classes:
            warn("class_code", f"класс {cc} не входит в продукт {p['product_code']} (по справочнику: {', '.join(classes)})")
        p["class_code"] = cc
    elif classes:
        p["class_code"] = classes[0]
    # регион
    if p.get("region"):
        r = region_of(p["region"], lists["regions"])
        if r:
            p["region"] = r
        elif template:
            err("region", "регион не из списка (лист «Справочники»)")
    # страхователь: только тип и ИНН юрлица
    ht = _pick(p["holder_type"], HOLDER_TYPES) if p.get("holder_type") else None
    if p.get("holder_type") and not ht:
        (err if template else warn)("holder_type", "нужно «юрлицо» или «физлицо»")
    p["holder_type"] = ht
    inn = p.get("holder_inn")
    if inn:
        if ht == "физлицо":
            # значение в сообщение не пишем: это персональные данные
            err("holder_inn", "ИНН физлица в отчёт не вносится (персональные данные) — очистите ячейку")
            p["holder_inn"] = None
        elif not re.fullmatch(r"\d{9}", inn.replace(" ", "")):
            err("holder_inn", "ИНН юрлица — ровно 9 цифр")
            p["holder_inn"] = None
        else:
            p["holder_inn"] = inn.replace(" ", "")
    elif ht == "юрлицо" and template:
        warn("holder_inn", "ИНН юрлица не указан")
    # даты
    if p.get("date_from") and p.get("date_to") and p["date_to"] <= p["date_from"]:
        err("date_to", "дата окончания не позже даты начала")
    if p.get("date_signed") and p.get("date_from") and p["date_signed"] > p["date_from"]:
        warn("date_signed", "договор заключён позже начала ответственности")
    # суммы
    for f in ("sum_insured", "value_amount", "premium", "premium_paid", "franchise",
              "claims_claimed", "claims_paid"):
        if p.get(f) is not None and p[f] < 0:
            err(f, "сумма не может быть отрицательной")
    if p.get("sum_insured") is not None and p["sum_insured"] == 0 and template:
        err("sum_insured", "страховая сумма — больше нуля")
    if p.get("franchise") and p.get("sum_insured") and p["franchise"] >= p["sum_insured"]:
        err("franchise", "франшиза не меньше страховой суммы")
    if p.get("premium_paid") is not None and p.get("premium") is not None and p["premium_paid"] > p["premium"] * 1.0001:
        warn("premium_paid", "оплачено больше начисленного")
    if p.get("rate") is not None and not (0 < p["rate"] <= 100):
        err("rate", "ставка в процентах: больше 0 и не больше 100")
    sz = p.get("seismic_zone")
    if sz is not None:
        if not float(sz).is_integer() or not 6 <= sz <= 10:
            err("seismic_zone", "сейсмозона — целое число от 6 до 10")
            p["seismic_zone"] = None
        else:
            p["seismic_zone"] = int(sz)
    # списки
    if p.get("currency"):
        c = _pick(p["currency"], CURRENCIES)
        if not c:
            err("currency", "валюта — UZS, USD, EUR или RUB")
        p["currency"] = c or p["currency"].upper()[:3]
    else:
        p["currency"] = "UZS"
    if p.get("status"):
        s = _pick(p["status"], STATUSES)
        if not s:
            (err if template else warn)("status", "статус — действует, расторгнут или истёк")
        p["status"] = s
        today = date.today().isoformat()
        if s == "действует" and p.get("date_to") and p["date_to"] < today:
            warn("status", "указано «действует», но срок договора уже закончился")
        if s == "истёк" and p.get("date_to") and p["date_to"] >= today:
            warn("status", "указано «истёк», но срок договора ещё идёт")
    # убытки одной строкой
    cnt = p.get("claims_count")
    if cnt is not None:
        if cnt < 0 or not float(cnt).is_integer():
            err("claims_count", "число убытков — целое, не меньше нуля")
            p["claims_count"] = None
        else:
            p["claims_count"] = int(cnt)
    if (p.get("claims_paid") or 0) > (p.get("claims_claimed") or 0) > 0:
        warn("claims_paid", "выплачено больше заявленного")
    if p.get("claims_count") == 0 and ((p.get("claims_claimed") or 0) > 0 or (p.get("claims_paid") or 0) > 0):
        warn("claims_count", "убытков 0, а суммы убытков указаны")
    if p.get("last_loss_date") and p.get("date_from") and p["last_loss_date"] < p["date_from"]:
        warn("last_loss_date", "убыток раньше начала ответственности")
    return out


def parse_claim(row, mapping, issues) -> dict:
    idx = mapping["index"]
    g = lambda f: _cell(row, idx, f)
    out = {}
    for f in ("contract_no", "claim_no", "status", "cause"):
        out[f] = _text(g(f))
    for f in ("claimed", "paid"):
        v = g(f)
        out[f] = to_number(v)
        if out[f] is None and not _blank(v):
            issues.append((f, "не число"))
    for f in ("event_date", "reported_date"):
        v = g(f)
        d = to_date(v)
        out[f] = d.isoformat() if d else None
        if d is None and not _blank(v):
            issues.append((f, "дата не распознана — нужна ДД.ММ.ГГГГ"))
    return out


def check_claim(c: dict, row_no: int, sheet: str, names: dict, contracts: dict, parse_issues: list) -> list:
    out = []
    col = lambda f: _col_title(f, CLAIM_COLUMNS, names)
    err = lambda f, m: out.append(_issue(sheet, row_no, col(f), m))
    warn = lambda f, m: out.append(_issue(sheet, row_no, col(f), m, "предупреждение"))
    for f, m in parse_issues:
        err(f, m)
    bad = {f for f, _ in parse_issues}
    for f, _ru, _uz, req, *_ in CLAIM_COLUMNS:
        if req and f not in bad and c.get(f) in (None, ""):
            err(f, "обязательная колонка не заполнена")
    k = c.get("contract_no")
    if k and k not in contracts:
        err("contract_no", "такого договора нет на листе «Портфель»")
    for f in ("claimed", "paid"):
        if c.get(f) is not None and c[f] < 0:
            err(f, "сумма не может быть отрицательной")
    if (c.get("paid") or 0) > (c.get("claimed") or 0) > 0:
        warn("paid", "выплачено больше заявленного")
    if c.get("status"):
        s = _pick(c["status"], CLAIM_STATUSES)
        if not s:
            err("status", "статус — заявлен, оплачен или отказ")
        c["status"] = s
        if s == "отказ" and (c.get("paid") or 0) > 0:
            warn("paid", "статус «отказ», а выплата есть")
    if c.get("event_date") and c.get("reported_date") and c["reported_date"] < c["event_date"]:
        err("reported_date", "заявление раньше события")
    p = contracts.get(k) if k else None
    if p and c.get("event_date"):
        if (p.get("date_from") and c["event_date"] < p["date_from"]) or (p.get("date_to") and c["event_date"] > p["date_to"]):
            warn("event_date", "событие вне срока договора")
    return out


# ---------- разбор файла целиком ----------

def _find_sheet(path, names: set) -> Optional[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return next((n for n in wb.sheetnames if _norm(n) in names), None)
    finally:
        wb.close()


def analyze_file(path, sheet: Optional[str] = None, con=None) -> dict:
    """Файл → договоры, убытки, ошибки по строкам. В базу ничего не пишет.
    Нет обязательной колонки — ValueError с объяснением по-русски."""
    tpl_sheet = None if sheet else _find_sheet(path, PORTFOLIO_SHEETS)
    template = tpl_sheet is not None or (sheet is not None and _norm(sheet) in PORTFOLIO_SHEETS)
    sh = read_sheet(path, sheet or tpl_sheet)
    required = ["sum_insured"] if not template else ["external_no", "sum_insured", "premium"]
    mapping = map_columns(sh["headers"], SYNONYMS, required)
    if mapping["missing"]:
        labels = dict(LABELS, **{c[0]: c[1] for c in PORTFOLIO_COLUMNS})
        raise ValueError(explain_missing(mapping["missing"], labels) + f" Заголовки на листе «{sh['sheet']}» "
                         f"(строка {sh['header_row']}): " + ", ".join(h for h in sh["headers"] if h))
    if "product" not in mapping["index"] and "external_no" not in mapping["index"]:
        raise ValueError("Не найдены ни колонка «Вид страхования», ни «Договор №» — код продукта взять негде.")

    def run(c):
        ref = db.load_reference(c)
        return ref, reference_lists(c)
    if con is None:
        with db.tx() as c:
            ref, lists = run(c)
    else:
        ref, lists = run(con)
    known = ref.product_classes

    issues, contracts, seen = [], [], {}
    for row_no, row in zip(sh["row_nos"], sh["rows"]):
        pi = []
        p = parse_row(row, mapping, sh["formats"], known, pi)
        if p["external_no"] is None and p["sum_insured"] is None and not template:
            continue
        if p["external_no"] and _norm(p["external_no"]).startswith(("итого", "всего", "jami")):
            continue
        row_issues = check_contract(p, row_no, sh["sheet"], mapping["names"], lists, known, template, pi)
        if p["external_no"]:
            if p["external_no"] in seen:
                row_issues.append(_issue(sh["sheet"], row_no, _col_title("external_no", PORTFOLIO_COLUMNS, mapping["names"]),
                                         f"номер договора повторяется (уже в строке {seen[p['external_no']]})"))
            else:
                seen[p["external_no"]] = row_no
        raw = {h: _json_safe(v) for h, v in zip(sh["headers"], row) if h}
        if "holder_inn" in mapping["index"] and p.get("holder_type") != "юрлицо":
            raw.pop(mapping["names"]["holder_inn"], None)     # ИНН физлица не храним даже в исходной строке
        issues += row_issues
        contracts.append({"row_no": row_no, "p": p, "raw": raw, "issues": row_issues})

    claims, claim_sheet, claim_mapping = [], None, None
    cs = _find_sheet(path, CLAIM_SHEETS) if template else None
    if cs:
        csh = read_sheet_safe(path, cs)
        if csh and csh["rows"]:
            claim_sheet = csh["sheet"]
            claim_mapping = map_columns(csh["headers"], CLAIM_SYNONYMS, ["contract_no", "claimed"])
            if claim_mapping["missing"]:
                labels = {c[0]: c[1] for c in CLAIM_COLUMNS}
                issues.append(_issue(claim_sheet, csh["header_row"], "—",
                                     explain_missing(claim_mapping["missing"], labels)))
            else:
                by_no = {c["p"]["external_no"]: c["p"] for c in contracts if c["p"]["external_no"]}
                for row_no, row in zip(csh["row_nos"], csh["rows"]):
                    pi = []
                    c = parse_claim(row, claim_mapping, pi)
                    ci = check_claim(c, row_no, claim_sheet, claim_mapping["names"], by_no, pi)
                    issues += ci
                    claims.append({"row_no": row_no, "c": c, "issues": ci})
    errors = [i for i in issues if i["level"] == "ошибка"]
    return {"mode": "шаблон" if template else "выгрузка", "template": template,
            "sheet": sh["sheet"], "header_row": sh["header_row"], "mapping": mapping["names"],
            "unmapped": [h for i, h in enumerate(sh["headers"]) if h and i not in mapping["index"].values()],
            "claims_sheet": claim_sheet, "claims_mapping": claim_mapping["names"] if claim_mapping else {},
            "contracts": contracts, "claims": claims, "errors": errors,
            "warnings": [i for i in issues if i["level"] != "ошибка"], "ref": ref, "lists": lists}


def read_sheet_safe(path, sheet):
    try:
        return read_sheet(path, sheet)
    except ValueError:
        return None


# ---------- проверка движком ----------

def _seismic_option(zone) -> Optional[str]:
    if zone is None:
        return None
    return "z7" if zone <= 7 else "z8" if zone == 8 else "z9"


def review_row(ref, p: dict) -> dict:
    """Один договор → результат проверки. Если данных не хватает — вердикт «не проверен» с причиной."""
    out = {k: p.get(k) for k in ("external_no", "product_code", "policyholder", "branch", "sum_insured",
                                 "value_amount", "region", "term_days")}
    out.update({"premium_file": p.get("premium"), "premium_calc": None, "applied_rate_pct": None,
                "technical_rate_pct": None, "min_rate_pct": None, "verdict": "не проверен",
                "violations": [], "note": None})
    if not p.get("sum_insured"):
        out["note"] = "нет страховой суммы"
        return out
    if not p.get("product_code"):
        out["note"] = "не распознан код продукта" + (f" («{p['product_text']}»)" if p.get("product_text") else "")
        return out
    classes = ref.product_classes.get(p["product_code"])
    if not classes:
        out["note"] = f"продукт {p['product_code']} не найден в справочнике"
        return out
    if p.get("currency") and p["currency"] != "UZS":
        # движок и минимальные премии — в сумах; курса в шаблоне нет, пересчитывать не берёмся
        out["note"] = f"валюта {p['currency']}: движок считает в сумах, курс в отчёте не задан"
        return out
    cls = p["class_code"] if p.get("class_code") in classes else classes[0]
    obj = p.get("object_type") if (cls, p.get("object_type")) in ref.base_rates else OBJECT_BY_CLASS.get(cls, "Прочее")
    value = p["value_amount"] if p.get("value_amount") else p["sum_insured"]   # нет стоимости — считаем равной сумме
    if not p.get("value_amount"):
        out["note"] = "страховая стоимость не указана, принята равной сумме"
    rate = p.get("rate")
    if rate is None and p.get("premium") and p["sum_insured"]:
        rate = p["premium"] / p["sum_insured"] * 100 * 365 / p["term_days"]
    factors = {**DEFAULT_FACTORS, "franchise": "f1" if (p.get("franchise") or 0) > 0 else "f0"}
    if _seismic_option(p.get("seismic_zone")):
        factors["seismic"] = _seismic_option(p["seismic_zone"])
    inp = Input(product_code=p["product_code"], class_code=cls,
                object_type=obj,
                value_amount=value, sum_insured=p["sum_insured"], term_days=p["term_days"],
                factors=factors, applied_rate_pct=rate if rate else None,
                manual_reason="из выгрузки" if rate else "",
                # в портфельном режиме проверяем цену и закон, а не документооборот: эти три
                # предупреждения (документы, оплата, раскрытие) к выгрузке неприменимы
                docs_received=[c["doc_name"] for c in ref.checklists], premium_paid=True, disclosure_done=True)
    res = calculate(ref, inp)
    r = res["rates"]
    prem_calc = round(premium_of(r["auto_pct"], p["sum_insured"], p["term_days"]))
    prem_file = p.get("premium") if p.get("premium") is not None else res["premium"]
    out.update({
        "class_code": cls, "object_type": inp.object_type,
        "applied_rate_pct": round(r["applied_pct"], 4), "technical_rate_pct": round(r["technical_pct"], 4),
        "min_rate_pct": r["min_pct"], "premium_file": prem_file, "premium_calc": prem_calc,
        "premium_diff": round(prem_calc - prem_file), "verdict": res["verdict"],
        # ручная ставка «из выгрузки» — не нарушение, а способ передачи ставки; в список не включаем
        "violations": [{"rule": c["rule"], "status": c["status"], "title": c["title"]}
                       for c in res["checks"] if c["status"] != "ok" and c["rule"] != "manual_rate"],
    })
    return out


def review_rows(ref, sheet: dict, mapping: dict) -> list:
    """Все строки листа → список результатов. Строки без номера и без суммы (итоги, подписи) пропускаются."""
    known = ref.product_classes
    out = []
    row_nos = sheet.get("row_nos") or list(range(sheet["header_row"] + 1, sheet["header_row"] + 1 + len(sheet["rows"])))
    for i, row in zip(row_nos, sheet["rows"]):
        p = parse_row(row, mapping, sheet["formats"], known)
        if p["external_no"] is None and p["sum_insured"] is None:
            continue
        if p["external_no"] and _norm(p["external_no"]).startswith(("итого", "всего", "jami")):
            continue
        r = review_row(ref, p)
        r["row_no"] = i
        r["raw"] = {h: _json_safe(v) for h, v in zip(sheet["headers"], row) if h}
        out.append(r)
    return out


def _json_safe(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


# ---------- сохранение и сводка ----------

CONTRACT_FIELDS = ["external_no", "date_signed", "date_from", "date_to", "product_code", "class_code", "branch",
                   "region", "holder_type", "holder_inn", "object_type", "sum_insured", "value_amount", "premium",
                   "premium_paid", "franchise", "rate_pct", "seismic_zone", "currency", "status", "claims_count",
                   "claims_claimed", "claims_paid", "last_loss_date"]
CLAIM_FIELDS = ["contract_no", "claim_no", "event_date", "reported_date", "claimed", "paid", "status", "cause"]


def save_batch(con, file_name: str, mapping: dict, reviews: list, claims: list = None) -> int:
    ts = db.now()
    cnt = lambda v: sum(1 for r in reviews if r["verdict"] == v)
    cur = con.execute(
        "INSERT INTO portfolio_batches (imported_at, file_name, rows_total, rows_ok, rows_warn, rows_stop, mapping)"
        " VALUES (?,?,?,?,?,?,?)",
        (ts, file_name, len(reviews), cnt("ок"), cnt("на утверждение"), cnt("отклонено"),
         json.dumps(mapping["names"], ensure_ascii=False)))
    bid = cur.lastrowid
    for r in reviews:
        con.execute(
            "INSERT INTO portfolio_reviews (batch_id, imported_at, file_name, external_no, product_code, policyholder,"
            " branch, sum_insured, value_amount, premium_file, premium_calc, applied_rate_pct, technical_rate_pct,"
            " min_rate_pct, verdict, violations, row_json) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (bid, ts, file_name, r["external_no"], r["product_code"], r["policyholder"], r["branch"],
             r["sum_insured"], r["value_amount"], r["premium_file"], r["premium_calc"], r["applied_rate_pct"],
             r["technical_rate_pct"], r["min_rate_pct"], r["verdict"],
             json.dumps(r["violations"], ensure_ascii=False),
             json.dumps({"raw": r.get("raw"), "row_no": r.get("row_no"), "note": r.get("note"),
                         "premium_diff": r.get("premium_diff"), "issues": r.get("issues") or []},
                        ensure_ascii=False)))
        p = r.get("contract")
        if p:
            vals = dict(p, rate_pct=p.get("rate"))
            con.execute(f"INSERT INTO portfolio_contracts (batch_id, row_no, {', '.join(CONTRACT_FIELDS)})"
                        f" VALUES ({', '.join(['?'] * (len(CONTRACT_FIELDS) + 2))})",
                        [bid, r.get("row_no")] + [vals.get(f) for f in CONTRACT_FIELDS])
    for c in claims or []:
        con.execute(f"INSERT INTO portfolio_claims (batch_id, row_no, {', '.join(CLAIM_FIELDS)})"
                    f" VALUES ({', '.join(['?'] * (len(CLAIM_FIELDS) + 2))})",
                    [bid, c["row_no"]] + [c["c"].get(f) for f in CLAIM_FIELDS])
    db.audit(con, "backend", "портфельный аудит: загрузка", f"portfolio_batch:{bid}",
             {"file": file_name, "rows": len(reviews), "stop": cnt("отклонено"), "claims": len(claims or [])})
    return bid


def _review_from_db(r: dict) -> dict:
    extra = json.loads(r.pop("row_json") or "{}")
    r["violations"] = json.loads(r["violations"] or "[]")
    r["note"] = extra.get("note")
    r["row_no"] = extra.get("row_no")
    r["premium_diff"] = extra.get("premium_diff")
    r["raw"] = extra.get("raw")
    r["issues"] = extra.get("issues") or []
    return r


def batch_rows(con, bid: int, verdict: Optional[str] = None) -> list:
    sql, args = "SELECT * FROM portfolio_reviews WHERE batch_id=?", [bid]
    if verdict:
        sql += " AND verdict=?"; args.append(verdict)
    sql += " ORDER BY id"
    return [_review_from_db(r) for r in db.rows(con, sql, *args)]


def summarize(reviews: list) -> dict:
    """Сводка по списку результатов: счётчики нарушений, недополученная премия, топы по филиалам и продуктам."""
    has = lambda r, code: any(v["rule"] == code for v in r["violations"])
    verdicts = {}
    for r in reviews:
        verdicts[r["verdict"]] = verdicts.get(r["verdict"], 0) + 1
    lost = sum(max(r.get("premium_diff") or 0, 0) for r in reviews)
    tops = {"branch": {}, "product_code": {}}
    for r in reviews:
        if r["verdict"] in ("ок", "не проверен"):
            continue
        for key in tops:
            k = r.get(key) or "—"
            t = tops[key].setdefault(k, {"name": k, "rows": 0, "stop": 0, "lost": 0})
            t["rows"] += 1
            t["stop"] += r["verdict"] == "отклонено"
            t["lost"] += max(r.get("premium_diff") or 0, 0)
    top = lambda d: sorted(d.values(), key=lambda t: (-t["stop"], -t["rows"], -t["lost"]))[:10]
    return {
        "rows_total": len(reviews),
        "by_verdict": verdicts,
        "below_min": sum(has(r, "min_rate") for r in reviews),
        "below_technical": sum(has(r, "below_technical") for r in reviews),
        "sum_over_value": sum(has(r, "sum_over_value") for r in reviews),
        "no_franchise": sum(has(r, "no_franchise") for r in reviews),
        "risk_limit_20": sum(any(v["rule"] == "risk_limit_20" and v["status"] == "stop" for v in r["violations"])
                             for r in reviews),
        "unchecked": sum(r["verdict"] == "не проверен" for r in reviews),
        "sum_insured_total": sum(r["sum_insured"] or 0 for r in reviews),
        "premium_file_total": sum(r["premium_file"] or 0 for r in reviews),
        "premium_calc_total": sum(r["premium_calc"] or 0 for r in reviews if r["premium_calc"] is not None),
        "lost_premium": round(lost),
        "top_branches": top(tops["branch"]),
        "top_products": top(tops["product_code"]),
    }


def overview(contracts: list, claims: list, class_names: dict = None) -> dict:
    """Сводка портфеля: число договоров, суммы, премии и убытки — всего, по классам и по регионам.
    Суммы складываются только по договорам в сумах: другие валюты без курса не смешиваем."""
    class_names = class_names or {}
    uzs = [c for c in contracts if (c.get("currency") or "UZS") == "UZS"]
    uzs_nos = {c.get("external_no") for c in uzs}
    # убытки: по листу «Убытки», если он заполнен; иначе — по колонкам на листе «Портфель»
    detail = [k for k in claims if k.get("contract_no") in uzs_nos]
    paid_by_no, claimed_by_no, count_by_no = {}, {}, {}
    for k in detail:
        n = k["contract_no"]
        paid_by_no[n] = paid_by_no.get(n, 0) + (k.get("paid") or 0)
        claimed_by_no[n] = claimed_by_no.get(n, 0) + (k.get("claimed") or 0)
        count_by_no[n] = count_by_no.get(n, 0) + 1
    use_detail = bool(claims)

    def claims_of(c):
        if use_detail:
            n = c.get("external_no")
            return count_by_no.get(n, 0), claimed_by_no.get(n, 0), paid_by_no.get(n, 0)
        return c.get("claims_count") or 0, c.get("claims_claimed") or 0, c.get("claims_paid") or 0

    def bucket(key_fn, name_fn):
        d = {}
        for c in uzs:
            k = key_fn(c) or "не указан"
            b = d.setdefault(k, {"key": k, "name": name_fn(k), "contracts": 0, "sum_insured": 0, "premium": 0,
                                 "premium_paid": 0, "claims": 0, "claims_paid": 0})
            n, _cl, pd = claims_of(c)
            b["contracts"] += 1
            b["sum_insured"] += c.get("sum_insured") or 0
            b["premium"] += c.get("premium") or 0
            b["premium_paid"] += c.get("premium_paid") or 0
            b["claims"] += n
            b["claims_paid"] += pd
        for b in d.values():
            b["loss_ratio"] = round(b["claims_paid"] / b["premium"] * 100, 1) if b["premium"] else None
        return sorted(d.values(), key=lambda b: -b["premium"])

    tot = {"contracts": len(contracts), "contracts_uzs": len(uzs), "other_currency": len(contracts) - len(uzs),
           "sum_insured": sum(c.get("sum_insured") or 0 for c in uzs),
           "premium": sum(c.get("premium") or 0 for c in uzs),
           "premium_paid": sum(c.get("premium_paid") or 0 for c in uzs)}
    cl = [claims_of(c) for c in uzs]
    tot["claims"] = sum(x[0] for x in cl)
    tot["claims_claimed"] = sum(x[1] for x in cl)
    tot["claims_paid"] = sum(x[2] for x in cl)
    tot["loss_ratio"] = round(tot["claims_paid"] / tot["premium"] * 100, 1) if tot["premium"] else None
    tot["paid_share"] = round(tot["premium_paid"] / tot["premium"] * 100, 1) if tot["premium"] else None
    by = lambda f: {k: sum(1 for c in contracts if (c.get(f) or "не указан") == k)
                    for k in sorted({c.get(f) or "не указан" for c in contracts})}
    return {"totals": tot, "claims_source": "лист «Убытки»" if use_detail else "колонки листа «Портфель»",
            "by_status": by("status"), "by_holder": by("holder_type"), "by_currency": by("currency"),
            "by_class": bucket(lambda c: c.get("class_code"),
                               lambda k: f"{k} — {class_names[k]}" if k in class_names else k),
            "by_region": bucket(lambda c: c.get("region"), lambda k: k)}


def _class_names(con) -> dict:
    return {r["code"]: r["name"] for r in db.rows(con, "SELECT code, name FROM classes")}


def batch_overview(con, bid: int) -> dict:
    contracts = db.rows(con, "SELECT * FROM portfolio_contracts WHERE batch_id=? ORDER BY id", bid)
    claims = db.rows(con, "SELECT * FROM portfolio_claims WHERE batch_id=? ORDER BY id", bid)
    return overview(contracts, claims, _class_names(con))


def _preview_out(a: dict) -> dict:
    """Ответ предпросмотра: без справочников движка, ошибок — не больше 500 строк."""
    contracts = [c["p"] for c in a["contracts"]]
    claims = [c["c"] for c in a["claims"]]
    with db.tx() as con:
        names = _class_names(con)
    show = [dict(c["p"], row_no=c["row_no"], issues=c["issues"]) for c in a["contracts"][:200]]
    for s in show:
        s.pop("policyholder", None)      # в предпросмотре имена страхователей не показываем
    return {"mode": a["mode"], "template": a["template"], "sheet": a["sheet"], "header_row": a["header_row"],
            "mapping": a["mapping"], "unmapped": a["unmapped"], "claims_sheet": a["claims_sheet"],
            "claims_mapping": a["claims_mapping"], "rows_total": len(contracts), "claims_total": len(claims),
            "errors_total": len(a["errors"]), "warnings_total": len(a["warnings"]),
            "errors": a["errors"][:500], "warnings": a["warnings"][:500],
            "can_import": not (a["template"] and a["errors"]) and bool(contracts),
            "overview": overview(contracts, claims, names), "rows": show}


class RowErrors(ValueError):
    """Файл-шаблон с ошибками в строках: загрузка не делается, список ошибок — в .errors."""

    def __init__(self, message: str, errors: list):
        super().__init__(message)
        self.errors = errors


def import_file(path, file_name: str, sheet: Optional[str] = None) -> dict:
    """Полный цикл без сервера: файл → сопоставление → проверка → база. Ошибки — ValueError по-русски."""
    a = analyze_file(path, sheet)
    if not a["contracts"]:
        raise ValueError("В файле нет строк с договорами после строки заголовков.")
    if a["template"] and a["errors"]:
        raise RowErrors(f"В файле {len(a['errors'])} ошибок в строках — загрузка не сделана. "
                        "Исправьте ячейки и загрузите снова.", a["errors"])
    ref = a["ref"]
    reviews = []
    for c in a["contracts"]:
        r = review_row(ref, c["p"])
        r["row_no"] = c["row_no"]
        r["raw"] = c["raw"]
        r["issues"] = c["issues"]
        if c["issues"] and not a["template"]:
            msg = "; ".join(f"{i['column']}: {i['message']}" for i in c["issues"])
            r["note"] = (r["note"] + "; " if r.get("note") else "") + msg
        r["contract"] = {k: v for k, v in c["p"].items() if k not in ("policyholder", "product_text")}
        reviews.append(r)
    mapping = {"names": a["mapping"]}
    with db.tx() as con:
        bid = save_batch(con, file_name, mapping, reviews, a["claims"])
        ov = batch_overview(con, bid)
    return {"batch_id": bid, "file_name": file_name, "mode": a["mode"], "sheet": a["sheet"],
            "header_row": a["header_row"], "mapping": a["mapping"], "unmapped": a["unmapped"],
            "claims_sheet": a["claims_sheet"], "claims_total": len(a["claims"]),
            "errors": a["errors"], "warnings": a["warnings"],
            "summary": summarize(reviews), "overview": ov, "rows": reviews}


# ---------- шаблон ----------

def template_bytes(lang: str = "ru") -> bytes:
    """Пустой шаблон отчётности портфеля: «Портфель», «Убытки», «Инструкция», «Справочники»."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    lang = "uz" if lang == "uz" else "ru"
    names = SHEETS[lang]
    with db.tx() as con:
        lists = reference_lists(con)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = names["portfolio"]
    wc = wb.create_sheet(names["claims"])
    wh = wb.create_sheet(names["help"])
    wr = wb.create_sheet(names["ref"])
    thin = Side(style="thin", color="B7C0CC")
    bold = Font(bold=True)

    # --- «Справочники»: колонки со списками, на них ссылаются выпадающие списки ---
    ref_cols = [
        ("Код продукта", [p["code"] for p in lists["products"]]),
        ("Продукт", [p["name"] for p in lists["products"]]),
        ("Классы продукта", [(p["classes"] or "") for p in lists["products"]]),
        ("Класс", [c["code"] for c in lists["classes"]]),
        ("Название класса", [c["name"] for c in lists["classes"]]),
        ("Регион", [r["name"] for r in lists["regions"]]),
        ("Тип объекта", lists["objects"]),
        ("Причина / риск", lists["causes"]),
        ("Тип страхователя", list(HOLDER_TYPES)),
        ("Валюта", list(CURRENCIES)),
        ("Статус договора", list(STATUSES)),
        ("Статус убытка", list(CLAIM_STATUSES)),
    ]
    ranges = {}
    for j, (title, vals) in enumerate(ref_cols, start=1):
        L = get_column_letter(j)
        c = wr.cell(row=1, column=j, value=title)
        c.font = bold
        c.fill = PatternFill("solid", fgColor=FILL_OPT)
        for i, v in enumerate(vals, start=2):
            wr.cell(row=i, column=j, value=v).number_format = "@"
        wr.column_dimensions[L].width = 44 if title == "Продукт" else 18 if "Название" not in title else 36
        ranges[title] = f"'{names['ref']}'!${L}$2:${L}${max(len(vals), 1) + 1}"
    wr.freeze_panes = "A2"
    list_ref = {"product": "Код продукта", "class": "Класс", "region": "Регион", "object": "Тип объекта",
                "holder": "Тип страхователя", "currency": "Валюта", "status": "Статус договора",
                "claim_status": "Статус убытка", "cause": "Причина / риск"}
    soft = {"object", "cause"}       # здесь можно вписать своё — список только подсказывает

    def sheet(wsx, cols):
        for j, (f, ru, uz, req, kind, width, hint, _ex) in enumerate(cols, start=1):
            L = get_column_letter(j)
            title = (uz if lang == "uz" else ru) + (" *" if req else "")
            c = wsx.cell(row=1, column=j, value=title)
            c.font = bold
            c.fill = PatternFill("solid", fgColor=FILL_REQ if req else FILL_OPT)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = Border(bottom=thin, right=thin)
            c.comment = Comment(hint, "INSON")
            wsx.column_dimensions[L].width = width
            rng = f"{L}2:{L}{MAX_ROWS + 1}"
            dv = None
            if kind in list_ref:
                dv = DataValidation(type="list", formula1="=" + ranges[list_ref[kind]], allow_blank=True,
                                    errorStyle="warning" if kind in soft else "stop")
                dv.error = "Значение не из списка на листе «Справочники»."
            elif kind == "date":
                dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                                    formula2="DATE(2100,12,31)", allow_blank=True)
                dv.error = "Нужна дата ДД.ММ.ГГГГ."
            elif kind == "money":
                dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
                dv.error = "Нужно число не меньше нуля, без пробелов и букв."
            elif kind == "rate":
                dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True)
                dv.error = "Ставка в процентах: от 0 до 100."
            elif kind == "seismic":
                dv = DataValidation(type="whole", operator="between", formula1="6", formula2="10", allow_blank=True)
                dv.error = "Целое число от 6 до 10."
            elif kind == "int":
                dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
                dv.error = "Целое число не меньше нуля."
            elif kind == "inn":
                dv = DataValidation(type="textLength", operator="equal", formula1="9", allow_blank=True)
                dv.error = "ИНН юрлица — 9 цифр. У физлица оставьте пустым."
            if dv is not None:
                dv.showErrorMessage = True
                dv.errorTitle = "Проверьте значение"
                dv.prompt = hint
                dv.showInputMessage = True
                wsx.add_data_validation(dv)
                dv.add(rng)
            fmt = {"date": "DD.MM.YYYY", "money": "#,##0.00", "rate": "0.000", "seismic": "0", "int": "0",
                   "inn": "@", "text": "@", "product": "@", "class": "@"}.get(kind)
            if fmt:
                for i in range(2, 202):      # формат на первые 200 строк: дальше Excel продолжит сам
                    wsx.cell(row=i, column=j).number_format = fmt
        wsx.row_dimensions[1].height = 32
        wsx.freeze_panes = "B2"
        wsx.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    sheet(ws, PORTFOLIO_COLUMNS)
    sheet(wc, CLAIM_COLUMNS)

    # --- «Инструкция» ---
    wh.column_dimensions["A"].width = 26
    wh.column_dimensions["B"].width = 14
    wh.column_dimensions["C"].width = 70
    wh.column_dimensions["D"].width = 22
    lines = [
        ("Отчётность портфеля INSON — как заполнять", None),
        ("1. Одна строка листа «%s» — один договор. Одна строка листа «%s» — один убыток." % (names["portfolio"], names["claims"]), None),
        ("2. Колонки с тёплой заливкой и звёздочкой (*) — обязательные. Голубые — по возможности.", None),
        ("3. Даты — ДД.ММ.ГГГГ. Суммы — числом, без пробелов и слова «сум», в валюте договора.", None),
        ("4. Ставка — в процентах годовых: 0,3 означает 0,3 %.", None),
        ("5. Страхователь: только тип (юрлицо / физлицо) и ИНН юрлица. ФИО, паспорт и ПИНФЛ не вносить — "
         "это персональные данные; ИНН физлица система отклонит.", None),
        ("6. Выпадающие списки берут значения с листа «%s»: продукты, классы и регионы — из базы сюрвейера." % names["ref"], None),
        ("7. Заголовки не переименовывайте и строки над ними не добавляйте. Колонки можно менять местами.", None),
        ("8. Загрузка: админка → «Рынок и данные» → «Портфель» → «Загрузить заполненный». Сначала покажем "
         "предпросмотр с ошибками по строкам; сохранить можно, когда ошибок нет (предупреждения не мешают).", None),
        ("9. Убытки: если заполнен лист «%s», сводка берёт убытки оттуда; иначе — из колонок «Число убытков», "
         "«Заявлено», «Выплачено» на листе «%s»." % (names["claims"], names["portfolio"]), None),
        ("", None),
    ]
    for i, (t, _) in enumerate(lines, start=1):
        c = wh.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=14)
        wh.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i > 1 and t:
            wh.row_dimensions[i].height = 30
    r0 = len(lines) + 1

    def help_table(r, title, cols):
        wh.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
        r += 1
        for j, h in enumerate(["Колонка", "Обязательная", "Что вписать", "Пример"], start=1):
            c = wh.cell(row=r, column=j, value=h)
            c.font = bold
            c.fill = PatternFill("solid", fgColor=FILL_OPT)
        for f, ru, uz, req, kind, _w, hint, ex in cols:
            r += 1
            vals = [f"{ru} / {uz}", "да" if req else "нет", hint, ex]
            for j, v in enumerate(vals, start=1):
                c = wh.cell(row=r, column=j, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if req:
                    c.fill = PatternFill("solid", fgColor=FILL_REQ)
        return r + 2
    r0 = help_table(r0, f"Лист «{names['portfolio']}»", PORTFOLIO_COLUMNS)
    help_table(r0, f"Лист «{names['claims']}»", CLAIM_COLUMNS)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- API ----------

def _save_upload(file: UploadFile, folder: Path) -> Path:
    name = Path(file.filename or "выгрузка.xlsx").name
    if not name.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Нужен файл Excel (.xlsx). Старый формат .xls сохраните как .xlsx.")
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{name}"
    with dest.open("wb") as f:
        shutil.copyfileobj(file.file, f)
    if dest.read_bytes()[:2] != b"PK":
        dest.unlink(missing_ok=True)
        raise HTTPException(400, "Файл не похож на xlsx (нет сигнатуры PK) — возможно, это .xls или html.")
    return dest


@router.get("/portfolio/template.xlsx")
def portfolio_template(lang: str = "ru"):
    fn = "portfolio_template_uz.xlsx" if lang == "uz" else "portfolio_template.xlsx"
    return Response(template_bytes(lang),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fn}"})


@router.post("/portfolio/preview")
async def portfolio_preview(file: UploadFile = File(...), sheet: Optional[str] = None):
    """Разбор без сохранения: ошибки по строкам, предупреждения и сводка портфеля."""
    tmp = Path(tempfile.mkdtemp(prefix="portfolio-preview-"))
    try:
        dest = _save_upload(file, tmp)
        try:
            a = analyze_file(dest, sheet)
        except ValueError as e:
            raise HTTPException(400, str(e))
        out = _preview_out(a)
        out["file_name"] = Path(file.filename or "").name
        return out
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@router.post("/portfolio/import")
async def portfolio_import(file: UploadFile = File(...), sheet: Optional[str] = None):
    name = Path(file.filename or "выгрузка.xlsx").name
    dest = _save_upload(file, UPLOAD_DIR)
    try:
        out = import_file(dest, name, sheet)
    except RowErrors as e:
        dest.unlink(missing_ok=True)
        raise HTTPException(400, {"message": str(e), "errors": e.errors[:500]})
    except ValueError as e:
        raise HTTPException(400, str(e))
    out["stored"] = db.stored_path(dest)
    for r in out["rows"]:
        r.pop("contract", None)
    return out


@router.get("/portfolio/batches")
def portfolio_batches(limit: int = 50):
    with db.tx() as con:
        out = db.rows(con, """SELECT b.*, (SELECT COUNT(*) FROM portfolio_claims c WHERE c.batch_id = b.id) claims_rows
                              FROM portfolio_batches b ORDER BY b.id DESC LIMIT ?""", limit)
    for b in out:
        b["mapping"] = json.loads(b["mapping"] or "{}")
    return out


def _batch(con, bid: int) -> dict:
    b = db.rows(con, "SELECT * FROM portfolio_batches WHERE id=?", bid)
    if not b:
        raise HTTPException(404, "Загрузка не найдена")
    b[0]["mapping"] = json.loads(b[0]["mapping"] or "{}")
    return b[0]


@router.get("/portfolio/batches/{bid}")
def portfolio_batch(bid: int, verdict: Optional[str] = None):
    with db.tx() as con:
        b = _batch(con, bid)
        return {"batch": b, "rows": batch_rows(con, bid, verdict)}


@router.get("/portfolio/batches/{bid}/summary")
def portfolio_summary(bid: int):
    with db.tx() as con:
        b = _batch(con, bid)
        return {"batch": b, "summary": summarize(batch_rows(con, bid))}


@router.get("/portfolio/batches/{bid}/overview")
def portfolio_batch_overview(bid: int):
    with db.tx() as con:
        b = _batch(con, bid)
        return {"batch": b, "overview": batch_overview(con, bid)}


@router.get("/portfolio/overview")
def portfolio_overview():
    """Сводка портфеля по последней загрузке."""
    with db.tx() as con:
        last = db.rows(con, "SELECT id FROM portfolio_batches ORDER BY id DESC LIMIT 1")
        if not last:
            return {"batch": None, "overview": None, "reason": "загрузок портфеля ещё не было"}
        return {"batch": _batch(con, last[0]["id"]), "overview": batch_overview(con, last[0]["id"])}


@router.get("/portfolio", response_class=HTMLResponse)
def portfolio_page(embed: int = 0):
    return _page(web.read_text(ROOT / "app" / "portfolio.html"), "/portfolio", embed)


def _page(html: str, active: str, embed: int) -> str:
    """Общая раскладка из app/main.py; импорт отложенный — main.py сам подключает этот модуль."""
    from .ui.pages import page
    return page(html, active, bool(embed))
