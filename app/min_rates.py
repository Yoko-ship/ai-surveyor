"""
Минимальные ставки страховщика (задание заказчика 01.10.2026: «минимальную ставку пусть назначает сам страховщик»).

Где хранится. Таблица min_rates (та же, из которой минимум читают калькулятор и акт, app/db.load_reference): у каждой
правки администратора — своя версия тарифа (tariff_versions, уровень «компания», своя дата начала действия) и одна
строка min_rates; сведения о правке (тип ставки, кто, когда, примечание, откуда — админка или Excel) — в таблице
min_rate_versions по id версии. Старые версии не удаляются и не меняются (правило проекта № 9): акт берёт минимум,
действующий на свою дату; сохранённые акты хранят минимум и его источник в своих данных.

Какая версия действует. На дату D для продукта действует строка версии с наибольшей датой начала (≤ D, и
effective_to пусто или ≥ D); при равных датах — более поздняя версия (больший id). Так же выбирает
db._load_reference: правка администратора заменяет ставку тарифной политики (приказ 54-П) с даты начала, а не
складывается с ней по максимуму. Поставляемые значения тарифной политики остаются версией «из тарифной политики,
приказ 54-П» и видны в истории.

Тип ставки (rate_type): annual — годовая (по умолчанию: премия = сумма × ставка × дни / 365); fixed — на весь срок
(премия = сумма × ставка, без деления на срок). У строк тарифной политики тип annual; администратор меняет тип новой
версией. Правка задаёт одну ставку для всех типов клиента (payer_type = NULL).

Правка — только с сегодняшней или будущей даты: прошлые расчёты должны воспроизводиться (правило № 9).
У обязательных видов (pricing_mode «нормативный акт») ставку устанавливает нормативный акт — правка не принимается.

Модуль: таблица, выбор версии на дату, история, правка, разбор Excel (предпросмотр и применение) и маршруты
  PUT  /reference/min-rates/{product_code}      — новая версия (администратор)
  GET  /reference/min-rates                     — продукты с действующей минимальной ставкой и типом
  GET  /reference/min-rates/{product_code}/history
  GET  /reference/min-rates/template.xlsx       — шаблон импорта
  POST /reference/min-rates/import?apply=0|1    — импорт из Excel: apply=0 — предпросмотр, 1 — записать (администратор)
"""
import io
import re
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Body, Depends, File, Query, Request, UploadFile
from fastapi.responses import JSONResponse, Response

from . import db
from .auth import optional_user, require

ADMIN = "админ"
RATE_TYPES = ("annual", "fixed")
RATE_TYPE_WORDS = {"annual": "annual", "годовая": "annual", "годовой": "annual", "год": "annual",
                   "yillik": "annual", "fixed": "fixed", "фиксированная": "fixed", "фиксированный": "fixed",
                   "на весь срок": "fixed", "за весь срок": "fixed", "qatʼiy": "fixed", "qat'iy": "fixed"}
RATE_TYPE_LABELS = {"annual": {"ru": "годовая", "uz": "yillik", "en": "annual"},
                    "fixed": {"ru": "фиксированная (на весь срок)", "uz": "qatʼiy (butun muddatga)",
                              "en": "fixed (for the whole term)"}}
SRC_POLICY, SRC_ADMIN, SRC_IMPORT = "policy", "admin", "import"
POLICY_LEVEL = "компания"
MAX_NOTE = 300
MAX_PCT = 100.0
FILE_MAX = 2 * 1024 * 1024
MAX_ROWS = 2000
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_HEAD = ["Код продукта", "Минимальная ставка, %", "Тип ставки (annual / fixed)",
                 "Дата начала (ДД.ММ.ГГГГ)", "Примечание"]

SCHEMA_SQL = [
    """CREATE TABLE IF NOT EXISTS min_rate_versions (
        tariff_version_id INTEGER PRIMARY KEY REFERENCES tariff_versions(id),
        product_code      TEXT NOT NULL,
        min_rate_pct      REAL NOT NULL,
        rate_type         TEXT NOT NULL DEFAULT 'annual',   -- annual | fixed
        source            TEXT NOT NULL,                    -- admin | import
        created_at        TEXT NOT NULL,
        created_by        TEXT,
        note              TEXT
    )""",
    "CREATE INDEX IF NOT EXISTS idx_min_rate_versions_product ON min_rate_versions (product_code)",
]

router = APIRouter()


def ensure(con) -> None:
    for sql in SCHEMA_SQL:
        con.execute(sql)


# --------------------------------------------------------------------------- #
#  Чистые функции
# --------------------------------------------------------------------------- #

def parse_date(v) -> Optional[date]:
    """«2026-10-02», «02.10.2026», date/datetime → date; иначе None."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_pct(v) -> Optional[float]:
    """«0,08», «0.08 %», 0.08 → 0.08; иначе None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").replace("%", "").replace(" ", "").replace(" ", "").replace(",", ".").strip()
    try:
        return float(s) if s else None
    except ValueError:
        return None


def parse_rate_type(v) -> Optional[str]:
    """annual | fixed по слову (ru/uz/en); пусто — annual; непонятное — None."""
    s = str(v or "").strip().lower()
    if not s:
        return "annual"
    return RATE_TYPE_WORDS.get(s)


def check_edit(pct, rate_type, effective_from, today: Optional[date] = None, note: str = "") -> tuple:
    """(чистые значения, ошибки {поле: текст}) для одной правки."""
    today = today or date.today()
    errs, out = {}, {}
    p = parse_pct(pct)
    if p is None or not (0 < p <= MAX_PCT):
        errs["min_rate_pct"] = f"число больше 0 и не больше {MAX_PCT:g} (%)"
    else:
        out["min_rate_pct"] = round(p, 6)
    rt = parse_rate_type(rate_type)
    if rt is None:
        errs["rate_type"] = "annual (годовая) или fixed (на весь срок)"
    else:
        out["rate_type"] = rt
    d = parse_date(effective_from) if effective_from not in (None, "") else today
    if d is None:
        errs["effective_from"] = "дата ГГГГ-ММ-ДД или ДД.ММ.ГГГГ"
    elif d < today:
        errs["effective_from"] = ("не раньше сегодняшней даты: прошлые расчёты и акты должны воспроизводиться "
                                  "(правило проекта № 9)")
    elif d.year > today.year + 10:
        errs["effective_from"] = "не дальше чем на 10 лет вперёд"
    else:
        out["effective_from"] = d.isoformat()
    n = str(note or "").strip()
    if len(n) > MAX_NOTE:
        errs["note"] = f"не длиннее {MAX_NOTE} знаков"
    out["note"] = n[:MAX_NOTE]
    return out, errs


def premium(rate_pct: float, sum_insured: float, term_days: int, rate_type: str = "annual") -> float:
    """Премия по типу ставки: annual — сумма × ставка × дни / 365; fixed — сумма × ставка (без деления на срок).
    100 млн, 0,5 %, 1 095 дней: annual — 1 500 000; fixed — 500 000."""
    if rate_type == "fixed":
        return rate_pct / 100 * sum_insured
    return rate_pct / 100 * sum_insured * term_days / 365


def annual_equivalent(rate_pct: Optional[float], term_days: int, rate_type: str = "annual") -> Optional[float]:
    """Годовой эквивалент ставки для сравнения с рынком: fixed × 365 / дни; annual — как есть."""
    if rate_pct is None:
        return None
    if rate_type == "fixed":
        return round(float(rate_pct) * 365 / max(int(term_days or 365), 1), 6)
    return float(rate_pct)


def pick(rows: list, as_of: str, payer_type: Optional[str] = None) -> Optional[dict]:
    """Строка, действующая на дату as_of (ISO): наибольшая дата начала, затем больший id версии; внутри версии —
    ставка типа клиента, иначе общая, иначе наибольшая (как engine._pick)."""
    act = [r for r in rows if str(r["effective_from"]) <= as_of
           and (not r.get("effective_to") or str(r["effective_to"]) >= as_of)]
    if not act:
        return None
    best = max((str(r["effective_from"]), int(r["version_id"])) for r in act)
    same = [r for r in act if (str(r["effective_from"]), int(r["version_id"])) == best]
    if payer_type is not None:
        hit = [r for r in same if r.get("payer_type") == payer_type]
        if hit:
            return hit[0]
    hit = [r for r in same if r.get("payer_type") is None]
    if hit:
        return hit[0]
    return max(same, key=lambda r: float(r["min_rate_pct"]))


def source_label(row: Optional[dict], lang: str = "ru") -> Optional[str]:
    """Источник минимума словами: «минимальная ставка страховщика (установлена администратором 02.10.2026,
    примечание …)» или «из тарифной политики, приказ 54-П»."""
    if not row:
        return None
    src = row.get("source") or SRC_POLICY
    d = _ddmmyyyy(row.get("effective_from"))
    note = (row.get("note") or "").strip()
    if src in (SRC_ADMIN, SRC_IMPORT):
        how = {"ru": {"admin": "установлена администратором", "import": "загружена из Excel"},
               "uz": {"admin": "administrator tomonidan belgilangan", "import": "Excel'dan yuklangan"},
               "en": {"admin": "set by the administrator", "import": "imported from Excel"}}[_lang(lang)][src]
        head = {"ru": "минимальная ставка страховщика", "uz": "sugʻurtalovchining eng kam stavkasi",
                "en": "the insurer's minimum rate"}[_lang(lang)]
        frm = {"ru": "действует с", "uz": "amal qilish sanasi", "en": "effective from"}[_lang(lang)]
        nt = {"ru": "примечание", "uz": "izoh", "en": "note"}[_lang(lang)]
        inner = f"{how}, {frm} {d}" + (f", {nt}: {note}" if note else "")
        return f"{head} ({inner})"
    order = _order_of(row.get("document_ref") or row.get("name") or "")
    if order and _lang(lang) != "ru":
        order = "".join(_LAT.get(ch.upper(), ch) for ch in order)
    if _lang(lang) == "uz":
        return "tarif siyosatidan" + (f", {order}-son buyruq" if order else "")
    if _lang(lang) == "en":
        return "from the tariff policy" + (f", Order No. {order}" if order else "")
    return "из тарифной политики" + (f", приказ {order}" if order else "")


def _lang(lang) -> str:
    s = str(lang or "ru").lower()[:2]
    return s if s in ("ru", "uz", "en") else "ru"


_LAT = {"А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Е": "E", "К": "K", "Л": "L", "М": "M", "Н": "N", "О": "O",
        "П": "P", "Р": "R", "С": "S", "Т": "T", "У": "U", "Ф": "F", "Х": "X"}


def _order_of(ref: str) -> Optional[str]:
    m = re.search(r"№\s*([0-9A-Za-zА-Яа-яЁё\-/]+)", str(ref or ""))
    return m.group(1) if m else None


def _ddmmyyyy(iso) -> str:
    d = parse_date(iso)
    return d.strftime("%d.%m.%Y") if d else str(iso or "")


# --------------------------------------------------------------------------- #
#  База
# --------------------------------------------------------------------------- #

def _rows(con, product_code: str) -> list:
    """Все строки минимальной ставки компании по продукту — с версией и сведениями о правке."""
    ensure(con)
    out = []
    for r in db.rows(con, """SELECT m.id AS row_id, m.min_rate_pct, m.payer_type, m.class_code,
                                    v.id AS version_id, v.name, v.document_ref, v.effective_from, v.effective_to,
                                    x.rate_type, x.source, x.created_at, x.created_by, x.note
                             FROM min_rates m JOIN tariff_versions v ON v.id = m.tariff_version_id
                             LEFT JOIN min_rate_versions x ON x.tariff_version_id = v.id
                             WHERE m.product_code = ? AND v.level = ?
                             ORDER BY v.effective_from, v.id""", product_code, POLICY_LEVEL):
        r["rate_type"] = r.get("rate_type") or "annual"
        r["source"] = r.get("source") or SRC_POLICY
        out.append(r)
    return out


def on_date(con, product_code: Optional[str], as_of=None, payer_type: Optional[str] = None) -> Optional[dict]:
    """Минимальная ставка страховщика по продукту на дату (по умолчанию — сегодня) с источником и типом ставки."""
    if not product_code:
        return None
    d = parse_date(as_of) if as_of else date.today()
    row = pick(_rows(con, product_code), (d or date.today()).isoformat(), payer_type)
    if not row:
        return None
    return {"pct": row["min_rate_pct"], "rate_type": row["rate_type"], "source": row["source"],
            "version_id": row["version_id"], "effective_from": row["effective_from"],
            "created_at": row.get("created_at"), "note": row.get("note"),
            "document_ref": row.get("document_ref"), "name": row.get("name"), "payer_type": row.get("payer_type"),
            "label_ru": source_label(row, "ru")}


def rate_type_on(con, product_code: Optional[str], as_of=None) -> str:
    """Тип ставки продукта на дату: из действующей версии минимума; нет версии — annual."""
    try:
        return (on_date(con, product_code, as_of) or {}).get("rate_type") or "annual"
    except Exception:
        return "annual"


def history(con, product_code: str, admin: bool = True) -> list:
    """Все версии минимальной ставки продукта (старые не удаляются): от ранней к поздней, с пометкой действующей."""
    rows = _rows(con, product_code)
    today = date.today().isoformat()
    cur = pick(rows, today)
    out = []
    for r in rows:
        out.append({"version_id": r["version_id"], "min_rate_pct": r["min_rate_pct"], "payer_type": r["payer_type"],
                    "rate_type": r["rate_type"], "effective_from": r["effective_from"],
                    "effective_to": r["effective_to"], "source": r["source"],
                    "source_label": source_label(r, "ru"), "document_ref": r["document_ref"],
                    "created_at": r.get("created_at"),
                    "created_by": (r.get("created_by") if admin else ("администратор" if r.get("created_by") else None)),
                    "note": r.get("note"),
                    "active": bool(cur) and cur["version_id"] == r["version_id"]
                    and cur.get("payer_type") == r.get("payer_type"),
                    "future": str(r["effective_from"]) > today})
    return out


def products(con) -> list:
    """Продукты с действующей минимальной ставкой страховщика, типом и датой (для админки)."""
    ensure(con)
    out = []
    for p in db.rows(con, "SELECT code, name, pricing_mode, rate_text FROM products ORDER BY code"):
        cur = on_date(con, p["code"])
        rows = _rows(con, p["code"])
        future = [r for r in rows if str(r["effective_from"]) > date.today().isoformat()]
        out.append({"code": p["code"], "name": p["name"], "pricing_mode": p["pricing_mode"],
                    "rate_text": p["rate_text"], "editable": p["pricing_mode"] != "нормативный акт",
                    "min_rate_pct": (cur or {}).get("pct"), "rate_type": (cur or {}).get("rate_type") or "annual",
                    "effective_from": (cur or {}).get("effective_from"), "source": (cur or {}).get("source"),
                    "source_label": (cur or {}).get("label_ru"), "note": (cur or {}).get("note"),
                    "versions": len(rows),
                    "next": ({"min_rate_pct": future[-1]["min_rate_pct"], "rate_type": future[-1]["rate_type"],
                              "effective_from": future[-1]["effective_from"]} if future else None)})
    return out


def save(con, product_code: str, clean: dict, who: str, source: str = SRC_ADMIN) -> dict:
    """Новая версия минимальной ставки: версия тарифа + строка min_rates + сведения о правке. Ничего не удаляет."""
    ensure(con)
    pr = db.rows(con, "SELECT code, pricing_mode FROM products WHERE code=?", product_code)
    if not pr:
        raise ValueError("продукт не найден в справочнике")
    if pr[0]["pricing_mode"] == "нормативный акт":
        raise ValueError("обязательный вид: ставку устанавливает нормативный акт, минимум страховщика не задаётся")
    cls = db.rows(con, "SELECT class_code FROM product_classes WHERE product_code=? ORDER BY part_no", product_code)
    cls_code = cls[0]["class_code"] if cls else None
    what = "правка администратора" if source == SRC_ADMIN else "импорт из Excel"
    cur = con.execute("INSERT INTO tariff_versions (level, name, document_ref, effective_from) VALUES (?,?,?,?)",
                      (POLICY_LEVEL, f"Минимальная ставка страховщика — продукт {product_code}",
                       what, clean["effective_from"]))
    vid = cur.lastrowid
    con.execute("INSERT INTO min_rates (tariff_version_id, product_code, class_code, payer_type, min_rate_pct) "
                "VALUES (?,?,?,?,?)", (vid, product_code, cls_code, None, clean["min_rate_pct"]))
    con.execute("INSERT INTO min_rate_versions (tariff_version_id, product_code, min_rate_pct, rate_type, source, "
                "created_at, created_by, note) VALUES (?,?,?,?,?,?,?,?)",
                (vid, product_code, clean["min_rate_pct"], clean["rate_type"], source, db.now(), who,
                 clean.get("note") or ""))
    # журнал: «новая версия тарифов» version:<id> и «минимальная ставка» — признаки ручной правки для app/refsync.py
    # (обновление справочников из образа их не затирает)
    db.audit(con, who, "новая версия тарифов", f"version:{vid}",
             {"level": POLICY_LEVEL, "name": f"Минимальная ставка страховщика — продукт {product_code}",
              "effective_from": clean["effective_from"], "source": source})
    db.audit(con, who, "минимальная ставка", product_code,
             {"tariff_version_id": vid, "product_code": product_code, "class_code": cls_code, "payer_type": None,
              "min_rate_pct": clean["min_rate_pct"], "rate_type": clean["rate_type"],
              "effective_from": clean["effective_from"], "note": clean.get("note") or "", "source": source})
    db.reference_changed(con)          # калькулятор и акт сразу видят новую версию
    return {"version_id": vid, "product_code": product_code, **clean, "source": source}


# --------------------------------------------------------------------------- #
#  Excel
# --------------------------------------------------------------------------- #

def template_bytes() -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Минимальные ставки"
    ws.append(TEMPLATE_HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3F5")
    ws.append(["0832", 0.08, "annual", date.today().strftime("%d.%m.%Y"), "пример — удалите строку"])
    for col, w in zip("ABCDE", (16, 24, 28, 26, 50)):
        ws.column_dimensions[col].width = w
    ins = wb.create_sheet("Инструкция")
    ins.column_dimensions["A"].width = 110
    for line in ("Импорт минимальных ставок страховщика",
                 "",
                 "1. Одна строка — один продукт. Код продукта — как в справочнике (например, 0832).",
                 "2. Минимальная ставка — в процентах: 0,08 или 0.08 (не 0,0008).",
                 "3. Тип ставки: annual (годовая, премия = сумма × ставка × дни / 365) или fixed (на весь срок,",
                 "   премия = сумма × ставка). Пусто — annual.",
                 "4. Дата начала — не раньше сегодняшней: прошлые расчёты должны воспроизводиться.",
                 "5. Каждая строка — новая версия; прежние версии сохраняются и видны в истории.",
                 "6. Сначала загрузка показывает предпросмотр (было → станет, ошибки по строкам);",
                 "   записывается только файл без ошибок."):
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header_map(row) -> Optional[dict]:
    names = [str(c or "").strip().lower() for c in row]
    out = {}
    for i, n in enumerate(names):
        if "код" in n or n in ("product", "product_code", "code"):
            out.setdefault("code", i)
        elif "ставк" in n and "тип" not in n or n in ("min_rate_pct", "rate"):
            out.setdefault("pct", i)
        elif "тип" in n or n in ("rate_type", "type"):
            out.setdefault("type", i)
        elif "дата" in n or n in ("effective_from", "date"):
            out.setdefault("date", i)
        elif "примеч" in n or n in ("note",):
            out.setdefault("note", i)
    return out if {"code", "pct"} <= set(out) else None


def parse_workbook(data: bytes) -> dict:
    """Строки файла → [{"row", "product_code", "min_rate_pct", "rate_type", "effective_from", "note"}] сырыми
    значениями; ошибки чтения файла — {"error"}."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        return {"error": "файл не читается как Excel (.xlsx): " + type(e).__name__, "rows": []}
    ws = wb.worksheets[0]
    head, rows = None, []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if head is None:
            head = _header_map(r)
            if head is None and i >= 5:
                return {"error": "не найдена строка заголовка: «Код продукта», «Минимальная ставка, %», …", "rows": []}
            continue
        if not any(c not in (None, "") for c in r):
            continue
        if len(rows) >= MAX_ROWS:
            return {"error": f"строк больше {MAX_ROWS}", "rows": []}
        g = lambda k: r[head[k]] if k in head and head[k] < len(r) else None   # noqa: E731
        code = g("code")
        if isinstance(code, float) and code == int(code):
            code = str(int(code)).zfill(4)
        elif isinstance(code, int):
            code = str(code).zfill(4)
        rows.append({"row": i, "product_code": str(code or "").strip()[:10], "min_rate_pct": g("pct"),
                     "rate_type": g("type"), "effective_from": g("date"), "note": g("note")})
    wb.close()
    if head is None:
        return {"error": "файл пустой", "rows": []}
    return {"error": None, "rows": rows}


def preview(con, raw_rows: list, today: Optional[date] = None) -> dict:
    """Проверка строк импорта: ошибки по строкам и «было → станет». Ничего не пишет."""
    items, seen, n_err = [], set(), 0
    for r in raw_rows:
        clean, errs = check_edit(r.get("min_rate_pct"), r.get("rate_type"), r.get("effective_from"), today,
                                 r.get("note") or "")
        code = r.get("product_code") or ""
        pr = db.rows(con, "SELECT code, name, pricing_mode FROM products WHERE code=?", code) if code else []
        if not code:
            errs["product_code"] = "пустой код продукта"
        elif not pr:
            errs["product_code"] = "продукт не найден в справочнике"
        elif pr[0]["pricing_mode"] == "нормативный акт":
            errs["product_code"] = "обязательный вид: ставку устанавливает нормативный акт"
        key = (code, clean.get("effective_from"))
        if key in seen:
            errs["product_code"] = "повтор: тот же продукт и та же дата начала уже есть в файле"
        seen.add(key)
        cur = on_date(con, code) if pr else None
        n_err += 1 if errs else 0
        items.append({"row": r.get("row"), "product_code": code, "product_name": pr[0]["name"] if pr else None,
                      "now": {"min_rate_pct": (cur or {}).get("pct"), "rate_type": (cur or {}).get("rate_type"),
                              "source_label": (cur or {}).get("label_ru")},
                      "new": clean if not errs else None, "errors": errs, "ok": not errs})
    return {"rows": len(items), "errors": n_err, "items": items, "can_apply": bool(items) and n_err == 0}


# --------------------------------------------------------------------------- #
#  Маршруты
# --------------------------------------------------------------------------- #

def _who(user: dict) -> str:
    return (user or {}).get("login") or ADMIN


@router.get("/reference/min-rates")
def list_min_rates():
    """Продукты с действующей минимальной ставкой страховщика: ставка, тип, дата, источник."""
    with db.tx() as con:
        return {"ok": True, "items": products(con), "rate_types": list(RATE_TYPES)}


@router.get("/reference/min-rates/template.xlsx")
def min_rates_template():
    return Response(template_bytes(), media_type=XLSX_MIME,
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''min_rates_template.xlsx"})


@router.get("/reference/min-rates/{product_code}/history")
def min_rates_history(request: Request, product_code: str):
    """Все версии минимальной ставки продукта (старые сохраняются); кто правил — видит только администратор."""
    admin = (optional_user(request) or {}).get("role") == ADMIN
    with db.tx() as con:
        if not db.rows(con, "SELECT code FROM products WHERE code=?", product_code):
            return JSONResponse({"ok": False, "detail": "продукт не найден"}, status_code=404)
        cur = on_date(con, product_code)
        return {"ok": True, "product_code": product_code, "current": cur,
                "history": history(con, product_code, admin)}


@router.put("/reference/min-rates/{product_code}")
def min_rates_put(product_code: str, body: dict = Body(...), user: dict = Depends(require(ADMIN))):
    """Новая версия минимальной ставки страховщика (администратор). Поля: min_rate_pct, rate_type (annual | fixed),
    effective_from (ГГГГ-ММ-ДД, не раньше сегодня; пусто — сегодня), note. min_rate_pct можно не передавать —
    тогда берётся действующая ставка (например, меняется только тип ставки)."""
    body = body if isinstance(body, dict) else {}
    with db.tx() as con:
        ensure(con)
        if not db.rows(con, "SELECT code FROM products WHERE code=?", product_code):
            return JSONResponse({"ok": False, "detail": "продукт не найден"}, status_code=404)
        cur = on_date(con, product_code)
        pct = body.get("min_rate_pct")
        if pct in (None, "") and cur:
            pct = cur["pct"]
        rt = body.get("rate_type")
        if rt in (None, "") and cur:
            rt = cur["rate_type"]
        clean, errs = check_edit(pct, rt, body.get("effective_from"), note=body.get("note") or "")
        if errs:
            return JSONResponse({"ok": False, "errors": errs}, status_code=422)
        try:
            saved = save(con, product_code, clean, _who(user), SRC_ADMIN)
        except ValueError as e:
            return JSONResponse({"ok": False, "errors": {"product_code": str(e)}}, status_code=422)
        return {"ok": True, "saved": saved, "current": on_date(con, product_code),
                "history": history(con, product_code)}


@router.post("/reference/min-rates/import")
async def min_rates_import(file: UploadFile = File(...), apply: int = Query(0),
                           user: dict = Depends(require(ADMIN))):
    """Импорт из Excel: apply=0 — предпросмотр (ничего не пишет), apply=1 — записать, если ошибок нет."""
    data = await file.read(FILE_MAX + 1)
    if len(data) > FILE_MAX:
        return JSONResponse({"ok": False, "detail": "файл больше 2 МБ"}, status_code=400)
    if data[:2] != b"PK":
        return JSONResponse({"ok": False, "detail": "нужен файл Excel (.xlsx)"}, status_code=400)
    parsed = parse_workbook(data)
    if parsed["error"]:
        return JSONResponse({"ok": False, "detail": parsed["error"]}, status_code=422)
    with db.tx() as con:
        ensure(con)
        pv = preview(con, parsed["rows"])
        if not apply:
            return {"ok": True, "applied": False, "preview": pv}
        if not pv["can_apply"]:
            return JSONResponse({"ok": False, "applied": False, "preview": pv,
                                 "detail": "в файле есть ошибки — ничего не записано"}, status_code=422)
        saved = [save(con, it["product_code"], it["new"], _who(user), SRC_IMPORT) for it in pv["items"]]
        return {"ok": True, "applied": True, "saved": saved, "preview": pv}
