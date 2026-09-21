"""
Выгрузки по запросу: анализ в PDF и в XLSX, а также отправка файла в Telegram.

    GET  /requests/{rid}/analysis.pdf     — PDF: карточка запроса, объект, стоимость с источниками
                                            и ссылками, цепочка формирования ставки, проверки,
                                            вероятность подтверждения с разбором, документы,
                                            кто рассматривает и их решения.
    GET  /requests/{rid}/analysis.xlsx    — та же выгрузка листами: «Запрос», «Ставка», «Проверки»,
                                            «Стоимость», «Документы», «Решения», «Вероятность».
    POST /requests/{rid}/analysis/send-telegram?format=pdf|xlsx|both
                                          — бот присылает файл в чат вошедшего: во встроенном
                                            браузере Telegram обычное скачивание часто не работает.

Доступ строго: участники запроса (инициатор и назначенные рассматривающие) и администраторы.
Остальным 403 — в выгрузке есть сведения о страхователе, а они составляют тайну страхования
(ЗРУ-730, ст. 62).

Почему отдельный модуль, а не продолжение app/proposal.py: там документ ДЛЯ КЛИЕНТА (что покрыто,
сколько стоит, что можно улучшить), здесь — документ ДЛЯ КОМПАНИИ (откуда цифры, что проверено,
кто решает). Разные читатели и разная ответственность, но вёрстка общая: класс _Pdf и форматирование
берутся из app/proposal.py, поэтому оба документа выглядят одинаково.

Вероятность подтверждения берётся из app/outcomes.py (считает её app/analysis.py при отправке
запроса на согласование): число, вердикт, «что снижает» и «что повысит», пометка «оценка
экспертная, не калибрована». В XLSX под это отведён отдельный лист «Вероятность».
Запрос на согласование ещё не отправляли — в документе честно написано, что числа пока нет.
"""
import io
import json
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import approvals, db, outcomes, proposal
from .auth import current_user
from .proposal import CONTENT_W, GRAY, money, pct, ru_date

router = APIRouter()

ADMIN = "админ"
FORMATS = ("pdf", "xlsx", "both")

HEAD_FILL = PatternFill("solid", fgColor="EEF2F4")
HEAD_FONT = Font(bold=True)


# --------------------------------------------------------------------------- #
#  Доступ
# --------------------------------------------------------------------------- #

def participants(con, rid: int) -> set:
    """Кто участвует в запросе: инициатор и назначенные рассматривающие."""
    ids = set(approvals.initiator_user_ids(con, rid))
    ids |= {r["user_id"] for r in db.rows(con, "SELECT user_id FROM request_reviewers WHERE request_id=?", rid)}
    return ids


def check_access(con, user: dict, rid: int):
    if not db.rows(con, "SELECT 1 FROM requests WHERE id=?", rid):
        raise HTTPException(404, "Запрос не найден")
    if user["role"] == ADMIN:
        return
    if user["id"] not in participants(con, rid):
        raise HTTPException(403, "Выгрузка доступна участникам запроса и администраторам: "
                                 "сведения о страхователе не разглашаются (ЗРУ-730, ст. 62)")


# --------------------------------------------------------------------------- #
#  Сбор данных
# --------------------------------------------------------------------------- #

def collect(con, rid: int) -> dict:
    req = db.rows(con, "SELECT * FROM requests WHERE id=?", rid)
    if not req:
        raise HTTPException(404, "Запрос не найден")
    req = req[0]
    obj = db.rows(con, "SELECT * FROM objects WHERE request_id=? ORDER BY id LIMIT 1", rid)
    calc = db.rows(con, "SELECT * FROM calculations WHERE request_id=? ORDER BY id DESC LIMIT 1", rid)
    calc = calc[0] if calc else None
    checks = db.rows(con, "SELECT rule_code, status, detail FROM check_results WHERE calculation_id=?",
                     calc["id"]) if calc else []
    rules = {r["code"]: r for r in db.rows(con, "SELECT code, name, legal_ref FROM rules")}
    val = db.rows(con, "SELECT * FROM valuations WHERE request_id=? ORDER BY id DESC LIMIT 1", rid)
    val = val[0] if val else None
    sources = db.rows(con, "SELECT * FROM valuation_sources WHERE valuation_id=? ORDER BY id",
                      val["id"]) if val else []
    docs = db.rows(con, "SELECT doc_name, received, file_path FROM documents WHERE request_id=? ORDER BY id", rid)
    files = db.rows(con, "SELECT filename, doc_kind, mime, size_bytes, uploaded_at, parse_status"
                         " FROM photos WHERE request_id=? ORDER BY id", rid)
    prod = db.rows(con, "SELECT code, name FROM products WHERE code=?", req.get("product_code") or "")
    agent = db.rows(con, "SELECT name FROM agents WHERE id=?", req["agent_id"]) if req.get("agent_id") else []
    partner = db.rows(con, "SELECT partner FROM general_agreements WHERE id=?",
                      req["general_agreement_id"]) if req.get("general_agreement_id") else []
    author = db.rows(con, "SELECT full_name, department, position FROM users WHERE id=?",
                     req["created_by_user_id"]) if req.get("created_by_user_id") else []
    expl, attrs = {}, {}
    if calc:
        try:
            expl = json.loads(calc["explanation"] or "{}")
        except ValueError:
            expl = {}
    if obj:
        try:
            attrs = json.loads(obj[0]["attributes"] or "{}")
        except ValueError:
            attrs = {}
    return {"request": req, "object": obj[0] if obj else None, "calculation": calc, "checks": checks,
            "rules": rules, "valuation": val, "sources": sources, "documents": docs, "files": files,
            "product": prod[0] if prod else None, "agent": agent[0] if agent else None,
            "partner": partner[0]["partner"] if partner else None,
            "author": author[0] if author else None,
            "reviewers": approvals.reviewers(con, rid), "probability": outcomes.summary(con, rid),
            "explanation": expl, "attributes": attrs}


def _rate_rows(data: dict) -> list:
    """Цепочка формирования ставки — та же таблица, что в предложении клиенту."""
    calc, expl = data["calculation"], data["explanation"]
    if not calc:
        return []
    return proposal._chain_rows(expl.get("chain") or [], calc, expl)


def _source_rows(data: dict) -> list:
    out = []
    for s in data["sources"]:
        out.append([s["source"], s["status"], str(s["ads_count"] or "—"),
                    money(s["median"]) if s["median"] else "—", s["url"] or "ссылки нет",
                    ru_date(s["fetched_at"] or "")])
    return out


# --------------------------------------------------------------------------- #
#  PDF
# --------------------------------------------------------------------------- #

def build_analysis_pdf(con, rid: int) -> bytes:
    if not proposal.FONT_REGULAR.exists():
        raise HTTPException(503, "На сервере нет шрифта для PDF (Arial). Пока доступна выгрузка XLSX; "
                                 "для PDF нужно положить шрифт рядом с приложением")
    data = collect(con, rid)
    req, obj, calc = data["request"], data["object"], data["calculation"]
    pdf = proposal._Pdf()
    pdf.footer = "СО АО «INSON» — Анализ запроса (для компании)"

    # 1. Карточка запроса
    pdf.para("СО АО «INSON»", size=10, bold=True, color=GRAY, gap=2)
    pdf.para("Анализ запроса", size=20, bold=True, gap=2)
    pdf.para(f"Запрос № {req.get('external_no') or rid} · сформировано "
             f"{datetime.now().strftime('%d.%m.%Y %H:%M')}", size=10, color=GRAY, gap=2)
    pdf.space(4)
    pdf.rule()
    pdf.heading("Запрос")
    author = data["author"]
    pdf.kv([
        ("Номер в системе", str(rid)),
        ("Номер договора", req.get("external_no") or "—"),
        ("Филиал", req.get("branch") or "—"),
        ("Продукт", f"{data['product']['code']} — {data['product']['name']}" if data["product"]
         else (req.get("product_code") or "—")),
        ("Страхователь", req.get("policyholder") or "—"),
        ("Подал", (author["full_name"] if author else None)
         or (data["agent"]["name"] if data["agent"] else "не указан")),
        ("Департамент, должность", f"{author['department'] or '—'}, {author['position'] or '—'}"
         if author else "—"),
        ("Генеральное соглашение", data["partner"] or "нет"),
        ("Создан", ru_date(req.get("created_at") or "")),
        ("Состояние согласования", req.get("approval_status") or "не требуется"),
    ])

    # 2. Объект
    pdf.heading("Объект")
    if obj:
        pdf.kv([
            ("Тип объекта", obj["object_type"]),
            ("Адрес, регион", ", ".join(x for x in (obj.get("region"), obj.get("address")) if x) or "—"),
            ("Страховая стоимость", money(obj.get("value_amount"))),
            ("Страховая сумма", money(obj.get("sum_insured"))),
            ("Франшиза", str(obj.get("franchise") or "не установлена")),
        ])
    else:
        pdf.para("Объект по запросу не сохранён.")

    # 3. Стоимость и источники
    pdf.heading("Стоимость объекта и откуда она взялась")
    val = data["valuation"]
    if val:
        pdf.kv([
            ("Заявлено страхователем", money(val.get("declared_value"))),
            ("Оценка системы", money(val.get("ai_value"))),
            ("Метод", f"{val.get('method') or '—'} (версия методики {val.get('method_version') or '—'})"),
            ("Подтвердил андеррайтер", val.get("confirmed_by_underwriter") or "нет"),
        ])
        if data["sources"]:
            pdf.para("Источники — у каждого ссылка, по которой цифру можно проверить:",
                     size=9.5, color=GRAY, gap=4)
            pdf.table([("Источник", 110, "l"), ("Итог", 90, "l"), ("Объявлений", 60, "r"),
                       ("Медиана", 95, "r"), ("Ссылка", CONTENT_W - 415, "l"), ("Дата", 60, "l")],
                      _source_rows(data))
        else:
            pdf.para("Источники по этой оценке не сохранены.", size=9.5, color=GRAY)
    else:
        pdf.para("Оценка стоимости по запросу не проводилась: сумма принята со слов страхователя "
                 "или из документов. Проверить цифру по объявлениям и нормам износа — раздел «Оценка».")

    # 4. Ставка
    pdf.heading("Как сложилась ставка")
    rows = _rate_rows(data)
    if rows:
        pdf.table([("Составляющая", CONTENT_W - 190, "l"), ("Действие", 95, "r"), ("Ставка", 95, "r")], rows)
        pdf.kv([
            ("Техническая ставка", pct(calc.get("gross_rate_pct"))),
            ("Минимум по политике", pct(calc.get("min_rate_pct"))),
            ("Применённая ставка", pct(calc.get("applied_rate_pct"))),
            ("Премия", money(calc.get("premium"))),
            ("Вердикт", calc.get("verdict") or "—"),
            ("Версия тарифов", str(calc.get("tariff_version_id") or "—")),
        ])
    else:
        pdf.para("Расчёт по запросу не сохранён.")

    # 5. Проверки
    pdf.heading("Проверки")
    if data["checks"]:
        rows = []
        for ch in data["checks"]:
            rule = data["rules"].get(ch["rule_code"]) or {}
            rows.append([ch["rule_code"] + (f" — {rule.get('legal_ref')}" if rule.get("legal_ref") else ""),
                         ch["status"], ch["detail"] or ""])
        pdf.table([("Правило и норма", 190, "l"), ("Итог", 70, "l"),
                   ("Что именно", CONTENT_W - 260, "l")], rows)
    else:
        pdf.para("Проверки по запросу не сохранены.")

    # 6. Вероятность подтверждения
    pdf.heading("Вероятность подтверждения")
    prob = data["probability"]
    if prob["ready"]:
        pdf.big_number("Вероятность, что запрос подтвердят", f"{prob['probability']}%",
                       f"Методика версии {prob.get('model_version') or '—'}, расчёт "
                       f"{ru_date(prob.get('sent_at') or '')} · {prob.get('verdict') or ''}")
        pdf.para(prob.get("summary") or "", size=10)
        pdf.para("Что снижает вероятность:", size=10, bold=True, gap=2)
        if prob["minus"]:
            for x in prob["minus"]:
                # у стоп-пунктов delta = −100: это не вклад в цифру, а жёсткое нарушение нормы
                tail = " — жёсткое нарушение нормы" if x.get("stop") else f" (−{abs(x['delta'])} п.п.)"
                pdf.bullet(x["text"] + tail + (f". Норма: {x['norm']}" if x.get("norm") else ""))
        else:
            pdf.bullet("снижающих обстоятельств не нашлось")
        pdf.para("Что идёт в плюс:", size=10, bold=True, gap=2)
        if prob["plus"]:
            for x in prob["plus"]:
                pdf.bullet(f"{x['text']} (+{x['delta']} п.п.)")
        else:
            pdf.bullet("плюсов по правилам не набралось")
        if prob["how_to_raise"]:
            pdf.para("Что повысит вероятность:", size=10, bold=True, gap=2)
            for x in prob["how_to_raise"]:
                pdf.bullet(x["text"])
        if (prob.get("stat") or {}).get("text"):
            pdf.para(prob["stat"]["text"], size=9.5, color=GRAY)
        if not prob.get("calibrated"):
            pdf.para("Оценка экспертная, не калибрована: веса правил статистикой компании пока "
                     "не подтверждены (правило проекта № 7).", size=9, color=GRAY)
        if prob.get("decision"):
            pdf.para(f"Фактическое решение: {prob['decision']}"
                     + (f" · {prob['decided_by']}" if prob.get("decided_by") else "")
                     + (f" · {ru_date(prob['decided_at'])}" if prob.get("decided_at") else ""),
                     size=9.5, color=GRAY)
    else:
        pdf.para(prob["text"] + ". Вероятность считается, когда запрос уходит на согласование: "
                 "тогда здесь появятся число, вердикт и разбор — что снижает и что повысит.",
                 size=10, color=GRAY)

    # 7. Документы
    pdf.heading("Документы")
    rows = [[d["doc_name"], "получен" if d["received"] else "не получен",
             d["file_path"] or ""] for d in data["documents"]]
    rows += [[f["filename"] or f["doc_kind"], f["doc_kind"],
              f"{round((f['size_bytes'] or 0) / 1024)} КБ, загружен {ru_date(f['uploaded_at'] or '')}"]
             for f in data["files"]]
    if rows:
        pdf.table([("Документ", 200, "l"), ("Состояние", 110, "l"),
                   ("Файл", CONTENT_W - 310, "l")], rows)
    else:
        pdf.para("По запросу не загружено ни одного документа.")

    # 8. Кто рассматривает
    pdf.heading("Кто рассматривает и что решил")
    if data["reviewers"]:
        pdf.table([("Рассматривающий", 190, "l"), ("Должность", 120, "l"), ("Решение", 90, "l"),
                   ("Когда и примечание", CONTENT_W - 400, "l")],
                  [[r["full_name"], r.get("position") or "—", r["status"],
                    (ru_date(r["decided_at"]) if r["decided_at"] else "ещё не решил")
                    + (f" · {r['comment']}" if r.get("comment") else "")]
                   for r in data["reviewers"]])
    else:
        pdf.para("Рассматривающие по запросу не назначены.")

    pdf.space(6)
    pdf.rule()
    pdf.para("Документ для внутреннего пользования: содержит сведения, составляющие тайну страхования "
             "(ЗРУ-730, ст. 62). Передача третьим лицам запрещена.", size=9, color=GRAY)
    return pdf.finish()


# --------------------------------------------------------------------------- #
#  XLSX
# --------------------------------------------------------------------------- #

def _sheet(wb, title: str, header: list, rows: list, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    widths = widths or [max(14, min(60, len(str(h)) + 10)) for h in header]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _probability_sheet(wb, prob: dict):
    """Отдельный лист «Вероятность»: число, вердикт, «что снижает», «что повысит», пометка о калибровке."""
    rows = []
    if prob["ready"]:
        rows += [["Вероятность подтверждения, %", prob["probability"], "", ""],
                 ["Вердикт", prob.get("verdict") or "", "", ""],
                 ["Коротко", prob.get("summary") or "", "", ""],
                 ["Версия методики", prob.get("model_version") or "", "", ""],
                 ["Рассчитано", prob.get("sent_at") or "", "", ""],
                 ["", "", "", ""]]
        for x in prob["minus"]:
            # стоп-пункт показываем словами: это не «минус столько-то», а нарушение нормы
            rows.append(["Что снижает", x["text"],
                         "жёсткое нарушение" if x.get("stop") else -abs(x["delta"]),
                         x.get("norm") or ""])
        for x in prob["plus"]:
            rows.append(["Что в плюс", x["text"], abs(x["delta"]), x.get("norm") or ""])
        for x in prob["how_to_raise"]:
            rows.append(["Что повысит", x["text"], x.get("delta") if x.get("delta") else "", ""])
        if (prob.get("stat") or {}).get("text"):
            rows.append(["Статистика решений", prob["stat"]["text"], "", ""])
        if not prob.get("calibrated"):
            rows.append(["Пометка", "оценка экспертная, не калибрована", "", "правило проекта № 7"])
        if prob.get("decision"):
            rows.append(["Фактическое решение", prob["decision"], prob.get("decided_by") or "",
                         prob.get("decided_at") or ""])
    else:
        rows.append(["Вероятность", prob["text"], "", ""])
    _sheet(wb, "Вероятность", ["Раздел", "Что именно", "Вклад, п.п.", "Норма или примечание"],
           rows, widths=[24, 78, 16, 60])


def build_analysis_xlsx(con, rid: int) -> bytes:
    data = collect(con, rid)
    req, obj, calc = data["request"], data["object"], data["calculation"]
    wb = Workbook()
    wb.remove(wb.active)                      # лист по умолчанию не нужен: свои листы называем сами

    author = data["author"]
    prob = data["probability"]
    _sheet(wb, "Запрос", ["Показатель", "Значение"], [
        ["Номер в системе", rid],
        ["Номер договора", req.get("external_no") or ""],
        ["Филиал", req.get("branch") or ""],
        ["Продукт", f"{data['product']['code']} — {data['product']['name']}" if data["product"]
         else (req.get("product_code") or "")],
        ["Страхователь", req.get("policyholder") or ""],
        ["Подал", (author["full_name"] if author else None)
         or (data["agent"]["name"] if data["agent"] else "")],
        ["Департамент", (author or {}).get("department") or ""],
        ["Должность", (author or {}).get("position") or ""],
        ["Генеральное соглашение", data["partner"] or ""],
        ["Создан", req.get("created_at") or ""],
        ["Состояние согласования", req.get("approval_status") or "не требуется"],
        ["Тип объекта", obj["object_type"] if obj else ""],
        ["Адрес", ", ".join(x for x in ((obj or {}).get("region"), (obj or {}).get("address")) if x)],
        ["Страховая стоимость", (obj or {}).get("value_amount")],
        ["Страховая сумма", (obj or {}).get("sum_insured")],
        ["Вероятность подтверждения, %", prob["probability"] if prob["ready"] else prob["text"]],
        ["Вердикт по вероятности", prob.get("verdict") or ""],
        ["Версия методики вероятности", prob.get("model_version") or ""],
    ], widths=[36, 60])

    _sheet(wb, "Ставка", ["Составляющая", "Действие", "Ставка нарастающим итогом"],
           _rate_rows(data) or [["Расчёт по запросу не сохранён", "", ""]], widths=[62, 24, 26])
    if calc:
        ws = wb["Ставка"]
        ws.append([])
        for label, value in (("Техническая ставка, %", calc.get("gross_rate_pct")),
                             ("Минимум по политике, %", calc.get("min_rate_pct")),
                             ("Применённая ставка, %", calc.get("applied_rate_pct")),
                             ("Премия, сум", calc.get("premium")),
                             ("Вердикт", calc.get("verdict")),
                             ("Версия тарифов (id)", calc.get("tariff_version_id"))):
            ws.append([label, value])

    _sheet(wb, "Проверки", ["Правило", "Норма", "Итог", "Что именно"],
           [[ch["rule_code"], (data["rules"].get(ch["rule_code"]) or {}).get("legal_ref") or "",
             ch["status"], ch["detail"] or ""] for ch in data["checks"]]
           or [["", "", "", "проверки не сохранены"]], widths=[24, 34, 16, 70])

    val = data["valuation"]
    rows = []
    if val:
        rows += [["Заявлено страхователем", "", "", val.get("declared_value"), "", ""],
                 ["Оценка системы", val.get("method") or "", val.get("method_version") or "",
                  val.get("ai_value"), "", val.get("created_at") or ""]]
    for s in data["sources"]:
        rows.append([s["source"], s["status"], s["ads_count"], s["median"], s["url"] or "", s["fetched_at"] or ""])
    _sheet(wb, "Стоимость", ["Источник", "Итог", "Объявлений", "Значение", "Ссылка", "Дата"],
           rows or [["Оценка стоимости не проводилась", "", "", "", "", ""]],
           widths=[28, 22, 14, 20, 60, 20])

    rows = [[d["doc_name"], "получен" if d["received"] else "не получен", d["file_path"] or "", "", ""]
            for d in data["documents"]]
    rows += [[f["filename"] or "", f["doc_kind"], "", f["size_bytes"], f["uploaded_at"] or ""]
             for f in data["files"]]
    _sheet(wb, "Документы", ["Документ", "Вид", "Файл", "Размер, байт", "Загружен"],
           rows or [["документов нет", "", "", "", ""]], widths=[36, 22, 46, 16, 22])

    _sheet(wb, "Решения", ["Рассматривающий", "Должность", "Решение", "Когда", "Примечание", "Назначил"],
           [[r["full_name"], r.get("position") or "", r["status"], r["decided_at"] or "",
             r.get("comment") or "", r.get("assigned_by") or ""] for r in data["reviewers"]]
           or [["рассматривающие не назначены", "", "", "", "", ""]], widths=[30, 24, 16, 22, 40, 20])

    _probability_sheet(wb, prob)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
#  API
# --------------------------------------------------------------------------- #

@router.get("/requests/{rid}/analysis.pdf")
def analysis_pdf(rid: int, user: dict = Depends(current_user)):
    with db.tx() as con:
        check_access(con, user, rid)
        blob = build_analysis_pdf(con, rid)
        db.audit(con, user["login"], "выгрузка анализа PDF", f"request:{rid}", None)
    return Response(content=blob, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="analysis_{rid}.pdf"'})


@router.get("/requests/{rid}/analysis.xlsx")
def analysis_xlsx(rid: int, user: dict = Depends(current_user)):
    with db.tx() as con:
        check_access(con, user, rid)
        blob = build_analysis_xlsx(con, rid)
        db.audit(con, user["login"], "выгрузка анализа XLSX", f"request:{rid}", None)
    return Response(content=blob,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="analysis_{rid}.xlsx"'})


@router.post("/requests/{rid}/analysis/send-telegram")
def analysis_to_telegram(rid: int, format: str = "pdf", user: dict = Depends(current_user)):
    """Присылает выгрузку в чат вошедшего: во встроенном браузере Telegram скачивание часто не работает."""
    fmt = (format or "pdf").strip().lower()
    if fmt not in FORMATS:
        raise HTTPException(422, "format: pdf, xlsx или both")
    if not user.get("telegram_id"):
        raise HTTPException(409, "К вашей учётной записи не привязан Telegram — откройте приложение "
                                 "через бота, тогда файл будет куда прислать")
    from . import tgbot
    sent, errors = [], []
    with db.tx() as con:
        check_access(con, user, rid)
        jobs = []
        if fmt in ("pdf", "both"):
            jobs.append((f"анализ_запроса_{rid}.pdf", build_analysis_pdf(con, rid), "application/pdf"))
        if fmt in ("xlsx", "both"):
            jobs.append((f"анализ_запроса_{rid}.xlsx", build_analysis_xlsx(con, rid),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        for name, blob, mime in jobs:
            res = tgbot.send_file(user["telegram_id"], name, blob, mime,
                                  caption=f"Анализ запроса № {rid}", con=con,
                                  user_id=user["id"], request_id=rid)
            (sent if res.get("ok") else errors).append(name if res.get("ok") else
                                                       f"{name}: {res.get('reason')}")
        db.audit(con, user["login"], "анализ отправлен в Telegram", f"request:{rid}", {"формат": fmt})
    if errors and not sent:
        raise HTTPException(502, "Файл не отправлен: " + "; ".join(errors))
    return {"ok": True, "sent": sent, "errors": errors,
            "message": "Файл отправлен вам в чат с ботом" if sent else "Отправить не удалось"}
