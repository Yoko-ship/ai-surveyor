"""
Финансы компании: отчёты о страховых резервах по периодам (задача заказчика 21.09.2026).

Период — год + «1–4 квартал» или «Годовая отчётность». Данные вносит только админ:
вручную (форма) или файлом Excel по шаблону. Файл сначала разбирается в предпросмотр
(POST /finance/reserves/parse — ничего не сохраняет), сохранение — одним путём для обоих
источников (POST /finance/reserves), строки проверяются заново: клиенту не доверяем.

Состав резервов — Положение 1882 (виды совпадают с колонками reserve_reports):
  РНП — пп. 5 «а», 9–17, по учётным группам п. 10; РЗУ — пп. 5 «г», 21–22; РПНУ — пп. 5 «г», 21, 23;
  стабилизационные ОСГО/ОСГОР/ОСГОП — пп. 5 «в», «в-1», «в2», 24–25⁴; РК — пп. 6, 28;
  прочие (РКУ, резерв премий по жизни и др.) — пп. 5 «б», 6, 29, 31. РПМ компания не ведёт — не принимаем.

Хранение: reserve_periods + reserve_period_lines (db/schema.sql). Замена и удаление не стирают
строки, а меняют status — старая версия остаётся в истории, каждое действие пишется в audit.
Ёмкость (app/capacity.py) берёт последний действующий период через latest_reserve_rows().
"""
import io
import math
import re
from datetime import date
from typing import List, Optional

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel

from . import db
from .auth import require

router = APIRouter()
admin_only = require("админ")

# (код, название, норма, допустимые группы; '' — без группы)
KINDS = [
    ("rnp", "РНП — резерв незаработанной премии", "Положение 1882, пп. 5 «а», 9–17; группы — п. 10",
     ["1", "2", "3", "4"]),
    ("rzu", "РЗУ — резерв заявленных, но неурегулированных убытков", "Положение 1882, пп. 5 «г», 21–22", [""]),
    ("rpnu", "РПНУ — резерв произошедших, но незаявленных убытков",
     "Положение 1882, пп. 5 «г», 21, 23 (не меньше 10% базовой премии за 12 мес.)", [""]),
    ("stab", "Стабилизационный резерв", "Положение 1882, пп. 5 «в», «в-1», «в2», 24–25⁴",
     ["ОСГО", "ОСГОР", "ОСГОП"]),
    ("cat_reserve", "РК — резерв катастроф", "Положение 1882, пп. 6, 28", [""]),
    ("other", "Прочие резервы (РКУ, резерв премий по жизни и др.)", "Положение 1882, пп. 5 «б», 6, 29, 31", [""]),
]
KIND = {k[0]: k for k in KINDS}
GROUPS = {
    "1": "1-я учётная группа",
    "2": "2-я учётная группа",
    "3": "3-я учётная группа",
    "4": "4-я учётная группа",
    "ОСГО": "ОСГО", "ОСГОР": "ОСГОР", "ОСГОП": "ОСГОП", "": "—",
}
GROUP_HINT = {
    "1": "классы 1–13 и 17 (кроме ответственности заёмщиков), непропорциональное перестрахование",
    "2": "классы 14, 15, 16 (кроме урожая), ответственность заёмщиков за непогашение кредитов",
    "3": "договоры с открытыми датами начала и окончания",
    "4": "страхование урожая сельхозкультур (класс 16)",
}
PERIODS = {"Q1": ("1 квартал", "03-31"), "Q2": ("2 квартал", "06-30"), "Q3": ("3 квартал", "09-30"),
           "Q4": ("4 квартал", "12-31"), "Y": ("Годовая отчётность", "12-31")}
YEAR_MIN, YEAR_MAX = 2000, 2100
AMOUNT_MAX = 1e16                  # больше — почти наверняка ошибка ввода (лишние нули)
FILE_MAX = 5 * 1024 * 1024
BASE_LABEL = "Базовая премия за 12 мес. (справочно, в итог не входит)"
ACTIVE, REPLACED, DELETED = "действует", "заменён", "удалён"


# ---------------------------------------------------------------- проверки

def report_date(year: int, ptype: str) -> str:
    return f"{year:04d}-{PERIODS[ptype][1]}"


def check_period(year, ptype) -> str:
    """Возвращает отчётную дату или бросает 400 с понятной причиной."""
    if not isinstance(year, int) or isinstance(year, bool) or not (YEAR_MIN <= year <= YEAR_MAX):
        raise HTTPException(400, f"Год — целое число от {YEAR_MIN} до {YEAR_MAX}")
    if ptype not in PERIODS:
        raise HTTPException(400, "Период — 1, 2, 3, 4 квартал или годовая отчётность (Q1, Q2, Q3, Q4, Y)")
    d = report_date(year, ptype)
    if d > date.today().isoformat():
        raise HTTPException(400, f"Отчётная дата {d[8:10]}.{d[5:7]}.{d[:4]} ещё не наступила — отчёта за этот период быть не может")
    return d


def to_amount(v):
    """Число из ячейки или поля формы. None — пусто; ValueError — не число."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise ValueError("ожидается число, а не да/нет")
    if isinstance(v, (int, float)):
        x = float(v)
    else:
        s = str(v).strip()
        if not s or s in ("-", "—", "–"):
            return None
        s = re.sub(r"[\s\u00a0\u202f\u2009']", "", s).replace("сум", "").replace("сўм", "").replace("som", "")
        if "," in s and "." in s:
            # разделитель тысяч — тот, что стоит первым
            s = s.replace(",", "") if s.index(",") < s.index(".") else s.replace(".", "").replace(",", ".")
        elif "," in s:
            head, _, tail = s.rpartition(",")
            s = f"{head.replace(',', '')}.{tail}" if s.count(",") == 1 and len(tail) <= 2 else s.replace(",", "")
        if not re.fullmatch(r"-?\d+(\.\d+)?", s):
            raise ValueError(f"«{v}» — не число")
        x = float(s)
    if not math.isfinite(x):
        raise ValueError("не число")
    if x < 0:
        raise ValueError("сумма меньше нуля")
    if x >= AMOUNT_MAX:
        raise ValueError("сумма слишком большая — проверьте лишние нули")
    return x


def validate_lines(lines) -> tuple:
    """Строки из формы или предпросмотра: [{reserve_code, group_code, amount}] → (строки, ошибки)."""
    out, errors, seen = [], [], set()
    for i, ln in enumerate(lines or [], 1):
        code, grp = str(ln.get("reserve_code") or ""), str(ln.get("group_code") or "")
        if code not in KIND:
            errors.append(f"строка {i}: неизвестный вид резерва «{code}»")
            continue
        if grp not in KIND[code][3]:
            allowed = ", ".join(GROUPS[g] for g in KIND[code][3])
            errors.append(f"строка {i}: для «{KIND[code][1]}» группа «{grp or '—'}» не подходит; можно: {allowed}")
            continue
        try:
            amt = to_amount(ln.get("amount"))
        except ValueError as e:
            errors.append(f"строка {i} ({KIND[code][1]}, {GROUPS[grp]}): {e}")
            continue
        if amt is None:
            continue
        if (code, grp) in seen:
            errors.append(f"строка {i}: «{KIND[code][1]}, {GROUPS[grp]}» указана дважды")
            continue
        seen.add((code, grp))
        out.append({"reserve_code": code, "group_code": grp, "amount": amt})
    if not out and not errors:
        errors.append("Не введено ни одной суммы")
    return out, errors


def summarize(lines) -> dict:
    by_kind = {k[0]: 0.0 for k in KINDS}
    for ln in lines:
        by_kind[ln["reserve_code"]] += ln["amount"]
    return {"total": sum(by_kind.values()), "by_kind": by_kind}


# ---------------------------------------------------------------- разбор Excel

def _norm(s) -> str:
    s = str(s or "").lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-яўқғҳ]", "", s)


def _col_role(h: str) -> Optional[str]:
    n = _norm(h)
    if not n:
        return None
    if any(x in n for x in ("сумм", "summa", "miqdor", "миқдор", "микдор")):
        return "amount"
    if any(x in n for x in ("групп", "guruh", "гуруҳ", "гурух")):
        return "group"
    if any(x in n for x in ("вид", "резерв", "zaxira", "захира", "rezerv", "наименован", "nomi", "turi")):
        return "kind"
    return None


def detect_kind(text) -> Optional[str]:
    n = _norm(text)
    if not n:
        return None
    if n.startswith(("итого", "всего", "jami", "жами")):
        return "total"
    if "рпм" in n or "предупредит" in n or "ogohlantir" in n:
        return "rpm"
    if "базов" in n or "bazaviy" in n or "базавий" in n:
        return "base"
    if "рпну" in n or "незаявл" in n or "ibnr" in n or "qilinmagan" in n or "қилинмаган" in n:
        return "rpnu"
    if "рзу" in n or "неурегул" in n or "rbns" in n or "tartibgasolinmagan" in n or "тартибгасолинмаган" in n:
        return "rzu"
    if "рнп" in n or "незаработ" in n or "ishlabtopilmagan" in n or "ишлабтопилмаган" in n:
        return "rnp"
    if "стабилиз" in n or "barqaror" in n or "барқарор" in n:
        return "stab"
    if n.startswith("рку") or "колебан" in n or "tebran" in n or "резервпремий" in n \
            or n.startswith(("прочие", "другие", "boshqa", "бошқа")):
        return "other"
    if n == "рк" or "катастроф" in n or "falokat" in n or "фалокат" in n:
        return "cat_reserve"
    return None


_STAB = re.compile(r"(ОСГОР|ОСГОП|ОСГО|OSGOR|OSGOP|OSGO)(?![А-ЯA-Z])")
_RNP_GROUP = re.compile(r"(?<!\d)([1-4])(?!\d)\s*-?\s*(?:я\s*)?(?:уч[её]тн|групп|guruh)", re.I)
_ROMAN = {"I": "1", "II": "2", "III": "3", "IV": "4"}
_WORDS = {"перв": "1", "втор": "2", "трет": "3", "четв": "4", "birinchi": "1", "ikkinchi": "2",
          "uchinchi": "3", "tortinchi": "4"}


def detect_group(text) -> str:
    """'1'..'4', 'ОСГО'/'ОСГОР'/'ОСГОП', '' — пусто, None — не распознано."""
    if text is None:
        return ""
    if isinstance(text, (int, float)) and not isinstance(text, bool):
        return str(int(text)) if float(text) in (1, 2, 3, 4) else None
    s = str(text).strip()
    if not s or s in ("-", "—", "–"):
        return ""
    m = _STAB.search(s.upper())
    if m:
        return {"OSGO": "ОСГО", "OSGOR": "ОСГОР", "OSGOP": "ОСГОП"}.get(m.group(1), m.group(1))
    m = re.match(r"\s*([1-4])(?!\d)", s)
    if m:
        return m.group(1)
    m = re.match(r"\s*(IV|I{1,3})(?![A-Za-z])", s)
    if m:
        return _ROMAN[m.group(1)]
    n = _norm(s)
    for w, g in _WORDS.items():
        if n.startswith(w):
            return g
    return None


def _find_header(rows) -> Optional[tuple]:
    for i, row in enumerate(rows[:15]):
        roles = {}
        for j, v in enumerate(row):
            r = _col_role(v) if isinstance(v, str) else None
            if r and r not in roles:
                roles[r] = j
        if "kind" in roles and "amount" in roles:
            return i, roles
    return None


def parse_workbook(data: bytes, file_name: str = "") -> dict:
    """Разбор файла в предпросмотр. Ничего не сохраняет. Ошибки и предупреждения — списком, не молча."""
    errors, warnings = [], []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Файл не читается как Excel (.xlsx): {type(e).__name__}")
    try:
        sheets = sorted(wb.worksheets, key=lambda ws: 0 if _norm(ws.title) in ("резервы", "zaxiralar", "захиралар") else 1)
        found = None
        for ws in sheets:
            if _norm(ws.title) in ("инструкция", "yoriqnoma", "йўриқнома"):
                continue
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            h = _find_header(rows)
            if h:
                found = (ws.title, rows, h)
                break
    finally:
        wb.close()
    if not found:
        raise HTTPException(400, "В файле не найден заголовок таблицы: нужны колонки «Вид резерва», «Учётная группа» "
                                 "и «Сумма, сум» (как в шаблоне). Скачайте шаблон и заполните его.")
    sheet, rows, (hi, cols) = found
    if "group" not in cols:
        warnings.append("Нет колонки «Учётная группа» — группы берём из названия резерва, если они там есть")

    acc = {}              # (вид, группа) -> {amount, rows, merged}
    base, file_total = None, None
    for r_i in range(hi + 1, len(rows)):
        row = rows[r_i]
        rn = r_i + 1       # номер строки как в Excel
        cell = lambda role: row[cols[role]] if role in cols and cols[role] < len(row) else None  # noqa: E731
        kind_txt, grp_txt, amt_raw = cell("kind"), cell("group"), cell("amount")
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        kind = detect_kind(kind_txt)
        try:
            amt = to_amount(amt_raw)
        except ValueError as e:
            errors.append(f"строка {rn} («{kind_txt or ''}»): {e}")
            continue
        if kind is None:
            if amt:
                warnings.append(f"строка {rn}: «{kind_txt or 'без названия'}» — вид резерва не распознан, "
                                f"сумма {amt:,.0f} не учтена".replace(",", " "))
            continue
        if kind == "total":
            file_total = amt
            continue
        if kind == "rpm":
            if amt:
                warnings.append(f"строка {rn}: РПМ компания не ведёт — сумма {amt:,.0f} не учтена".replace(",", " "))
            continue
        if kind == "base":
            base = amt
            continue
        if amt is None:
            continue           # незаполненная строка шаблона
        grp = detect_group(grp_txt)
        if grp is None:
            errors.append(f"строка {rn}: учётная группа «{grp_txt}» не распознана (нужно 1–4 или ОСГО / ОСГОР / ОСГОП)")
            continue
        if grp == "" and kind == "stab":
            grp = detect_group(kind_txt) or ""
        if grp == "" and kind == "rnp":
            m = _RNP_GROUP.search(str(kind_txt))
            grp = m.group(1) if m else ""
        allowed = KIND[kind][3]
        merged = False
        if grp not in allowed:
            if allowed == [""]:
                warnings.append(f"строка {rn}: для «{KIND[kind][1]}» учётная группа не нужна "
                                f"(группы п. 10 — только для РНП) — сумма добавлена к общему {KIND[kind][1].split(' ')[0]}")
                grp, merged = "", True
            elif grp == "":
                need = "учётную группу 1–4" if kind == "rnp" else "вид: ОСГО, ОСГОР или ОСГОП"
                errors.append(f"строка {rn}: для «{KIND[kind][1]}» укажите {need}")
                continue
            else:
                errors.append(f"строка {rn}: для «{KIND[kind][1]}» группа «{GROUPS.get(grp, grp)}» не подходит")
                continue
        if kind == "other" and (kind, grp) in acc:
            merged = True
            warnings.append(f"строка {rn}: ещё одна строка прочих резервов — суммы сложены")
        key = (kind, grp)
        if key in acc:
            if not (merged and acc[key]["merged"] or kind == "other"):
                errors.append(f"строка {rn}: «{KIND[kind][1]}, {GROUPS[grp]}» уже была в строке {acc[key]['rows'][0]}")
                continue
            acc[key]["amount"] += amt
            acc[key]["rows"].append(rn)
        else:
            acc[key] = {"amount": amt, "rows": [rn], "merged": merged}

    order = {(k[0], g): (i, j) for i, k in enumerate(KINDS) for j, g in enumerate(k[3])}
    lines = [{"reserve_code": k, "group_code": g, "amount": v["amount"], "rows": v["rows"],
              "label": KIND[k][1], "group_label": GROUPS[g]} for (k, g), v in sorted(acc.items(), key=lambda x: order[x[0]])]
    s = summarize(lines)
    if not lines and not errors:
        errors.append("В файле не найдено ни одной суммы резерва")
    if file_total is not None and abs(file_total - s["total"]) > 1:
        warnings.append(f"Итог в файле {file_total:,.0f} не совпадает с суммой строк {s['total']:,.0f} — "
                        "сохраняется сумма строк".replace(",", " "))
    return {"file_name": file_name, "sheet": sheet, "lines": lines, "base_premium_12m": base,
            "total": s["total"], "by_kind": s["by_kind"], "errors": errors, "warnings": warnings}


def template_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Резервы"
    bold = Font(bold=True)
    ws.append(["Вид резерва", "Учётная группа", "Сумма, сум"])
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDE3F5")
    for code, label, _norm_ref, groups in KINDS:
        for g in groups:
            ws.append([label, GROUPS[g], None])
    ws.append([BASE_LABEL, "—", None])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.freeze_panes = "A2"

    ins = wb.create_sheet("Инструкция")
    ins.column_dimensions["A"].width = 110
    text = [
        "Как заполнить шаблон отчёта о страховых резервах",
        "",
        "1. Заполняйте только колонку «Сумма, сум» на листе «Резервы». Суммы — в сумах, не меньше нуля.",
        "   Можно числом или текстом с пробелами: 1 250 000 000 или 1 250 000 000,50.",
        "2. Пустая сумма — резерв не заполнен, в отчёт не попадает. Ноль — резерв есть и равен нулю.",
        "3. Период (год и квартал или годовая отчётность) выбирается на экране перед загрузкой — в файле его нет.",
        "4. Названия строк и колонок не меняйте. Лишние строки допустимы: неизвестные будут показаны",
        "   в предупреждениях и не войдут в итог. Строку «Итого» можно оставить — её сверят с суммой строк.",
        "5. Базовая премия за 12 месяцев — справочно, в итог резервов не входит: по ней проверяется норма",
        "   РПНУ не меньше 10% (Положение 1882, п. 23).",
        "",
        "Состав резервов (Положение 1882):",
    ]
    text += [f"   • {label} — {ref}" for _c, label, ref, _g in KINDS]
    text += ["", "Учётные группы для РНП (Положение 1882, п. 10):"]
    text += [f"   • {GROUPS[g]}: {GROUP_HINT[g]}" for g in ("1", "2", "3", "4")]
    text += ["", "Стабилизационные резервы — отдельно по ОСГО, ОСГОР и ОСГОП (пп. 24–25⁴).",
             "РПМ (резерв предупредительных мероприятий) компания не ведёт — строку для него не добавляйте."]
    for t in text:
        ins.append([t])
    ins["A1"].font = Font(bold=True, size=13)
    for row in ins.iter_rows():
        row[0].alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- хранение

def _period_out(p: dict) -> dict:
    p = dict(p)
    p["period_label"] = PERIODS.get(p["period_type"], (p["period_type"],))[0]
    return p


def _lines(con, pid: int) -> list:
    out = db.rows(con, "SELECT reserve_code, group_code, amount FROM reserve_period_lines WHERE period_id=?", pid)
    order = {(k[0], g): (i, j) for i, k in enumerate(KINDS) for j, g in enumerate(k[3])}
    out.sort(key=lambda x: order.get((x["reserve_code"], x["group_code"]), (99, 99)))
    for ln in out:
        ln["label"] = KIND[ln["reserve_code"]][1] if ln["reserve_code"] in KIND else ln["reserve_code"]
        ln["group_label"] = GROUPS.get(ln["group_code"], ln["group_code"])
    return out


def active_period(con, year: int, ptype: str) -> Optional[dict]:
    r = db.rows(con, "SELECT * FROM reserve_periods WHERE period_year=? AND period_type=? AND status=?",
                year, ptype, ACTIVE)
    return r[0] if r else None


def save_period(con, who: str, year, ptype, lines, base_premium=None, source="вручную",
                file_name=None, replace=False) -> dict:
    rdate = check_period(year, ptype)
    if source not in ("вручную", "файл"):
        raise HTTPException(400, "Источник — «вручную» или «файл»")
    good, errors = validate_lines(lines)
    try:
        base = to_amount(base_premium)
    except ValueError as e:
        errors.append(f"базовая премия за 12 мес.: {e}")
        base = None
    if errors:
        raise HTTPException(400, {"message": "Отчёт не сохранён: исправьте ошибки", "errors": errors})
    s = summarize(good)
    old = active_period(con, year, ptype)
    if old and not replace:
        raise HTTPException(409, {"message": f"За {PERIODS[ptype][0].lower()} {year} года отчёт уже есть. Заменить?",
                                  "existing": _period_out(old)})
    ts = db.now()
    if old:
        # сначала снимаем старую версию: уникальный индекс держит один действующий отчёт на период
        con.execute("UPDATE reserve_periods SET status=?, closed_by=?, closed_at=? WHERE id=?",
                    (REPLACED, who, ts, old["id"]))
    cur = con.execute(
        "INSERT INTO reserve_periods (period_year, period_type, report_date, total, base_premium_12m, source,"
        " file_name, status, created_by, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
        (year, ptype, rdate, s["total"], base, source, (file_name or None) and str(file_name)[:200], ACTIVE, who, ts))
    pid = cur.lastrowid
    con.executemany("INSERT INTO reserve_period_lines (period_id, reserve_code, group_code, amount) VALUES (?,?,?,?)",
                    [(pid, ln["reserve_code"], ln["group_code"], ln["amount"]) for ln in good])
    entity = f"резервы:{year}-{ptype}"
    detail = {"id": pid, "total": s["total"], "source": source, "file": file_name, "lines": len(good)}
    if old:
        con.execute("UPDATE reserve_periods SET replaced_by=? WHERE id=?", (pid, old["id"]))
        db.audit(con, who, "резервы: отчёт заменён", entity, dict(detail, old_id=old["id"], old_total=old["total"]))
    else:
        db.audit(con, who, "резервы: отчёт внесён", entity, detail)
    return {"ok": True, "id": pid, "replaced": old["id"] if old else None, "report_date": rdate, **s}


def latest_reserve_rows(con) -> tuple:
    """Последний действующий период в форме строк reserve_reports: (период, строки) или (None, [])."""
    # при одной дате 31.12 годовая отчётность ('Y') важнее 4 квартала: 'Y' > 'Q4'
    try:
        p = db.rows(con, "SELECT * FROM reserve_periods WHERE status=? ORDER BY report_date DESC, period_type DESC, id DESC"
                         " LIMIT 1", ACTIVE)
    except Exception as e:
        # база, на которой ещё не прогнан schema.sql (старые скрипты, копии): периодов просто нет
        if "reserve_periods" not in str(e):
            raise
        return None, []
    if not p:
        return None, []
    p = _period_out(p[0])
    cols = ("rnp", "rzu", "rpnu", "stab", "cat_reserve", "other")
    blank = lambda st, sc: {"report_date": p["report_date"], "scope_type": st, "scope_code": sc,  # noqa: E731
                            **{c: 0.0 for c in cols}, "base_premium_12m": None, "source": f"финансы: период {p['id']}"}
    total = blank("итого", "итого")
    total["base_premium_12m"] = p["base_premium_12m"]
    scoped = {}
    for ln in _lines(con, p["id"]):
        total[ln["reserve_code"]] += ln["amount"]
        g = ln["group_code"]
        if g in ("1", "2", "3", "4"):
            key = ("группа", g)
        elif g:
            key = ("вид", g)
        else:
            continue          # без группы — только в итоговой строке
        scoped.setdefault(key, blank(*key))[ln["reserve_code"]] += ln["amount"]
    return p, [scoped[k] for k in sorted(scoped)] + [total]


# ---------------------------------------------------------------- маршруты

class LineIn(BaseModel):
    reserve_code: str
    group_code: str = ""
    amount: Optional[float] = None


class PeriodIn(BaseModel):
    period_year: int
    period_type: str
    lines: List[LineIn]
    base_premium_12m: Optional[float] = None
    source: str = "вручную"
    file_name: Optional[str] = None
    replace: bool = False


@router.get("/finance/reserves/structure")
def get_structure(user: dict = Depends(admin_only)):
    """Состав формы: виды резервов, группы, периоды — экран строит форму отсюда, а не держит копию."""
    return {"kinds": [{"code": c, "label": l, "norm": n, "groups": [{"code": g, "label": GROUPS[g],
                                                                        "hint": GROUP_HINT.get(g, "")} for g in gs]}
                      for c, l, n, gs in KINDS],
            "periods": [{"code": k, "label": v[0], "report_day": v[1]} for k, v in PERIODS.items()],
            "year_min": YEAR_MIN, "year_max": YEAR_MAX, "base_label": BASE_LABEL}


@router.get("/finance/reserves/template.xlsx")
def get_template(user: dict = Depends(admin_only)):
    return Response(template_bytes(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''reserves_template.xlsx"})


@router.post("/finance/reserves/parse")
async def post_parse(file: UploadFile = File(...), user: dict = Depends(admin_only)):
    name = (file.filename or "файл.xlsx").replace("\\", "/").rsplit("/", 1)[-1]
    data = await file.read(FILE_MAX + 1)
    if len(data) > FILE_MAX:
        raise HTTPException(400, "Файл больше 5 МБ — для отчёта о резервах это слишком много; проверьте, тот ли файл")
    if not name.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Нужен файл Excel (.xlsx). Старый формат .xls сохраните как .xlsx.")
    if data[:2] != b"PK":
        raise HTTPException(400, "Файл не похож на xlsx (нет сигнатуры PK) — возможно, это .xls или страница сайта.")
    return parse_workbook(data, name)


@router.post("/finance/reserves")
def post_period(body: PeriodIn, user: dict = Depends(admin_only)):
    with db.tx() as con:
        return save_period(con, user["login"], body.period_year, body.period_type,
                           [ln.model_dump() for ln in body.lines], body.base_premium_12m, body.source,
                           body.file_name, body.replace)


@router.get("/finance/reserves")
def list_periods(history: int = 0, user: dict = Depends(admin_only)):
    with db.tx() as con:
        sql = "SELECT * FROM reserve_periods" + ("" if history else " WHERE status=?") + \
              " ORDER BY period_year DESC, report_date DESC, period_type DESC, id DESC"
        items = db.rows(con, sql, *(() if history else (ACTIVE,)))
        latest, _ = latest_reserve_rows(con)
    return {"items": [_period_out(p) for p in items], "latest_id": latest["id"] if latest else None}


@router.get("/finance/reserves/{pid:int}")
def get_period(pid: int, user: dict = Depends(admin_only)):
    with db.tx() as con:
        p = db.rows(con, "SELECT * FROM reserve_periods WHERE id=?", pid)
        if not p:
            raise HTTPException(404, "Отчёт не найден")
        p = _period_out(p[0])
        p["lines"] = _lines(con, pid)
        p["by_kind"] = summarize(p["lines"])["by_kind"]
        p["history"] = [_period_out(h) for h in db.rows(
            con, "SELECT * FROM reserve_periods WHERE period_year=? AND period_type=? ORDER BY id DESC",
            p["period_year"], p["period_type"])]
    return p


@router.delete("/finance/reserves/{pid:int}")
def delete_period(pid: int, user: dict = Depends(admin_only)):
    with db.tx() as con:
        p = db.rows(con, "SELECT * FROM reserve_periods WHERE id=?", pid)
        if not p:
            raise HTTPException(404, "Отчёт не найден")
        p = p[0]
        if p["status"] != ACTIVE:
            raise HTTPException(409, f"Эта версия уже {p['status']} — удалить можно только действующий отчёт")
        con.execute("UPDATE reserve_periods SET status=?, closed_by=?, closed_at=? WHERE id=?",
                    (DELETED, user["login"], db.now(), pid))
        db.audit(con, user["login"], "резервы: отчёт удалён", f"резервы:{p['period_year']}-{p['period_type']}",
                 {"id": pid, "total": p["total"], "source": p["source"]})
    return {"ok": True, "id": pid}
