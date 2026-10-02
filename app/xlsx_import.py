"""
Общее для импорта справочников из Excel в два шага: «предпросмотр → применить».

  * check_upload — размер и сигнатура PK (xlsx — это zip; правило проекта № 10);
  * read_sheet   — первая таблица листа: строка заголовка по известным названиям колонок;
  * Previews     — хранилище предпросмотров по токену: 30 минут, один раз, только тому же администратору.
                   Хранится в памяти процесса: сервер работает одним процессом (Procfile, --workers 1).
                   При применении строки проверяются ЗАНОВО по текущей базе — между предпросмотром
                   и «применить» справочник мог измениться.
  * template     — шаблон с заголовком, примером и листом «Инструкция».
"""
import io
import secrets
import threading
import time
from datetime import date, datetime
from typing import Callable, Optional

from fastapi import HTTPException

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TOKEN_TTL_SEC = 30 * 60
TOKENS_MAX = 50
MAX_COLS = 64
HEAD_ROWS = 10          # строку заголовка ищем в первых строках листа


async def check_upload(file, limit: int) -> bytes:
    data = await file.read(limit + 1)
    if len(data) > limit:
        raise HTTPException(413, f"Файл больше {limit // (1024 * 1024)} МБ")
    if not data:
        raise HTTPException(400, "Файл пустой")
    if data[:2] != b"PK":
        raise HTTPException(400, "Нужен файл Excel (.xlsx)")
    return data


def norm_head(v) -> str:
    return " ".join(str(v or "").replace("\n", " ").split()).strip().lower()


def read_sheet(data: bytes, match: Callable[[str], Optional[str]], required: set, max_rows: int,
               head_hint: str) -> dict:
    """
    Строки первого листа. Вызывать в пуле потоков (run_in_threadpool): разбор xlsx — работа процессора,
    в цикле событий он остановил бы остальные запросы. Пределы: max_rows строк листа (пустые строки
    тоже считаются — иначе файл из миллиона пустых строк разбирался бы до конца) и MAX_COLS колонок;
    размер листа сначала берётся из его заголовка (dimension), до чтения строк. Возвращает
    {"rows": [{"row": № строки, "values": {ключ: значение}}], "columns": {ключ: заголовок},
     "unknown": [заголовки, которые не узнали]}.
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(422, "Файл не читается как Excel (.xlsx): " + type(e).__name__)
    try:
        ws = wb.worksheets[0]
        limit_rows = max_rows + HEAD_ROWS
        if (ws.max_row or 0) > limit_rows:
            raise HTTPException(422, f"Строк больше {max_rows} — разделите файл")
        if (ws.max_column or 0) > MAX_COLS:
            raise HTTPException(422, f"Колонок больше {MAX_COLS} — оставьте колонки шаблона")
        head, names, unknown, rows = None, {}, [], []
        for i, r in enumerate(ws.iter_rows(values_only=True, max_col=MAX_COLS), start=1):
            if i > limit_rows:         # заголовок листа мог не указать размер или солгать
                raise HTTPException(422, f"Строк больше {max_rows} — разделите файл")
            if head is None:
                m, unk = {}, []
                for j, c in enumerate(r):
                    h = norm_head(c)
                    if not h:
                        continue
                    key = match(h)
                    if key and key not in m.values():
                        m[j] = key
                        names[key] = str(c).strip()
                    else:
                        unk.append(str(c).strip())
                if required <= set(m.values()):
                    head, unknown = m, unk
                elif i >= HEAD_ROWS:
                    break
                else:
                    names = {}
                continue
            if not any(c not in (None, "") for c in r):
                continue
            if len(rows) >= max_rows:
                raise HTTPException(422, f"Строк больше {max_rows} — разделите файл")
            rows.append({"row": i, "values": {k: (r[j] if j < len(r) else None) for j, k in head.items()}})
    finally:
        wb.close()
    if head is None:
        raise HTTPException(422, "Не найдена строка заголовка. Нужны колонки: " + head_hint)
    return {"rows": rows, "columns": names, "unknown": unknown}


# ---------- значения ячеек ----------

def as_text(v, limit: int = 200) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return " ".join(str(v).split())[:limit]


def as_number(v) -> tuple:
    """(число | None, ошибка | None). Пусто — (None, None). «1 500 000,50» → 1500000.5."""
    if v is None or v == "":
        return None, None
    if isinstance(v, bool):
        return None, "не число"
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).replace(" ", "").replace(" ", "").replace("%", "").replace(",", ".").strip()
    if not s:
        return None, None
    try:
        return float(s), None
    except ValueError:
        return None, "не число"


def as_date(v) -> tuple:
    """(date | None, ошибка | None). Принимает дату Excel, ГГГГ-ММ-ДД, ДД.ММ.ГГГГ, ДД/ММ/ГГГГ."""
    if v is None or v == "":
        return None, None
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    return None, "дата ДД.ММ.ГГГГ или ГГГГ-ММ-ДД"


def product_code(v) -> str:
    """Код продукта: 832 (число из Excel) → «0832», как в справочнике."""
    if isinstance(v, bool):
        return ""
    if isinstance(v, (int, float)) and float(v) == int(v):
        return str(int(v)).zfill(4)
    return str(v or "").strip()[:20]


# ---------- предпросмотры ----------

class Previews:
    def __init__(self, kind: str):
        self.kind = kind
        self._d = {}
        self._lock = threading.Lock()

    def _gc(self):
        now = time.monotonic()
        for k in [k for k, v in self._d.items() if now - v["t"] > TOKEN_TTL_SEC]:
            self._d.pop(k, None)
        while len(self._d) >= TOKENS_MAX:
            self._d.pop(min(self._d, key=lambda k: self._d[k]["t"]))

    def put(self, who: str, payload: dict) -> str:
        token = secrets.token_urlsafe(18)
        with self._lock:
            self._gc()
            self._d[token] = {"t": time.monotonic(), "who": who, "payload": payload}
        return token

    def take(self, token: str, who: str) -> dict:
        """Забирает предпросмотр (одноразово). Чужой, устаревший или неизвестный токен — 404."""
        with self._lock:
            self._gc()
            v = self._d.get(token or "")
            if not v or v["who"] != who:
                raise HTTPException(404, "Предпросмотр не найден или устарел (30 минут) — загрузите файл заново")
            self._d.pop(token, None)
        return v["payload"]


# ---------- шаблон ----------

def template(sheet: str, head: list, example: list, instructions: list, widths: Optional[list] = None) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3F5")
    if example:
        ws.append(example)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("Инструкция")
    ins.column_dimensions["A"].width = 120
    for line in instructions:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
